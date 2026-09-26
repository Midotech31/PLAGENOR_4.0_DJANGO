from __future__ import annotations

import argparse
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document


NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def _shared_strings(book: ZipFile) -> list[str]:
    try:
        root = ET.fromstring(book.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    out = []
    for si in root.findall("m:si", NS):
        out.append("".join(node.text or "" for node in si.iterfind(".//m:t", NS)))
    return out


def _cell_value(cell, shared):
    kind = cell.get("t")
    if kind == "inlineStr":
        return "".join(node.text or "" for node in cell.iterfind(".//m:t", NS))
    value = cell.find("m:v", NS)
    raw = value.text if value is not None else ""
    if kind == "s":
        return shared[int(raw)] if raw else ""
    if kind == "b":
        return raw == "1"
    if kind in ("str", "e"):
        return raw
    if raw == "":
        return ""
    try:
        number = float(raw)
    except ValueError:
        return raw
    return int(number) if number.is_integer() else number


def xlsx_rows(path: Path) -> dict[str, list[dict]]:
    with ZipFile(path) as book:
        shared = _shared_strings(book)
        wb = ET.fromstring(book.read("xl/workbook.xml"))
        rels = ET.fromstring(book.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.get("Id"): rel.get("Target") for rel in rels.findall("pr:Relationship", NS)}
        result = {}
        for sheet in wb.findall("m:sheets/m:sheet", NS):
            name = sheet.get("name")
            target = rel_map[sheet.get(f"{{{NS['r']}}}id")]
            part = target if target.startswith("xl/") else "xl/" + target.lstrip("/")
            root = ET.fromstring(book.read(part))
            rows = []
            for row in root.findall("m:sheetData/m:row", NS):
                cells = {}
                for cell in row.findall("m:c", NS):
                    ref = cell.get("r", "")
                    match = re.match(r"([A-Z]+)", ref)
                    if not match:
                        continue
                    col = 0
                    for ch in match.group(1):
                        col = col * 26 + ord(ch) - 64
                    cells[col] = _cell_value(cell, shared)
                if cells:
                    width = max(cells)
                    values = [cells.get(i, "") for i in range(1, width + 1)]
                else:
                    values = []
                if any(value not in ("", None) for value in values):
                    rows.append({"source_row": int(row.get("r")), "values": values})
            result[name] = rows
        return result


def room_rows(path: Path):
    room_files = []
    rows = []
    with ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.lower().endswith(".docx"))
        for name in names:
            payload = archive.read(name)
            doc = Document(BytesIO(payload))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            label = paragraphs[1] if len(paragraphs) > 1 else (paragraphs[0] if paragraphs else "")
            match = re.search(r"(?:room|salle)\s*0*([0-9]{1,2})", label, re.I)
            room = f"Room {int(match.group(1)):02d}" if match else ""
            filename = Path(name).name
            room_files.append({
                "filename": filename,
                "sha256": digest(payload),
                "room": room,
                "label": label,
            })
            for table_number, table in enumerate(doc.tables, 1):
                for row_number, row in enumerate(table.rows[1:], 2):
                    values = [cell.text.strip() for cell in row.cells]
                    if not any(values):
                        continue
                    rows.append({
                        "source_file": filename,
                        "room": room,
                        "room_label": label,
                        "table": table_number,
                        "source_row": row_number,
                        "values": values,
                    })
    return room_files, rows


def build(workbook: Path, room_archive: Path):
    wb_bytes = workbook.read_bytes()
    archive_bytes = room_archive.read_bytes()
    room_files, rooms = room_rows(room_archive)
    return {
        "schema": 2,
        "title": "Inventaire PLAGENOR 2026 — source institutionnelle de reprise",
        "sources": {
            "workbook": {"filename": workbook.name, "sha256": digest(wb_bytes)},
            "room_archive": {
                "filename": room_archive.name,
                "sha256": digest(archive_bytes),
                "files": room_files,
            },
        },
        "xlsx": xlsx_rows(workbook),
        "room_lists": rooms,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("room_archive", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    data = build(args.workbook, args.room_archive)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(f"Wrote {args.output}")
    print({name: max(len(rows) - 1, 0) for name, rows in data["xlsx"].items()})
    print(f"room files={len(data['sources']['room_archive']['files'])} rows={len(data['room_lists'])}")


if __name__ == "__main__":
    main()
