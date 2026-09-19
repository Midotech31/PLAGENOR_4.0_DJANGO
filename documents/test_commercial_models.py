import hashlib
import io
import os
from pathlib import Path
import shutil
import tempfile
import zipfile
from decimal import Decimal
from unittest.mock import patch

from docx import Document
from django.test import TestCase, SimpleTestCase, override_settings
from django.utils import timezone

from accounts.models import User
from core.commercial import document_identity
from core.exceptions import FinancialValidationError
from core.financial import compute_invoice_totals
from core.models import Invoice, Request, Service, PlatformContent
from documents.generators import generate_quote, generate_invoice_document
from documents.genoclab_layout import (
    ESSBO_LOGO, _GENOCLAB_LOGO, add_prestation_table, amount_in_words_fr,
)


def document_text(path):
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    def tables(items):
        for table in items:
            for row in table.rows:
                for cell in row.cells:
                    chunks.append(cell.text)
                    tables(cell.tables)
    tables(document.tables)
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
        tables(section.footer.tables)
    return '\n'.join(chunks)


class CommercialDocumentModelsTests(TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(prefix='plagenor-commercial-')
        self.addCleanup(self.folder.cleanup)
        storage = override_settings(MEDIA_ROOT=self.folder.name, DOCUMENT_PDF_ENABLED=False)
        storage.enable(); self.addCleanup(storage.disable)
        self.user = User.objects.create_user('document-client', role='CLIENT',
                    first_name='Client', last_name='Fictif', email='client@example.test',
                    organization='Établissement de recette', phone='+213 000 000 000')
        self.service = Service.objects.create(code='MODEL-TEST', name='Analyse de recette')
        self.items = [{'label':'Identification microbienne — prestation de recette',
                       'quantity':3,'unit_price':'2500.00','total':'7500.00'},
                      {'label':'Contrôle qualité des acides nucléiques',
                       'quantity':2,'unit_price':'1250.25','total':'2500.50'}]

    def make_request(self, channel, suffix=''):
        self.assertLessEqual(len('MDL-' + channel + suffix), Request._meta.get_field('display_id').max_length)
        request = Request.objects.create(
            display_id='MDL-' + channel + suffix, title='Prestation de recette',
            requester=self.user, service=self.service, channel='GENOCLAB',
            billing_channel=channel, status='QUOTE_DRAFT',
        )
        rate = 0 if channel == 'OHB' else 0.19
        request.quote_detail = {'items': self.items, 'vat_rate':rate, 'admin_fees':0, 'report_fees':0,
                                'identity':document_identity(request)}
        request.quote_amount = compute_invoice_totals(self.items, vat_rate=rate)['total_ttc']
        request.save()
        return request

    def keep_preview(self, path, name):
        target = os.environ.get('PLAGENOR_DOCUMENT_PREVIEW_DIR')
        if target:
            destination = Path(target); destination.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, destination/name)

    def test_quote_and_invoice_share_model_with_correct_tax_and_logo(self):
        for channel in ('OHB', 'GENOCLAB'):
            with self.subTest(channel=channel):
                req = self.make_request(channel)
                rate = Decimal('0') if channel == 'OHB' else Decimal('0.19')
                totals = compute_invoice_totals(self.items, vat_rate=rate)
                quote = generate_quote(req)
                invoice = Invoice.objects.create(request=req, client=self.user,
                    invoice_number='TEST-' + channel + '-001', line_items=self.items,
                    document_snapshot=document_identity(req), subtotal_ht=totals['subtotal_before_tax'],
                    vat_rate=rate, vat_amount=totals['vat_amount'], total_ttc=totals['total_ttc'])
                bill = generate_invoice_document(invoice)
                for kind, path in (('DEVIS',quote), ('FACTURE',bill)):
                    text = document_text(path)
                    self.assertIn('Client Fictif',text)
                    self.assertIn('Établissement de recette',text)
                    self.assertIn('7 500,00',text)
                    self.assertIn('2 500,50',text)
                    self.assertIn('Montant DA',text)
                    self.assertNotIn('/100', text)
                    expected_words = ('dix mille dinars algériens et cinquante centimes' if channel == 'OHB'
                                      else 'onze mille neuf cents dinars algériens et soixante centimes')
                    self.assertIn(expected_words, text)
                    self.assertIn('Siège social',text)
                    self.assertIn('Coordonnées',text)
                    self.assertNotIn('_____ ',text)
                    with zipfile.ZipFile(path) as archive:
                        hashes = {hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('word/media/')}
                    expected = ESSBO_LOGO if channel == 'OHB' else _GENOCLAB_LOGO
                    other = _GENOCLAB_LOGO if channel == 'OHB' else ESSBO_LOGO
                    self.assertIn(hashlib.sha256(expected.read_bytes()).hexdigest(),hashes)
                    self.assertNotIn(hashlib.sha256(other.read_bytes()).hexdigest(),hashes)
                    if channel == 'OHB':
                        self.assertIn('10 000,50',text)
                        self.assertIn('Non assujetti à la TVA',text)
                        self.assertNotIn('TVA (0 %)',text)
                        self.assertNotIn('TVA (19 %)',text)
                    else:
                        self.assertIn('TVA (19 %)',text)
                        self.assertIn('1 900,10',text)
                        self.assertIn('11 900,60',text)
                        self.assertNotIn('Non assujetti',text)
                    self.keep_preview(path, channel + '_' + kind + '_EXEMPLE_FICTIF.docx')
                self.assertIn('Arrêté le présent devis',document_text(quote))
                self.assertIn('Validité du devis : 30 jours',document_text(quote))
                self.assertIn('Arrêtée la présente facture',document_text(bill))
                self.assertIn('Validité de la facture : 30 jours',document_text(bill))

    def test_exact_fractional_total_in_all_commercial_documents(self):
        for channel in ('OHB', 'GENOCLAB'):
            with self.subTest(channel=channel):
                price = '11900.50' if channel == 'OHB' else '10000.42'
                self.items = [{'label': 'Prestation de recette', 'quantity': 1,
                               'unit_price': price, 'total': price}]
                req = self.make_request(channel, '-WORDS')
                rate = Decimal('0') if channel == 'OHB' else Decimal('0.19')
                totals = compute_invoice_totals(self.items, vat_rate=rate)
                self.assertEqual(Decimal(str(totals['total_ttc'])), Decimal('11900.50'))
                invoice = Invoice.objects.create(request=req, client=self.user,
                    invoice_number='TEST-WORDS-' + channel, line_items=self.items,
                    document_snapshot=document_identity(req), subtotal_ht=totals['subtotal_before_tax'],
                    vat_rate=rate, vat_amount=totals['vat_amount'], total_ttc=totals['total_ttc'])
                quote = generate_quote(req)
                bill = generate_invoice_document(invoice)
                for path in (quote, bill):
                    text = document_text(path)
                    self.assertIn('onze mille neuf cents dinars algériens et cinquante centimes', text)
                    self.assertIn('11 900,50', text)
                    self.assertNotIn('/100', text)
                if channel == 'GENOCLAB':
                    self.keep_preview(bill, 'GENOCLAB_FACTURE_11900_50_EXEMPLE_FICTIF.docx')

    def test_ohb_missing_explicit_rate_does_not_default_to_nineteen_percent(self):
        req = self.make_request('OHB', '-DEFAULT')
        req.quote_detail.pop('vat_rate')
        req.save()
        text = document_text(generate_quote(req))
        self.assertIn('10 000,50',text)
        self.assertNotIn('TVA (19 %)',text)

    def test_missing_invoice_snapshot_preserves_request_billing_identity(self):
        req = self.make_request('OHB', '-LEGACY')
        totals = compute_invoice_totals(self.items, vat_rate=0)
        invoice = Invoice.objects.create(request=req, client=self.user, invoice_number='TEST-OHB-LEGACY',
                    line_items=self.items, vat_rate=0, subtotal_ht=totals['subtotal_ht'], total_ttc=totals['total_ttc'])
        text = document_text(generate_invoice_document(invoice))
        self.assertIn('Non assujetti',text)

    def test_issued_invoice_original_and_snapshot_are_immutable(self):
        req = self.make_request('GENOCLAB', '-FROZEN')
        identity = document_identity(req)
        invoice = Invoice.objects.create(request=req, client=self.user, invoice_number='TEST-FROZEN',
                    line_items=self.items, vat_rate='0.19', document_snapshot=identity)
        path = generate_invoice_document(invoice)
        before = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        self.user.first_name='Changed'; self.user.save()
        PlatformContent.objects.update_or_create(key='genoclab_issuer_name',lang='fr', defaults={'value':'Changed issuer'})
        second = generate_invoice_document(invoice)
        self.assertEqual(hashlib.sha256(Path(second).read_bytes()).hexdigest(),before)
        self.assertIn('Client Fictif',document_text(second))
        self.assertNotIn('Changed issuer',document_text(second))

    def test_long_table_arabic_client_and_french_accents_remain_intact(self):
        self.user.first_name='عميل'; self.user.last_name='تجريبي'; self.user.save()
        req = self.make_request('OHB', '-LONG')
        req.quote_detail['identity']=document_identity(req)
        req.quote_detail['items']=[{'label':f'Prestation {i:02d} — contrôle qualité et séquençage',
                                    'quantity':i,'unit_price':'12.34','total':str(Decimal(i)*Decimal('12.34'))}
                                   for i in range(1,31)]
        req.save(); path=generate_quote(req)
        text=document_text(path)
        self.assertIn('عميل تجريبي',text)
        self.assertIn('Prestation 30',text)
        with zipfile.ZipFile(path) as archive:
            xml=archive.read('word/document.xml')
            self.assertIn(b'w:tblHeader',xml)
        self.keep_preview(path,'OHB_DEVIS_30_LIGNES_EXEMPLE_FICTIF.docx')


class CommercialArithmeticTests(SimpleTestCase):
    def test_french_money_words(self):
        for amount, words in [(71,'soixante et onze dinars algériens'),
                              (80,'quatre-vingts dinars algériens'),
                              (80000,'quatre-vingt mille dinars algériens'),
                              (280000,'deux cent quatre-vingt mille dinars algériens'),
                              (200000000,'deux cents millions de dinars algériens'),
                              (1000.5,'mille dinars algériens et cinquante centimes')]:
            self.assertEqual(amount_in_words_fr(amount),words)
        self.assertEqual(amount_in_words_fr(10**15),'')

    def test_invalid_amounts_and_ohb_tax_mismatch_fail_closed(self):
        for row in ({'quantity':-1}, {'unit_price':'NaN'}, {'total':'Infinity'}):
            with self.assertRaises(FinancialValidationError):
                add_prestation_table(Document(),[row],vat_rate=0.19)
        with self.assertRaises(FinancialValidationError):
            add_prestation_table(Document(),[],vat_rate=0.19,non_taxable=True)
