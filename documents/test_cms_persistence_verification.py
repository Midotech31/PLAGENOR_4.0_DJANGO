import io
import tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.base import ContentFile
from django.core.files.storage import Storage
from lxml import html
from accounts.models import User
from core.models import Service, Request, ServiceFormField
from documents.models import ServiceTemplate, DocumentBlock
from documents.generators import _get_uploaded_template
from documents.views import _cached_doc_path, _block_signature, _service_fields_signature


@override_settings(SECURE_SSL_REDIRECT=False, RATE_LIMIT_BACKEND='cache',
    EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
              'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class DocumentPersistenceVerificationTests(TestCase):
    def setUp(self):
        folder=tempfile.TemporaryDirectory();self.addCleanup(folder.cleanup)
        setting=override_settings(MEDIA_ROOT=folder.name);setting.enable();self.addCleanup(setting.disable)
        self.admin=User.objects.create_user(username='document-verifier',role='SUPER_ADMIN',password=None)
        self.client.force_login(self.admin)
        self.service=Service.objects.create(code='AUDIT-DOCUMENT',name_fr='Documents',name_en='Documents',name_ar='وثائق')
        self.req=Request.objects.create(display_id='AUDIT-001',title='Test',channel='IBTIKAR',service=self.service)

    def upload(self,text='Stored template'):
        output=io.BytesIO();doc=Document();doc.add_paragraph(text);doc.save(output)
        return SimpleUploadedFile('template.docx',output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    def create(self,name):
        response=self.client.post('/documents/templates/create/',{'service':str(self.service.pk),'template_type':'PLATFORM_NOTE','name':name,'file':self.upload(name)})
        self.assertEqual(response.status_code,302,response.content[:300])
        return ServiceTemplate.objects.get(name=name)

    def test_three_versions_preserve_history_and_can_be_reactivated(self):
        versions=[self.create('Version '+str(i)) for i in range(3)]
        self.assertEqual(ServiceTemplate.objects.filter(is_active=True).count(),1)
        self.assertEqual(ServiceTemplate.objects.filter(is_active=False).count(),2)
        for item in versions:
            item.refresh_from_db()
            with item.file.open('rb') as source:self.assertEqual(Document(source).paragraphs[0].text,item.name)
        self.client.post('/documents/templates/'+str(versions[0].pk)+'/toggle/')
        versions[0].refresh_from_db();self.assertTrue(versions[0].is_active)
        self.assertEqual(ServiceTemplate.objects.filter(is_active=True).count(),1)

    def test_invalid_upload_keeps_values_and_existing_active_template(self):
        previous=self.create('Previous')
        response=self.client.post('/documents/templates/create/',{'service':str(self.service.pk),'template_type':'PLATFORM_NOTE','name':'Retained title','description':'Retained description','file':SimpleUploadedFile('bad.docx',b'not a docx')})
        self.assertEqual(response.status_code,400)
        page=html.fromstring(response.content)
        self.assertEqual(page.xpath('//input[@name="name"]/@value'),['Retained title'])
        self.assertEqual(page.xpath('//textarea[@name="description"]/text()'),['Retained description'])
        previous.refresh_from_db();self.assertTrue(previous.is_active);self.assertEqual(ServiceTemplate.objects.count(),1)

    def test_changing_service_type_and_name_is_persisted_and_reopened(self):
        obj=self.create('Original');other=Service.objects.create(code='AUDIT-OTHER',name='Other')
        response=self.client.post('/documents/templates/'+str(obj.pk)+'/edit/',{'service':str(other.pk),'template_type':'RECEPTION_FORM','name':'Moved','description':'Updated','is_active':'on'})
        self.assertEqual(response.status_code,302);obj.refresh_from_db()
        self.assertEqual(obj.service_id,other.pk);self.assertEqual(obj.template_type,'RECEPTION_FORM')
        self.client.logout();self.client.force_login(self.admin)
        self.assertContains(self.client.get('/documents/templates/'+str(obj.pk)+'/edit/'),'Moved')
        self.client.post('/documents/templates/'+str(obj.pk)+'/delete/')
        self.assertFalse(ServiceTemplate.objects.filter(pk=obj.pk).exists())

    def test_cache_invalidates_on_template_revision_and_older_block_deletion(self):
        first=self.create('First');before=_cached_doc_path(self.req,'PLATFORM_NOTE')
        self.create('Second');self.assertNotEqual(before,_cached_doc_path(self.req,'PLATFORM_NOTE'))
        older=DocumentBlock.objects.create(template_type='PLATFORM_NOTE',language='fr',body='Oldest')
        DocumentBlock.objects.create(template_type='PLATFORM_NOTE',language='fr',body='Latest')
        before=_block_signature(self.req,'PLATFORM_NOTE');older.delete()
        self.assertNotEqual(before,_block_signature(self.req,'PLATFORM_NOTE'))

    def test_in_place_custom_field_edit_invalidates_cache_without_id_change(self):
        field=ServiceFormField.objects.create(service=self.service,name='field',label_fr='Original')
        before=_service_fields_signature(self.req)
        ServiceFormField.objects.filter(pk=field.pk).update(label_fr='Changed')
        self.assertNotEqual(before,_service_fields_signature(self.req))

    def test_missing_stored_file_is_explicit_not_an_unnoticed_fallback(self):
        obj=self.create('Missing');obj.file.storage.delete(obj.file.name)
        with self.assertRaises(FileNotFoundError):_get_uploaded_template(self.service,'PLATFORM_NOTE')

    def test_object_storage_without_local_path_is_used_by_generator(self):
        class ObjectStorage(Storage):
            def __init__(self):self.data={}
            def exists(self,name):return name in self.data
            def _save(self,name,content):self.data[name]=content.read();return name
            def _open(self,name,mode='rb'):return ContentFile(self.data[name])
            def path(self,name):raise NotImplementedError('Object storage has no local path')
        storage=ObjectStorage();field=ServiceTemplate._meta.get_field('file')
        with patch.object(field,'storage',storage):
            obj=ServiceTemplate.objects.create(service=self.service,name='Object storage',template_type='PLATFORM_NOTE',file=self.upload('REMOTE SOURCE'))
            source=_get_uploaded_template(self.service,'PLATFORM_NOTE')
            self.assertTrue(source.is_file());self.assertEqual(Document(source).paragraphs[0].text,'REMOTE SOURCE')
            self.assertEqual(_get_uploaded_template(self.service,'PLATFORM_NOTE'),source)
            source.write_bytes(b'corrupt cache')
            restored=_get_uploaded_template(self.service,'PLATFORM_NOTE')
            self.assertEqual(Document(restored).paragraphs[0].text,'REMOTE SOURCE')
            self.assertIn(obj.file.name,storage.data)

    def test_active_constraint_does_not_limit_archived_versions(self):
        from django.db import IntegrityError,transaction
        self.create('Active')
        for i in range(4):ServiceTemplate.objects.create(service=self.service,name='Archive '+str(i),template_type='PLATFORM_NOTE',file=self.upload(),is_active=False)
        self.assertEqual(ServiceTemplate.objects.filter(is_active=False).count(),4)
        with self.assertRaises(IntegrityError),transaction.atomic():
            ServiceTemplate.objects.create(service=self.service,name='Conflict',template_type='PLATFORM_NOTE',file=self.upload(),is_active=True)

    def test_failed_template_write_keeps_previous_active_file(self):
        from django.db import DatabaseError
        from django.conf import settings
        previous=self.create('Previous durable model')
        before={str(p) for p in Path(settings.MEDIA_ROOT).rglob('*') if p.is_file()}
        with patch('documents.views.ServiceTemplate.save',side_effect=DatabaseError('Simulated failure')):
            response=self.client.post('/documents/templates/create/',{'service':str(self.service.pk),'template_type':'PLATFORM_NOTE','name':'Not committed','file':self.upload()})
        self.assertEqual(response.status_code,503)
        previous.refresh_from_db();self.assertTrue(previous.is_active)
        self.assertEqual(ServiceTemplate.objects.count(),1)
        self.assertEqual({str(p) for p in Path(settings.MEDIA_ROOT).rglob('*') if p.is_file()},before)

    def test_failed_replacement_preserves_original_document_and_entered_name(self):
        from django.db import DatabaseError
        previous=self.create('Previous')
        with patch('documents.views.ServiceTemplate.save',side_effect=DatabaseError('Simulated failure')):
            response=self.client.post('/documents/templates/'+str(previous.pk)+'/edit/',{'service':str(self.service.pk),'template_type':'PLATFORM_NOTE','name':'Retained name','is_active':'on','file':self.upload('REPLACEMENT')})
        self.assertEqual(response.status_code,503);self.assertContains(response,'Retained name',status_code=503)
        previous.refresh_from_db();self.assertEqual(previous.name,'Previous')
        with previous.file.open('rb') as source:self.assertEqual(Document(source).paragraphs[0].text,'Previous')

    def test_block_validation_retains_values_and_relations_are_atomic(self):
        from django.db import DatabaseError
        original=DocumentBlock.objects.create(template_type='QUOTE',body='Previous block')
        payload={'template_type':'QUOTE','language':'fr','position':'BOTTOM','title':'Retained title','body':'Retained body','priority':'invalid','services':[str(self.service.pk)]}
        response=self.client.post('/documents/blocks/create/',payload)
        self.assertEqual(response.status_code,400);self.assertContains(response,'Retained body',status_code=400)
        payload['priority']='2'
        with patch.object(type(original.services),'set',side_effect=DatabaseError('Simulated relation failure')):
            response=self.client.post('/documents/blocks/create/',payload)
        self.assertEqual(response.status_code,503)
        self.assertEqual(DocumentBlock.objects.count(),1)
        self.assertContains(response,'Retained body',status_code=503)
        payload['title']='X'*256
        self.assertEqual(self.client.post('/documents/blocks/create/',payload).status_code,400)

    def test_invalid_template_metadata_does_not_deactivate_previous(self):
        previous=self.create('Previous')
        base={'service':str(self.service.pk),'template_type':'PLATFORM_NOTE','name':'Retained'}
        for changes in ({'service':'invalid'},{'service':'00000000-0000-0000-0000-000000000000'},{'template_type':'invalid'},{}):
            response=self.client.post('/documents/templates/create/',{**base,**changes})
            self.assertEqual(response.status_code,400)
        self.service.active=False;self.service.save()
        self.assertEqual(self.client.post('/documents/templates/create/',{**base,'file':self.upload()}).status_code,400)
        previous.refresh_from_db();self.assertTrue(previous.is_active)

    def test_template_create_and_toggle_get_do_not_write(self):
        self.assertEqual(self.client.get('/documents/templates/create/').status_code,200)
        item=self.create('Stable')
        response=self.client.get('/documents/templates/'+str(item.pk)+'/toggle/')
        self.assertEqual(response.status_code,302)
        item.refresh_from_db();self.assertTrue(item.is_active)

    def test_empty_stored_file_and_failed_cleanup_are_explicit(self):
        from unittest.mock import Mock
        from documents.template_storage import materialize_template
        from documents.views import _cleanup_template_file
        item=self.create('Empty source')
        with item.file.open('wb') as target:target.write(b'')
        with self.assertRaises(ValueError):materialize_template(item)
        storage=Mock();storage.delete.side_effect=OSError('Test cleanup failure')
        with self.assertLogs('plagenor.documents',level='ERROR'):_cleanup_template_file((storage,'test.docx'))
        storage.delete.assert_called_once_with('test.docx')

    def test_concurrent_template_move_is_not_silently_overwritten(self):
        item=self.create('Stable');other=Service.objects.create(code='OTHER-CONCURRENT',name='Other')
        moved=ServiceTemplate.objects.get(pk=item.pk);moved.service=other
        with patch('documents.views.ServiceTemplate.objects.select_for_update') as locked:
            locked.return_value.get.return_value=moved
            response=self.client.post('/documents/templates/'+str(item.pk)+'/edit/',{'service':str(self.service.pk),'template_type':'PLATFORM_NOTE','name':'Pending edit','is_active':'on'})
        self.assertEqual(response.status_code,400)
        item.refresh_from_db();self.assertEqual(item.name,'Stable')
