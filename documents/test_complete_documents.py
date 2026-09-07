"""Document integration boundaries, conversion recovery and legacy inputs."""
import io
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from django.test import TestCase, SimpleTestCase, override_settings
from docx import Document
from core.models import Service
from documents.models import ServiceTemplate, TemplatePlaceholder, DocumentBlock
from documents import pdf_converter as pc


class CompleteConversionContracts(SimpleTestCase):
    def test_uno_connection_retries_caches_and_reports_unavailable_bridge(self):
        daemon=pc._LibreOfficeDaemon();uno=MagicMock()
        resolver=uno.getComponentContext.return_value.ServiceManager.createInstanceWithContext.return_value
        desktop=MagicMock();ctx=MagicMock();ctx.ServiceManager.createInstanceWithContext.return_value=desktop
        with patch.dict(sys.modules,{'uno':uno}),patch.object(pc.time,'sleep'):
            resolver.resolve.side_effect=[RuntimeError('starting'),ctx]
            self.assertIs(daemon._connect_desktop(),desktop)
            self.assertIs(daemon._connect_desktop(),desktop)
            self.assertEqual(resolver.resolve.call_count,2)
            daemon._desktop=None;resolver.resolve.side_effect=RuntimeError('unavailable')
            with self.assertRaises(pc._UnoConversionError):daemon._connect_desktop()
        with patch.object(daemon,'_alive',return_value=False),patch.object(daemon,'_start') as start:
            daemon._ensure_running();start.assert_called_once()
        daemon._proc=MagicMock();daemon._proc.terminate.side_effect=RuntimeError('gone');daemon._proc.kill.side_effect=RuntimeError('gone')
        with self.assertLogs(pc.logger,level='ERROR'):daemon.shutdown()
        self.assertIsNone(daemon._proc)

    def test_uno_conversion_closes_documents_and_restarts_only_once(self):
        daemon=pc._LibreOfficeDaemon();desktop=MagicMock();uno=MagicMock()
        beans=SimpleNamespace(PropertyValue=lambda:SimpleNamespace())
        with tempfile.TemporaryDirectory() as temp,patch.dict(sys.modules,{'uno':uno,'com.sun.star.beans':beans}),patch.object(daemon,'_ensure_running'),patch.object(daemon,'_connect_desktop',return_value=desktop):
            src=Path(temp)/'contract.docx';src.write_bytes(b'docx')
            with self.assertRaisesRegex(pc._UnoConversionError,'no PDF'):daemon.convert(src,Path(temp))
            desktop.loadComponentFromURL.return_value.close.assert_called_once_with(False)
            desktop.loadComponentFromURL.side_effect=RuntimeError('bridge lost')
            with patch.object(daemon,'shutdown') as shutdown:
                with self.assertRaisesRegex(pc._UnoConversionError,'bridge lost'):daemon.convert(src,Path(temp))
                self.assertEqual(shutdown.call_count,2)
        with patch.object(pc,'_daemon',None),patch.object(pc.atexit,'register') as register:
            first=pc._get_daemon();self.assertIs(pc._get_daemon(),first);register.assert_called_once()

    def test_cancellation_watermark_survives_missing_system_font(self):
        from documents.watermark import add_cancellation_watermark
        from PIL import ImageFont
        fallback=ImageFont.load_default()
        doc=Document()
        with patch('documents.watermark.ImageFont.truetype',side_effect=OSError),patch('documents.watermark.ImageFont.load_default',return_value=fallback):
            add_cancellation_watermark(doc.add_paragraph())
        self.assertIn('behindDoc="1"',doc._element.xml)
        buf=io.BytesIO();doc.save(buf);self.assertTrue(buf.getvalue().startswith(b'PK'))


class CompleteDocumentModelContracts(TestCase):
    def test_template_labels_and_empty_or_populated_file_urls(self):
        service=Service.objects.create(code='DOC-COMP',name='Document')
        template=ServiceTemplate(service=service,name='Devis',template_type='QUOTE')
        self.assertIn('DOC-COMP',str(template));self.assertIsNone(template.file_url)
        template.file.name='templates/quote.docx';self.assertTrue(template.file_url.endswith('templates/quote.docx'))
        self.assertIn('client',str(TemplatePlaceholder(placeholder='client',description='Client identity')))
        block=DocumentBlock.objects.create(template_type='QUOTE',position='before_body',language='fr')
        self.assertIn('GLOBAL',str(block));self.assertEqual(block.scope_label(),'Global')
        block.services.add(service);self.assertIn('DOC-COMP',str(block))

    def test_french_amount_words_and_invalid_legacy_amounts(self):
        from documents import genoclab_layout as gl
        from core.models import PlatformContent
        for value,expected in [(21,'vingt et un'),(30,'trente'),(0,'zéro'),(-1,'moins un'),(2000000,'deux millions'),(1000000000,'un milliard')]:
            self.assertEqual(gl.amount_in_words_fr(value),expected)
        for value in ('invalid','NaN','Infinity',None):self.assertEqual(gl.amount_in_words_fr(value),'')
        self.assertEqual(gl._three_digits(0),'')
        self.assertEqual(gl._money(12.5),'12,50 DA');self.assertEqual(gl._money('bad'),'bad');self.assertEqual(gl._money_int('bad'),'bad')
        PlatformContent.objects.create(key='custom-contract',lang='fr',value='Custom')
        self.assertEqual(gl.cms_get('custom-contract'),'Custom')
        with patch.object(PlatformContent.objects,'filter',side_effect=RuntimeError):
            self.assertEqual(gl.cms_get('missing','Fallback'),'Fallback')
        doc=Document()
        with patch.object(gl,'cms_get',return_value='invalid'):
            gl.add_prestation_table(doc,[{'label':'A','quantity':2,'unit_price':10},
                                          {'label':'B','quantity':'bad','unit_price':10},
                                          {'label':'C','total':'bad'}])
        self.assertEqual(len(doc.tables),1)
        gl.add_genoclab_footer(doc,total_amount=21,identity={'values':{'genoclab_footer_legal':'Montant arrêté'}})
        self.assertIn('Vingt et un',' '.join(p.text for p in doc.paragraphs))

    def test_spreadsheet_numeric_formats_and_period_labels(self):
        from documents import stats_excel as se
        from openpyxl import Workbook
        section={'title':'Revenue','columns':['Group','Count','Percent','IBTIKAR (DA)','GenoClab (DA)'],
                 'rows':[['A',2,50,120,240]],'total_row':['Total',2,50,120,240]}
        ws=Workbook().active
        se._write_section(ws,section,1)
        self.assertEqual(ws.cell(3,3).number_format,se._PCT_FMT)
        self.assertEqual(ws.cell(3,4).number_format,se._MONEY_FMT)
        for filters,expected in [({},'toutes périodes confondues'),({'date_to':'2026-12-31'},"jusqu'au 2026-12-31"),({'date_from':'2026-01-01','date_to':'2026-12-31'},'du 2026-01-01 au 2026-12-31')]:
            self.assertEqual(se._period_label(filters),expected)

    def test_document_helpers_handle_sparse_imported_documents(self):
        from documents import docx_helpers as dh
        from docx.oxml import OxmlElement
        doc=Document();p=doc.add_paragraph('{{MISSING}}')
        dh._strip_in_subtree(p._p);self.assertEqual(p.text,'')
        with patch.object(type(p._p),'iter',side_effect=ValueError):self.assertIsNone(dh._strip_in_subtree(p._p))
        self.assertFalse(dh._is_section_heading(p))
        p.text='1. Title';self.assertTrue(dh._is_section_heading(p))
        empty=Document();self.assertEqual(dh._find_anchor_for_position(empty,'TOP'),(None,'end'))
        p=empty.add_paragraph('plain');self.assertEqual(dh._find_anchor_for_position(empty,'TOP')[1],'before')
        p.text='';self.assertEqual(dh._find_anchor_for_position(empty,'BEFORE_SIGNATURE'),(None,'end'))
        for position in dh._ANCHOR_KEYWORDS:
            self.assertEqual(dh._find_anchor_for_position(empty,position),(None,'end'))
        self.assertIsNone(dh._find_section_end(empty,p))
        table=doc.add_table(rows=2,cols=3);cell=table.cell(0,0)
        cell.paragraphs[0].add_run('old');cell.paragraphs[0].add_run(' extra')
        dh._set_cell_text(cell,'new');self.assertEqual(cell.text,'new')
        for child in list(cell._tc):cell._tc.remove(child)
        dh._set_cell_text(cell,'replacement');self.assertEqual(cell.text,'replacement')
        self.assertEqual(dh._build_param_label_index({},''),[])
        with patch('core.registry.get_service_def',side_effect=RuntimeError):
            self.assertTrue(dh._build_param_label_index({'mode':'A'},'missing'))
        self.assertIsNone(dh._fuzzy_pick('question',[('', 'value','label','key')],set()))
        dh.populate_legacy_param_questions(doc,SimpleNamespace(service_params={},service=None))
        self.assertIsNone(dh.populate_legacy_sample_table(doc,SimpleNamespace(sample_table=[])))

    def test_question_and_sample_legacy_tables_fill_without_placeholders(self):
        from documents import docx_helpers as dh
        service=Service.objects.create(code='LEGACYDOC',name='Legacy')
        definition={'parameters':[{'name':'mode','label':'Analysis mode','label_fr':'Analysis mode'}],
                    'sample_table':{'columns':[{}, {'name':'sample','label':'Sample'}, {'name':'strain','label':'Strain'}, {'name':'  ','label':''}]}}
        doc=Document();p=doc.add_paragraph();p.add_run('Analysis mode : ');p.add_run('Cliquez ou appuyez ici pour entrer du texte.')
        doc.add_paragraph('')
        req=SimpleNamespace(service=service,service_params={'mode':'Complete'},sample_table=[{'sample':'A','strain':'B','empty':None}])
        with patch('core.registry.get_service_def',return_value=definition):
            dh.populate_legacy_param_questions(doc,req)
            self.assertIn('Complete',p.text)
            table=doc.add_table(rows=2,cols=4)
            for c,t in zip(table.rows[0].cells,['N°','Sample','Strain','']):c.text=t
            dh.populate_legacy_sample_table(doc,req)
            self.assertIn('A',' '.join(c.text for c in table.rows[1].cells))
            req.sample_table=[{'empty':None}];dh.populate_legacy_sample_table(doc,req)
        with patch('core.registry.get_service_def',side_effect=RuntimeError):
            dh.populate_legacy_sample_table(Document(),req)
        tiny=Document();tiny.add_table(rows=1,cols=1)
        dh.populate_legacy_sample_table(tiny,req)


class CompleteGeneratorContracts(TestCase):
    def setUp(self):
        from core.models import Request
        from accounts.models import User
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        override=override_settings(BASE_DIR=self.root,MEDIA_ROOT=self.root/'media');override.enable();self.addCleanup(override.disable)
        self.service=Service.objects.create(code='GEN-COMP',name='Analysis')
        self.user=User.objects.create_user('doc-complete',role='REQUESTER',first_name='National',last_name='Student')
        self.req=Request.objects.create(display_id='GEN-COMP',channel='IBTIKAR',service=self.service,requester=self.user,service_params={'mode':'complete'},sample_table=[{'sample':'S-1'}])
        self.templates=self.root/'documents'/'docx_templates';self.templates.mkdir(parents=True)

    def read_text(self,path):
        doc=Document(path)
        return '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])

    def test_uploaded_legacy_generic_and_programmatic_document_variants(self):
        from documents import generators as g
        from django.core.files import File
        for kind,fn in [('IBTIKAR_FORM',g.generate_ibtikar_form),('PLATFORM_NOTE',g.generate_platform_note),('RECEPTION_FORM',g.generate_reception_form)]:
            doc=Document();doc.add_paragraph('Uploaded {{REQUEST_ID}}')
            stream=io.BytesIO();doc.save(stream);stream.seek(0)
            template=ServiceTemplate.objects.create(service=self.service,template_type=kind,name=kind,file=File(stream,name=kind+'.docx'))
            self.assertIn('Uploaded',self.read_text(fn(self.req)))
            template.delete()
        legacy=self.templates/'ibtikar';legacy.mkdir()
        doc=Document();doc.add_paragraph('Legacy Nom et prénom : * Nom complet du demandeur');doc.save(legacy/'contract.docx')
        with patch.dict(g.IBTIKAR_TEMPLATE_MAP,{'GEN-COMP':'contract.docx'}):
            self.assertIn('National Student',self.read_text(g.generate_ibtikar_form(self.req)))
        doc=Document();doc.add_paragraph('Generic {{REQUEST_ID}}');doc.add_paragraph('Tableau des échantillons');doc.add_paragraph('[Tableau des échantillons à remplir]');doc.add_paragraph('Signature du demandeur');doc.save(self.templates/'ibtikar_form_template.docx')
        text=self.read_text(g.generate_ibtikar_form(self.req));self.assertIn('S-1',text);self.assertNotIn('[Tableau',text)
        self.assertIn('S-1',self.read_text(g.generate_reception_form(self.req)))
        self.req.sample_table=[{}]
        from accounts.models import User
        from django.utils import timezone
        self.req.assigned_to=User.objects.create_user('doc-analyst',role='MEMBER').member_profile
        self.req.appointment_date=timezone.now()
        self.assertIn('Assignation',self.read_text(g.generate_platform_note(self.req)))
        with patch.object(ServiceTemplate.objects,'filter',side_effect=RuntimeError('storage metadata unavailable')):
            with self.assertLogs(g.logger,level='ERROR'):self.assertIsNone(g._get_uploaded_template(self.service,'QUOTE'))

    def test_defensive_field_rendering_and_custom_labels(self):
        from documents import generators as g
        from core.models import ServiceFormField
        from django.core.exceptions import ObjectDoesNotExist
        from unittest.mock import PropertyMock
        for value in (None,'bad'):
            self.assertEqual(g._money_2dp(value),'N/A')
        self.assertEqual(g._money('bad'),'N/A');self.assertEqual(g._format_date('bad'),'Non défini')
        class Orphan:pass
        orphan=Orphan()
        for exc in (ObjectDoesNotExist,RuntimeError):
            with patch.object(Orphan,'assigned_to',new_callable=PropertyMock,create=True,side_effect=exc):
                self.assertEqual(g._assigned_name(orphan),'Non assigné');self.assertEqual(g._assigned_email(orphan),'')
                self.assertEqual(g._safe_attr(orphan,'assigned_to','Missing'),'Missing')
        self.assertEqual(g._format_sample_table_text([None,{}]),'')
        self.assertIn('1 échantillon',g._format_sample_summary([{}]))
        doc=Document();g._render_sample_table(doc,[None,{}]);g._render_service_params(doc,{'empty':None});self.assertEqual(len(doc.tables),0)
        ServiceFormField.objects.create(service=self.service,name='custom',label='Custom label')
        definition={'parameters':[{'name':'mode','label_fr':'Mode choisi'}],'sample_table':{'columns':[{'name':'sample','label':'Échantillon'}]}}
        with patch('core.registry.get_service_def',return_value=definition):
            self.assertEqual(g._field_label_map(self.req),{'mode':'Mode choisi','sample':'Échantillon','custom':'Custom label'})
        with patch('core.registry.get_service_def',side_effect=RuntimeError),patch.object(type(self.service.custom_fields),'all',side_effect=RuntimeError):
            with self.assertLogs(g.logger,level='ERROR'):self.assertEqual(g._field_label_map(self.req),{})
        block=DocumentBlock.objects.create(template_type='IBTIKAR_FORM',position='TOP',body='Notice')
        g._inject_document_blocks(doc,'IBTIKAR_FORM',self.req);self.assertIn('Notice',' '.join(p.text for p in doc.paragraphs))
        anchor=doc.add_paragraph('Anchor');empty=DocumentBlock(title='',body='')
        self.assertIs(g._insert_block_relative(anchor,'after',empty,{}),anchor)
        # Defensive placement fallback is part of this reusable helper's API.
        g._insert_block_relative(anchor,'end',block,{});self.assertEqual(doc.paragraphs[-1].text,'Notice')
        table=doc.add_table(rows=0,cols=2);g._apply_brand_table_style_everywhere(doc)
        doc.add_table(rows=1,cols=1)
        with patch.object(g,'style_brand_table',side_effect=ValueError):g._apply_brand_table_style_everywhere(doc)
        self.assertEqual(len(doc.tables),2)

    def test_cancelled_invoice_first_render_keeps_original_and_marks_copy(self):
        from core.models import Invoice,IssuedDocument
        from documents.generators import generate_invoice_document
        from django.utils import timezone
        inv=Invoice.objects.create(request=self.req,invoice_number='CANCEL-COMP',cancelled_at=timezone.now(),cancellation_reason='Duplicate entry',line_items=[{'description':'Analysis','quantity':1,'unit_price':100,'total':100}])
        path=generate_invoice_document(inv)
        self.assertIn('ANNUL',Document(path).sections[0].header._element.xml)
        self.assertTrue(IssuedDocument.objects.filter(kind='INVOICE',number='CANCEL-COMP').exists())

    def test_partial_headers_missing_styles_and_incomplete_imported_tables(self):
        from documents import docx_helpers as dh,generators as g,build_default_templates as bt
        from docx.oxml import OxmlElement
        # DOCX documents can omit named styles or header/footer paragraphs.
        doc=Document();doc.styles['Normal'].delete();doc.styles['Heading 1'].delete()
        dh.apply_house_style(doc)
        self.assertNotIn('Heading 1',doc.styles)
        footer=doc.sections[0].footer
        for p in list(footer.paragraphs):p._p.getparent().remove(p._p)
        bt._add_footer(doc);self.assertIn('PLAGENOR',footer.paragraphs[0].text)
        header=doc.sections[0].header
        for p in list(header.paragraphs):p._p.getparent().remove(p._p)
        corrupt=self.root/'corrupt.png';corrupt.write_bytes(b'broken image')
        dh.ensure_institutional_header(doc,corrupt);self.assertEqual(len(header.paragraphs),1)
        sparse=SimpleNamespace(sections=[SimpleNamespace(header=None,footer=None)])
        dh.ensure_institutional_header(sparse,corrupt);dh.add_brand_footer(sparse)
        blank=Document();blank.add_paragraph('')
        self.assertEqual(dh._find_anchor_for_position(blank,'BEFORE_FOOTER'),(None,'end'))
        blank.add_heading('Informations du demandeur',level=2)
        self.assertEqual(dh._find_anchor_for_position(blank,'AFTER_REQUESTER'),(None,'end'))
        p=blank.add_paragraph();p.add_run('Nom et prénom : ');p.add_run('* Nom complet du demandeur')
        dh.apply_legacy_label_substitution(blank,{'FULL_NAME':'National Student'})
        self.assertEqual(p.runs[-1].text,'');self.assertIn('National Student',p.text)
        table=doc.add_table(rows=2,cols=3)
        for c,text in zip(table.rows[0].cells,['N°','Sample','Strain']):c.text=text
        table.rows[1]._tr.remove(table.rows[1].cells[-1]._tc)
        table.rows[1]._tr.remove(table.rows[1].cells[-1]._tc)
        request=SimpleNamespace(service=self.service,sample_table=[{'sample':'S1','strain':'X'}])
        definition={'sample_table':{'columns':[{'name':'sample','label':'Sample'},{'name':'strain','label':'Strain'}]}}
        with patch('core.registry.get_service_def',return_value=definition):dh.populate_legacy_sample_table(doc,request)
        self.assertEqual(table.rows[1].cells[0].text,'01')
        detached=Document().add_table(rows=1,cols=1);detached._element.getparent().remove(detached._element)
        g._detableize(SimpleNamespace(tables=[detached]));self.assertEqual(len(detached.rows),1)

    def test_template_builder_entrypoint_and_empty_statistical_sections(self):
        import runpy
        from contextlib import redirect_stdout
        from core.stats import stats_for_user
        from documents import generators as g
        with redirect_stdout(io.StringIO()):runpy.run_module('documents.build_default_templates',run_name='__main__')
        self.assertEqual(len(list(self.templates.glob('*.docx'))),3)
        from accounts.models import User
        admin=User.objects.create_user('stats-doc',role='SUPER_ADMIN')
        bundle=stats_for_user(admin)
        for key in ('by_service','by_status','by_wilaya','by_channel','by_month','by_member'):bundle[key]=[]
        self.assertIn('Statistiques',self.read_text(g.generate_stats_report(bundle,{},admin)))
        self.req.service_params={'analysis_mode':'missing'}
        with patch('core.registry.get_service_def',return_value={'pricing':{'base_price':{'default':10},'multipliers':{}}}):
            doc=Document();g._render_tariff_breakdown(doc,self.req)
        self.assertTrue(doc.paragraphs)
