from pathlib import Path
from types import SimpleNamespace
import os
import shutil
import tempfile

from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document

from core.ibtikar.schema import definitions, label
from documents.document_contracts import (
    FINANCIAL_SOURCE_FIELDS, FINANCIAL_SOURCE_VALUES, IBTIKAR_CONTROL_FIELDS,
    IBTIKAR_VALIDATION_FIELDS, PLATFORM_NOTE_FIELDS, RECEPTION_FIELDS, STATS_FIELDS,
)
from documents.generators import (
    _build_platform_note_programmatic, _build_reception_form_programmatic,
    generate_stats_report,
)
from documents.genoclab_layout import CMS_DEFAULTS, render_commercial_document
from documents.ibtikar_canonical import build_document
from documents.ibtikar_reference import reference_content


def _keep_preview(source, name):
    target = os.environ.get("PLAGENOR_DOCUMENT_PREVIEW_DIR")
    if not target:
        return
    destination = Path(target)
    destination.mkdir(parents=True, exist_ok=True)
    path = destination / name
    if isinstance(source, (str, Path)):
        shutil.copy2(source, path)
    else:
        source.save(path)


def _collect_document_text(document):
    values = []

    def visit_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    values.append(cell.text)
                    visit_tables(cell.tables)

    values.extend(paragraph.text for paragraph in document.paragraphs)
    visit_tables(document.tables)
    for section in document.sections:
        values.extend(paragraph.text for paragraph in section.header.paragraphs)
        visit_tables(section.header.tables)
        values.extend(paragraph.text for paragraph in section.footer.paragraphs)
        visit_tables(section.footer.tables)
    return "\n".join(values)


def _row(spec, language="fr", marker="Valeur de recette"):
    return {
        "name": spec["name"], "label": label(spec["label"], language),
        "display": marker, "value": marker, "options": [], "all_options": False,
    }


def _project_from_schema(schema, language="fr"):
    staff = [
        _row(spec, language)
        for spec in schema.get("staff", [])
        if spec["name"] not in ("validated_price", "price_justification")
    ]
    return {
        "service_code": schema["service_code"],
        "title": label(schema["title"], language),
        "source_version": schema.get("source_version", ""),
        "applicant": [_row(spec, language) for spec in schema.get("applicant", [])],
        "parameters": [_row(spec, language) for spec in schema.get("parameters", [])],
        "samples": [[_row(spec, language) for spec in schema.get("samples", [])]],
        "staff": staff,
        "sample_count": 1,
        "read_count": None,
        "notices": [label(item, language) for item in schema.get("notices", [])],
    }


class IbtikarPreservationContractTests(SimpleTestCase):
    maxDiff = None

    def test_all_ten_services_preserve_schema_and_source_content(self):
        services = definitions()["services"]
        self.assertEqual(len(services), 10)
        for code, schema in services.items():
            with self.subTest(service=code):
                project = _project_from_schema(schema)
                document = build_document(
                    project,
                    {
                        "number": "IBK-TEST-001",
                        "date": "19/09/2026",
                        "external_reference": "IBTIKAR-REFERENCE",
                        "revision": 3,
                        "draft": False,
                        "operator_name": "Opérateur de recette",
                    },
                    "fr",
                    attachment_rows=[
                        {"label": label(spec["label"], "fr"), "display": "piece-test.pdf"}
                        for spec in schema.get("attachments", [])
                    ],
                )
                text = _collect_document_text(document)
                self.assertIn(code, text)
                self.assertIn(schema.get("source_version", ""), text)
                for group in ("applicant", "parameters", "samples"):
                    for spec in schema.get(group, []):
                        self.assertIn(label(spec["label"], "fr"), text)
                for spec in schema.get("staff", []):
                    if spec["name"] not in ("validated_price", "price_justification"):
                        self.assertIn(label(spec["label"], "fr"), text)
                for spec in schema.get("attachments", []):
                    self.assertIn(label(spec["label"], "fr"), text)
                source = reference_content(code)
                for paragraph in source["guidance"]:
                    self.assertIn(paragraph, text)
                for source_table in source["tables"]:
                    for row in source_table:
                        for cell in row:
                            if cell:
                                self.assertIn(cell, text)
                if source["ethics"]:
                    self.assertIn(source["ethics"], text)
                for expected in IBTIKAR_CONTROL_FIELDS + IBTIKAR_VALIDATION_FIELDS:
                    self.assertIn(expected, text)
                section = document.sections[0]
                self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
                self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
                self.assertIn("NUMPAGES", section.footer._element.xml)
                preview_name = f"IBTIKAR_{code.replace('-', '_').upper()}_EXEMPLE_FICTIF.docx"
                _keep_preview(document, preview_name)

    def test_long_source_tables_repeat_headers_and_signature_zones_remain(self):
        schema = definitions()["services"]["EGTP-Lyoph"]
        document = build_document(
            _project_from_schema(schema),
            {"number": "IBK-LYOPH", "date": "19/09/2026", "external_reference": "",
             "revision": 1, "draft": False, "operator_name": ""},
            "fr",
        )
        xml = document._element.xml
        text = _collect_document_text(document)
        self.assertIn("Récipients recommandés pour lyophilisation", text)
        self.assertIn("Substances interdites", text)
        self.assertIn("Beta 2-8 LSCplus", text)
        self.assertNotIn("Alpha 3-4 LSCbasic", text)
        self.assertIn("w:tblHeader", xml)
        self.assertLess(xml.index("Récipients recommandés"), xml.index("50 ml"))
        self.assertLess(xml.index("50 ml"), xml.index("Critères de remplissage"))
        self.assertLess(xml.index("Autres récipients"), xml.index("Microtubes (1.5 ml)"))
        self.assertLess(xml.index("Microtubes (1.5 ml)"), xml.index("Tout échantillon biologique"))
        self.assertLess(xml.index("strictement interdite"), xml.index("Substances interdites"))
        self.assertLess(xml.index("Substances interdites"), xml.index("Seuils de concentration"))
        for label_text in ("Signature du demandeur", "Signature de l’opérateur",
                           "Visa du Chef du Service Commun", "Visa du Directeur de l’ESSBO"):
            self.assertIn(label_text, text)
        _keep_preview(document, "IBTIKAR_LYOPH_EXEMPLE_FICTIF.docx")

    def test_sanger_current_acknowledgment_clause_is_preserved(self):
        schema = definitions()["services"]["EGTP-SeqS"]
        document = build_document(
            _project_from_schema(schema),
            {"number": "IBK-SANGER", "date": "19/09/2026", "external_reference": "",
             "revision": 1, "draft": False, "operator_name": ""},
            "fr",
        )
        text = _collect_document_text(document)
        self.assertIn("Acknowledgment and Citation Clause", text)
        self.assertIn("This research was conducted at the Genomics Technology Platform", text)
        _keep_preview(document, "IBTIKAR_SANGER_EXEMPLE_FICTIF.docx")


class FinancialPreservationContractTests(TestCase):
    def _identity(self, channel):
        return {
            "billing_channel": channel,
            "values": dict(CMS_DEFAULTS),
            "client_name": "Client Fictif",
            "client_organization": "Université de recette",
            "client_laboratory": "Laboratoire de recette",
            "client_phone": "+213 555 000 001",
            "client_fax": "+213 41 00 00 00",
            "client_email": "client@example.test",
            "client_lines": [],
            "request_reference": "GCL-TEST-001",
        }

    def _document(self, channel, kind):
        title = "Devis" if kind == "quote" else "Facture"
        return render_commercial_document(
            title=title, number="TEST-001", date="19/09/2026",
            identity=self._identity(channel),
            items=[
                {"label": "Analyse génomique", "quantity": 2,
                 "unit_price": "5000.00", "total": "10000.00"},
            ],
            vat_rate=0 if channel == "OHB" else 0.19,
            document_kind=kind,
        )

    def test_original_financial_fields_are_preserved_in_both_channels_and_kinds(self):
        for channel in ("OHB", "GENOCLAB"):
            for kind in ("quote", "invoice"):
                with self.subTest(channel=channel, kind=kind):
                    document = self._document(channel, kind)
                    text = _collect_document_text(document)
                    for expected in FINANCIAL_SOURCE_FIELDS:
                        self.assertIn(expected, text)
                    for expected in FINANCIAL_SOURCE_VALUES:
                        self.assertIn(expected, text)
                    for expected in ("Client Fictif", "Université de recette",
                                     "Laboratoire de recette", "+213 555 000 001",
                                     "+213 41 00 00 00", "client@example.test",
                                     "GCL-TEST-001"):
                        self.assertIn(expected, text)
                    self.assertIn("dinar", text.lower())
                    self.assertNotRegex(text, r"\b\d{1,2}/100\b")
                    if channel == "GENOCLAB":
                        self.assertIn("TVA (19 %)", text)
                        self.assertIn("Total TTC", text)
                        self.assertNotIn("Non assujetti à la TVA", text)
                    else:
                        self.assertIn("Total DA", text)
                        self.assertIn("Non assujetti à la TVA", text)
                        self.assertNotIn("TVA (19 %)", text)
                    self.assertIn("NUMPAGES", document.sections[0].footer._element.xml)


class OtherGeneratedDocumentContracts(TestCase):
    def field_map(self):
        return {
            "DISPLAY_ID": "REQ-TEST-001", "DATETIME": "19/09/2026 à 14:00",
            "FULL_NAME": "Demandeur Fictif", "ETABLISSEMENT": "ESSBO",
            "LABORATORY": "Laboratoire de recette", "STUDENT_LEVEL": "Doctorant",
            "SUPERVISOR": "Encadrant Fictif", "EMAIL": "test@example.test",
            "PHONE": "+213 555 000 002", "SERVICE_CODE": "EGTP-CAN",
            "SERVICE_NAME": "Contrôle qualité", "SERVICE_DESCRIPTION": "Description test",
            "SERVICE_TURNAROUND": "7", "CHANNEL": "IBTIKAR", "URGENCY": "Normale",
            "TITLE": "Projet fictif", "DESCRIPTION": "Description de demande",
            "BUDGET_AMOUNT": "10 000 DA", "IBTIKAR_BALANCE": "190 000 DA",
            "ASSIGNED_ANALYST": "Analyste Fictif", "ANALYST_EMAIL": "analyst@example.test",
            "APPOINTMENT_DATE": "20/09/2026", "TRACKING_CODE": "TRACK-001",
            "SUBMISSION_DATE": "19/09/2026",
        }

    def test_platform_note_keeps_legacy_information_contract(self):
        request = SimpleNamespace(
            service_params={}, sample_table=[], declared_ibtikar_balance=190000,
            assigned_to=None, service=None, budget_amount=10000, channel="IBTIKAR",
            urgency="Normal", admin_validated_price=None,
        )
        document = _build_platform_note_programmatic(request, self.field_map())
        text = _collect_document_text(document)
        for expected in PLATFORM_NOTE_FIELDS:
            self.assertIn(expected, text)
        _keep_preview(document, "NOTE_PLATEFORME_EXEMPLE_FICTIF.docx")

    def test_reception_form_keeps_original_information_and_signature_contract(self):
        request = SimpleNamespace(
            sample_table=[{"sample_code": "S01", "sample_type": "Sang"}],
            service=SimpleNamespace(code="EGTP-CAN", custom_fields=SimpleNamespace(all=lambda: [])),
        )
        document = _build_reception_form_programmatic(request, self.field_map())
        text = _collect_document_text(document)
        for expected in RECEPTION_FIELDS:
            self.assertIn(expected, text)
        self.assertIn("S01", text)
        self.assertIn("Sang", text)
        _keep_preview(document, "RECEPTION_EXEMPLE_FICTIF.docx")

    def test_statistics_report_keeps_all_report_sections(self):
        with tempfile.TemporaryDirectory() as folder, override_settings(MEDIA_ROOT=folder):
            actor = SimpleNamespace(get_full_name=lambda: "Agent Fictif", username="agent")
            bundle = {
                "kpis": {
                    "total": 3, "completed": 1, "in_progress": 1, "rejected": 1,
                    "completion_rate": 33.3, "ibtikar_count": 2, "genoclab_count": 1,
                    "ibtikar_virtual_revenue": 10000, "genoclab_revenue": 11900,
                },
                "by_service": [{"label": "PCR", "count": 1}],
                "by_status": [{"label": "Terminé", "count": 1}],
                "by_wilaya": [{"label": "Oran", "count": 3}],
                "by_organization": [{"label": "ESSBO", "count": 3}],
                "by_analysis_frame": [{"label": "Doctorat", "count": 2}],
                "by_gender": [{"label": "F", "count": 2}],
                "trend": [{"month": "2026-09", "count": 3}],
            }
            path = generate_stats_report(bundle, {"channel": "IBTIKAR"}, actor)
            text = _collect_document_text(Document(path))
            for expected in STATS_FIELDS:
                self.assertIn(expected, text)
            _keep_preview(path, "STATISTIQUES_EXEMPLE_FICTIF.docx")
