import base64
import io
import tempfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

from core.exceptions import FinancialValidationError
from documents import ibtikar_reference as reference
from documents.document_design import (
    GENOCLAB_THEME, PLAGENOR_THEME, add_identity_header, style_data_table,
)
from documents.docx_helpers import add_brand_footer, style_brand_table
from documents.genoclab_layout import (
    CMS_DEFAULTS, add_genoclab_footer, add_genoclab_header,
    add_prestation_table, amount_in_words_fr,
)
from documents.ibtikar_canonical import (
    _direction, _kv_table, _render_reference, _render_samples,
    _rtl_document, _signature_image, generate_canonical_form,
)


def all_text(doc):
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    for section in doc.sections:
        chunks.extend(p.text for p in section.header.paragraphs)
        for table in section.header.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells)
        chunks.extend(p.text for p in section.footer.paragraphs)
        for table in section.footer.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


class DocumentDesignEdgeCoverageTests(SimpleTestCase):
    def test_identity_header_falls_back_to_text_when_plagenor_image_fails(self):
        doc = Document()
        with patch.object(Run, "add_picture", side_effect=RuntimeError("image failure")):
            add_identity_header(doc, PLAGENOR_THEME)
        text = all_text(doc)
        self.assertIn("ESSBO", text)

    def test_identity_header_falls_back_to_brand_name_when_logo_fails(self):
        doc = Document()
        with patch.object(Run, "add_picture", side_effect=RuntimeError("image failure")):
            add_identity_header(doc, GENOCLAB_THEME)
        self.assertIn("GENOCLAB", all_text(doc))

    def test_empty_data_table_is_a_noop(self):
        doc = Document()
        table = doc.add_table(rows=0, cols=2)
        style_data_table(table)
        self.assertEqual(len(table.rows), 0)


    def test_legacy_first_column_style_and_footer_replacement_paths(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Libellé"
        table.cell(0, 1).text = "Valeur"
        table.cell(1, 0).text = "Second"
        table.cell(1, 1).text = "Contenu"
        style_brand_table(table, accent="first-col")
        self.assertIn("F8FAFC", table.cell(0, 0)._tc.xml)

        footer = doc.sections[0].footer
        footer.paragraphs[0].add_run("Ancien pied")
        add_brand_footer(doc)
        self.assertNotIn("Ancien pied", footer.paragraphs[0].text)
        self.assertIn("ESSBO", footer.paragraphs[0].text)


class GeneratorStyleFailureCoverageTests(SimpleTestCase):
    def test_table_style_failure_is_logged_and_does_not_abort_document_generation(self):
        from documents import generators
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Libellé"
        with patch(
            "documents.generators.style_key_value_table",
            side_effect=RuntimeError("style failure"),
        ), patch.object(generators.logger, "exception") as logged:
            generators._apply_brand_table_style_everywhere(doc)
        logged.assert_called_once_with("Unable to style document table")


class CommercialLayoutEdgeCoverageTests(SimpleTestCase):
    def identity(self, **updates):
        values = dict(CMS_DEFAULTS)
        identity = {
            "billing_channel": "GENOCLAB",
            "values": values,
            "client_name": "Client de recette",
            "client_organization": "ESSBO",
            "client_laboratory": "Laboratoire",
            "client_phone": "041000000",
            "client_fax": "041000001",
            "client_email": "client@example.test",
            "client_lines": [],
            "request_reference": "GCL-EDGE",
            "commercial_terms": "Conditions standard de recette.",
        }
        identity.update(updates)
        return identity

    def test_hundred_eighty_thousand_suppresses_plural_s_before_mille(self):
        self.assertEqual(
            amount_in_words_fr(180000),
            "cent quatre-vingt mille dinars algériens",
        )

    def test_header_preserves_legal_details_and_extra_client_lines(self):
        identity = self.identity(
            client_lines=["ESSBO", "Adresse postale complémentaire"],
        )
        identity["values"]["genoclab_issuer_legal_details"] = "RC : TEST-001"
        doc = Document()
        add_genoclab_header(
            doc, title="Devis", doc_number="D-EDGE", doc_date="19/09/2026",
            client_name=identity["client_name"],
            client_lines=identity["client_lines"],
            identity=identity,
        )
        text = all_text(doc)
        self.assertIn("RC : TEST-001", text)
        self.assertIn("Informations complémentaires", text)
        self.assertIn("Adresse postale complémentaire", text)

    def test_ohb_rejects_non_zero_vat(self):
        with self.assertRaises(FinancialValidationError):
            add_prestation_table(
                Document(),
                [{"label": "Test", "quantity": 1, "unit_price": 100}],
                vat_rate="0.19",
                non_taxable=True,
            )

    def test_footer_preserves_custom_payment_terms(self):
        doc = Document()
        identity = self.identity(payment_terms="Paiement par virement à réception.")
        add_genoclab_footer(
            doc, total_amount=100, identity=identity, document_kind="invoice",
        )
        text = all_text(doc)
        self.assertIn("Conditions de paiement", text)
        self.assertIn("Paiement par virement à réception.", text)


class IbtikarReferenceEdgeCoverageTests(SimpleTestCase):
    def test_dynamic_prompt_handles_blank_text(self):
        self.assertTrue(reference._is_dynamic_prompt("EGTP-PCR", "   "))

    def test_missing_reference_file_returns_empty_contract(self):
        with tempfile.TemporaryDirectory() as folder, \
             patch.object(reference, "_TEMPLATE_DIR", Path(folder)), \
             patch.dict(reference._TEMPLATE_MAP, {"EDGE": "missing.docx"}, clear=False):
            self.assertEqual(
                reference.extract_reference_content("EDGE"),
                {"guidance": [], "tables": [], "blocks": [], "ethics": ""},
            )

    def test_validation_table_stops_collection_and_signature_after_ethics_is_safe(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "edge.docx"
            doc = Document()
            doc.add_paragraph("Avant échantillons")
            sample = doc.add_table(rows=1, cols=2)
            sample.cell(0, 0).text = "Code échantillon"
            sample.cell(0, 1).text = "Séquence"
            doc.add_paragraph("Guide fixe")
            doc.add_paragraph("Déclaration de responsabilité éthique")
            doc.add_paragraph("Signature du demandeur")
            validation = doc.add_table(rows=1, cols=1)
            validation.cell(0, 0).text = "Validation de la demande"
            doc.add_paragraph("NE DOIT PAS APPARAÎTRE")
            doc.save(path)
            with patch.object(reference, "_TEMPLATE_DIR", Path(folder)), \
                 patch.dict(reference._TEMPLATE_MAP, {"EDGE": "edge.docx"}, clear=False):
                result = reference.extract_reference_content("EDGE")
            self.assertIn("Guide fixe", result["guidance"])
            self.assertNotIn("NE DOIT PAS APPARAÎTRE", result["guidance"])
            self.assertEqual(result["ethics"], "")


class IbtikarCanonicalEdgeCoverageTests(SimpleTestCase):
    def test_rtl_helpers_and_empty_kv_table(self):
        doc = Document()
        p = doc.add_paragraph("Texte")
        _direction(p, False)
        self.assertIsNone(p._p.get_or_add_pPr().find(qn("w:bidi")))
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "خلية"
        _rtl_document(doc, "ar")
        self.assertIsNotNone(doc.paragraphs[0]._p.get_or_add_pPr().find(qn("w:bidi")))
        self.assertIsNotNone(
            table.cell(0, 0).paragraphs[0]._p.get_or_add_pPr().find(qn("w:bidi"))
        )
        self.assertIsNone(_kv_table(Document(), [], "fr"))
        ar_doc = Document()
        ar_table = _kv_table(
            ar_doc,
            [{"label": "الاسم", "display": "قيمة", "options": [], "all_options": False}],
            "ar",
        )
        self.assertIsNotNone(
            ar_table.cell(0, 0).paragraphs[0]._p.get_or_add_pPr().find(qn("w:bidi"))
        )

    def test_sample_renderer_covers_empty_read_count_and_compact_layout(self):
        doc = Document()
        _render_samples(doc, {"samples": []}, "fr")
        project = {
            "samples": [[{
                "name": "sample_code", "label": "Code", "display": "S01",
                "options": [], "all_options": False,
            }]],
            "read_count": 2,
        }
        _render_samples(doc, project, "fr")
        self.assertIn("Nombre de lectures demandé : 2", all_text(doc))
        self.assertIn("S01", all_text(doc))

    def test_reference_renderer_covers_non_french_note_and_empty_source_table(self):
        source = {
            "guidance": ["Important : test", "1. Rubrique", "Texte libre"],
            "tables": [[]],
            "blocks": [
                {"type": "table", "rows": []},
                {"type": "paragraph", "text": "Important : test"},
                {"type": "paragraph", "text": "1. Rubrique"},
                {"type": "paragraph", "text": "Texte libre"},
            ],
            "ethics": "",
        }
        doc = Document()
        with patch("documents.ibtikar_canonical.reference_content", return_value=source):
            ethics = _render_reference(
                doc, {"service_code": "EDGE", "notices": ["Notice supplémentaire"]}, "en"
            )
        text = all_text(doc)
        self.assertIn("official French source wording", text)
        self.assertIn("Important", text)
        self.assertIn("Notice supplémentaire", text)
        self.assertTrue(ethics)

    def test_signature_renderer_uses_valid_image_and_falls_back_on_invalid_bytes(self):
        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
            "+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        valid = Document()
        _signature_image(valid, png)
        self.assertEqual(len(valid.inline_shapes), 1)
        invalid = Document()
        _signature_image(invalid, b"not-an-image")
        self.assertTrue(invalid.tables)

    @override_settings(MEDIA_ROOT=tempfile.gettempdir())
    def test_generate_canonical_form_covers_submission_attachment_filtering(self):
        schema = {
            "attachments": [
                {"name": "applicant_signature", "label": {"fr": "Signature"}},
            ]
        }
        fake_file = MagicMock()
        fake_file.open.return_value = io.BytesIO(b"signature-bytes")
        attachments = [
            SimpleNamespace(field_name="inactive", original_name="inactive.pdf", file=fake_file),
            SimpleNamespace(field_name="orphan", original_name="orphan.pdf", file=fake_file),
            SimpleNamespace(
                field_name="applicant_signature",
                original_name="signature.png",
                file=fake_file,
            ),
        ]
        manager = MagicMock()
        manager.filter.return_value = attachments
        submission = SimpleNamespace(
            schema=schema,
            applicant={},
            parameters={},
            samples=[{}],
            staff={"operator_name": "Opérateur"},
            revision=2,
            attachments=manager,
        )
        query = MagicMock()
        query.first.return_value = submission
        request = SimpleNamespace(
            pk="edge-request",
            display_id="IBK-EDGE",
            created_at=datetime(2026, 9, 19),
            ibtikar_external_code="IBT-EDGE",
            status="REQUEST_CREATED",
            service=SimpleNamespace(code="EDGE"),
        )
        projected = {
            "title": "Service edge",
            "service_code": "EDGE",
            "source_version": "V1",
            "applicant": [],
            "parameters": [],
            "samples": [],
            "staff": [
                {"name": "validated_price", "label": "Prix", "display": ""},
                {"name": "price_justification", "label": "Justification", "display": ""},
                {"name": "operator_name", "label": "Opérateur", "display": "Opérateur"},
            ],
            "sample_count": 0,
            "notices": [],
        }
        output = Document()
        with tempfile.TemporaryDirectory() as folder, \
             override_settings(MEDIA_ROOT=folder), \
             patch("core.ibtikar.models.IbtikarSubmission.objects.filter", return_value=query), \
             patch("core.ibtikar.schema.projection", return_value=projected), \
             patch("core.ibtikar.schema.active_data", return_value={}), \
             patch("core.ibtikar.schema.active_names",
                   return_value={"orphan", "applicant_signature"}), \
             patch("core.ibtikar.schema.label",
                   side_effect=lambda value, language: value.get(language, "") if isinstance(value, dict) else value), \
             patch("documents.ibtikar_canonical.build_document", return_value=output), \
             patch("documents.generators._inject_document_blocks"):
            path = generate_canonical_form(request)
        self.assertTrue(Path(path).name.startswith("IBTIKAR_edge-request_"))
        self.assertEqual(
            [item["display"] for item in projected.get("_attachments", [])], []
        )
