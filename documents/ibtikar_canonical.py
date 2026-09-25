from copy import deepcopy
import io
from pathlib import Path
import re
import uuid

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from documents.document_design import (
    IBTIKAR_CONTENT_INDENT_CM, IBTIKAR_CONTENT_WIDTH_CM, IBTIKAR_FONT,
    PLAGENOR_THEME, add_callout, add_document_footer, add_ibtikar_master_header,
    add_ibtikar_section_heading, add_ibtikar_signature_grid, add_ibtikar_subheading,
    add_invisible_header_drawing,
    add_section_heading, align_ibtikar_content_table,
    apply_document_style, set_cant_split, style_data_table, style_ibtikar_key_value_table,
)
from documents.ibtikar_reference import reference_content


TEXT = {
    'form': ('FICHE DE DEMANDE IBTIKAR', 'IBTIKAR REQUEST FORM', 'استمارة طلب إبتكار'),
    'service_requested': ('Service demandé', 'Requested service', 'الخدمة المطلوبة'),
    'general': ('INFORMATIONS GÉNÉRALES', 'GENERAL INFORMATION', 'معلومات عامة'),
    'requester': ('Demandeur et projet', 'Applicant and project', 'صاحب الطلب والمشروع'),
    'parameters': ('Paramètres de la prestation', 'Service parameters', 'معلمات الخدمة'),
    'samples': ('Échantillons / amorces', 'Samples / primers', 'العينات / البادئات'),
    'sample': ('Échantillon / amorce', 'Sample / primer', 'العينة / البادئ'),
    'guidance': ('Exigences, sécurité et conditions', 'Requirements, safety and conditions', 'المتطلبات والسلامة والشروط'),
    'staff': ('Cadre réservé à PLAGENOR', 'PLAGENOR section', 'إطار مخصص للأرضية'),
    'attachments': ('Documents joints applicables', 'Applicable attachments', 'الوثائق المرفقة المعنية'),
    'ethics': ('Déclaration de responsabilité éthique', 'Ethical responsibility statement', 'إقرار المسؤولية الأخلاقية'),
    'signature': ('Signature du demandeur', 'Applicant signature', 'توقيع صاحب الطلب'),
    'operator_signature': ('Signature de l’opérateur', 'Operator signature', 'توقيع الموظف المكلف'),
    'head': ('Visa du Chef du Service Commun', 'Common Service Head endorsement', 'تأشيرة رئيس المصلحة المشتركة'),
    'director': ('Visa du Directeur de l’ESSBO', 'ESSBO Director endorsement', 'تأشيرة مدير المدرسة'),
    'staff_help': ('À compléter par le personnel habilité lors de la réception et du traitement, dans l’application ou sur la fiche imprimée.',
                   'To be completed by authorised staff during receipt and processing, in the application or on the printed form.',
                   'يستكملها الموظفون المخولون عند الاستلام والمعالجة، في التطبيق أو على الاستمارة المطبوعة.'),
    'unknown': ('Non renseigné', 'Not provided', 'غير مذكور'),
    'source': ('Version du formulaire source', 'Source form version', 'نسخة الاستمارة المرجعية'),
    'revision': ('Révision numérique', 'Digital revision', 'المراجعة الرقمية'),
    'legacy_revision': ('Non applicable — demande historique', 'Not applicable — historical request', 'غير مطبق — طلب تاريخي'),
    'date': ('Date de la demande', 'Request date', 'تاريخ الطلب'),
    'number': ('Numéro de demande', 'Request number', 'رقم الطلب'),
    'count': ('Nombre de lignes enregistrées', 'Number of recorded rows', 'عدد الصفوف المسجلة'),
    'reads': ('Nombre de lectures demandé', 'Requested number of reads', 'عدد القراءات المطلوبة'),
    'reference': ('Référence IBTIKAR-DGRSDT', 'IBTIKAR-DGRSDT reference', 'مرجع إبتكار للمديرية العامة للبحث العلمي'),
    'operator': ('Opérateur ayant enregistré la réception', 'Operator recording receipt', 'الموظف الذي سجل الاستلام'),
    'legacy': ('Données historiques conservées — aucune valeur absente n’est déduite du modèle papier.',
               'Historical data retained — no missing value is inferred from the paper template.',
               'بيانات تاريخية محفوظة — لا تستنتج القيم الغائبة من النموذج الورقي.'),
    'unsigned': ('Un nom ou un visa saisi ne constitue pas une signature manuscrite. Les emplacements non signés restent à compléter.',
                 'A typed name or endorsement is not a handwritten signature. Unsigned areas remain to be completed.',
                 'لا يعد الاسم أو التأشير المكتوب توقيعا بخط اليد. تبقى مواضع التوقيع الفارغة لاستكمالها.'),
    'draft': ('BROUILLON — demande non soumise', 'DRAFT — request not submitted', 'مسودة — لم يرسل الطلب'),
    'ethics_body': (
        "La signature de cette fiche atteste que les échantillons soumis ont été collectés, manipulés et transférés conformément aux exigences éthiques et réglementaires applicables. Le demandeur assume la responsabilité de leur nature, de leur origine et de leur utilisation.",
        "Signing this form certifies that the submitted samples were collected, handled and transferred in accordance with applicable ethical and regulatory requirements. The applicant accepts responsibility for their nature, origin and use.",
        "يقر توقيع هذه الاستمارة بأن العينات المقدمة جُمعت وعولجت ونُقلت وفقا للمتطلبات الأخلاقية والتنظيمية المعمول بها، ويتحمل صاحب الطلب مسؤولية طبيعتها ومصدرها واستعمالها.",
    ),
}


def text(key, language):
    return TEXT[key][{'fr': 0, 'en': 1, 'ar': 2}.get(language, 0)]


def _direction(paragraph, rtl):
    if not rtl:
        return
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn('w:bidi')) is None:
        ppr.append(OxmlElement('w:bidi'))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _rtl_document(doc, language):
    if language != 'ar':
        return
    for paragraph in doc.paragraphs:
        _direction(paragraph, True)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _direction(paragraph, True)


def _display(row):
    selected = [
        ('☑ ' if option['selected'] else '☐ ') + option['label']
        for option in row.get('options', [])
        if option['selected'] or row.get('all_options')
    ]
    return '\n'.join(selected) if selected else str(row.get('display', ''))


def _kv_table(doc, rows, language, *, dense=False, writable=False):
    if not rows:
        return None
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    align_ibtikar_content_table(table, [6.10, 11.70])
    for item in rows:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(6.10), Cm(11.70)
        cells[0].text = str(item['label'])
        cells[1].text = _display(item)
        if writable:
            row = table.rows[-1]
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row.height = Cm(1.25 if item.get('field_type') == 'textarea' else 0.85)
        set_cant_split(table.rows[-1])
    style_ibtikar_key_value_table(table, dense=dense)
    if language == 'ar':
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _direction(paragraph, True)
    return table


def _control_table(doc, project, metadata, language):
    rows = [
        {'label': text('number', language), 'display': metadata['number']},
        {'label': text('date', language), 'display': metadata['date']},
        {'label': text('source', language), 'display': project.get('source_version') or text('unknown', language)},
        {'label': text('revision', language), 'display': str(metadata.get('revision') or text('unknown', language))},
        {'label': 'Code', 'display': project['service_code']},
        {'label': text('reference', language), 'display': metadata.get('external_reference') or ''},
        {'label': text('count', language), 'display': project.get('sample_count', 0)},
    ]
    return _kv_table(doc, rows, language, dense=True)


def _sample_columns(samples):
    names, labels = [], {}
    for rows in samples:
        for row in rows:
            name = row.get('name') or row.get('label')
            if name not in names:
                names.append(name)
                labels[name] = row.get('label') or name
    return names, labels


def _render_samples(doc, project, language):
    samples = project.get('samples') or []
    if not samples:
        return
    add_ibtikar_section_heading(doc, text('samples', language).upper(), theme=PLAGENOR_THEME)
    if project.get('read_count') is not None:
        p = doc.add_paragraph(f"{text('reads', language)} : {project['read_count']}")
        p.paragraph_format.space_after = Pt(5)
    names, labels = _sample_columns(samples)
    compact = len(names) <= 6 and sum(len(str(labels[n])) for n in names) <= 130
    if compact:
        table = doc.add_table(rows=1, cols=len(names) + 1)
        table.autofit = False
        table.rows[0].cells[0].text = 'N°'
        for index, name in enumerate(names, 1):
            table.rows[0].cells[index].text = str(labels[name])
        for number, rows in enumerate(samples, 1):
            values = {row.get('name') or row.get('label'): _display(row) for row in rows}
            cells = table.add_row().cells
            cells[0].text = f'{number:02d}'
            for index, name in enumerate(names, 1):
                cells[index].text = str(values.get(name, '—'))
        usable = IBTIKAR_CONTENT_WIDTH_CM
        first = 0.90
        remaining = max(usable - first, 1.0)
        widths = [first] + [remaining / max(len(names), 1)] * len(names)
        align_ibtikar_content_table(table, widths)
        style_data_table(table, theme=PLAGENOR_THEME, dense=True, font_name=IBTIKAR_FONT)
        return
    for number, rows in enumerate(samples, 1):
        add_ibtikar_subheading(
            doc, f"{text('sample', language)} {number:02d}", theme=PLAGENOR_THEME,
        )
        _kv_table(doc, rows, language, dense=True)


def _looks_heading(value):
    clean = value.strip()
    if not clean or len(clean) > 95:
        return False
    return bool(re.match(r'^\d+\.\s', clean)) or clean.endswith(':') or clean.lower().startswith(
        ('conditions ', 'sécurité ', 'critères ', 'autres récipients', 'acknowledgment ')
    )


def _render_reference(doc, project, language):
    source = reference_content(project['service_code'])
    if language != 'fr' and (source.get('guidance') or source.get('tables')):
        note = {
            'en': 'The operational requirements below are preserved in their official French source wording.',
            'ar': 'تُحفظ المتطلبات التشغيلية أدناه بصياغتها الرسمية الأصلية باللغة الفرنسية.',
        }.get(language)
        if note:
            table = add_callout(doc, note, theme=PLAGENOR_THEME, font_name=IBTIKAR_FONT)
            align_ibtikar_content_table(table, [IBTIKAR_CONTENT_WIDTH_CM])
    guidance = list(source.get('guidance') or [])
    for notice in project.get('notices') or []:
        # EGTP-IMT already carries the complete official MALDI-TOF guidance
        # extracted from the source form. Its schema notices are condensed
        # reminders of the same ethics/transport requirements and would
        # otherwise be printed a second time.
        if project.get('service_code') == 'EGTP-IMT':
            continue
        if notice and not any(notice.casefold() in value.casefold() or value.casefold() in notice.casefold()
                              for value in guidance if value):
            guidance.append(notice)
    blocks = list(source.get('blocks') or [])
    source_paragraphs = {
        block.get('text') for block in blocks if block.get('type') == 'paragraph'
    }
    for notice in guidance:
        if notice not in source_paragraphs:
            blocks.append({'type': 'paragraph', 'text': notice})
    if blocks:
        add_ibtikar_section_heading(doc, text('guidance', language).upper(), theme=PLAGENOR_THEME)
        for block in blocks:
            if block.get('type') == 'table':
                rows = block.get('rows') or []
                if not rows:
                    continue
                table = doc.add_table(rows=0, cols=max(len(row) for row in rows))
                for row_values in rows:
                    cells = table.add_row().cells
                    for index, value in enumerate(row_values):
                        cells[index].text = str(value)
                col_count = max(len(row) for row in rows)
                align_ibtikar_content_table(
                    table, [IBTIKAR_CONTENT_WIDTH_CM / max(col_count, 1)] * col_count,
                )
                style_data_table(table, theme=PLAGENOR_THEME, dense=True, font_name=IBTIKAR_FONT)
                continue
            value = block.get('text') or ''
            if value.lower().startswith(('très important', 'important')):
                table = add_callout(
                    doc, value, title='Important', theme=PLAGENOR_THEME, kind='warning',
                    font_name=IBTIKAR_FONT,
                )
                align_ibtikar_content_table(table, [IBTIKAR_CONTENT_WIDTH_CM])
            elif _looks_heading(value):
                add_ibtikar_subheading(doc, value, theme=PLAGENOR_THEME)
            else:
                p = doc.add_paragraph(value)
                p.paragraph_format.space_after = Pt(3)
    return source.get('ethics') or text('ethics_body', language)


def _render_attachments(doc, rows, language):
    if not rows:
        return
    add_ibtikar_section_heading(doc, text('attachments', language).upper(), theme=PLAGENOR_THEME)
    _kv_table(doc, rows, language, dense=True)


def _signature_image(doc, signature_bytes):
    if signature_bytes:
        try:
            doc.add_picture(io.BytesIO(signature_bytes), width=Cm(4.0))
            return
        except Exception:
            pass
    table = doc.add_table(rows=1, cols=1)
    align_ibtikar_content_table(table, [IBTIKAR_CONTENT_WIDTH_CM])
    cell = table.cell(0, 0)
    cell.text = '\n\n'
    from documents.document_design import set_cell_border, set_cell_margins
    set_cell_border(cell, top=(PLAGENOR_THEME.line, 3), bottom=(PLAGENOR_THEME.line, 3),
                    start=(PLAGENOR_THEME.line, 3), end=(PLAGENOR_THEME.line, 3))
    set_cell_margins(cell, top=120, bottom=360, start=120, end=120)


def build_document(project, metadata, language='fr', attachment_rows=None,
                   signature_bytes=None, legacy=None):
    doc = Document()
    apply_document_style(doc, PLAGENOR_THEME, dense=True)
    for section in doc.sections:
        section.top_margin = Cm(0.45)
        section.left_margin = Cm(0.85)
        section.right_margin = Cm(0.85)
        section.bottom_margin = Cm(1.25)
        section.header_distance = Cm(0.25)
    normal = doc.styles['Normal']
    normal.font.name = IBTIKAR_FONT
    normal.font.size = Pt(9.5)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        fonts.set(qn(f'w:{key}'), IBTIKAR_FONT)
    for style_name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[style_name]
        style.font.name = IBTIKAR_FONT
        style_rpr = style.element.get_or_add_rPr()
        style_fonts = style_rpr.get_or_add_rFonts()
        for key in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            style_fonts.set(qn(f'w:{key}'), IBTIKAR_FONT)
    add_invisible_header_drawing(doc)
    add_ibtikar_master_header(
        doc,
        form_title=text('form', language),
        service_label=text('service_requested', language),
        service_title=project['title'],
        service_code=project['service_code'],
    )
    if metadata.get('draft'):
        table = add_callout(
            doc, text('draft', language), theme=PLAGENOR_THEME, kind='warning',
            font_name=IBTIKAR_FONT,
        )
        align_ibtikar_content_table(table, [IBTIKAR_CONTENT_WIDTH_CM])
    add_ibtikar_section_heading(
        doc, f"1. {text('general', language)}", icon='general', theme=PLAGENOR_THEME,
    )
    _control_table(doc, project, metadata, language)
    if legacy:
        table = add_callout(
            doc, text('legacy', language), theme=PLAGENOR_THEME, font_name=IBTIKAR_FONT,
        )
        align_ibtikar_content_table(table, [IBTIKAR_CONTENT_WIDTH_CM])
    if project.get('applicant'):
        add_ibtikar_section_heading(
            doc, f"2. {text('requester', language).upper()}", icon='user',
            theme=PLAGENOR_THEME,
        )
        _kv_table(doc, project['applicant'], language)
    if project.get('parameters'):
        add_ibtikar_section_heading(doc, text('parameters', language).upper(), theme=PLAGENOR_THEME)
        _kv_table(doc, project['parameters'], language)
    _render_samples(doc, project, language)
    _render_attachments(doc, attachment_rows or [], language)
    ethics = _render_reference(doc, project, language)
    add_ibtikar_section_heading(doc, text('ethics', language).upper(), theme=PLAGENOR_THEME)
    table = add_callout(doc, ethics, theme=PLAGENOR_THEME, font_name=IBTIKAR_FONT)
    align_ibtikar_content_table(table, [IBTIKAR_CONTENT_WIDTH_CM])
    add_ibtikar_section_heading(doc, text('signature', language).upper(), theme=PLAGENOR_THEME)
    _signature_image(doc, signature_bytes)
    page_break = doc.add_paragraph()
    page_break.paragraph_format.page_break_before = True
    heading = add_ibtikar_section_heading(doc, text('staff', language).upper(), theme=PLAGENOR_THEME)
    p = doc.add_paragraph(text('staff_help', language))
    p.paragraph_format.left_indent = Cm(0)
    if metadata.get('operator_name'):
        p = doc.add_paragraph(f"{text('operator', language)} : {metadata['operator_name']}")
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.left_indent = Cm(0)
    staff_rows = project.get('staff') or []
    _kv_table(doc, staff_rows, language, writable=True)
    add_ibtikar_signature_grid(
        doc, [text('operator_signature', language), text('head', language), text('director', language)],
        theme=PLAGENOR_THEME,
    )
    p = doc.add_paragraph(text('unsigned', language))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.left_indent = Cm(IBTIKAR_CONTENT_INDENT_CM)
    if legacy:
        add_section_heading(
            doc, text('legacy', language), theme=PLAGENOR_THEME, font_name=IBTIKAR_FONT,
        )
        _kv_table(doc, legacy, language, dense=True)
    add_document_footer(
        doc, theme=PLAGENOR_THEME, reference=metadata['number'], font_name=IBTIKAR_FONT,
    )
    _rtl_document(doc, language)
    return doc


def generate_canonical_form(req):
    from django.conf import settings
    from django.utils.translation import get_language
    from core.ibtikar.models import IbtikarSubmission
    from core.ibtikar.schema import (
        active_data, active_names, label, projection, reference_projection, schema_for_service,
    )
    from core.ibtikar.legacy import document_initial
    language = (get_language() or 'fr').split('-')[0]
    submission = IbtikarSubmission.objects.filter(request=req).first()
    attachments, signature, legacy = [], None, None
    if submission:
        schema = submission.schema
        project = reference_projection(submission, language)
        active_samples = [
            active_data(schema, 'samples', row, submission.parameters)
            for row in submission.samples
        ]
        active = active_names(
            schema['attachments'], {}, submission.parameters, active_samples,
        )
        for obj in submission.attachments.filter(active=True):
            if obj.field_name not in active:
                continue
            spec = next((field for field in schema['attachments']
                         if field['name'] == obj.field_name), None)
            if not spec:
                continue
            attachments.append({
                'label': label(spec['label'], language),
                'display': obj.original_name,
            })
            if obj.field_name == 'applicant_signature':
                with obj.file.open('rb') as stream:
                    signature = stream.read()
    else:
        schema = schema_for_service(req.service)
        old = document_initial(req, schema)
        project = projection(
            schema, old['document_applicant'], old['parameters'],
            old['document_samples'], language=language, print_blank_staff=True,
        )
        legacy = old['legacy_display']
    project['staff'] = [
        row for row in project['staff']
        if row['name'] not in ('validated_price', 'price_justification')
    ]
    metadata = {
        'number': req.display_id,
        'date': req.created_at.strftime('%d/%m/%Y'),
        'external_reference': req.ibtikar_external_code,
        'revision': submission.revision if submission else text('legacy_revision', language),
        'draft': req.status == 'DRAFT',
        'operator_name': submission.staff.get('operator_name') if submission else None,
    }
    doc = build_document(project, metadata, language, attachments, signature, legacy)
    from documents.generators import _inject_document_blocks
    _inject_document_blocks(doc, 'IBTIKAR_FORM', req)
    directory = Path(settings.MEDIA_ROOT) / 'documents'
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f'IBTIKAR_{req.pk}_{uuid.uuid4().hex}.docx'
    doc.save(path)
    return str(path)
