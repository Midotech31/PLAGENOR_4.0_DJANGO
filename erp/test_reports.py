from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import uuid

from django.contrib.auth.models import AnonymousUser
from django.core.exceptions import PermissionDenied,ValidationError
from django.test import TestCase,override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from erp.models import Article,Capability,InventoryLine,StockContainer,WorkItem
from erp.services.biobank import receive_sample
from erp.services.inventory import create_inventory,count_inventory
from erp.services.reports import export_report,period,render_value,report,stock_value_estimate
from erp.services.stock import remove_stock,reconcile_stock,reverse_stock,transfer_stock
from erp.services.storage import save_location
from erp.services.work import create_work
from erp.test_operations import OperationFixtures


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],SECURE_SSL_REDIRECT=False,
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
        'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ReportTests(OperationFixtures,TestCase):
    def rows(self,kind,user=None,**filters):
        dataset=report(user or self.ops,kind,{'year':timezone.localdate().year,**filters})
        return [dataset.serialize(obj) for obj in dataset.queryset]

    def test_stock_available_quantities_and_filters_keep_incompatible_units_separate(self):
        first,_,_=self.receive('A',quantity=10)
        second,_,_=self.receive('B',quantity=20,accepted=False)
        rows=self.rows('STOCK')
        self.assertEqual(len(rows),2)
        bycode={row[2]:row for row in rows}
        self.assertEqual(bycode[first.code][7],10)
        self.assertEqual(bycode[second.code][7],0)
        self.assertEqual(len(self.rows('STOCK',state='USABLE')),1)
        self.assertEqual(len(self.rows('STOCK',state='BLOCKED')),1)
        self.assertEqual(len(self.rows('STOCK',state='EMPTY')),0)
        self.assertEqual(len(self.rows('STOCK',state='EXPIRING',days=10)),0)
        self.assertEqual(len(self.rows('STOCK',q=first.code)),1)
        self.assertEqual(len(self.rows('STOCK',article=self.liquid)),0)
        self.assertEqual(len(self.rows('STOCK',location=self.lab)),2)

    def test_consumption_monthly_yearly_and_loss_counts_exclude_reversals(self):
        container,_,_=self.receive(quantity=100)
        first=remove_stock(self.ops,container.pk,key=uuid.uuid4(),amount=10,unit=self.unit,reason='Analyse')
        cancelled=remove_stock(self.ops,container.pk,key=uuid.uuid4(),amount=5,unit=self.unit,reason='Saisie erronée')
        reverse_stock(self.ops,cancelled.pk,key=uuid.uuid4(),reason='Annulation de la saisie')
        loss=remove_stock(self.ops,container.pk,key=uuid.uuid4(),amount=3,unit=self.unit,kind='LOSS',reason='Perte')
        monthly=self.rows('CONSUMPTION')
        yearly=self.rows('CONSUMPTION',granularity='YEAR')
        self.assertEqual(len(monthly),1)
        self.assertEqual(monthly[0][3],10)
        self.assertEqual(yearly[0][3],10)
        self.assertEqual(yearly[0][2],timezone.localdate().year)
        self.assertEqual(len(self.rows('LOSSES')),1)
        self.assertEqual(self.rows('LOSSES')[0][5],-3)
        self.assertEqual(len(self.rows('MOVEMENTS')),5)
        self.assertEqual(reconcile_stock(self.ops),[])

    def test_receipt_history_uses_event_location_after_container_transfer(self):
        container,_,_=self.receive(ordered_on=timezone.localdate()-timedelta(days=15),order_reference='ORDER-TRACE')
        other=save_location(self.ops,{'code':'REPORT-OTHER','name':'Autre stockage','kind':self.storage_kind,'parent':self.lab})
        transfer_stock(self.ops,container.pk,key=uuid.uuid4(),destination=other,reason='Déplacement')
        self.grant(Capability.VIEW_STOCK,location=self.freezer,category=self.category)
        receipt=self.rows('RECEIPTS',user=self.operator)
        self.assertEqual(len(receipt),1)
        self.assertEqual(receipt[0][9],15)
        self.assertEqual(receipt[0][7],'ORDER-TRACE')
        self.assertEqual(self.rows('STOCK',user=self.operator),[])

    def test_inventory_export_does_not_reveal_blind_counts(self):
        container,_,_=self.receive(quantity=10)
        campaign=create_inventory(self.ops,title='Inventaire à l’aveugle',assignee=self.operator,location=self.freezer,blind=True)
        line=campaign.lines.get()
        count_inventory(self.operator,line.pk,expected=line.version,container_version=container.version,amount=8)
        dataset=report(self.operator,'INVENTORY',{'year':timezone.localdate().year,'campaign':campaign})
        rows=[dataset.serialize(row) for row in dataset.queryset]
        self.assertIsNone(rows[0][5])
        self.assertEqual(rows[0][6],8)
        self.assertIsNone(rows[0][7])
        manager=self.rows('INVENTORY',campaign=campaign)[0]
        self.assertEqual(manager[5:8],[10,8,-2])
        output=BytesIO()
        export_report(dataset,output)
        book=load_workbook(BytesIO(output.getvalue()),data_only=False)
        self.assertEqual(book['Données']['F5'].value,'—')
        self.assertEqual(book['Données']['G5'].value,8)
        self.assertEqual(book['Données']['H5'].value,'—')

    def test_samples_and_freezer_capacity_reports_use_anonymous_codes(self):
        box=save_location(self.ops,{'code':'REPORT-BOX','name':'Boîte','kind':self.storage_kind,
            'parent':self.freezer,'grid_rows':2,'grid_columns':3})
        sample=receive_sample(self.ops,key=uuid.uuid4(),code='CODE-ONLY',amount=20,unit=self.ul,location=box,
            position=box.positions.first(),received_on=timezone.localdate(),reason='Stockage de recette')
        self.assertEqual(self.rows('SAMPLES')[0][0],'CODE-ONLY')
        self.assertEqual(self.rows('SAMPLES')[0][5],20)
        rows=self.rows('STORAGE')
        freezer=next(row for row in rows if row[0]==self.freezer.code)
        boxed=next(row for row in rows if row[0]==box.code)
        self.assertEqual(freezer[3:7],[6,1,0,5])
        self.assertEqual(boxed[3:7],[6,1,0,5])
        self.assertNotIn(self.outsider.email,str(rows)) if self.outsider.email else None

    def test_indicative_stock_value_keeps_currencies_separate_and_costs_private(self):
        self.receive('DZD',quantity=10,unit_price=Decimal('12.50'),currency='DZD')
        self.receive('EUR',quantity=2,unit_price=Decimal('2.50'),currency='EUR')
        self.receive('UNKNOWN',quantity=3)
        estimate=stock_value_estimate(self.ops,{})
        self.assertEqual(estimate['currencies'],{'DZD':Decimal('125.00'),'EUR':Decimal('5.00')})
        self.assertEqual(estimate['unpriced_containers'],1)
        self.grant(Capability.VIEW_STOCK,category=self.category)
        limited=stock_value_estimate(self.operator,{})
        self.assertEqual(limited['currencies'],{})
        self.assertEqual(limited['hidden_containers'],3)
        self.grant(Capability.VIEW_COST,category=self.category)
        self.assertEqual(stock_value_estimate(self.operator,{})['currencies'],estimate['currencies'])

    def test_excel_text_is_literal_and_native_views_export_same_filtered_records(self):
        self.article.name='=HYPERLINK(1)'
        self.article.save(update_fields=['name'])
        self.receive()
        self.client.force_login(self.ops)
        url=reverse('erp:reports')
        self.assertEqual(self.client.get(url).status_code,200)
        for kind in ('MOVEMENTS','CONSUMPTION','LOSSES','RECEIPTS','INVENTORY','SAMPLES','STORAGE','TRACE','TASKS'):
            with self.subTest(kind=kind):
                self.assertEqual(self.client.get(url,{'kind':kind,'year':timezone.localdate().year}).status_code,200)
        response=self.client.get(url,{'kind':'STOCK','year':timezone.localdate().year,'export':'xlsx'})
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(b''.join(response.streaming_content)),data_only=False)
        value=book['Données']['B5']
        self.assertEqual(value.data_type,'s')
        self.assertTrue(value.value.startswith(chr(39)+'='))
        self.assertEqual(book['Données']['F5'].value,10)
        self.assertEqual(self.client.get(url,{'year':'bad'}).status_code,400)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(url).status_code,403)

    def test_task_only_member_has_a_usable_report_without_stock_permission(self):
        create_work(self.ops,kind='CONTROL',title='Tâche du membre',assignee=self.operator,due_on=timezone.localdate())
        self.client.force_login(self.operator)
        response=self.client.get(reverse('erp:reports'))
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'Tâche du membre')
        self.assertNotIn('STOCK',dict(response.context['form'].fields['kind'].choices))
        for kind in ('STOCK','RECEIPTS','STORAGE','SAMPLES','TRACE'):
            with self.subTest(kind=kind),self.assertRaises(PermissionDenied):
                report(self.operator,kind,{'year':timezone.localdate().year})
        with self.assertRaises(PermissionDenied):
            report(AnonymousUser(),'TASKS',{'year':timezone.localdate().year})
        with self.assertRaises(ValidationError):
            report(self.ops,'INVALID',{'year':timezone.localdate().year})
        with self.assertRaises(ValidationError):
            period(1999)
        self.assertEqual(render_value(None),'—')
        self.assertEqual(render_value(Decimal('1.200000')),'1.2')
        self.assertEqual(render_value(timezone.localdate()),timezone.localdate().strftime('%d/%m/%Y'))

    def test_streamed_export_preserves_a4_print_settings_and_numeric_cells(self):
        self.receive(quantity=Decimal('10.125'))
        dataset=report(self.ops,'STOCK',{'year':timezone.localdate().year})
        output=BytesIO()
        self.assertEqual(export_report(dataset,output),1)
        book=load_workbook(BytesIO(output.getvalue()))
        sheet=book['Données']
        self.assertEqual(str(sheet.page_setup.paperSize),'9')
        self.assertEqual(sheet.page_setup.orientation,'landscape')
        self.assertEqual(sheet.page_setup.fitToWidth,1)
        self.assertEqual(sheet.page_setup.fitToHeight,0)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(sheet.freeze_panes,'A5')
        self.assertEqual(sheet.print_title_rows,'$1:$4')
        self.assertEqual(sheet.auto_filter.ref,'A4:L5')
        self.assertEqual(sheet['F5'].value,10.125)
        self.assertEqual(sheet['F5'].data_type,'n')
        self.assertTrue(sheet['A4'].font.bold)
        book.close()

    def test_empty_streamed_export_keeps_headers_and_valid_filter(self):
        output=BytesIO()
        self.assertEqual(export_report(report(self.ops,'STOCK',{}),output),0)
        book=load_workbook(BytesIO(output.getvalue()))
        sheet=book['Données']
        self.assertEqual(sheet.auto_filter.ref,'A4:L4')
        self.assertEqual(sheet.max_row,4)
        self.assertEqual(sheet['A4'].value,'Article')
        book.close()
