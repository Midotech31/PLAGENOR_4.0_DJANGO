import copy
from decimal import Decimal
import io
import zipfile

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils.translation import override
from openpyxl import load_workbook

from erp.cdc.lot_catalog import fingerprint
from erp.models import Capability, CdcCriterion, CdcRequirement, WorkItem
from erp.services.cdc import create_dossier, save_cdc_item, save_criterion, save_requirement
from erp.services.cdc_exports import criteria_workbook
from erp.services.cdc_exchange import set_lot_active
from erp.test_operations import OperationFixtures


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],
    SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class CdcGridExportTests(OperationFixtures, TestCase):
    def setUp(self):
        self.dossier = create_dossier(self.ops, family='equipment',
            reference='93/SME/SDFM/SG/ESSBO/2026', title='Grille technique',
            assignee=self.operator, allow_costs=True)
        self.item = self.dossier.lots.filter(active=True).first().items.filter(active=True).first()
        save_cdc_item(self.operator, self.item.lot_id, expected=self.dossier.version, pk=self.item.pk,
            values={'estimate_supplier': self.party, 'estimated_price': Decimal('987654.32'),
                'tax_rate': Decimal('19'), 'price_source': 'Estimation interne confidentielle',
                'currency': 'DZD'}, reason='Estimation séparée')
        self.dossier.refresh_from_db()
        save_requirement(self.operator, self.item.pk, expected=self.dossier.version,
            values={'kind': CdcRequirement.Kind.ELIMINATORY, 'position': 1,
                'statement': 'Pureté ≥ 99 %', 'evidence': 'Certificat d’analyse',
                'verification_method': 'Contrôle à réception', 'justification': 'Méthode validée',
                'active': True}, reason='Exigence technique')
        self.dossier.refresh_from_db()
        self.revision = save_criterion(self.operator, self.dossier.pk, expected=self.dossier.version,
            values={'code': 'TECH', 'lot': self.item.lot, 'title': 'Conformité technique',
                'method': CdcCriterion.Method.BINARY, 'weight': Decimal('12.50'),
                'threshold': Decimal('7.25'), 'eliminatory': True,
                'evidence': 'Mémoire technique', 'position': 1, 'active': True}, reason='Grille initiale')
        self.dossier.refresh_from_db()
        self.url = reverse('erp:cdc-criteria-export', args=[self.dossier.pk])

    def workbook(self, revision=None):
        return load_workbook(io.BytesIO(criteria_workbook(revision or self.revision)))

    def changed_snapshot(self):
        revision = copy.copy(self.revision)
        revision.data = copy.deepcopy(revision.data)
        return revision

    def seal(self, revision):
        revision.sha256 = fingerprint({'document': revision.data, 'estimates': revision.estimates})
        return revision

    def test_download_preserves_numeric_scores_provenance_and_excludes_estimates(self):
        self.client.force_login(self.operator)
        count = self.dossier.revisions.count()
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
        self.assertIn(f'CDC-{self.dossier.pk}-R{self.revision.number}-grille.xlsx', response['Content-Disposition'])
        payload = b''.join(response.streaming_content)
        book = load_workbook(io.BytesIO(payload))
        criteria, requirements, trace = book.worksheets
        self.assertEqual(criteria['A2'].value, 'TECH')
        self.assertEqual(criteria['E2'].value, 12.5)
        self.assertEqual(criteria['F2'].value, 7.25)
        self.assertEqual(criteria['E2'].data_type, 'n')
        self.assertEqual(requirements['D2'].value, self.item.designation)
        self.assertEqual(requirements['F2'].value, 'Pureté ≥ 99 %')
        self.assertEqual(trace['B2'].value, self.dossier.reference)
        self.assertEqual(trace['B3'].value, self.revision.number)
        self.assertEqual(trace['B4'].value, str(self.revision.pk))
        self.assertEqual(trace['B5'].value, self.revision.sha256)
        self.assertEqual(trace['B6'].value, self.revision.created_at.isoformat())
        self.assertEqual(trace['B7'].value, str(self.operator.pk))
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            source = b''.join(archive.read(name) for name in archive.namelist() if name.endswith('.xml'))
        for private in (b'987654.32', b'Estimation interne confidentielle'):
            self.assertNotIn(private, source)
        for sheet in book:
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
            self.assertEqual(sheet.page_setup.paperSize, int(sheet.PAPERSIZE_A4))
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(self.dossier.revisions.count(), count)
        self.assertEqual(self.client.post(self.url).status_code, 405)

    def test_export_uses_frozen_lot_item_and_criteria_after_live_changes(self):
        frozen = self.revision.data['lot_catalog']['lots'][0]
        self.item.lot.__class__.objects.filter(pk=self.item.lot_id).update(name='Lot modifié')
        self.item.__class__.objects.filter(pk=self.item.pk).update(designation='Article modifié')
        CdcCriterion.objects.filter(dossier=self.dossier).update(title='Critère modifié', weight=99)
        CdcRequirement.objects.filter(item=self.item).update(statement='Exigence modifiée')
        self.dossier.reference = '94/SME/SDFM/SG/ESSBO/2026'
        self.dossier.save(update_fields=['reference'])
        with self.assertNumQueries(0):
            book = self.workbook()
        self.assertEqual(book.worksheets[0]['B2'].value, frozen['name'])
        self.assertEqual(book.worksheets[0]['C2'].value, 'Conformité technique')
        self.assertEqual(book.worksheets[0]['E2'].value, 12.5)
        self.assertEqual(book.worksheets[1]['D2'].value, self.item.designation)
        self.assertEqual(book.worksheets[1]['F2'].value, 'Pureté ≥ 99 %')
        self.assertEqual(book.worksheets[2]['B2'].value, self.revision.data['reference'])

    def test_formula_like_text_is_exact_literal_in_all_sheets(self):
        revision = self.changed_snapshot()
        revision.data['reference'] = '=1+1'
        row = revision.data['criteria'][0]
        row.update(code='=SUM(1,2)', title='+addition', evidence='@command')
        revision.data['requirements'][0]['statement'] = '-negative text'
        revision.data['lot_catalog']['lots'][0]['name'] = '=HYPERLINK("https://example.invalid","lot")'
        book = self.workbook(self.seal(revision))
        self.assertEqual(book.worksheets[0]['A2'].value, '=SUM(1,2)')
        self.assertEqual(book.worksheets[0]['C2'].value, '+addition')
        self.assertEqual(book.worksheets[0]['H2'].value, '@command')
        self.assertEqual(book.worksheets[1]['F2'].value, '-negative text')
        self.assertEqual(book.worksheets[2]['B2'].value, '=1+1')
        for sheet in book:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f')
                    self.assertIsNone(cell.hyperlink)

    def test_incomplete_global_and_legacy_grids_can_be_reviewed(self):
        revision = self.changed_snapshot()
        revision.data['criteria'][0].update(lot=None, threshold=None, eliminatory=False)
        book = self.workbook(self.seal(revision))
        self.assertEqual(book.worksheets[0]['B2'].value, 'Global')
        self.assertIsNone(book.worksheets[0]['F2'].value)
        self.assertEqual(book.worksheets[0]['G2'].value, 'Non')
        for key in ('criteria', 'requirements', 'lot_catalog'):
            revision.data.pop(key)
        book = self.workbook(self.seal(revision))
        self.assertEqual([s.max_row for s in book], [1, 1, 7])

    def test_criterion_for_removed_lot_preserves_its_snapshot_identity(self):
        revision = set_lot_active(self.operator, self.item.lot_id, expected=self.dossier.version,
            active=False, reason='Lot retiré pour réexamen de la grille')
        book = self.workbook(revision)
        self.assertEqual(book.worksheets[0]['B2'].value, str(self.item.lot_id))
        self.assertEqual(book.worksheets[0]['C2'].value, 'Conformité technique')
        self.assertEqual(book.worksheets[1].max_row, 1)

    def test_reader_permissions_and_revocation(self):
        self.assertEqual(self.client.get(self.url).status_code, 302)
        for user in (self.outsider, self.second):
            self.client.force_login(user)
            self.assertEqual(self.client.get(self.url).status_code, 404)
        self.grant(Capability.REVIEW_CDC_TECHNICAL, user=self.second)
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.dossier.work.status = WorkItem.Status.SUBMITTED
        self.dossier.work.save(update_fields=['status'])
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(reverse('erp:cdc-criteria-export', args=[self.dossier.pk]),
            self.client.get(reverse('erp:cdc-governance', args=[self.dossier.pk])).content.decode())
        self.dossier.work.status = WorkItem.Status.DRAFT
        self.dossier.work.assignee = self.second
        self.dossier.work.save(update_fields=['status', 'assignee'])
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(self.url).status_code, 404)

    def test_translated_headers_and_arabic_sheet_direction(self):
        for language, title, header in (('en', 'Criteria', 'Criterion'), ('ar', 'المعايير', 'المعيار')):
            with self.subTest(language=language), override(language):
                book = self.workbook()
                self.assertEqual(book.worksheets[0].title, title)
                self.assertEqual(book.worksheets[0]['C1'].value, header)
                self.assertTrue(all(bool(sheet.sheet_view.rightToLeft) == (language == 'ar') for sheet in book))

    def test_corrupt_or_unrepresentable_snapshots_fail_without_silent_loss(self):
        revision = self.changed_snapshot()
        revision.data['reference'] = 'Altérée'
        with self.assertRaisesRegex(ValidationError, 'intégrité'):
            self.workbook(revision)
        for change in ('method', 'nan', 'number', 'control', 'long'):
            with self.subTest(change=change):
                revision = self.changed_snapshot()
                row = revision.data['criteria'][0]
                if change == 'method':
                    row['method'] = 'METHODE_INCONNUE'
                elif change == 'nan':
                    row['weight'] = 'NaN'
                elif change == 'number':
                    row['weight'] = 'invalide'
                elif change == 'control':
                    row['title'] = 'Contrôle\x00'
                else:
                    row['evidence'] = 'x' * 32768
                with self.assertRaises(ValidationError):
                    self.workbook(self.seal(revision))
        self.client.force_login(self.operator)
        from unittest.mock import patch
        with patch('erp.cdc_views.criteria_workbook', side_effect=ValidationError('Export impossible')):
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.content, b'Export impossible')

    def test_missing_current_revision_is_not_replaced_by_live_data(self):
        self.dossier.revision_number += 1
        self.dossier.save(update_fields=['revision_number'])
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(self.url).status_code, 404)
