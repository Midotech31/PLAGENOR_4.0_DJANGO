import io
from pathlib import PurePath

from django.core.exceptions import ValidationError
from django.utils import translation
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

import threading

_VALIDATION_SLOT=threading.BoundedSemaphore(1)


MAX_FILE=10*1024*1024
MAX_ROWS=500
MAX_COLUMNS=40

LABELS={
    'code':_('Code interne'),'name':_('Désignation française'),'name_en':_('Désignation anglaise'),
    'name_ar':_('Désignation arabe'),'category_code':_('Code catégorie'),'base_unit_code':_('Code unité de gestion'),
    'purchase_unit_code':_('Code unité d’achat'),'manufacturer_code':_('Code fabricant'),
    'manufacturer_reference':_('Référence fabricant'),'supplier_code':_('Code fournisseur'),
    'cas':_('Numéro CAS'),'packaging':_('Conditionnement décrit'),'specifications':_('Spécifications techniques'),
    'minimum_stock':_('Stock minimal'),'safety_stock':_('Stock de sécurité'),'reorder_point':_('Seuil de réapprovisionnement'),
    'target_stock':_('Stock cible'),'order_multiple':_('Multiple de commande'),'minimum_order_quantity':_('Commande minimale (unité d’achat)'),
    'lead_time_days':_('Délai fournisseur (jours)'),'criticality':_('Criticité'),
    'article_code':_('Code article'),'location_code':_('Code emplacement'),'lot_code':_('Code interne du lot'),
    'manufacturer_lot':_('Lot fabricant'),'container_code':_('Code interne du contenant'),'amount':_('Quantité'),
    'unit_code':_('Code unité'),'received_on':_('Date de réception'),'expires_on':_('Date de péremption'),
    'manufactured_on':_('Date de fabrication'),'condition':_('État à la réception'),'cold_chain_ok':_('Chaîne du froid respectée'),
    'unit_price_base':_('Prix HT par unité de gestion'),'currency':_('Devise'),'order_reference':_('Commande / marché'),
    'kind_code':_('Code type d’emplacement'),'parent_code':_('Code emplacement parent'),
    'grid_rows':_('Nombre de lignes'),'grid_columns':_('Nombre de colonnes'),'capacity':_('Capacité'),
    'temperature_target':_('Température cible (°C)'),'temperature_min':_('Température minimale (°C)'),
    'temperature_max':_('Température maximale (°C)'),'position_row':_('Ligne de la position'),'position_column':_('Colonne de la position'),
    'request_reference':_('Référence de la demande PLAGENOR'),'source_code':_('Code échantillon déclaré'),
    'sample_type':_('Type d’échantillon'),'matrix':_('Matrice'),'preservation':_('Conservation documentée'),
    'observed_on':_('Date de référence du prix'),'price':_('Prix dans l’unité indiquée'),'source':_('Source documentaire'),
    'lot_name':_('Lot d’achat'),'retained_quantity':_('Quantité retenue en unité d’achat'),
    'estimated_price':_('Prix unitaire estimé hors taxes'),'tax_rate':_('Taux de taxe (%)'),
    'decision_reason':_('Justification de la décision'),'measured_at':_('Date et heure du relevé'),
    'value':_('Température mesurée (°C)'),'comment':_('Observation'),
    'new_name':_('Désignation d’un nouvel article'),'new_category_code':_('Catégorie d’un nouvel article'),
    'new_base_unit_code':_('Unité de gestion d’un nouvel article')}

SCHEMAS={
    'CATALOG':{'required':['code','name','category_code','base_unit_code'],
        'optional':['name_en','name_ar','purchase_unit_code','manufacturer_code','manufacturer_reference','supplier_code','cas',
            'packaging','specifications','minimum_stock','safety_stock','reorder_point','target_stock','order_multiple',
            'minimum_order_quantity','lead_time_days','criticality']},
    'INITIAL':{'required':['article_code','location_code','lot_code','manufacturer_lot','container_code','amount','unit_code','received_on','condition'],
        'optional':['expires_on','manufactured_on','supplier_code','cold_chain_ok','unit_price_base','currency','new_name','new_category_code','new_base_unit_code']},
    'RECEIPTS':{'required':['article_code','location_code','lot_code','manufacturer_lot','container_code','amount','unit_code','received_on','condition'],
        'optional':['expires_on','manufactured_on','supplier_code','cold_chain_ok','unit_price_base','currency','order_reference']},
    'LOCATIONS':{'required':['code','name','kind_code'],
        'optional':['parent_code','name_en','name_ar','grid_rows','grid_columns','capacity','temperature_target','temperature_min','temperature_max']},
    'SAMPLES':{'required':['code','location_code','amount','unit_code','received_on'],
        'optional':['position_row','position_column','request_reference','source_code','sample_type','matrix','preservation']},
    'PRICES':{'required':['article_code','unit_code','price','currency','observed_on','source'],'optional':['supplier_code']},
    'PLAN':{'required':['article_code','lot_name','retained_quantity','decision_reason'],
        'optional':['estimated_price','tax_rate','currency','source','supplier_code']},
    'TEMPERATURE':{'required':['location_code','measured_at','value'],'optional':['comment']}}


def schema(kind):
    if kind not in SCHEMAS:
        raise ValidationError(_('Domaine d’import inconnu.'))
    return SCHEMAS[kind]


def aliases(kind):
    fields=schema(kind)
    names=fields['required']+fields['optional']
    mapping={name.casefold():name for name in names}
    for language in ('fr','en','ar'):
        with translation.override(language):
            mapping.update({str(LABELS[name]).strip().casefold():name for name in names})
    return mapping


def read_matrix(data,extension):
    import os
    import subprocess
    import sys
    import json
    from pathlib import Path
    if extension not in ('.xlsx','.csv'):
        raise ValidationError(_('Utilisez un fichier XLSX sans macros ou un CSV UTF-8.'))
    if not _VALIDATION_SLOT.acquire(timeout=1):
        raise ValidationError(_('Validation occupée. Réessayez dans quelques instants.'))
    try:
        environment={key:value for key,value in os.environ.items() if key.upper() in {'SYSTEMROOT','WINDIR','PATH','TEMP','TMP'}}
        result=subprocess.run([sys.executable,'-I',str(Path(__file__).resolve().parents[1]/'table_probe.py'),extension],
            input=data,stdout=subprocess.PIPE,stderr=subprocess.DEVNULL,timeout=8,env=environment,
            creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        if len(result.stdout)>6*1024*1024:
            raise ValidationError(_('Le tableau dépasse les limites de validation.'))
        value=json.loads(result.stdout)
        if result.returncode!=0 or 'rows' not in value:
            code=value.get('error','INVALID')
            messages={'FORMULA':_('Remplacez les formules par des valeurs avant l’import.'),
                'EXTERNAL':_('Les liens externes ne sont pas admis dans un import.'),
                'ACTIVE':_('Les macros et objets incorporés ne sont pas admis.'),
                'SHEETS':_('Le classeur doit comporter une seule feuille de données et, éventuellement, une feuille Guide.'),
                'SIZE':_('Le tableau dépasse les limites : 500 lignes, 40 colonnes ou 4 Mo de texte utile.'),
                'EMPTY':_('Le tableau ne contient aucune ligne de données.')}
            raise ValidationError(messages.get(code,_('Le fichier contient une structure invalide ou non autorisée.')))
        return value['rows']
    except (OSError,subprocess.TimeoutExpired,ValueError,TypeError) as exc:
        raise ValidationError(_('Le fichier ne peut pas être validé dans les limites de temps et de mémoire autorisées.')) from exc
    finally:
        _VALIDATION_SLOT.release()


def parse_table(kind,filename,data):
    expected=schema(kind)
    if not isinstance(data,bytes) or not data or len(data)>MAX_FILE:
        raise ValidationError(_('Le fichier doit contenir des données et ne pas dépasser 10 Mo.'))
    extension=PurePath(filename).suffix.lower()
    matrix=read_matrix(data,extension)
    while matrix and not any(matrix[-1]):
        matrix.pop()
    if len(matrix)<2:
        raise ValidationError(_('Le tableau ne contient aucune ligne de données.'))
    if len(matrix)>MAX_ROWS+1:
        raise ValidationError(_('Le tableau dépasse 500 lignes de données.'))
    mapping=aliases(kind)
    raw_headers=matrix[0]
    while raw_headers and not raw_headers[-1]:
        raw_headers.pop()
    headers=[mapping.get(value.strip().casefold()) for value in raw_headers]
    if None in headers or len(headers)!=len(set(headers)):
        raise ValidationError(_('Une colonne est inconnue, vide ou dupliquée. Utilisez le modèle fourni.'))
    missing=set(expected['required'])-set(headers)
    if missing:
        raise ValidationError(_('Colonnes obligatoires manquantes : %(columns)s') % {'columns':', '.join(str(LABELS[name]) for name in sorted(missing))})
    rows=[]
    for index,values in enumerate(matrix[1:],2):
        if not any(values):
            continue
        if any(values[len(headers):]):
            raise ValidationError(_('La ligne %(row)s contient des cellules sans en-tête.') % {'row':index})
        padded=values+['']*max(0,len(headers)-len(values))
        rows.append({'row':index,'data':dict(zip(headers,padded))})
    if not rows:
        raise ValidationError(_('Le tableau ne contient aucune ligne de données.'))
    return rows


def import_template(kind):
    fields=schema(kind)
    names=fields['required']+fields['optional']
    book=Workbook()
    sheet=book.active
    sheet.title='Données'
    for index,name in enumerate(names,1):
        cell=sheet.cell(1,index,str(LABELS[name]))
        cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='203F50' if name in fields['required'] else '496778')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.column_dimensions[get_column_letter(index)].width=26
    sheet.row_dimensions[1].height=40
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=f'A1:{get_column_letter(len(names))}1'
    guide=book.create_sheet('Guide')
    content=[(_('Mode d’emploi'),_('Conservez les en-têtes. Saisissez vos données réelles dans la feuille Données. Aucune ligne d’exemple n’est importée.')),
        (_('Colonnes obligatoires'),', '.join(str(LABELS[name]) for name in fields['required'])),
        (_('Références'),_('Les codes doivent correspondre aux référentiels PLAGENOR. Les données inconnues sont signalées avant toute application.')),
        (_('Quantités'),_('Utilisez des nombres avec au plus six décimales ; l’unité est donnée dans la colonne Code unité.')),
        (_('Dates'),_('Dates : AAAA-MM-JJ. Relevés : AAAA-MM-JJ HH:MM, dans le fuseau horaire de la plateforme.')),
        (_('Formules'),_('Collez uniquement des valeurs. Les macros, formules et liens externes sont refusés.')),
        (_('Aperçu'),_('Le premier dépôt produit un aperçu et un rapport d’erreurs. L’import n’est appliqué qu’après confirmation explicite.')),
        (_('Mise à jour'),_('Les articles du catalogue et les emplacements existants sont mis à jour uniquement si leurs versions n’ont pas changé depuis l’aperçu.')),
        (_('Nouveaux articles'),_('Pour le stock initial, renseignez ensemble la désignation, la catégorie et l’unité de gestion du nouvel article ; leur création sera soumise aux droits du catalogue.')),
        (_('Prix des réceptions'),_('Le prix de réception est exprimé par unité de gestion, pas par conditionnement. Les prix dans une autre unité peuvent être importés comme observations de prix.'))]
    for index,(label,text) in enumerate(content,1):
        guide.cell(index,1,str(label)).font=Font(name='Calibri',bold=True,color='203F50')
        guide.cell(index,2,str(text)).alignment=Alignment(wrap_text=True,vertical='center')
        guide.row_dimensions[index].height=48
    guide.column_dimensions['A'].width=30
    guide.column_dimensions['B'].width=110
    buffer=io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
