import copy
import io
import json
from datetime import timedelta
from decimal import Decimal

from django import forms
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.cell.cell import IllegalCharacterError
from openpyxl.styles.numbers import is_date_format

from erp.cdc.docengine import DocumentError, safe_zip, sha
from erp.cdc.financial import excel_number, summarize
from erp.cdc.lot_catalog import MAX_ITEMS, diff_catalog, fingerprint
from erp.cdc.lot_workbook import MAX_FILE, Q, _sheet_paths, read_cells, xml
from erp.models import CdcItem, CdcWorkbookPreview
from .cdc import _dossier, cdc_cost_allowed
from .cdc_exports import _sheet
from .common import check_version

SALT = 'plagenor.cdc.finance.v1'


def financial_snapshot(dossier):
    revision = dossier.revisions.get(number=dossier.revision_number)
    if fingerprint({'document': revision.data, 'estimates': revision.estimates}) != revision.sha256:
        raise ValidationError(_('L’intégrité de la révision CDC ne peut pas être confirmée.'))
    legacy_ids = [entry['id'] for entry in revision.estimates if 'source_key' not in entry]
    keys = {str(item.pk): item.source_key for item in CdcItem.objects.filter(pk__in=legacy_ids).only('id', 'source_key')}
    estimates = {(entry['lot'], entry.get('source_key', keys.get(entry['id']))): entry for entry in revision.estimates}
    rows = []
    for lot in revision.data['lot_catalog']['lots']:
        for item in lot['items']:
            entry = estimates[(lot['id'], item['key'])]
            rows.append({'id': entry['id'], 'key': item['key'], 'lot': lot['id'], 'lot_name': lot['name'],
                'designation': item['designation'], 'unit': item['unit'], 'quantity': Decimal(item['quantity']),
                'price': Decimal(entry['price']) if entry['price'] is not None else None,
                'tax_rate': Decimal(entry['tax_rate']) if entry['tax_rate'] is not None else None,
                'currency': entry['currency'], 'source': entry['source']})
    return revision, summarize(rows)


def export_financial(user, dossier):
    if not cdc_cost_allowed(user, dossier):
        raise PermissionDenied
    revision, summary = financial_snapshot(dossier)
    if any(row['currency'] != 'DZD' for row in summary['lines']):
        raise ValidationError(_('Le classeur financier est en DZD ; aucune conversion implicite n’est autorisée.'))
    headers = [_('Lot'), _('Désignation'), _('Unité documentaire'), _('Quantité prévue'),
        _('Prix unitaire estimé hors taxes'), _('Taux de taxe (%)'), _('Hors taxes'), _('Taxes'), _('Total TTC'), _('Identifiant interne')]
    names = [_('Estimations'), _('Totaux par lot'), _('Traçabilité')]
    meta = {'schema': 1, 'dossier': str(dossier.pk), 'revision': revision.number, 'revision_id': str(revision.pk),
        'sha256': revision.sha256, 'sheets': names, 'headers': headers}
    token = signing.dumps(meta, salt=SALT, compress=True)
    book = Workbook(); book.remove(book.active)
    try:
        estimates = _sheet(book, names[0], headers,
            [tuple(excel_number(value) for value in (row['lot_name'], row['designation'], row['unit'],
                row['quantity'], row['price'], row['tax_rate'], row['net'], row['tax'], row['gross'], row['id']))
                for row in summary['lines']], (36, 48, 20, 18, 22, 18, 22, 22, 22, 38))
        estimates.column_dimensions['J'].hidden = True
        for row in estimates.iter_rows(min_row=2, min_col=5, max_col=9):
            for cell in row:
                cell.number_format = '#,##0.00'
        totals = [(lot['name'], currency, *(excel_number(values[key]) for key in ('net', 'tax', 'gross')),
                   lot['incomplete_lines']) for lot in summary['lots'] for currency, values in lot['currencies'].items()]
        totals += [(lot['name'], 'DZD', None, None, None, lot['incomplete_lines'])
                   for lot in summary['lots'] if not lot['currencies']]
        totals += [(_('Global'), currency, *(excel_number(values[key]) for key in ('net', 'tax', 'gross')),
                    summary['incomplete_lines']) for currency, values in summary['currencies'].items()]
        _sheet(book, names[1], (_('Lot'), _('Devise'), _('Hors taxes'), _('Taxes'), _('Total TTC'), _('Articles incomplets')),
            totals, (36, 14, 22, 22, 22, 22))
        _sheet(book, names[2], (_('Champ'), _('Valeur')), [(_('Référence CDC'), dossier.reference),
            (_('Révision'), revision.number), (_('Empreinte SHA-256'), revision.sha256),
            (_('Mode d’emploi'), _('Modifiez uniquement les quantités, prix unitaires et taux. Les totaux sont recalculés après import. Une cellule de prix ou de taux vide efface cette valeur.')),
            *[(row['id'], row['source']) for row in summary['lines']]], (38, 90))
        identity = book.create_sheet('_CDC')
        identity['A1'] = 'PLAGENOR_CDC_FINANCE_V1'; identity['B2'] = token
        identity.sheet_state = 'veryHidden'
        output = io.BytesIO(); book.save(output)
        return output.getvalue()
    except (ValueError, IllegalCharacterError) as error:
        raise ValidationError(_('Le classeur financier contient une valeur incompatible avec Excel.')) from error


def parse_financial(payload, filename, dossier):
    if not filename.lower().endswith('.xlsx') or len(payload) > MAX_FILE:
        raise DocumentError(_('Choisissez un fichier XLSX de 20 Mo maximum.'))
    revision, summary = financial_snapshot(dossier)
    if any(row['currency'] != 'DZD' for row in summary['lines']):
        raise DocumentError(_('Le classeur financier est en DZD ; aucune conversion implicite n’est autorisée.'))
    with safe_zip(payload, 64 * 1024 * 1024) as archive:
        for name in archive.namelist():
            if any(value in name.lower() for value in ('vbaproject', 'externallinks/', 'embeddings/', 'activex/', 'connections.xml')):
                raise DocumentError(_('Macros, connexions ou objets incorporés interdits.'))
            if name.endswith('.rels') and any(rel.get('TargetMode') == 'External' for rel in xml(archive.read(name))):
                raise DocumentError(_('Les liens externes ne sont pas acceptés.'))
        paths = _sheet_paths(archive)
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            strings = [''.join(node.text or '' for node in value.iter(Q+'t'))
                for value in xml(archive.read('xl/sharedStrings.xml')).findall(Q+'si')]
            if len(strings) > 10000 or sum(map(len, strings)) > 8000000:
                raise DocumentError(_('Chaînes du classeur trop volumineuses.'))
        cells = {name: read_cells(archive, path, strings, maximum_column='J', maximum_cells=(MAX_ITEMS + 2) * 10)
            for name, path in paths.items()}
        val = lambda values, address: values.get(address, {}).get('value', '')
        try:
            identity = cells['_CDC']
            if val(identity, 'A1') != 'PLAGENOR_CDC_FINANCE_V1':
                raise ValueError
            meta = signing.loads(val(identity, 'B2'), salt=SALT)
            if (meta['schema'] != 1 or meta['dossier'] != str(dossier.pk) or meta['revision'] != revision.number
                    or meta['revision_id'] != str(revision.pk) or meta['sha256'] != revision.sha256
                    or set(cells) != {'_CDC', *meta['sheets']}):
                raise ValueError
            sheet = cells[meta['sheets'][0]]
            if [val(sheet, column+'1') for column in 'ABCDEFGHIJ'] != meta['headers']:
                raise ValueError
        except (KeyError, ValueError, TypeError, signing.BadSignature) as error:
            raise DocumentError(_('L’identité du classeur est invalide ou le dossier a changé. Téléchargez sa révision actuelle.')) from error
        try:
            styles = xml(archive.read('xl/styles.xml'))
        except KeyError as error:
            raise DocumentError(_('Le classeur financier est incomplet. Téléchargez un nouveau modèle.')) from error
        formats = {node.get('numFmtId'): node.get('formatCode', '') for node in styles.iter(Q+'numFmt')}
        cell_styles = styles.find(Q+'cellXfs')
        if cell_styles is None or not len(cell_styles):
            raise DocumentError(_('Style numérique Excel invalide.'))
        xfs = list(cell_styles)
        originals = {row['id']: row for row in summary['lines']}
        candidate = copy.deepcopy(revision.data)
        items = {item['key']: item for lot in candidate['lot_catalog']['lots'] for item in lot['items']}
        financial, taxes, seen = {}, {}, set()
        numbers = sorted({int(address[1:]) for address in sheet if int(address[1:]) > 1})
        for number in numbers:
            identifier = val(sheet, 'J'+str(number))
            if not any(val(sheet, column+str(number)) for column in 'ABCDEFGHIJ'):
                continue
            if identifier not in originals or identifier in seen:
                raise DocumentError(_('Identifiant d’article inconnu ou dupliqué dans le classeur financier.'))
            seen.add(identifier)
            original = originals[identifier]
            if [val(sheet, column+str(number)) for column in 'ABC'] != [original['lot_name'], original['designation'], original['unit']]:
                raise DocumentError(_('Les lots, désignations et unités se modifient dans le dossier, pas dans ce classeur.'))
            parsed = []
            for column, precision, scale, minimum, maximum in [('D', 18, 6, Decimal('0.000001'), None),
                ('E', 18, 2, 0, None), ('F', 5, 2, 0, 100)]:
                entry = sheet.get(column+str(number), {})
                try:
                    style = int(entry.get('style', '0'))
                except ValueError as error:
                    raise DocumentError(_('Style numérique Excel invalide.')) from error
                if not 0 <= style < len(xfs):
                    raise DocumentError(_('Style numérique Excel invalide.'))
                fmt = xfs[style].get('numFmtId', '0')
                if fmt in {str(value) for value in (*range(14, 23), 45, 46, 47)} or is_date_format(formats.get(fmt, '')):
                    raise DocumentError(_('Une quantité, un prix ou un taux ne peut pas être une date Excel.'))
                try:
                    value = forms.DecimalField(required=column == 'D', max_digits=precision, decimal_places=scale,
                        min_value=minimum, max_value=maximum).clean(entry.get('value', '').replace(',', '.'))
                except ValidationError as error:
                    raise DocumentError('%s!%s%s : %s' % (meta['sheets'][0], column, number, error.messages[0])) from error
                parsed.append(value)
            quantity, price, rate = parsed
            items[original['key']]['quantity'] = str(quantity)
            financial[original['key']] = str(price) if price is not None else None
            taxes[original['key']] = str(rate) if rate is not None else None
            original.update(quantity=quantity, price=price, tax_rate=rate)
        if seen != set(originals):
            raise DocumentError(_('Des articles sont absents du classeur financier. Aucun retrait implicite n’est autorisé.'))
        return {'data': candidate, 'financial': financial, 'taxes': taxes, 'file_sha256': sha(payload),
            'finance_totals': json.loads(json.dumps(summarize(list(originals.values())), default=str)),
            'diff': diff_catalog(revision.data['lot_catalog'], candidate['lot_catalog'])}


@transaction.atomic
def preview_financial(user, pk, *, expected, upload, reason):
    dossier = _dossier(user, pk, edit=True)
    check_version(dossier, expected)
    if not cdc_cost_allowed(user, dossier):
        raise PermissionDenied
    if not reason.strip():
        raise ValidationError(_('Justifiez l’import et indiquez la source des prix.'))
    if upload.size > MAX_FILE:
        raise ValidationError(_('Le fichier dépasse la limite de 20 Mo.'))
    payload = parse_financial(upload.read(MAX_FILE + 1), upload.name, dossier)
    return CdcWorkbookPreview.objects.create(dossier=dossier, actor=user, base_version=dossier.version,
        filename=upload.name[:180], payload=payload, import_prices=True, reason=reason.strip()[:500],
        expires_at=timezone.now() + timedelta(hours=24))
