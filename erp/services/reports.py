from dataclasses import dataclass
from datetime import date,datetime,time,timedelta
from decimal import Decimal
import io

from django.core.exceptions import PermissionDenied,ValidationError
from django.db.models import Case,When,Count,F,Max,Min,OuterRef,Q,Subquery,Sum
from django.db.models.functions import TruncMonth
from django.utils import timezone
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from erp.models import (Article,BiologicalSample,Capability,InventoryLine,Location,LocationClosure,
    PurchaseReceiptLink,RunConsumption,StockContainer,StockEntry,StockLot,StockReceipt)
from erp.permissions import grants,is_manager,is_team,operational_scope,permitted,storage_scope
from .biobank import biobank_scope
from .report_exports import literal
from .stock import usable_filter
from .work import work_scope


ZERO=Decimal(0)
KINDS=[('TASKS',_('Tâches et responsabilités')),('STOCK',_('État du stock')),('MOVEMENTS',_('Journal des mouvements')),('CONSUMPTION',_('Consommation mensuelle et annuelle')),
    ('LOSSES',_('Pertes et péremptions')),('RECEIPTS',_('Réceptions et délais')),('INVENTORY',_('Rapport d’inventaire physique')),
    ('SAMPLES',_('Échantillons et aliquots')),('STORAGE',_('Occupation du stockage froid')),('TRACE',_('Traçabilité lot → analyses'))]


@dataclass
class Report:
    title: str
    headers: list
    queryset: object
    serialize: object
    notes: str=''


def period(year):
    if type(year) is not int or not 2000<=year<=timezone.localdate().year+2:
        raise ValidationError(_('Année du rapport non valide.'))
    return timezone.make_aware(datetime(year,1,1)),timezone.make_aware(datetime(year+1,1,1))


def _stock(user):
    if not is_manager(user) and not grants(user,Capability.VIEW_STOCK).exists():
        raise PermissionDenied
    return operational_scope(StockContainer.objects.all(),user)


def _filter(qs,filters,*,article_field=None,location_field=None,search_fields=()):
    if filters.get('article') and article_field:
        qs=qs.filter(**{article_field:filters['article'].pk})
    if filters.get('location') and location_field:
        qs=qs.filter(**{location_field+'__in':LocationClosure.objects.filter(ancestor=filters['location']).values('descendant_id')})
    text=filters.get('q','')
    if text and search_fields:
        condition=Q()
        for field in search_fields:
            condition|=Q(**{field+'__icontains':text})
        qs=qs.filter(condition)
    return qs


def report(user,kind,filters):
    if not is_team(user):
        raise PermissionDenied
    if kind not in dict(KINDS):
        raise ValidationError(_('Type de rapport inconnu.'))
    start,end=period(filters.get('year',timezone.localdate().year))
    title=str(dict(KINDS)[kind])
    if kind=='TASKS':
        qs=_filter(work_scope(user),filters,location_field='location_id',search_fields=('title','instructions')).select_related('schedule')
        return Report(title,[_('Tâche'),_('Type'),_('Responsable'),_('Statut'),_('Priorité'),_('Échéance'),_('Début prévu'),_('Fin prévue')],
            qs.order_by('due_on','title'),lambda obj:[obj.title,obj.get_kind_display(),
                obj.assignee.get_full_name() or obj.assignee.username if obj.assignee else '',obj.get_status_display(),
                obj.get_priority_display(),obj.due_on,obj.schedule.starts_at if hasattr(obj,'schedule') else None,
                obj.schedule.ends_at if hasattr(obj,'schedule') else None])
    if kind=='STOCK':
        qs=_filter(_stock(user),filters,article_field='lot__article_id',location_field='location_id',
            search_fields=('code','lot__code','lot__manufacturer_lot','lot__article__name','lot__article__cas','location__code')).select_related('lot__article__base_unit','lot__article__category','location')
        state=filters.get('state')
        if state=='USABLE':
            qs=qs.filter(usable_filter(),quantity__gt=F('reserved'))
        elif state=='BLOCKED':
            qs=qs.filter(quantity__gt=0).exclude(usable_filter())
        elif state=='EXPIRING':
            limit=timezone.localdate()+timedelta(days=filters.get('days') or 30)
            qs=qs.filter(quantity__gt=0).filter(Q(use_by__lte=limit)|Q(lot__expires_on__lte=limit))
        elif state=='EMPTY':
            qs=qs.filter(quantity=0)
        qs=qs.annotate(available_flag=Case(
            When(usable_filter(),then=True),default=False))
        headers=[_('Article'),_('Désignation'),_('Contenant'),_('Lot fabricant'),_('Emplacement'),_('Stock physique'),_('Réservé'),
            _('Disponible'),_('Unité'),_('Péremption / limite'),_('Contrôle'),_('Criticité')]
        def row(obj):
            expiry=min((day for day in (obj.use_by,obj.lot.expires_on) if day),default=None)
            return [obj.lot.article.code,obj.lot.article.name,obj.code,obj.lot.manufacturer_lot,obj.location.code,
                obj.quantity,obj.reserved,max(ZERO,obj.quantity-obj.reserved) if obj.available_flag else ZERO,
                obj.lot.article.base_unit.code,expiry,obj.get_status_display(),obj.lot.article.get_criticality_display()]
        return Report(title,headers,qs.order_by('lot__article__code','code'),row,str(_('Les quantités de produits ayant des unités différentes ne sont pas additionnées.')))
    if kind in ('MOVEMENTS','CONSUMPTION','LOSSES'):
        _stock(user)
        qs=operational_scope(StockEntry.objects.all(),user,category_field='container__lot__article__category_id')
        qs=_filter(qs,filters,article_field='container__lot__article_id',location_field='location_id',
            search_fields=('container__code','container__lot__manufacturer_lot','container__lot__article__name','movement__request__display_id'))
        qs=qs.filter(created_at__gte=start,created_at__lt=end)
        if kind=='CONSUMPTION':
            qs=qs.filter(quantity_delta__lt=0,movement__kind__in=['CONSUMPTION','PREPARATION'],movement__reversal__isnull=True)
            if filters.get('granularity')=='YEAR':
                grouped=qs.values('container__lot__article__code','container__lot__article__name',
                    'container__lot__article__base_unit__code').annotate(quantity=Sum('quantity_delta')).order_by('container__lot__article__code')
                return Report(title,[_('Article'),_('Désignation'),_('Année'),_('Quantité consommée'),_('Unité')],grouped,
                    lambda obj:[obj['container__lot__article__code'],obj['container__lot__article__name'],start.year,-obj['quantity'],
                        obj['container__lot__article__base_unit__code']],str(_('Total annuel des consommations et prélèvements internes non contre-passés.')))
            grouped=qs.annotate(month=TruncMonth('created_at')).values('month','container__lot__article__code',
                'container__lot__article__name','container__lot__article__base_unit__code').annotate(quantity=Sum('quantity_delta')).order_by('container__lot__article__code','month')
            return Report(title,[_('Article'),_('Désignation'),_('Mois'),_('Quantité consommée'),_('Unité')],grouped,
                lambda obj:[obj['container__lot__article__code'],obj['container__lot__article__name'],obj['month'].strftime('%Y-%m'),
                    -obj['quantity'],obj['container__lot__article__base_unit__code']],
                str(_('Les sorties réellement consommées et les prélèvements de préparations internes sont inclus. Les mouvements contre-passés sont exclus.')))
        if kind=='LOSSES':
            qs=qs.filter(quantity_delta__lt=0,movement__kind__in=['LOSS','BREAKAGE','CONTAMINATION','EXPIRY','DESTRUCTION'],movement__reversal__isnull=True)
        qs=qs.select_related('container__lot__article__base_unit','movement__actor','movement__request','location').order_by('-created_at','-id')
        headers=[_('Date'),_('Article'),_('Contenant'),_('Lot fabricant'),_('Mouvement'),_('Quantité'),_('Réservation'),_('Unité'),
            _('Emplacement au mouvement'),_('Demande'),_('Opérateur'),_('Justification')]
        return Report(title,headers,qs,lambda obj:[obj.created_at,obj.container.lot.article.code,obj.container.code,
            obj.container.lot.manufacturer_lot,obj.movement.get_kind_display(),obj.quantity_delta,obj.reserved_delta,
            obj.container.lot.article.base_unit.code,obj.location.code,obj.movement.request.display_id if obj.movement.request_id else '',
            obj.movement.actor.get_full_name() or obj.movement.actor.username,obj.movement.reason],
            str(_('Le journal conserve les écritures et les corrections. Les chiffres de consommation et de pertes excluent leurs mouvements contre-passés.')))
    if kind=='RECEIPTS':
        _stock(user)
        receipts=StockReceipt.objects.filter(received_on__gte=start.date(),received_on__lt=end.date())
        qs=operational_scope(receipts,user,category_field='container__lot__article__category_id',location_field='movement__entries__location_id').distinct()
        qs=_filter(qs,filters,article_field='container__lot__article_id',location_field='movement__entries__location_id',
            search_fields=('container__code','container__lot__manufacturer_lot','order_reference','supplier__name')).select_related('container__lot__article__base_unit','supplier','movement__actor')
        headers=[_('Date de réception'),_('Article'),_('Contenant'),_('Lot fabricant'),_('Quantité reçue'),_('Unité de gestion'),
            _('Fournisseur'),_('Commande'),_('Date de commande'),_('Délai observé (jours)'),_('État reçu')]
        return Report(title,headers,qs.order_by('-received_on','id'),lambda obj:[obj.received_on,obj.container.lot.article.code,
            obj.container.code,obj.container.lot.manufacturer_lot,obj.received_quantity,obj.container.lot.article.base_unit.code,
            str(obj.supplier) if obj.supplier else '',obj.order_reference,obj.ordered_on,
            (obj.received_on-obj.ordered_on).days if obj.ordered_on else None,obj.condition])
    if kind=='INVENTORY':
        qs=InventoryLine.objects.filter(campaign__work__in=work_scope(user))
        qs=_filter(qs,filters,article_field='container__lot__article_id',location_field='location_id',
            search_fields=('container__code','campaign__work__title')).select_related('campaign__work','container__lot__article__base_unit','location','counted_by')
        if filters.get('campaign'):
            qs=qs.filter(campaign=filters['campaign'])
        def row(obj):
            reveal=is_manager(user) or not obj.campaign.blind or obj.campaign.work.status=='APPROVED'
            difference=obj.counted_quantity-obj.theoretical_quantity if reveal and obj.counted_quantity is not None else None
            return [obj.campaign.work.title,obj.campaign.work.get_status_display(),obj.container.code,obj.container.lot.article.code,
                obj.location.code,obj.theoretical_quantity if reveal else None,obj.counted_quantity,difference,
                obj.container.lot.article.base_unit.code,obj.counted_at,obj.counted_by.get_full_name() or obj.counted_by.username if obj.counted_by else '',obj.note]
        return Report(title,[_('Campagne'),_('Statut'),_('Contenant'),_('Article'),_('Emplacement'),_('Stock théorique'),_('Compté'),
            _('Écart'),_('Unité'),_('Date du comptage'),_('Opérateur'),_('Observation')],qs.order_by('campaign__created_at','container__code'),row,
            str(_('Les quantités théoriques et les écarts restent masqués dans les inventaires à l’aveugle non encore approuvés.')))
    if kind=='SAMPLES':
        if not is_manager(user) and not grants(user,Capability.VIEW_BIOBANK).exists():
            raise PermissionDenied
        qs=_filter(biobank_scope(user),filters,location_field='location_id',search_fields=('code','parent__code','origin_request__display_id'))
        qs=qs.select_related('origin_request','root_sample__origin_request')
        return Report(title,[_('Code échantillon'),_('Parent'),_('Demande'),_('Réception'),_('Quantité initiale'),_('Quantité restante'),
            _('Unité'),_('Emplacement'),_('Position'),_('Statut'),_('Cycles de décongélation'),_('Secondes hors stockage')],qs.order_by('code'),
            lambda obj:[obj.code,obj.parent.code if obj.parent else '',obj.shared_metadata.origin_request.display_id if obj.shared_metadata.origin_request_id else '',
                obj.received_on,obj.initial_quantity,obj.remaining_quantity,obj.unit.code,obj.location.code,
                str(obj.position.row)+':'+str(obj.position.column) if obj.position else '',obj.get_status_display(),obj.freeze_thaw_cycles,obj.out_of_storage_seconds])
    if kind=='STORAGE':
        if not is_manager(user) and not grants(user,Capability.VIEW_BIOBANK).exists():
            raise PermissionDenied
        qs=storage_scope(Location.objects.filter(active=True),user,Capability.VIEW_BIOBANK)
        qs=_filter(qs,filters,location_field='pk',search_fields=('code','name')).filter(Q(kind__cold_storage=True)|Q(grid_rows__isnull=False)).select_related('kind')
        qs=qs.annotate(position_count=Count('descendant_links__descendant__positions',distinct=True),
            occupied=Count('descendant_links__descendant__positions',filter=Q(descendant_links__descendant__positions__occupant__isnull=False),distinct=True),
            reserved_positions=Count('descendant_links__descendant__positions',filter=Q(descendant_links__descendant__positions__reservations__active=True,
                descendant_links__descendant__positions__reservations__until__gt=timezone.now()),distinct=True))
        return Report(title,[_('Code emplacement'),_('Nom'),_('Type'),_('Positions physiques'),_('Occupées'),_('Réservées'),_('Libres'),
            _('Occupation (%)'),_('Température minimale'),_('Température maximale')],qs.order_by('code'),lambda obj:[obj.code,obj.name,str(obj.kind),
                obj.position_count,obj.occupied,obj.reserved_positions,max(0,obj.position_count-obj.occupied-obj.reserved_positions),
                Decimal(100)*(obj.occupied+obj.reserved_positions)/obj.position_count if obj.position_count else None,obj.temperature_min,obj.temperature_max],
            str(_('Les équipements incluent les positions de leurs boîtes descendantes ; n’additionnez pas un équipement et ses propres boîtes.')))
    _stock(user)
    from .links import request_scope
    consumption=operational_scope(RunConsumption.objects.all(),user,category_field='container__lot__article__category_id',location_field='movement__entries__location_id').filter(run__request__in=request_scope(user),
        movement__reversal__isnull=True,created_at__gte=start,created_at__lt=end)
    consumption=_filter(consumption,filters,article_field='container__lot__article_id',location_field='movement__entries__location_id',
        search_fields=('container__lot__manufacturer_lot','container__lot__code','run__code','run__request__display_id'))
    return Report(title,[_('Lot fabricant'),_('Contenant'),_('Article'),_('Série analytique'),_('Demande'),_('Quantité réelle'),
        _('Unité'),_('Date de consommation'),_('Identifiants des échantillons')],consumption.select_related(
            'container__lot__article','run__request','requirement__unit').prefetch_related('run__inputs').order_by('container__lot__code','run__code'),
        lambda obj:[obj.container.lot.manufacturer_lot,obj.container.code,obj.container.lot.article.code,obj.run.code,obj.run.request.display_id,
            obj.quantity,obj.requirement.unit.code,obj.created_at,' ; '.join(item.label for item in obj.run.inputs.all())],
        str(_('Seules les séries dont les consommations réelles ont été confirmées sont présentées.')))


def render_value(value):
    if value is None:
        return '—'
    if isinstance(value,datetime):
        return timezone.localtime(value).strftime('%d/%m/%Y %H:%M')
    if isinstance(value,date):
        return value.strftime('%d/%m/%Y')
    if isinstance(value,Decimal):
        return format(value.normalize(),'f')
    return str(value)


def export_report(dataset,output):
    count=dataset.queryset.count()
    if count>100000:
        raise ValidationError(_('Cet export dépasse 100 000 lignes. Affinez les filtres ou la période.'))
    book=Workbook(write_only=True)
    sheet=book.create_sheet('Données')
    sheet.freeze_panes='A5'
    sheet.print_title_rows='1:4'
    sheet.auto_filter.ref='A4:'+get_column_letter(len(dataset.headers))+str(max(4,count+4))
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape'
    sheet.page_setup.paperSize=Worksheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1
    sheet.page_setup.fitToHeight=0
    for column in range(1,len(dataset.headers)+1):
        sheet.column_dimensions[get_column_letter(column)].width=24 if column!=2 else 38
    sheet.row_dimensions[1].height=28
    sheet.row_dimensions[2].height=38
    sheet.row_dimensions[4].height=30
    def styled(values,header=False):
        cells=[]
        for value in values:
            cell=WriteOnlyCell(sheet)
            if isinstance(value,Decimal):
                cell.value=value
                cell.number_format='#,##0.######'
            elif type(value) in (int,float):
                cell.value=value
            else:
                cell.value=literal(render_value(value))
            cell.font=Font(name='Calibri',size=10,bold=header,color='FFFFFF' if header else '203F50')
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if header:
                cell.fill=PatternFill('solid',fgColor='203F50')
            cells.append(cell)
        return cells
    sheet.append(styled(['PLAGENOR 4.0 — '+dataset.title]))
    sheet.append(styled([dataset.notes]))
    sheet.append([])
    sheet.append(styled([str(value) for value in dataset.headers],True))
    for obj in dataset.queryset.iterator(chunk_size=500):
        sheet.append(styled(dataset.serialize(obj)))
    book.save(output)
    return count


def stock_value_estimate(user,filters):
    stock=_filter(_stock(user),filters,article_field='lot__article_id',location_field='location_id').filter(quantity__gt=0).select_related('lot__article__base_unit','lot__article__category','location')
    latest=StockReceipt.objects.filter(container__lot_id=OuterRef('lot_id'),movement__reversal__isnull=True).filter(
        Q(unit_price__isnull=False)|Q(purchase_delivery__isnull=False)).order_by('-received_on','-created_at')
    stock=stock.annotate(price=Subquery(latest.values('unit_price')[:1]),
        purchase_price=Subquery(latest.values('purchase_delivery__actual_unit_price')[:1]),
        purchase_factor=Subquery(latest.values('purchase_delivery__line__factor')[:1]),
        price_currency=Subquery(latest.values('currency')[:1]))
    values={}
    unpriced=0
    hidden=0
    for row in stock.iterator(chunk_size=500):
        if not permitted(user,Capability.VIEW_COST,category=row.lot.article.category,location=row.location):
            hidden+=1
            continue
        price=row.price
        if price is None and row.purchase_price is not None and row.purchase_factor:
            price=row.purchase_price/row.purchase_factor
        if price is None:
            unpriced+=1
        else:
            values[row.price_currency]=values.get(row.price_currency,ZERO)+row.quantity*price
    return {'currencies':values,'unpriced_containers':unpriced,'hidden_containers':hidden,
        'basis':str(_('Estimation indicative au dernier prix de réception documenté de chaque lot, hors taxes. Il ne s’agit pas d’une valorisation comptable.'))}
