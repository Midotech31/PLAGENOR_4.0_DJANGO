from datetime import date
from decimal import Decimal
import io
from pathlib import Path
import tempfile

from django.core.exceptions import ValidationError
from django.utils import timezone
from django.utils.translation import gettext as _
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt
from openpyxl import Workbook
from openpyxl.styles import Font

from documents.pdf_converter import convert_docx_to_pdf
from .stock_reporting import with_availability


def stock_table(queryset, *, limit=100000):
    if queryset.count() > limit:
        raise ValidationError(_('Réduisez les filtres : cet export dépasse la limite de lignes.'))
    headers = [_('Code article'), _('Désignation'), _('Référence catalogue'), _('Référence fabricant'),
        _('Référence fournisseur'), _('Fabricant'), _('Marque'), _('Fournisseur préféré'), _('Catégorie'),
        _('Code lot'), _('Lot fabricant'), _('Numéro de série'), _('Code contenant'), _('Emplacement'),
        _('Stock physique'), _('Stock réservé'), _('Disponible'), _('Unité'), _('Date de réception initiale'),
        _('Ouverture'), _('Date limite d’utilisation'), _('Contrôle du contenant')]
    rows = []
    for container in with_availability(queryset).order_by('lot__article__code', 'code').iterator():
        article = container.lot.article
        rows.append([article.code, str(article), article.catalog_reference, article.manufacturer_reference,
            article.supplier_reference, str(article.manufacturer) if article.manufacturer else '', article.brand,
            str(article.preferred_supplier) if article.preferred_supplier else '', str(article.category),
            container.lot.code, container.lot.manufacturer_lot, container.lot.serial_number, container.code,
            str(container.location), container.quantity, container.reserved,
            container.quantity - container.reserved if container.stock_usable else Decimal(0), article.base_unit.code,
            container.fifo_received_on, container.opened_on, container.use_by, str(container.get_status_display())])
    return headers, rows


def stock_workbook(queryset):
    headers, rows = stock_table(queryset)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Inventaire'
    for number, row in enumerate([headers, *rows], 1):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(number, column)
            if isinstance(value, Decimal) and len(value.normalize().as_tuple().digits) <= 15:
                cell.value = value
                cell.number_format = '0.######'
            else:
                cell.value = value.isoformat() if isinstance(value, date) else '' if value is None else str(value)
                cell.data_type = 's'
            if number == 1:
                cell.font = Font(bold=True)
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def table_pdf(title, subtitle, headers, rows):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(297), Mm(210)
    section.left_margin = section.right_margin = Mm(10)
    section.top_margin = section.bottom_margin = Mm(12)
    doc.styles['Normal'].font.name = 'DejaVu Sans'
    doc.styles['Normal'].font.size = Pt(8)
    doc.add_heading(title, level=1)
    doc.add_paragraph(subtitle)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value.isoformat() if isinstance(value, date) else '' if value is None else str(value)
    with tempfile.TemporaryDirectory(prefix='plagenor-stock-pdf-') as directory:
        source = Path(directory) / 'inventaire.docx'
        doc.save(source)
        pdf = convert_docx_to_pdf(source)
        if pdf.suffix.lower() != '.pdf' or not pdf.is_file():
            raise ValidationError(_('La conversion PDF n’a pas abouti. Réessayez ou utilisez l’export Excel.'))
        return pdf.read_bytes()


def stock_pdf(queryset):
    headers, rows = stock_table(queryset, limit=2000)
    positions = (0, 1, 2, 3, 4, 10, 12, 13, 14, 16, 17, 20)
    return table_pdf(_('PLAGENOR 4.0 — Inventaire'), str(timezone.localdate()),
        [headers[index] for index in positions], [[row[index] for index in positions] for row in rows])


def dispatch_pdf(dispatch, lines):
    rows = [[line.snapshot['article'], line.snapshot['designation'], line.snapshot.get('catalog_reference', ''),
        line.snapshot.get('supplier_reference', ''), line.snapshot['lot'], line.snapshot['container'],
        line.snapshot['location'], line.quantity, line.snapshot['unit']] for line in lines]
    return table_pdf(_('PLAGENOR 4.0 — Bon de distribution'),
        ' | '.join((str(dispatch.pk), str(dispatch.get_mode_display()), dispatch.beneficiary,
            str(dispatch.distributed_on), str(dispatch.actor), dispatch.reason)),
        [_('Code'), _('Désignation'), _('Référence catalogue'), _('Référence fournisseur'),
         _('Lot fabricant'), _('Contenant'), _('Origine'), _('Quantité'), _('Unité')], rows)
