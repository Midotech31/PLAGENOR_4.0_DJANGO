import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .common import snapshot
from .work import require_work,work_allowed


HEAD_FILL=PatternFill('solid',fgColor='203F50')
WHITE=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
LINE=Side(style='hair',color='D8E2E8')


def literal(value):
    if value is None:
        return ''
    text=str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text


def workbook_table(title,subtitle,headers,rows,widths,*,numeric=(),formulas=None):
    workbook=Workbook()
    sheet=workbook.active
    sheet.title='Données'
    sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    sheet.cell(1,1,literal(title)).font=Font(name='Calibri',size=18,bold=True,color='203F50')
    sheet.row_dimensions[1].height=32
    sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
    sheet.cell(2,1,literal(subtitle)).alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[2].height=32
    for column,header in enumerate(headers,1):
        cell=sheet.cell(4,column,literal(header))
        cell.fill,cell.font,cell.alignment=HEAD_FILL,WHITE,Alignment(wrap_text=True,vertical='center')
        sheet.column_dimensions[get_column_letter(column)].width=widths[column-1]
    sheet.row_dimensions[4].height=32
    for index,row in enumerate(rows,5):
        sheet.row_dimensions[index].height=36
        for column,value in enumerate(row,1):
            cell=sheet.cell(index,column)
            cell.value=value if column in numeric and value is not None else literal(value)
            cell.font=Font(name='Calibri',size=10,color='243746')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
            cell.border=Border(bottom=LINE)
            if column in numeric:
                cell.number_format='#,##0.######'
        if formulas:
            for column,expression in formulas(index).items():
                sheet.cell(index,column,expression).number_format='#,##0.00'
    sheet.freeze_panes='C5'
    sheet.auto_filter.ref=f'A4:{get_column_letter(len(headers))}{max(4,len(rows)+4)}'
    sheet.print_title_rows='1:4'
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape'
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth,sheet.page_setup.fitToHeight=1,0
    sheet.page_margins=PageMargins(left=.3,right=.3,top=.4,bottom=.4,header=.15,footer=.15)
    sheet.oddFooter.center.text='PLAGENOR 4.0 — &P / &N'
    workbook.calculation.fullCalcOnLoad=True
    return workbook


def export_plan(user,plan):
    require_work(user,plan.work)
    costs=work_allowed(user,plan.work,costs=True)
    source=plan.approved_revision.data if plan.approved_revision_id else {
        'reference':plan.reference,'year':plan.year,'lines':[snapshot(row) for row in plan.lines.all()]}
    headers=['Lot','Code','Désignation','Unité d’achat','Proposition','Quantité retenue','Décision','Priorité','Justification']
    widths=[26,18,42,20,16,18,14,14,42]
    if costs:
        headers+=['Prix HT','Taxe (%)','Devise','Total HT','Taxes','Total TTC','Source du prix']
        widths +=[16,14,12,18,18,18,40]
    rows=[]
    for item in source['lines']:
        article=item['article_snapshot']
        row=[item['lot_name'],article['code'],article['name'],article['purchase_unit_name'],
            Decimal(item['proposed_quantity']) if item['proposed_quantity'] is not None else None,
            Decimal(item['retained_quantity']) if item['retained_quantity'] is not None else None,
            'Retenu' if item['included'] else 'Exclu',item['priority'],item['decision_reason']]
        if costs:
            row +=[Decimal(item['estimated_price']) if item['estimated_price'] is not None else None,
                Decimal(item['tax_rate']) if item['tax_rate'] is not None else None,item['currency'],None,None,None,item['price_source']]
        rows.append(row)
    def formulas(index):
        item=source['lines'][index-5]
        if not costs or not item['included'] or any(item[field] is None for field in ('retained_quantity','estimated_price','tax_rate')):
            return {}
        return {13:f'=ROUND(F{index}*J{index},2)',14:f'=ROUND(M{index}*K{index}/100,2)',15:f'=M{index}+N{index}'}
    workbook=workbook_table('PLAGENOR 4.0 — Plan d’approvisionnement',
        source['reference']+' — '+str(source['year'])+' — '+str(plan.work.get_status_display()),
        headers,rows,widths,numeric=(5,6,10,11,13,14,15),formulas=formulas)
    provenance=workbook.create_sheet('Traçabilité')
    entries=[('Plan',str(plan.pk)),('Révision',plan.revision_number),
        ('Empreinte de la révision approuvée',plan.approved_revision.sha256 if plan.approved_revision_id else 'Non approuvé'),
        ('Quantités','Les quantités sont exprimées dans l’unité d’achat indiquée.'),
        ('Calcul','Les totaux des lignes retenues sont arrondis à deux décimales. Les devises ne sont pas additionnées entre elles.'),
        ('Confidentialité','Les coûts ne figurent dans cet export que lorsque leur consultation est autorisée.')]
    for index,(key,value) in enumerate(entries,1):
        provenance.cell(index,1,key).font=Font(bold=True,color='203F50')
        provenance.cell(index,2,literal(value)).alignment=Alignment(wrap_text=True,vertical='top')
        provenance.row_dimensions[index].height=35
    provenance.column_dimensions['A'].width=36
    provenance.column_dimensions['B'].width=80
    data=io.BytesIO()
    workbook.save(data)
    return data.getvalue()
