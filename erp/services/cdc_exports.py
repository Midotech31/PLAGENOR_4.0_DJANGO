import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


HEADERS = (
    'Code', 'Catégorie', 'Critère', 'Description', 'Lot', 'Preuve attendue',
    'Note min.', 'Note max.', 'Pondération (%)', 'Seuil', 'Formule / méthode',
    'Règle d’arrondi', 'Éliminatoire', 'Source', 'Justification',
)


def criteria_workbook(revision):
    """Export the exact criteria snapshot frozen in a CDC revision."""
    rows = revision.governance.get('criteria', [])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Critères'
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([
            row.get('code', ''), row.get('category', ''), row.get('title', ''),
            row.get('description', ''), row.get('lot') or '', row.get('expected_evidence', ''),
            row.get('min_score'), row.get('max_score'), row.get('weight'),
            row.get('threshold'), row.get('formula', ''), row.get('rounding_rule', ''),
            'Oui' if row.get('eliminatory') else 'Non', row.get('source', ''),
            row.get('justification', ''),
        ])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((16, 18, 30, 45, 36, 40, 12, 12, 16, 12, 35, 20, 14, 40, 40), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    requirements = workbook.create_sheet('Exigences')
    requirement_headers = ('Article', 'Lot', 'Code', 'Type', 'Exigence', 'Preuve attendue',
                           'Méthode de vérification', 'Justification')
    requirements.append(list(requirement_headers))
    for cell in requirements[1]:
        cell.font = Font(bold=True)
    for row in revision.governance.get('requirements', []):
        requirements.append([
            row.get('item_label') or row.get('item_key', ''), row.get('lot_label') or row.get('lot', ''),
            row.get('code', ''), row.get('kind', ''),
            row.get('statement', ''), row.get('evidence', ''), row.get('verification', ''),
            row.get('justification', ''),
        ])
    requirements.freeze_panes = 'A2'
    requirements.auto_filter.ref = requirements.dimensions
    for index, width in enumerate((24, 36, 16, 18, 50, 40, 40, 40), 1):
        requirements.column_dimensions[get_column_letter(index)].width = width
    info = workbook.create_sheet('Traçabilité')
    info.append(['Référence CDC', revision.dossier.reference])
    info.append(['Révision', revision.number])
    info.append(['Empreinte SHA-256', revision.sha256])
    info.append(['Créée le', revision.created_at.isoformat()])
    info.append(['Auteur', revision.actor.get_username()])
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()
