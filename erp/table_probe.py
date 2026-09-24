import csv
from datetime import date,datetime
import io
import json
from pathlib import Path
import runpy
import sys
import zipfile


MAX_FILE=10*1024*1024
MAX_ROWS=500
MAX_COLUMNS=40
MAX_TEXT=4*1024*1024


def normalize(value):
    if isinstance(value,(date,datetime)):
        return value.isoformat()
    if value is None:
        return ''
    if isinstance(value,bool):
        return 'true' if value else 'false'
    text=str(value).strip()
    if len(text)>10000 or any(ord(char)<32 and char not in (chr(10),chr(13),chr(9)) for char in text):
        raise ValueError('CELL')
    return text


def csv_rows(data):
    text=data.decode('utf-8-sig')
    dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;'+chr(9))
    rows=[]
    for index,row in enumerate(csv.reader(io.StringIO(text),dialect=dialect)):
        if index>MAX_ROWS or len(row)>MAX_COLUMNS:
            raise ValueError('SIZE')
        rows.append([normalize(value) for value in row])
    return rows


def xlsx_rows(data,helpers):
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries=archive.infolist()
        if not 1<=len(entries)<=512:
            raise ValueError('SIZE')
        keys=set()
        total=0
        for info in entries:
            name=helpers['_entry_key'](info)
            if name.casefold() in keys:
                raise ValueError('STRUCTURE')
            keys.add(name.casefold())
            total+=info.file_size
            if info.file_size>8*1024*1024 or total>32*1024*1024:
                raise ValueError('SIZE')
            if info.file_size>1024*1024 and info.file_size>max(info.compress_size,1)*200:
                raise ValueError('SIZE')
            lowered=name.casefold()
            if any(part in lowered for part in ('vbaproject','embeddings/','externallinks/','activex/','connections.xml')):
                raise ValueError('ACTIVE')
            if name.endswith(('.xml','.rels')):
                element=helpers['_xml'](archive.read(info))
                if name.endswith('.rels') and any(node.get('TargetMode')=='External' for node in element):
                    raise ValueError('EXTERNAL')
                if lowered.startswith('xl/worksheets/'):
                    for node in element.iter():
                        if node.tag.endswith('}f'):
                            raise ValueError('FORMULA')
                        if node.tag.endswith('}row') and not 1<=int(node.get('r','0'))<=MAX_ROWS+1:
                            raise ValueError('SIZE')
    workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        candidates=[sheet for sheet in workbook if sheet.title not in ('Guide','Instructions')]
        if len(candidates)!=1:
            raise ValueError('SHEETS')
        sheet=candidates[0]
        if (sheet.max_row and sheet.max_row>MAX_ROWS+1) or (sheet.max_column and sheet.max_column>MAX_COLUMNS):
            raise ValueError('SIZE')
        return [[normalize(cell.value) for cell in row] for row in sheet.iter_rows(
            max_row=MAX_ROWS+2,max_col=min(sheet.max_column or MAX_COLUMNS,MAX_COLUMNS))]
    finally:
        workbook.close()


def read_matrix(data,extension,helpers):
    if not data or len(data)>MAX_FILE:
        raise ValueError('SIZE')
    if extension=='.csv':
        rows=csv_rows(data)
    elif extension=='.xlsx':
        rows=xlsx_rows(data,helpers)
    else:
        raise ValueError('EXTENSION')
    while rows and not any(rows[-1]):
        rows.pop()
    if len(rows)<2:
        raise ValueError('EMPTY')
    if len(rows)>MAX_ROWS+1:
        raise ValueError('SIZE')
    if sum(len(cell.encode('utf-8')) for row in rows for cell in row)>MAX_TEXT:
        raise ValueError('SIZE')
    return rows


def main():
    try:
        helpers=runpy.run_path(str(Path(__file__).resolve().parents[1]/'core/document_probe.py'))
        helpers['limit_resources']()
        data=sys.stdin.buffer.read(MAX_FILE+1)
        rows=read_matrix(data,sys.argv[1],helpers)
        sys.stdout.buffer.write(json.dumps({'rows':rows},ensure_ascii=False,separators=(',',':')).encode('utf-8'))
        return 0
    except Exception as exc:
        code=str(exc)
        if code not in ('SIZE','CELL','STRUCTURE','ACTIVE','EXTERNAL','FORMULA','SHEETS','EXTENSION','EMPTY'):
            code='INVALID'
        sys.stdout.buffer.write(json.dumps({'error':code}).encode('ascii'))
        return 1


if __name__=='__main__':
    sys.exit(main())
