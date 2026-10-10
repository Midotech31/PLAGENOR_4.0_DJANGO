from datetime import timedelta
from decimal import Decimal
import io
import uuid
from unittest.mock import patch

from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone, translation
from openpyxl import load_workbook

from erp.models import (Article, Capability, ImportMapping, Location, StockContainer,
    StockDispatch, StockEntry, StockLot, StockMovement)
from erp.services.catalog import save_article
from erp.services.common import Conflict
from erp.services.distributions import distribute_stock, return_stock, returnable_quantity
from erp.services.import_mapping import preview_mapping, require_mapping, stage_import
from erp.services.bulk_imports import apply_import, preview_import
from erp.services.stock import (control_container, fefo, open_container, reconcile_stock,
    remove_stock, reserve_stock, reverse_stock, transfer_stock)
from erp.services.stock_reporting import dashboard_data, entry_balances
from erp.stock_views import stock_queryset
from erp.test_operations import OperationFixtures


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ScientificStockTests(OperationFixtures, TestCase):
    def setUp(self):
        self.client.force_login(self.ops)
        self.article = save_article(self.ops, {'catalog_reference': 'CAT-000123', 'preferred_supplier': self.party,
            'supplier_reference': 'SUP-00123', 'brand': 'Science', 'pack_quantity': Decimal('96')},
            pk=self.article.pk, expected=self.article.version)
        self.destination = Location.objects.create(code='SCIENCE-DEST', name='Laboratoire suivi', kind=self.storage_kind)
        from erp.models import LocationClosure
        LocationClosure.objects.create(ancestor=self.destination, descendant=self.destination, depth=0)

    def distribute(self, container, **changes):
        values = {'key': uuid.uuid4(), 'lines': [{'container': container, 'amount': Decimal('2.125'), 'unit': self.unit,
            'expected': container.version}], 'mode': 'EXIT', 'beneficiary': 'Laboratoire de génomique',
            'distributed_on': timezone.localdate(), 'reason': 'Analyse du projet'}
        values.update(changes)
        return distribute_stock(self.ops, **values), values

    def test_receipt_exit_return_and_internal_transfer_keep_exact_global_balance(self):
        source, receipt, unused = self.receive('EXACT', quantity='10.125', delivery_reference='BL-SCI-1', order_reference='BC-SCI-1')
        self.assertEqual(receipt.receipt.delivery_reference, 'BL-SCI-1')
        dispatch, values = self.distribute(source)
        self.assertEqual(distribute_stock(self.ops, **values).pk, dispatch.pk)
        source.refresh_from_db()
        self.assertEqual(source.quantity, Decimal('8'))
        line = dispatch.lines.get()
        self.assertEqual(line.snapshot['catalog_reference'], 'CAT-000123')
        returned = return_stock(self.ops, line.pk, key=uuid.uuid4(), amount='1.125', unit=self.unit,
            destination=self.freezer, container_code='BACK-EXACT', returned_on=timezone.localdate(), condition='Flacon intact', reason='Reliquat')
        self.assertEqual(returned.container.status, 'PENDING')
        self.assertEqual(returned.container.source_container_id, source.pk)
        self.assertEqual(returnable_quantity(line), Decimal('1'))
        self.assertEqual(sum(StockContainer.objects.values_list('quantity', flat=True)), Decimal('9.125'))
        moved, unused = self.distribute(source, mode='INTERNAL', destination=self.destination,
            lines=[{'container': source, 'amount': '8', 'unit': self.unit, 'expected': source.version}])
        source.refresh_from_db()
        self.assertEqual(source.location_id, self.destination.pk)
        self.assertEqual(sum(StockContainer.objects.values_list('quantity', flat=True)), Decimal('9.125'))
        self.assertEqual(reconcile_stock(self.ops), [])
        self.assertContains(self.client.get(reverse('erp:dispatch-detail', args=[moved.pk])), '8')

    def test_multi_article_failure_rolls_back_every_line_and_stale_version_is_refused(self):
        first, unused, unused = self.receive('A', quantity='10')
        second, unused, unused = self.receive('B', quantity='1', article=self.liquid, unit=self.ml)
        with self.assertRaises(ValidationError):
            self.distribute(first, lines=[{'container': first, 'amount': '2', 'unit': self.unit},
                {'container': second, 'amount': '2', 'unit': self.ml}])
        first.refresh_from_db()
        self.assertEqual(first.quantity, Decimal('10'))
        self.assertFalse(StockDispatch.objects.exists())
        self.assertEqual(StockMovement.objects.filter(kind='EXIT').count(), 0)
        with self.assertRaises(Conflict):
            self.distribute(first, lines=[{'container': first, 'amount': '1', 'unit': self.unit, 'expected': first.version + 1}])
        dispatch, unused = self.distribute(first, lines=[{'container': first, 'amount': '2', 'unit': self.unit},
            {'container': second, 'amount': '.125', 'unit': self.ml}])
        self.assertEqual(dispatch.lines.count(), 2)
        self.assertEqual(reconcile_stock(self.ops), [])

    def test_dispatch_input_refusals_and_idempotency_conflict(self):
        source, unused, unused = self.receive('GUARDS')
        line = {'container': source, 'amount': '1', 'unit': self.unit}
        for changes in ({'lines': []}, {'lines': [line] * 101}, {'lines': [line, line]},
                {'mode': 'OTHER'}, {'beneficiary': ''}, {'reason': ''}, {'distributed_on': timezone.localdate() + timedelta(days=1)},
                {'mode': 'INTERNAL'}, {'destination': self.destination}, {'distributed_on': timezone.localdate() - timedelta(days=1)}):
            with self.subTest(changes=changes), self.assertRaises(ValidationError):
                self.distribute(source, **changes)
        ghost = StockContainer(pk=uuid.uuid4())
        with self.assertRaises(ValidationError):
            self.distribute(source, lines=[{'container': ghost, 'amount': '1', 'unit': self.unit}])
        dispatch, values = self.distribute(source)
        values['beneficiary'] = 'Autre laboratoire'
        with self.assertRaises(Conflict):
            distribute_stock(self.ops, **values)
        self.assertEqual(StockDispatch.objects.count(), 1)

    def test_expired_unaccepted_and_reserved_stock_cannot_leave_as_distribution(self):
        pending, unused, unused = self.receive('PENDING', accepted=False)
        with self.assertRaises(ValidationError):
            self.distribute(pending)
        expired, unused, unused = self.receive('EXPIRED', accepted=False, expires_on=timezone.localdate() - timedelta(days=1))
        with self.assertRaises(ValidationError):
            remove_stock(self.ops, expired.pk, key=uuid.uuid4(), amount='1', unit=self.unit, kind='EXIT', reason='Distribution')
        source, unused, unused = self.receive('RESERVE', quantity='3')
        reserve_stock(self.ops, source.pk, key=uuid.uuid4(), amount='2', unit=self.unit, reference='Projet réservé')
        with self.assertRaises(ValidationError):
            self.distribute(source)
        self.assertEqual(reconcile_stock(self.ops), [])

    def test_return_ceiling_idempotency_dates_and_reversal_guard(self):
        source, unused, unused = self.receive('RETURN')
        dispatch, unused = self.distribute(source)
        line = dispatch.lines.get()
        values = dict(key=uuid.uuid4(), amount='1', unit=self.unit, destination=self.freezer,
            container_code='SCI-BACK', returned_on=timezone.localdate(), condition='Intact', reason='Non utilisé')
        for changes in ({'amount': '3'}, {'returned_on': timezone.localdate() - timedelta(days=1)},
                {'condition': ''}, {'reason': ''}, {'returned_on': timezone.localdate() + timedelta(days=1)}):
            with self.subTest(changes=changes), self.assertRaises(ValidationError):
                return_stock(self.ops, line.pk, **{**values, **changes})
        returned = return_stock(self.ops, line.pk, **values)
        self.assertEqual(return_stock(self.ops, line.pk, **values).pk, returned.pk)
        with self.assertRaises(ValidationError):
            reverse_stock(self.ops, line.movement_id, key=uuid.uuid4(), reason='Correction du bon')
        reverse_stock(self.ops, returned.movement_id, key=uuid.uuid4(), reason='Annulation du retour')
        self.assertEqual(returnable_quantity(line), line.quantity)
        reverse_stock(self.ops, line.movement_id, key=uuid.uuid4(), reason='Annulation du bon')
        self.assertEqual(returnable_quantity(line), Decimal(0))
        self.assertEqual(reconcile_stock(self.ops), [])

    def test_fifo_old_receipt_opened_aliquot_and_exact_partial_consumption(self):
        source, unused, unused = self.receive('FIFO', quantity='10.125', expires_on=None,
            received_on=timezone.localdate() - timedelta(days=10))
        recent, unused, unused = self.receive('LATER', expires_on=None)
        StockContainer.objects.filter(pk=source.pk).update(fifo_received_on=None)
        self.assertEqual(fefo(self.ops, self.article).first().pk, source.pk)
        source = open_container(self.ops, source.pk, expected=source.version,
            opened_on=timezone.localdate() - timedelta(days=9))
        self.assertContains(self.client.get(reverse('erp:stock-action', args=[source.pk, 'aliquot'])), 'Créer une aliquote')
        with self.assertRaises(ValidationError):
            transfer_stock(self.ops, recent.pk, key=uuid.uuid4(), destination=self.freezer, amount='1',
                destination_code='INVALID-ALIQUOT', reason='Aliquote', aliquot=True)
        for amount in (None, source.quantity):
            with self.assertRaises(ValidationError):
                transfer_stock(self.ops, source.pk, key=uuid.uuid4(), destination=self.freezer, amount=amount,
                    destination_code='INVALID-ALIQUOT', reason='Aliquote', aliquot=True)
        transfer_stock(self.ops, source.pk, key=uuid.uuid4(), destination=self.freezer, amount='.125',
            destination_code='SCI-ALIQUOT', reason='Aliquote', aliquot=True)
        child = StockContainer.objects.get(code='SCI-ALIQUOT')
        self.assertEqual(child.fifo_received_on, timezone.localdate() - timedelta(days=10))
        self.assertEqual(child.opened_on, source.opened_on)
        remove_stock(self.ops, child.pk, key=uuid.uuid4(), amount='.025', unit=self.unit, reason='Analyse sur aliquote')
        child.refresh_from_db()
        self.assertEqual(child.quantity, Decimal('.100'))
        self.assertEqual(reconcile_stock(self.ops), [])

    def test_catalog_supplier_binding_and_native_receipt_date_costs(self):
        with self.assertRaises(ValidationError):
            save_article(self.ops, {'supplier_reference': 'INVALID', 'preferred_supplier': None},
                pk=self.article.pk, expected=self.article.version)
        self.assertEqual(Article.objects.get(pk=self.article.pk).supplier_reference, 'SUP-00123')
        for language in ('fr', 'en', 'ar'):
            with translation.override(language):
                from erp.stock_forms import ReceiptForm
                self.assertIn('value="' + timezone.localdate().isoformat() + '"', str(ReceiptForm(user=self.ops)['received_on']))
        values = dict(key=uuid.uuid4(), article=self.article.pk, location=self.freezer.pk,
            manufacturer_lot='HTTP-SCI', lot_code='HTTP-SCI', container_code='HTTP-SCI', amount='10.125',
            unit=self.unit.pk, received_on=timezone.localdate().isoformat(), condition='Intact', unit_price='12.50',
            currency='', order_reference='BC-HTTP', delivery_reference='BL-HTTP')
        response = self.client.post(reverse('erp:receipt-create'), values)
        self.assertEqual(response.status_code, 302, getattr(response, 'context', None))
        container = StockContainer.objects.get(code='HTTP-SCI')
        receipt = container.stockreceipt_set.get()
        self.assertEqual((receipt.currency, receipt.unit_price, receipt.delivery_reference), ('DZD', Decimal('12.50'), 'BL-HTTP'))

    def test_search_grouping_qr_filters_and_excel_reference_precision(self):
        source, unused, unused = self.receive('SEARCH', quantity='10.125')
        for q in ('CAT-000123', 'SUP-00123', 'R-123', str(source.pk), 'PLAGENOR|article|' + str(self.article.pk)):
            self.assertContains(self.client.get(reverse('erp:stock-list'), {'q': q}), source.code)
        for view in ('products', 'lots', 'locations'):
            response = self.client.get(reverse('erp:stock-list'), {'view': view, 'direction': 'desc'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.context['page'].object_list[0]['quantities'][0]['physical'], Decimal('10.125'))
        for parameter, value in (('location', self.freezer.pk), ('category', self.category.pk),
                ('manufacturer', self.party.pk), ('supplier', self.party.pk), ('kind', self.storage_kind.pk),
                ('article', self.article.pk), ('lot', source.lot_id)):
            self.assertContains(self.client.get(reverse('erp:stock-list'), {parameter: value}), source.code)
        for params in ({'view': 'bad'}, {'location': 'bad'}, {'q': 'PLAGENOR|container|bad'}, {'format': 'bad'}):
            route = 'erp:stock-export' if 'format' in params else 'erp:stock-list'
            self.assertEqual(self.client.get(reverse(route), params).status_code, 404)
        response = self.client.get(reverse('erp:stock-export'), {'q': 'CAT-000123'})
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active['C2'].value, 'CAT-000123')
        self.assertEqual(book.active['E2'].value, 'SUP-00123')
        self.assertEqual(book.active['O2'].value, 10.125)
        self.assertContains(self.client.get(reverse('erp:identify'), {'q': 'CAT-000123'}), self.article.code)

    def test_dashboard_and_legacy_balance_reconstruction_do_not_modify_history(self):
        source, unused, unused = self.receive('LEDGER', quantity='10.125', unit_price='12.50')
        remove_stock(self.ops, source.pk, key=uuid.uuid4(), amount='.125', unit=self.unit, reason='Analyse')
        reserve_stock(self.ops, source.pk, key=uuid.uuid4(), amount='1', unit=self.unit, reference='Projet')
        original = list(StockEntry.objects.values_list('pk', 'snapshot'))
        rows = entry_balances(StockEntry.objects.order_by('-pk'))
        self.assertEqual((rows[0].physical_after, rows[0].reserved_after), (Decimal('10'), Decimal('1')))
        self.assertEqual(list(StockEntry.objects.values_list('pk', 'snapshot')), original)
        self.assertEqual(entry_balances([]), [])
        data = dashboard_data(self.ops, stock_queryset(self.ops))
        self.assertEqual(data['units'][0]['available'], Decimal('9'))
        self.assertEqual(data['costs']['DZD'], Decimal('125'))
        self.assertContains(self.client.get(reverse('erp:stock-dashboard')), 'CAT-000123')
        self.assertEqual(self.client.get(reverse('erp:stock-dashboard'), {'days': 'bad'}).status_code, 404)
        self.assertEqual(self.client.get(reverse('erp:stock-dashboard'), {'days': '0'}).status_code, 404)
        self.assertContains(self.client.get(reverse('erp:ledger')), source.code)

    def test_import_mapping_preserves_blank_existing_fields_and_explicit_clear(self):
        data = b'Id;Nom;Categorie;Unite;Catalogue;Pack\nTIPS;;;;CAT-REVISED;96\n'
        staged = stage_import(self.ops, key=uuid.uuid4(), kind='CATALOG', filename='real.csv', data=data, reason='Inventaire réel')
        choices = dict(code=0, name=1, category_code=2, base_unit_code=3, catalog_reference=4, pack_quantity=5)
        batch = preview_mapping(self.ops, staged.pk, choices=choices)
        self.assertTrue(batch.report['valid'], batch.report)
        self.assertEqual(preview_mapping(self.ops, staged.pk, choices=choices).pk, batch.pk)
        apply_import(self.ops, batch.pk, expected=batch.version, confirmed=True)
        self.article.refresh_from_db()
        self.assertEqual((self.article.name, self.article.catalog_reference, self.article.pack_quantity), ('Pointes', 'CAT-REVISED', Decimal('96')))
        self.assertEqual(self.article.supplier_reference, 'SUP-00123')
        blank = stage_import(self.ops, key=uuid.uuid4(), kind='CATALOG', filename='clear.csv',
            data=b'code;name;category_code;base_unit_code;pack_quantity\nTIPS;;;;\n', reason='Effacement demandé')
        batch = preview_mapping(self.ops, blank.pk, choices=dict(code=0, name=1, category_code=2, base_unit_code=3, pack_quantity=4), clear_fields=['pack_quantity'])
        self.assertTrue(batch.report['valid'], batch.report)
        apply_import(self.ops, batch.pk, expected=batch.version, confirmed=True)
        self.assertIsNone(Article.objects.get(pk=self.article.pk).pack_quantity)
        ImportMapping.objects.filter(pk=blank.pk).update(expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.client.get(reverse('erp:import-mapping', args=[blank.pk])).status_code, 400)

    def test_scoped_stock_and_dispatch_documents_never_escape_delegation(self):
        source, unused, unused = self.receive('SCOPED')
        dispatch, values = self.distribute(source)
        self.grant(Capability.VIEW_STOCK, category=self.category, location=self.destination)
        self.client.force_login(self.operator)
        self.assertNotContains(self.client.get(reverse('erp:stock-list')), source.code)
        self.assertEqual(self.client.get(reverse('erp:dispatch-detail', args=[dispatch.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse('erp:dispatch-create')).status_code, 403)
        with self.assertRaises(PermissionDenied):
            distribute_stock(self.operator, **{**values, 'key': uuid.uuid4()})
        with self.assertRaises(PermissionDenied):
            require_mapping(self.operator, ImportMapping(actor=self.ops, kind='CATALOG'))

    def test_real_inventory_and_distribution_pdfs_contain_references_and_quantities(self):
        from pypdf import PdfReader
        source, _, _ = self.receive('PDF', quantity='10.125')
        dispatch, _ = self.distribute(source)
        for url in (reverse('erp:stock-export'), reverse('erp:dispatch-detail', args=[dispatch.pk])):
            response = self.client.get(url, {'format': 'pdf'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response['Content-Type'], 'application/pdf')
            self.assertTrue(response.content.startswith(b'%PDF-'))
            reader = PdfReader(io.BytesIO(response.content))
            text = '\n'.join(page.extract_text() for page in reader.pages)
            self.assertIn('CAT-000123', text)
            self.assertIn('SUP-00123', text)
            self.assertIn(source.code, text)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.width), 842, delta=2)

    @override_settings(DOCUMENT_PDF_ENABLED=False)
    def test_failed_pdf_conversion_is_a_clear_error_and_never_a_docx_download(self):
        from erp.services.stock_exports import stock_table
        source, _, _ = self.receive('PDF-ERROR')
        dispatch, _ = self.distribute(source)
        for url in (reverse('erp:stock-export'), reverse('erp:dispatch-detail', args=[dispatch.pk])):
            response = self.client.get(url, {'format': 'pdf'})
            self.assertEqual(response.status_code, 400)
            self.assertContains(response, 'conversion PDF', status_code=400)
        with self.assertRaises(ValidationError):
            stock_table(stock_queryset(self.ops), limit=0)

    def test_excel_retains_formula_like_references_and_eighteen_digit_quantities(self):
        self.article = save_article(self.ops, {'catalog_reference': '=1+1'}, pk=self.article.pk, expected=self.article.version)
        self.receive('XLSX-LITERAL', quantity='10.125')
        response = self.client.get(reverse('erp:stock-export'))
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual((sheet['C2'].value, sheet['C2'].data_type), ('=1+1', 's'))
        self.assertEqual(sheet['O2'].value, 10.125)
        from erp.services.stock_exports import stock_workbook
        with patch('erp.services.stock_exports.stock_table', return_value=(['Reference', 'Quantity'],
                [['0000123', Decimal('999999999999.123456')]])):
            sheet = load_workbook(io.BytesIO(stock_workbook(None))).active
        self.assertEqual(sheet['A2'].value, '0000123')
        self.assertEqual(sheet['B2'].value, '999999999999.123456')

    def test_dashboard_includes_unstocked_products_and_respects_purchase_conversion(self):
        from erp.services.catalog import save_conversion
        from erp.models import Unit
        pack = Unit.objects.create(code='SCI-PACK', name='Boîte scientifique', dimension='COUNT', factor=1)
        save_conversion(self.ops, self.article, {'unit': pack, 'factor': Decimal('96'), 'justification': '96 tests par boîte'})
        self.article = save_article(self.ops, {'purchase_unit': pack, 'minimum_stock': Decimal('96'),
            'target_stock': Decimal('100'), 'minimum_order_quantity': Decimal('3'), 'order_multiple': Decimal('2')},
            pk=self.article.pk, expected=self.article.version)
        data = dashboard_data(self.ops, stock_queryset(self.ops))
        row = next(row for row in data['products'] if row['article'].pk == self.article.pk)
        self.assertTrue(row['low'])
        self.assertEqual((row['purchase_suggested'], row['suggested']), (Decimal('4'), Decimal('384')))
        self.assertEqual(data['references'], 2)
        self.assertContains(self.client.get(reverse('erp:stock-dashboard')), 'CAT-000123')
        self.assertEqual(self.client.get(reverse('erp:stock-dashboard'), {'q': 'ABSENT'}).context['references'], 0)
        self.grant(Capability.VIEW_STOCK, location=self.freezer)
        scoped = dashboard_data(self.operator, stock_queryset(self.operator))
        self.assertEqual(scoped['references'], 0)

    def test_dashboard_purchase_conversion_missing_is_explicit_and_units_are_separate(self):
        from erp.models import Unit
        volume = Unit.objects.create(code='SCI-L', name='Litre scientifique', dimension='VOLUME', factor=1)
        Article.objects.filter(pk=self.article.pk).update(purchase_unit=volume, minimum_stock=Decimal('5'), target_stock=Decimal('8'))
        self.receive('SEPARATE-A', quantity='2')
        self.receive('SEPARATE-B', quantity='3', article=self.liquid, unit=self.ml)
        data = dashboard_data(self.ops, stock_queryset(self.ops))
        self.assertEqual(len(data['units']), 2)
        row = next(row for row in data['products'] if row['article'].pk == self.article.pk)
        self.assertTrue(row['purchase_missing'])
        self.assertEqual(row['suggested'], Decimal('6'))

    def test_native_multi_item_distribution_return_and_internal_split_forms(self):
        first, _, _ = self.receive('HTTP-D1', quantity='10.125')
        second, _, _ = self.receive('HTTP-D2', quantity='10.125')
        response = self.client.get(reverse('erp:dispatch-create'), {'container': first.pk})
        self.assertEqual(response.status_code, 200)
        values = dict(key=uuid.uuid4(), mode='EXIT', beneficiary='Laboratoire HTTP', distributed_on=timezone.localdate(),
            reason='Besoin confirmé', **{'lines-TOTAL_FORMS': 2, 'lines-INITIAL_FORMS': 1,
                'lines-0-container': first.pk, 'lines-0-unit': self.unit.pk, 'lines-0-amount': '2.125', 'lines-0-expected': first.version,
                'lines-1-container': second.pk, 'lines-1-unit': self.unit.pk, 'lines-1-amount': '.125'})
        response = self.client.post(reverse('erp:dispatch-create'), values)
        self.assertEqual(response.status_code, 302, response.context)
        dispatch = StockDispatch.objects.get()
        line = dispatch.lines.get(container=first)
        self.assertContains(self.client.get(reverse('erp:dispatch-list')), 'Laboratoire HTTP')
        url = reverse('erp:stock-return', args=[line.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        returned = dict(key=uuid.uuid4(), amount='1.125', unit=self.unit.pk, destination=self.freezer.pk,
            container_code='HTTP-RETURN', returned_on=timezone.localdate(), condition='Intact', reason='Reliquat')
        self.assertEqual(self.client.post(url, returned).status_code, 302)
        self.assertEqual(self.client.post(url, {**returned, 'key': uuid.uuid4(), 'container_code': 'TOO-MUCH', 'amount': '5'}).status_code, 400)
        self.assertEqual(self.client.post(reverse('erp:dispatch-create'), {**values, 'key': uuid.uuid4(), 'beneficiary': ''}).status_code, 400)
        self.assertEqual(self.client.post(reverse('erp:dispatch-create'), {**values, 'key': uuid.uuid4(), 'lines-0-expected': 999}).status_code, 400)
        first.refresh_from_db()
        internal, _ = self.distribute(first, mode='INTERNAL', destination=self.destination,
            lines=[{'container': first, 'amount': '.125', 'unit': self.unit, 'destination_code': 'HTTP-INTERNAL'}])
        self.assertEqual(returnable_quantity(internal.lines.get()), Decimal('.125'))
        with self.assertRaises(ValidationError):
            return_stock(self.ops, internal.lines.get().pk, **{**returned, 'key': uuid.uuid4(), 'container_code': 'INTERNAL-RETURN',
                'amount': '.125', 'unit': self.unit, 'destination': self.freezer})
        self.assertEqual(reconcile_stock(self.ops), [])

    def test_scoped_search_picker_is_bounded_and_posted_selections_remain_valid(self):
        from erp.stock_forms import DispatchLineForm
        source, _, _ = self.receive('PICKER')
        pending, _, _ = self.receive('PICKER-PENDING', accepted=False)
        response = self.client.get(reverse('erp:dispatch-sources'), {'q': 'SUP-00123'})
        self.assertEqual([row['id'] for row in response.json()['results']], [str(source.pk)])
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        form = DispatchLineForm(user=self.ops, initial={'container': source})
        self.assertIn('data-version', str(form['container']))
        self.assertIn('data-unit', str(form['container']))
        bad = DispatchLineForm({'container': 'invalid', 'unit': self.unit.pk, 'amount': '1'}, user=self.ops)
        self.assertFalse(bad.is_valid())
        self.assertNotIn(pending, form.fields['container'].queryset)
        self.grant(Capability.VIEW_STOCK, location=self.destination)
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(reverse('erp:dispatch-sources'), {'q': source.code}).json()['results'], [])

    def test_ledger_date_drilldown_hierarchical_filters_and_fifo_export(self):
        parent = self.category.__class__.objects.create(code='SCI-PARENT', name='Catégorie parente')
        self.category.parent = parent
        self.category.save()
        source, _, _ = self.receive('DATE', received_on=timezone.localdate() - timedelta(days=5))
        StockContainer.objects.filter(pk=source.pk).update(fifo_received_on=None)
        response = self.client.get(reverse('erp:stock-list'), {'category': parent.pk, 'location': self.lab.pk})
        self.assertContains(response, source.code)
        self.assertIn(parent, response.context['filter_options'][0]['choices'])
        self.assertIn(self.lab, response.context['filter_options'][1]['choices'])
        response = self.client.get(reverse('erp:stock-export'))
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active['S2'].value,
            (timezone.localdate() - timedelta(days=5)).isoformat())
        self.assertContains(self.client.get(reverse('erp:ledger'), {'article': self.article.pk,
            'kind': 'RECEIPT', 'from': timezone.localdate(), 'until': timezone.localdate()}), source.code)
        self.assertEqual(self.client.get(reverse('erp:ledger'), {'from': 'invalid'}).status_code, 404)
        self.assertContains(self.client.get(reverse('erp:ledger'), {'container': source.pk}), source.code)
        self.assertContains(self.client.get(reverse('erp:ledger'), {'q': 'SUP-00123'}), source.code)

    def test_mapping_guards_deduplication_and_empty_table_leave_business_data_untouched(self):
        from erp.services.procurement import create_plan
        from erp.services.table_intake import SCHEMAS
        data = b'code;name;category_code;base_unit_code\nTIPS;;;\n'
        values = dict(key=uuid.uuid4(), kind='CATALOG', filename='source.csv', data=data, reason='Fichier scientifique')
        for changes in ({'reason': ''}, {'data': b''}, {'reason': 'x' * 501}):
            with self.subTest(changes=changes), self.assertRaises(ValidationError):
                stage_import(self.ops, **{**values, **changes})
        plan = create_plan(self.ops, reference='SCI-MAPPING-PLAN', year=timezone.localdate().year + 1, title='Plan du test')
        with self.assertRaises(ValidationError):
            stage_import(self.ops, **{**values, 'plan': plan})
        staged = stage_import(self.ops, **values)
        self.assertEqual(stage_import(self.ops, **values).pk, staged.pk)
        with self.assertRaises(Conflict):
            stage_import(self.ops, **{**values, 'reason': 'Autre source'})
        choices = dict(code=0, name=1, category_code=2, base_unit_code=3)
        for mapping in ({**choices, 'unknown': 1}, {**choices, 'code': -1}, {**choices, 'code': '0'},
                {**choices, 'code': 4}, {'code': 0}, {**choices, 'name': 0}):
            with self.subTest(mapping=mapping), self.assertRaises(ValidationError):
                preview_mapping(self.ops, staged.pk, choices=mapping)
        with self.assertRaises(ValidationError):
            preview_mapping(self.ops, staged.pk, choices=choices, clear_fields=['code'])
        schema = SCHEMAS['INITIAL']['required']
        headers = ';'.join(schema)
        other = stage_import(self.ops, key=uuid.uuid4(), kind='INITIAL', filename='initial.csv',
            data=(headers + '\n' + ';'.join('x' for _ in schema) + '\n').encode(), reason='Stock initial')
        with self.assertRaises(ValidationError):
            preview_mapping(self.ops, other.pk, choices={name: i for i, name in enumerate(schema)}, clear_fields=['brand'])
        with self.assertRaises(ValidationError):
            stage_import(self.ops, **{**values, 'key': uuid.uuid4(), 'data': b'code;name;category_code;base_unit_code\n;;;\n'})
        empty = stage_import(self.ops, **{**values, 'key': uuid.uuid4()})
        ImportMapping.objects.filter(pk=empty.pk).update(matrix=[['code', 'name', 'category_code', 'base_unit_code'], ['', '', '', '']])
        with self.assertRaises(ValidationError):
            preview_mapping(self.ops, empty.pk, choices=choices)
        batch = preview_mapping(self.ops, staged.pk, choices=choices)
        with self.assertRaises(Conflict):
            preview_mapping(self.ops, staged.pk, choices=choices, clear_fields=['brand'])
        self.assertFalse(StockMovement.objects.exists())
        self.assertEqual(Article.objects.get(pk=self.article.pk).catalog_reference, 'CAT-000123')
        self.assertTrue(batch.report['valid'])

    def test_assisted_import_native_upload_mapping_confirmation_and_conflict(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        data = b'Identifiant;Nom;Categorie;Unite;Catalogue\nTIPS;;;;CAT-NATIVE\n'
        values = dict(key=uuid.uuid4(), kind='CATALOG', reason='Reprise vérifiée', assisted='on',
            file=SimpleUploadedFile('scientifique.csv', data, content_type='text/csv'))
        response = self.client.post(reverse('erp:imports'), values)
        self.assertEqual(response.status_code, 302)
        staged = ImportMapping.objects.get()
        url = reverse('erp:import-mapping', args=[staged.pk])
        self.assertContains(self.client.get(url), 'scientifique.csv')
        choices = dict(column_code='0', column_name='1', column_category_code='2', column_base_unit_code='3', column_catalog_reference='4')
        self.assertEqual(self.client.post(url, choices).status_code, 302)
        staged.refresh_from_db()
        self.assertEqual(Article.objects.get(pk=self.article.pk).catalog_reference, 'CAT-000123')
        self.assertEqual(self.client.post(url, {**choices, 'clear_fields': ['brand']}).status_code, 400)
        response = self.client.post(reverse('erp:import-detail', args=[staged.batch_id]),
            {'expected_version': staged.batch.version, 'confirmed': 'on'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Article.objects.get(pk=self.article.pk).catalog_reference, 'CAT-NATIVE')
        self.assertEqual(Article.objects.get(pk=self.article.pk).supplier_reference, 'SUP-00123')

    def test_expiry_filters_do_not_hide_unknown_dates_and_stock_reorder_excludes_reversed_demand(self):
        expired, _, _ = self.receive('EXPIRY-PAST', accepted=False, expires_on=timezone.localdate() - timedelta(days=1))
        soon, _, _ = self.receive('EXPIRY-SOON', expires_on=timezone.localdate() + timedelta(days=20), unit_price='4')
        no_date, _, _ = self.receive('EXPIRY-NONE', expires_on=None, unit_price='3')
        for value, obj in (('expired', expired), ('30', soon), ('none', no_date)):
            response = self.client.get(reverse('erp:stock-list'), {'expiry': value})
            self.assertContains(response, obj.code)
            self.assertEqual(response.context['page'].paginator.count, 1)
        self.article = save_article(self.ops, {'lead_time_days': 30, 'safety_stock': Decimal('1')},
            pk=self.article.pk, expected=self.article.version)
        movement = remove_stock(self.ops, no_date.pk, key=uuid.uuid4(), amount='1', unit=self.unit, reason='Analyse', expected=no_date.version)
        no_date.refresh_from_db()
        with self.assertRaises(Conflict):
            remove_stock(self.ops, no_date.pk, key=uuid.uuid4(), amount='1', unit=self.unit, reason='Analyse', expected=1)
        reverse_stock(self.ops, movement.pk, key=uuid.uuid4(), reason='Analyse annulée')
        data = dashboard_data(self.ops, stock_queryset(self.ops))
        row = next(row for row in data['products'] if row['article'].pk == self.article.pk)
        self.assertEqual(row['consumed'], Decimal(0))
        self.grant(Capability.VIEW_STOCK, location=self.freezer)
        self.assertEqual(dashboard_data(self.operator, stock_queryset(self.operator))['costs'], {})

    def test_single_column_csv_is_supported_and_malformed_delimited_csv_is_rejected(self):
        from erp.table_probe import csv_rows
        import csv
        self.assertEqual(csv_rows(b'code\nTIPS\n'), [['code'], ['TIPS']])
        with self.assertRaises(csv.Error):
            csv_rows(b'a,b,c\n1,2\n3\n')

    def test_currency_validation_and_native_inventory_form_remain_operational(self):
        from erp.stock_forms import ReceiptForm, InventoryForm
        for currency in ('D1Z', 'DZ', '€€€'):
            self.assertIn('currency', ReceiptForm({'currency': currency}, user=self.ops).errors)
        self.assertIn('blind', InventoryForm(user=self.ops).fields)

    def test_legacy_expiry_is_used_for_fefo_sorting_and_exports(self):
        early, _, _ = self.receive('LEGACY-EXPIRY', expires_on=timezone.localdate() + timedelta(days=5))
        later, _, _ = self.receive('LATER-EXPIRY', expires_on=timezone.localdate() + timedelta(days=10))
        unknown, _, _ = self.receive('UNKNOWN-EXPIRY', expires_on=None)
        StockContainer.objects.filter(pk=early.pk).update(use_by=None)
        self.assertEqual(fefo(self.ops, self.article).first().pk, early.pk)
        response = self.client.get(reverse('erp:stock-list'), {'sort': 'expiry'})
        self.assertEqual([obj.pk for obj in response.context['page']], [early.pk, later.pk, unknown.pk])
        response = self.client.get(reverse('erp:stock-list'), {'sort': 'expiry', 'direction': 'desc'})
        self.assertEqual([obj.pk for obj in response.context['page']], [later.pk, early.pk, unknown.pk])
        response = self.client.get(reverse('erp:stock-export'), {'q': early.code})
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active['U2'].value,
            early.lot.expires_on.isoformat())

    def test_browser_database_names_are_isolated_and_reject_production_paths(self):
        import importlib
        from django.core.exceptions import ImproperlyConfigured
        import plagenor.settings_e2e as isolated
        try:
            with patch.dict('os.environ', {'PLAGENOR_E2E_DATABASE_NAME': 'plagenor-e2e-Safe123.sqlite3'}):
                self.assertEqual(importlib.reload(isolated).DATABASES['default']['NAME'].name, 'plagenor-e2e-Safe123.sqlite3')
            for name in ('plagenor.sqlite3', '../plagenor-e2e.sqlite3'):
                with patch.dict('os.environ', {'PLAGENOR_E2E_DATABASE_NAME': name}), self.assertRaises(ImproperlyConfigured):
                    importlib.reload(isolated)
        finally:
            importlib.reload(isolated)
