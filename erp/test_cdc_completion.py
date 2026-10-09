import copy
import io
import os
import uuid
from datetime import timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone, translation
from openpyxl import load_workbook

from erp.cdc.catalog import document, generate_document, profile
from erp.cdc.docengine import DocumentError, sha
from erp.cdc.financial import excel_number, summarize
from erp.cdc.schedule_adapter import managed_ids
from erp.models import CdcItem, CdcReusePreview, CdcWorkbookPreview
from erp.services.cdc import _estimates, _revision, create_dossier, document_data, estimate_totals, save_cdc_item, save_requirement
from erp.services.cdc_exchange import _apply_catalog, apply_workbook, restore_revision, set_lot_active
from erp.services.cdc_finance import export_financial, financial_snapshot, parse_financial, preview_financial
from erp.services.cdc_reuse import apply_reuse, preview_reuse, source_rows
from erp.services.cdc_tables import add_table_row, editable_table, remove_table_row
from erp.test_operations import OperationFixtures


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class CdcCompletionTests(OperationFixtures, TestCase):
    def setUp(self):
        self.source = create_dossier(self.ops, family='equipment', reference='401/SME/SDFM/SG/ESSBO/2026',
            title='Source historique', assignee=self.operator)
        self.target = create_dossier(self.ops, family='equipment', reference='402/SME/SDFM/SG/ESSBO/2026',
            title='Nouvelle opération', assignee=self.operator)
        self.lot = self.target.lots.first()
        self.client.force_login(self.ops)

    def price(self, dossier=None, **values):
        dossier = dossier or self.target
        item = dossier.lots.first().items.first()
        save_cdc_item(self.ops, item.lot_id, expected=dossier.version, pk=item.pk,
            values={'quantity': 2, 'estimated_price': 100, 'tax_rate': 19, 'price_source': 'Devis fournisseur', **values})
        dossier.refresh_from_db(); item.refresh_from_db()
        return item

    def reuse(self, actor=None, count=2):
        revision = self.source.revisions.first()
        _, rows = source_rows(actor or self.ops, revision.pk, self.target.family)
        return preview_reuse(actor or self.ops, self.target.pk, expected=self.target.version,
            source_revision=revision.pk, target_lot=self.lot.pk,
            selections=[row['selection'] for row in rows[:count]], reason='Besoins repris et vérifiés')

    def edits(self, preview):
        return [{key: value for key, value in row.items() if key in
            {'selection', 'designation', 'specifications', 'unit', 'packaging', 'quantity', 'details'}}
            for row in copy.deepcopy(preview.payload['rows'])]

    def upload(self, changes=None, dossier=None):
        book = load_workbook(io.BytesIO(export_financial(self.ops, dossier or self.target)))
        for cell, value in (changes or {}).items():
            book.worksheets[0][cell] = value
        output = io.BytesIO(); book.save(output)
        return SimpleUploadedFile('estimations.xlsx', output.getvalue())

    def financial_preview(self, upload=None, actor=None):
        return preview_financial(actor or self.ops, self.target.pk, expected=self.target.version,
            upload=upload or self.upload({'D2': 4, 'E2': 150, 'F2': 19}), reason='Nouveau devis confirmé')

    def row_template(self):
        _, managed = managed_ids(self.target.family)
        for table in profile(self.target.family)['tables']:
            if table['id'] not in managed and not table['nested'] and table['part'] == 'word/document.xml':
                for row in table['rows'][1:]:
                    if row['cloneable']:
                        cells = [' '.join(cell['text'].splitlines()) for cell in row['cells']]
                        if cells and cells[0]:
                            cells[0] = 'Ligne ajoutée et vérifiée'
                            try:
                                document(self.target.family).generate(table_edits={table['id']:
                                    [{'source_row': row['index'], 'after_row': row['index'], 'cells': cells}]})
                            except DocumentError:
                                continue
                            return table, row, cells
        self.fail('No safely cloneable canonical row')

    def test_selected_reuse_preview_edits_and_retry_keep_sources_and_prices_independent(self):
        original = self.price(self.source)
        save_requirement(self.ops, original.pk, expected=self.source.version, values={
            'statement': 'Preuve technique de réception', 'position': 1, 'kind': 'MANDATORY',
            'evidence': 'Fiche technique du fabricant', 'verification_method': 'Contrôle à la réception'})
        self.source.refresh_from_db()
        original_revision = self.source.revisions.first()
        baseline = copy.deepcopy(original_revision.data)
        before = self.lot.items.count()
        preview = self.reuse()
        self.assertEqual(self.lot.items.count(), before)
        rows = self.edits(preview); rows[0]['quantity'] = '7'; rows[0]['designation'] = 'Copie indépendante'
        rows[1]['omit'] = True
        revision = apply_reuse(self.ops, preview.pk, rows=rows)
        self.assertEqual(apply_reuse(self.ops, preview.pk, rows=rows).pk, revision.pk)
        copied = self.lot.items.get(designation='Copie indépendante')
        self.assertEqual(copied.quantity, 7)
        self.assertIsNone(copied.estimated_price); self.assertIsNone(copied.tax_rate)
        self.assertEqual(copied.requirements.get().statement, 'Preuve technique de réception')
        self.assertNotEqual(copied.pk, original.pk)
        self.assertEqual(self.lot.items.count(), before + 1)
        original.refresh_from_db(); original_revision.refresh_from_db()
        self.assertEqual(original.estimated_price, 100); self.assertEqual(original.quantity, 2)
        self.assertEqual(original_revision.data, baseline)

    def test_reuse_retains_structured_links_and_rejects_unit_reinterpretation(self):
        original = self.source.lots.first().items.first()
        save_cdc_item(self.ops, original.lot_id, expected=self.source.version, pk=original.pk,
            article=self.article, purchase_unit=self.unit, values={'quantity': 2})
        self.source.refresh_from_db()
        preview = self.reuse(count=1)
        rows = self.edits(preview); rows[0]['unit'] = 'Autre unité'
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=rows)
        rows[0]['unit'] = self.unit.name
        apply_reuse(self.ops, preview.pk, rows=rows)
        copied = self.lot.items.get(source_key__startswith='new-')
        self.assertEqual(copied.article_id, self.article.pk)
        self.assertEqual(copied.purchase_unit_id, self.unit.pk)
        self.assertEqual(copied.base_factor, 1)

    def test_reuse_missing_catalogue_link_is_explicitly_unstructured(self):
        preview = self.reuse(count=1)
        apply_reuse(self.ops, preview.pk, rows=self.edits(preview))
        self.assertIsNone(self.lot.items.get(source_key__startswith='new-').article_id)

    def test_reuse_expiry_stale_version_actor_and_source_revocation_are_atomic(self):
        preview = self.reuse(self.operator)
        with self.assertRaises(PermissionDenied): apply_reuse(self.ops, preview.pk, rows=self.edits(preview))
        self.source.work.assignee = self.second; self.source.work.save()
        with self.assertRaises(PermissionDenied): apply_reuse(self.operator, preview.pk, rows=self.edits(preview))
        self.source.work.assignee = self.operator; self.source.work.save()
        preview.expires_at = timezone.now() - timedelta(seconds=1); preview.save()
        with self.assertRaises(ValidationError): apply_reuse(self.operator, preview.pk, rows=self.edits(preview))
        preview = self.reuse()
        self.price()
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=self.edits(preview))
        self.assertFalse(self.lot.items.filter(source_key__startswith='new-').exists())

    def test_reuse_selection_integrity_and_preview_tampering_are_rejected(self):
        revision, rows = source_rows(self.ops, self.source.revisions.first().pk, self.target.family)
        for selections, reason in [([], 'Motif'), ([rows[0]['selection']] * 2, 'Motif'),
            (['inconnu'], 'Motif'), ([rows[0]['selection']], '')]:
            with self.subTest(selections=selections), self.assertRaises(ValidationError):
                preview_reuse(self.ops, self.target.pk, expected=self.target.version, source_revision=revision.pk,
                    target_lot=self.lot.pk, selections=selections, reason=reason)
        with patch('erp.services.cdc_reuse.fingerprint', return_value='0'*64):
            with self.assertRaises(ValidationError): source_rows(self.ops, revision.pk, self.target.family)
        preview = self.reuse(count=1)
        edits = self.edits(preview)
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=[])
        edits[0]['selection'] = 'autre'
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=edits)
        edits = self.edits(preview); edits[0]['omit'] = True
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=edits)
        preview.payload['source_sha256'] = '0'*64; preview.save()
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=self.edits(preview))

    def test_reuse_http_selection_edit_confirmation_and_access(self):
        url = reverse('erp:cdc-catalogue', args=[self.target.pk])
        self.assertContains(self.client.get(url), 'Catalogue et réutilisation')
        self.assertEqual(self.client.get(url+'?source_revision=incorrect').status_code, 200)
        revision, rows = source_rows(self.ops, self.source.revisions.first().pk, self.target.family)
        self.assertEqual(self.client.get(url+'?source_revision='+str(revision.pk)+'&q=introuvable').status_code, 200)
        data = {'expected_version': self.target.version, 'source_revision': revision.pk,
            'target_lot': self.lot.pk, 'selections': [rows[0]['selection']], 'reason': 'Réutilisation confirmée'}
        self.assertEqual(self.client.post(url, {}).status_code, 400)
        self.assertEqual(self.client.post(url, {**data, 'source_revision': 'incorrect'}).status_code, 400)
        self.assertEqual(self.client.post(url, {**data, 'expected_version': 999}).status_code, 400)
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        preview = CdcReusePreview.objects.get()
        self.assertContains(self.client.get(response.url), 'Vérifier les articles à réutiliser')
        posted = {'form-TOTAL_FORMS': 1, 'form-INITIAL_FORMS': 1, 'form-MIN_NUM_FORMS': 0,
            'form-MAX_NUM_FORMS': 1000, **{'form-0-'+key: value for key, value in self.edits(preview)[0].items()}}
        posted['form-0-specifications'] = 'Première ligne\r\nDeuxième ligne'
        self.assertEqual(self.client.post(response.url, posted).status_code, 400)
        self.assertEqual(self.client.post(response.url, {**posted, 'confirm': 'on', 'form-0-quantity': 0}).status_code, 400)
        self.assertEqual(self.client.post(response.url, {**posted, 'confirm': 'on'}).status_code, 302)
        self.assertContains(self.client.get(response.url), 'déjà été appliqué')
        self.client.force_login(self.second)
        self.assertEqual(self.client.get(response.url).status_code, 404)
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_financial_export_and_import_share_line_lot_global_totals(self):
        item = self.price()
        initial = self.target.revisions.first()
        revision, summary = financial_snapshot(self.target)
        self.assertEqual(summary['currencies']['DZD'], {'net': Decimal(200), 'tax': Decimal(38), 'gross': Decimal(238)})
        self.assertEqual(estimate_totals(self.ops, self.target)['currencies'], summary['currencies'])
        book = load_workbook(io.BytesIO(export_financial(self.ops, self.target)))
        self.assertEqual(book.worksheets[0]['G2'].value, 200)
        self.assertEqual(book.worksheets[0]['H2'].value, 38)
        self.assertEqual(book.worksheets[0]['I2'].value, 238)
        self.assertTrue(book.worksheets[0].column_dimensions['J'].hidden)
        self.assertEqual(book['_CDC'].sheet_state, 'veryHidden')
        self.assertFalse(any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row))
        preview = self.financial_preview()
        item.refresh_from_db(); self.assertEqual(item.estimated_price, 100)
        self.assertEqual(preview.payload['finance_totals']['currencies']['DZD']['gross'], '714.00')
        applied = apply_workbook(self.ops, preview.pk)
        self.assertEqual(apply_workbook(self.ops, preview.pk).pk, applied.pk)
        item.refresh_from_db(); self.assertEqual((item.quantity, item.estimated_price, item.tax_rate), (4, 150, 19))
        initial.refresh_from_db(); self.assertEqual(initial.estimates[0]['price'], '100.00')
        self.target.refresh_from_db()
        self.assertEqual(estimate_totals(self.ops, self.target)['currencies']['DZD']['gross'], 714)

    def test_financial_blank_cells_clear_values_and_preview_shows_incomplete_lines(self):
        item = self.price()
        preview = self.financial_preview(self.upload({'E2': None, 'F2': None}))
        self.assertGreater(preview.payload['finance_totals']['incomplete_lines'], 0)
        apply_workbook(self.ops, preview.pk)
        item.refresh_from_db(); self.assertIsNone(item.estimated_price); self.assertIsNone(item.tax_rate)

    def test_financial_untrusted_workbooks_are_rejected_without_writes(self):
        self.price()
        for changes in [{'E2': '=1+2'}, {'E2': -1}, {'F2': 101}, {'D2': 0}, {'J2': str(uuid.uuid4())},
            {'J3': self.target.lots.first().items.first().pk.hex}, {'A2': 'Lot déplacé'}, {'A1': 'En-tête modifié'},
            {column+'2': None for column in 'ABCDEFGHIJ'}]:
            with self.subTest(changes=changes), self.assertRaises(DocumentError):
                self.financial_preview(self.upload(changes))
        with self.assertRaises(DocumentError): self.financial_preview(self.upload(dossier=self.source))
        book = load_workbook(io.BytesIO(export_financial(self.ops, self.target)))
        book.worksheets[0]['E2'].number_format = 'DD/MM/YYYY'
        output = io.BytesIO(); book.save(output)
        with self.assertRaises(DocumentError): self.financial_preview(SimpleUploadedFile('date.xlsx', output.getvalue()))
        book.worksheets[0]['E2'].number_format = '0.00'
        book.worksheets[0]['B2'].hyperlink = 'https://example.org/'
        output = io.BytesIO(); book.save(output)
        with self.assertRaises(DocumentError): self.financial_preview(SimpleUploadedFile('lien.xlsx', output.getvalue()))
        self.assertFalse(CdcWorkbookPreview.objects.exists())

    def test_financial_permissions_revocation_and_revision_binding(self):
        with self.assertRaises(PermissionDenied): export_financial(self.operator, self.target)
        with self.assertRaises(PermissionDenied): self.financial_preview(actor=self.operator)
        self.target.work.allow_costs = True; self.target.work.save()
        preview = self.financial_preview(actor=self.operator)
        self.target.work.allow_costs = False; self.target.work.save()
        with self.assertRaises(PermissionDenied): apply_workbook(self.operator, preview.pk)
        uploaded = self.upload()
        self.price()
        with self.assertRaises(DocumentError): self.financial_preview(uploaded)

    def test_financial_error_handling_currency_large_values_and_integrity(self):
        with patch('erp.services.cdc_finance.fingerprint', return_value='0'*64):
            with self.assertRaises(ValidationError): financial_snapshot(self.target)
        self.price(currency='EUR')
        with self.assertRaises(ValidationError): export_financial(self.ops, self.target)
        with self.assertRaises(DocumentError): parse_financial(b'bad', 'foreign.xlsx', self.target)
        url = reverse('erp:cdc-financial-download', args=[self.target.pk])
        self.assertEqual(self.client.get(url).status_code, 400)
        row = {'lot': '1', 'lot_name': 'Grand budget', 'quantity': Decimal('999999999999.999999'),
            'price': Decimal('9999999999999999.99'), 'tax_rate': Decimal(19), 'currency': 'DZD'}
        summary = summarize([row])
        self.assertIsInstance(excel_number(summary['currencies']['DZD']['gross']), str)
        self.assertEqual(excel_number(Decimal('1.25')), Decimal('1.25'))

    def test_financial_http_download_preview_confirmation_and_missing_fields(self):
        self.price()
        url = reverse('erp:cdc-finance', args=[self.target.pk])
        self.assertContains(self.client.get(url), 'Estimations et totaux financiers')
        download = self.client.get(reverse('erp:cdc-financial-download', args=[self.target.pk]))
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download['Cache-Control'], 'private, no-store')
        self.assertEqual(self.client.post(url, {}).status_code, 400)
        self.assertEqual(self.client.post(url, {'expected_version': self.target.version,
            'file': self.upload({'E2': '=1+2'}), 'reason': 'Contrôle'}).status_code, 400)
        response = self.client.post(url, {'expected_version': self.target.version,
            'file': self.upload({'D2': 4, 'E2': 150, 'F2': 19}), 'reason': 'Devis confirmé'})
        self.assertEqual(response.status_code, 302)
        self.assertContains(self.client.get(response.url), 'Totaux financiers après import')
        self.assertEqual(self.client.post(response.url).status_code, 302)
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.get(reverse('erp:cdc-financial-download', args=[self.target.pk])).status_code, 403)

    def test_table_rows_are_versioned_and_original_word_is_preserved(self):
        table, row, cells = self.row_template()
        original = self.target.revisions.first()
        before = sha(document(self.target.family).data)
        add_table_row(self.ops, self.target.pk, expected=self.target.version, table_id=table['id'],
            source_row=row['index'], after_row=row['index'], cells=cells, reason='Adaptation documentée')
        self.target.refresh_from_db()
        payload, _ = generate_document(document_data(self.target))
        with ZipFile(io.BytesIO(payload)) as archive:
            self.assertIn('Ligne ajoutée et vérifiée', archive.read('word/document.xml').decode())
        self.assertEqual(sha(document(self.target.family).data), before)
        self.assertEqual(original.data['rows'], {})
        remove_table_row(self.ops, self.target.pk, expected=self.target.version, table_id=table['id'], index=0,
            reason='Retrait confirmé')
        self.target.refresh_from_db(); self.assertEqual(self.target.data['rows'][table['id']], [])
        self.assertEqual(self.target.revisions.count(), 3)

    def test_structural_or_invalid_table_edits_are_atomic(self):
        table, row, cells = self.row_template()
        for table_id, source_row in [('inconnu', 1), (table['id'], 0), (table['id'], 999)]:
            with self.assertRaises(ValidationError): editable_table(self.target.family, table_id, source_row)
        for reason, after in [('', row['index']), ('Contrôle', 999)]:
            with self.assertRaises((ValidationError, DocumentError)):
                add_table_row(self.ops, self.target.pk, expected=self.target.version, table_id=table['id'],
                    source_row=row['index'], after_row=after, cells=cells, reason=reason)
        with self.assertRaises(ValidationError):
            remove_table_row(self.ops, self.target.pk, expected=self.target.version, table_id=table['id'], index=0, reason='')
        self.assertEqual(self.target.revisions.count(), 1)

    def test_table_http_navigation_add_and_remove_with_confirmation(self):
        table, row, cells = self.row_template()
        tables = reverse('erp:cdc-tables', args=[self.target.pk])
        self.assertContains(self.client.get(tables), 'Tableaux du modèle')
        self.assertContains(self.client.get(tables+'?q=introuvable'), 'Aucun résultat')
        add = reverse('erp:cdc-table-add', args=[self.target.pk])
        params = '?table_id='+table['id']+'&source_row='+str(row['index'])
        self.assertEqual(self.client.get(add+params).status_code, 200)
        self.assertEqual(self.client.get(add+'?source_row=incorrect').status_code, 404)
        data = {'expected_version': self.target.version, 'table_id': table['id'], 'source_row': row['index'],
            'after_row': row['index'], 'reason': 'Adaptation confirmée',
            **{'cell_%s' % index: value for index, value in enumerate(cells)}}
        self.assertEqual(self.client.post(add, {**data, 'reason': ''}).status_code, 400)
        self.assertEqual(self.client.post(add, {**data, 'expected_version': 999}).status_code, 400)
        self.assertEqual(self.client.post(add, data).status_code, 302)
        self.target.refresh_from_db()
        remove = reverse('erp:cdc-table-remove', args=[self.target.pk, 0])
        self.assertEqual(self.client.get(remove+'?table_id='+table['id']).status_code, 200)
        self.assertEqual(self.client.post(remove, {}).status_code, 400)
        self.assertEqual(self.client.post(remove, {'expected_version': 999, 'table_id': table['id'],
            'reason': 'Contrôle', 'confirm': 'on'}).status_code, 400)
        self.assertEqual(self.client.post(remove, {'expected_version': self.target.version, 'table_id': table['id'],
            'reason': 'Retrait confirmé', 'confirm': 'on'}).status_code, 302)

    def test_import_guide_and_native_navigation_respect_cost_permissions(self):
        for language in ('fr', 'en', 'ar'):
            self.client.cookies['django_language'] = language
            for name in ('cdc-detail', 'cdc-guide', 'cdc-import-home'):
                page = self.client.get(reverse('erp:'+name, args=[self.target.pk]))
                self.assertEqual(page.status_code, 200)
                self.assertEqual(page['Content-Language'], language)
                if language == 'ar':
                    self.assertContains(page, 'dir="rtl"')
        self.client.cookies['django_language'] = 'fr'
        page = self.client.get(reverse('erp:cdc-detail', args=[self.target.pk]))
        for name in ('cdc-catalogue', 'cdc-finance', 'cdc-tables', 'cdc-import-home', 'cdc-guide'):
            self.assertContains(page, reverse('erp:'+name, args=[self.target.pk]))
        self.client.force_login(self.operator)
        self.assertNotContains(self.client.get(reverse('erp:cdc-import-home', args=[self.target.pk])),
            reverse('erp:cdc-finance', args=[self.target.pk]))
        self.client.force_login(self.outsider)
        for name in ('cdc-guide', 'cdc-import-home', 'cdc-tables'):
            self.assertEqual(self.client.get(reverse('erp:'+name, args=[self.target.pk])).status_code, 404)

    def test_legacy_revisions_without_source_keys_remain_financially_readable_and_reusable(self):
        def legacy(dossier):
            return [{key: value for key, value in entry.items() if key != 'source_key'}
                for entry in _estimates(dossier)]
        with patch('erp.services.cdc._estimates', side_effect=legacy):
            item = self.price(self.source)
        revision, totals = financial_snapshot(self.source)
        self.assertNotIn('source_key', revision.estimates[0])
        self.assertEqual(totals['currencies']['DZD']['gross'], 238)
        preview = self.reuse(count=1)
        apply_reuse(self.ops, preview.pk, rows=self.edits(preview))
        self.assertEqual(item.quantity, 2)

    def test_reuse_corrupted_or_oversized_copies_rollback_and_http_expiry_is_clear(self):
        preview = self.reuse(count=1)
        rows = self.edits(preview); rows[0]['details'] = 'x' * 10001
        with self.assertRaises(DocumentError): apply_reuse(self.ops, preview.pk, rows=rows)
        self.assertFalse(self.lot.items.filter(source_key__startswith='new-').exists())
        preview.expires_at = timezone.now() - timedelta(seconds=1); preview.save()
        url = reverse('erp:cdc-reuse-preview', args=[preview.pk])
        data = {'form-TOTAL_FORMS': 1, 'form-INITIAL_FORMS': 1, 'form-MIN_NUM_FORMS': 0,
            'form-MAX_NUM_FORMS': 1000, 'confirm': 'on',
            **{'form-0-'+key: value for key, value in self.edits(preview)[0].items()}}
        self.assertContains(self.client.post(url, data), 'a expiré', status_code=400)
        preview.actor = self.operator; preview.save()
        self.client.force_login(self.operator)
        self.source.work.assignee = self.second; self.source.work.save()
        self.assertEqual(self.client.get(url).status_code, 403)

    def mutate_archive(self, payload, *, replace=None, extra=None, remove=()):
        result = io.BytesIO()
        with ZipFile(io.BytesIO(payload)) as source, ZipFile(result, 'w') as destination:
            for name in source.namelist():
                if name not in remove:
                    value = source.read(name)
                    if replace and name in replace:
                        value = replace[name](value)
                    destination.writestr(name, value)
            for name, value in (extra or {}).items():
                destination.writestr(name, value)
        return result.getvalue()

    def numeric_style(self, payload, value):
        root = ET.fromstring(payload)
        root.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r="D2"]').set('s', value)
        return ET.tostring(root)

    def test_financial_malformed_archives_styles_and_signed_identity_are_rejected(self):
        payload = export_financial(self.ops, self.target)
        variants = [
            self.mutate_archive(payload, extra={'xl/vbaProject.bin': b'macro'}),
            self.mutate_archive(payload, extra={'xl/sharedStrings.xml':
                ('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                 + '<si><t>x</t></si>' * 10001 + '</sst>').encode()}),
            self.mutate_archive(payload, remove=['xl/styles.xml']),
            self.mutate_archive(payload, replace={'xl/styles.xml': lambda data:
                b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>'}),
            self.mutate_archive(payload, replace={'xl/worksheets/sheet1.xml': lambda data:
                self.numeric_style(data, 'incorrect')}),
            self.mutate_archive(payload, replace={'xl/worksheets/sheet1.xml': lambda data:
                self.numeric_style(data, '999')}),
            self.mutate_archive(payload, remove=['xl/worksheets/sheet4.xml']),
        ]
        for value in variants:
            with self.subTest(size=len(value)), self.assertRaises(DocumentError):
                parse_financial(value, 'malforme.xlsx', self.target)
        for filename, value in [('macro.xlsm', payload), ('trop-grand.xlsx', b'x'*(20*1024*1024+1))]:
            with self.assertRaises(DocumentError): parse_financial(value, filename, self.target)
        book = load_workbook(io.BytesIO(payload)); book['_CDC']['B2'] = 'signature-fausse'
        output = io.BytesIO(); book.save(output)
        with self.assertRaises(DocumentError): parse_financial(output.getvalue(), 'identite.xlsx', self.target)
        book = load_workbook(io.BytesIO(payload)); book['_CDC']['A1'] = 'AUTRE_CLASSEUR'
        output = io.BytesIO(); book.save(output)
        with self.assertRaises(DocumentError): parse_financial(output.getvalue(), 'modele-inconnu.xlsx', self.target)
        book = load_workbook(io.BytesIO(payload)); book.worksheets[0]['J3'] = book.worksheets[0]['J2'].value
        output = io.BytesIO(); book.save(output)
        with self.assertRaises(DocumentError): parse_financial(output.getvalue(), 'doublon.xlsx', self.target)

    def test_financial_shared_strings_empty_rows_and_export_failure_are_controlled(self):
        payload = export_financial(self.ops, self.target)
        shared = b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>Lot</t></si></sst>'
        payload = self.mutate_archive(payload, replace={'xl/worksheets/sheet1.xml': lambda data:
            data.replace(b'<c r="A1" s="1" t="inlineStr"><is><t>Lot</t></is></c>',
                b'<c r="A1" s="1" t="s"><v>0</v></c>').replace(b'</sheetData>', b'<row r="999"><c r="D999" t="inlineStr"><is><t></t></is></c></row></sheetData>')},
            extra={'xl/sharedStrings.xml': shared})
        parsed = parse_financial(payload, 'chaines.xlsx', self.target)
        self.assertEqual(len(parsed['financial']), CdcItem.objects.filter(lot__dossier=self.target, active=True).count())
        with patch('erp.services.cdc_finance._sheet', side_effect=ValueError('invalid XML')):
            with self.assertRaises(ValidationError): export_financial(self.ops, self.target)
        with patch('erp.services.cdc_finance.fingerprint', return_value='0'*64):
            self.assertEqual(self.client.get(reverse('erp:cdc-finance', args=[self.target.pk])).status_code, 400)
        for reason, size in [('', 1), ('Contrôle', 20*1024*1024+1)]:
            upload = SimpleUploadedFile('budget.xlsx', b'x'); upload.size = size
            with self.assertRaises(ValidationError):
                preview_financial(self.ops, self.target.pk, expected=self.target.version, upload=upload, reason=reason)

    def sized_dossier(self, count):
        dossier = create_dossier(self.ops, family='reagents', reference=f'{500+count}/SME/SDFM/SG/ESSBO/2026',
            title=f'Qualification synthétique {count} articles', assignee=self.operator)
        lots = list(dossier.lots.all())
        CdcItem.objects.filter(lot__dossier=dossier).update(active=False)
        CdcItem.objects.bulk_create([CdcItem(lot=lots[(index-1)%len(lots)], source_key='new-'+str(uuid.uuid4()),
            position=(index-1)//len(lots)+1,
            designation=f'Consommable synthétique {index:03d}', unit_label='Unité', quantity=2,
            specifications='Spécification française et عربية, contrôle à la réception. '*3,
            estimated_price=Decimal('12.50'), tax_rate=19, price_source='Jeu de données synthétique')
            for index in range(1, count+1)])
        _revision(self.ops, dossier, 'Jeu de qualification synthétique')
        return dossier

    def test_57_and_74_item_roundtrips_deletion_history_and_docx_generation(self):
        query_counts = []
        for count in (57, 74):
            dossier = self.sized_dossier(count)
            prior = dossier.revisions.first()
            upload = self.upload({'D2': 3, 'E2': Decimal('15.25'), 'F2': 19}, dossier=dossier)
            with CaptureQueriesContext(connection) as queries:
                preview = preview_financial(self.ops, dossier.pk, expected=dossier.version,
                    upload=upload, reason='Budget volumineux confirmé')
                revision = apply_workbook(self.ops, preview.pk)
            query_counts.append(len(queries))
            dossier.refresh_from_db()
            self.assertEqual(len(financial_snapshot(dossier)[1]['lines']), count)
            self.assertEqual(prior.data['lot_catalog']['lots'][0]['items'][0]['quantity'], '2.000000')
            self.assertEqual(Decimal(revision.data['lot_catalog']['lots'][0]['items'][0]['quantity']), 3)
            payload, _ = generate_document(document_data(dossier))
            with ZipFile(io.BytesIO(payload)) as archive:
                text = archive.read('word/document.xml').decode()
                for index in range(1, count+1): self.assertIn(f'Consommable synthétique {index:03d}', text)
                self.assertNotIn('15.25', text)
                self.assertNotIn('Jeu de données synthétique', text)
            directory = os.getenv('PLAGENOR_DOCUMENT_PREVIEW_DIR', '')
            if directory:
                Path(directory).mkdir(parents=True, exist_ok=True)
                Path(directory, f'cdc-reagents-{count}-items.docx').write_bytes(payload)
            lot = dossier.lots.filter(active=True).first()
            removed = lot.items.filter(active=True).count()
            set_lot_active(self.ops, lot.pk, expected=dossier.version, active=False, reason='Retrait contrôlé')
            dossier.refresh_from_db()
            self.assertEqual(len(financial_snapshot(dossier)[1]['lines']), count-removed)
            with self.assertRaises(DocumentError): generate_document(document_data(dossier))
            restore_revision(self.ops, revision.pk, expected=dossier.version, reason='Reprise vérifiée')
            dossier.refresh_from_db()
            self.assertEqual(len(financial_snapshot(dossier)[1]['lines']), count)
            prior.refresh_from_db()
            self.assertEqual(sum(len(lot['items']) for lot in prior.data['lot_catalog']['lots']), count)
        # Additional rows are written in bounded batches rather than one query per item.
        self.assertLessEqual(query_counts[1], query_counts[0]+4)
        self.assertLess(query_counts[1], 65)

    def test_large_selective_copy_keeps_one_revision_and_bounded_database_queries(self):
        source = self.sized_dossier(74)
        target = create_dossier(self.ops, family='reagents', reference='674/SME/SDFM/SG/ESSBO/2026',
            title='Destination synthétique', assignee=self.operator)
        _, rows = source_rows(self.ops, source.revisions.first().pk, target.family)
        lot = target.lots.first()
        preview = preview_reuse(self.ops, target.pk, expected=target.version,
            source_revision=source.revisions.first().pk, target_lot=lot.pk,
            selections=[row['selection'] for row in rows], reason='74 copies contrôlées')
        with CaptureQueriesContext(connection) as queries:
            apply_reuse(self.ops, preview.pk, rows=self.edits(preview))
        self.assertEqual(lot.items.filter(source_key__startswith='new-').count(), 74)
        self.assertEqual(target.revisions.count(), 2)
        self.assertLess(len(queries), 55)

    def test_missing_lot_in_catalog_apply_cannot_write_an_item(self):
        catalog = copy.deepcopy(document_data(self.target)['lot_catalog'])
        catalog['lots'][0]['id'] = str(uuid.uuid4())
        with self.assertRaises(ValidationError): _apply_catalog(self.target, catalog)

    def test_works_selective_append_is_rejected_before_an_incompatible_revision_is_created(self):
        dossier = create_dossier(self.ops, family='works', reference='799/SME/SDFM/SG/ESSBO/2026',
            title='Postes travaux protégés', assignee=self.operator)
        revision, rows = source_rows(self.ops, dossier.revisions.first().pk, dossier.family)
        lot = dossier.lots.first()
        with self.assertRaises(ValidationError):
            preview_reuse(self.ops, dossier.pk, expected=dossier.version, source_revision=revision.pk,
                target_lot=lot.pk, selections=[rows[0]['selection']], reason='Copie non admissible')
        preview = CdcReusePreview.objects.create(dossier=dossier, actor=self.ops, source_revision=revision,
            target_lot=lot, base_version=dossier.version, payload={'rows': rows[:1], 'source_sha256': revision.sha256},
            reason='Contrôle des aperçus restaurés', expires_at=timezone.now()+timedelta(hours=1))
        with self.assertRaises(ValidationError): apply_reuse(self.ops, preview.pk, rows=self.edits(preview))
        url = reverse('erp:cdc-catalogue', args=[dossier.pk])
        self.assertContains(self.client.get(url), '249 postes structurés')
        self.assertEqual(self.client.post(url, {'source_revision': revision.pk}).status_code, 400)
        self.client.force_login(self.operator)
        self.assertNotContains(self.client.get(url), reverse('erp:cdc-duplicate', args=[dossier.pk]))
        self.assertEqual(lot.items.count(), 249)
        self.assertEqual(dossier.revisions.count(), 1)
