import io
from datetime import timedelta
from decimal import Decimal

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from erp.cdc.docengine import DocumentError
from erp.models import CdcItem, CdcWorkbookPreview
from erp.services.cdc import create_dossier, document_data, save_cdc_item
from erp.services.cdc_exchange import (add_lot, apply_workbook, export_workbook, preview_workbook,
                                     restore_revision, set_lot_active)
from erp.test_operations import OperationFixtures


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class CdcWorkbookTests(OperationFixtures, TestCase):
    def setUp(self):
        self.dossier = create_dossier(self.ops, family='equipment', reference='81/SME/SDFM/SG/ESSBO/2026',
            title='CDC Excel', assignee=self.operator)
        self.item = self.dossier.lots.first().items.first()
        self.initial = self.dossier.revisions.first()

    def upload(self, changes=None, *, prices=False, family=None):
        d = family or self.dossier
        book = load_workbook(io.BytesIO(export_workbook(self.ops, d, prices=prices)))
        if changes is None:
            changes = {'F7': 12}
        for cell, value in changes.items():
            book['Lot_01'][cell] = value
        output = io.BytesIO(); book.save(output)
        return SimpleUploadedFile('lots.xlsx', output.getvalue())

    def preview(self, user=None, **kwargs):
        return preview_workbook(user or self.ops, self.dossier.pk, expected=self.dossier.version,
            upload=kwargs.pop('upload', None) or self.upload(), mode=kwargs.pop('mode', 'merge'),
            reason=kwargs.pop('reason', 'Besoins confirmés, devis fournisseur'), **kwargs)

    def test_full_http_export_preview_confirmation_and_retry(self):
        self.client.force_login(self.operator)
        url = reverse('erp:cdc-workbook', args=[self.dossier.pk])
        page = self.client.get(url)
        self.assertContains(page, 'Analyser et prévisualiser')
        self.assertNotIn('import_prices', page.context['form'].fields)
        self.assertEqual(self.client.get(reverse('erp:cdc-workbook-download', args=[self.dossier.pk])).status_code, 200)
        response = self.client.post(url, {'expected_version': self.dossier.version, 'file': self.upload(),
            'mode': 'merge', 'reason': 'Besoins validés', 'import_prices': 'on'})
        self.assertEqual(response.status_code, 302)
        self.item.refresh_from_db(); self.assertNotEqual(self.item.quantity, 12)
        p = CdcWorkbookPreview.objects.get(); self.assertFalse(p.import_prices)
        self.assertContains(self.client.get(response.url), 'Vérifier l’import Excel')
        self.assertEqual(self.client.post(response.url).status_code, 302)
        self.assertEqual(self.client.post(response.url).status_code, 302)
        self.item.refresh_from_db(); self.assertEqual(self.item.quantity, 12)
        self.assertEqual(self.dossier.revisions.count(), 2)
        self.assertContains(self.client.get(response.url), 'déjà été appliqué')
        self.initial.refresh_from_db(); self.assertNotEqual(self.initial.data['lot_catalog']['lots'][0]['items'][0]['quantity'], '12')

    def test_no_access_to_foreign_dossier_preview_or_costs(self):
        p = self.preview()
        with self.assertRaises(PermissionDenied): apply_workbook(self.operator, p.pk)
        with self.assertRaises(PermissionDenied): export_workbook(self.operator, self.dossier, prices=True)
        with self.assertRaises(PermissionDenied): self.preview(self.operator, import_prices=True)
        for user in (self.operator, self.second, self.outsider):
            self.client.force_login(user)
            self.assertEqual(self.client.get(reverse('erp:cdc-workbook-preview', args=[p.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse('erp:cdc-workbook', args=[self.dossier.pk])).status_code, 404)

    def test_stale_or_expired_import_is_atomic(self):
        p = self.preview()
        save_cdc_item(self.ops, self.item.lot_id, expected=self.dossier.version, pk=self.item.pk, values={'quantity': 3})
        self.client.force_login(self.ops)
        self.assertEqual(self.client.post(reverse('erp:cdc-workbook-preview', args=[p.pk])).status_code, 400)
        self.item.refresh_from_db(); self.assertEqual(self.item.quantity, 3)
        self.dossier.refresh_from_db()
        p = self.preview(); p.expires_at = timezone.now() - timedelta(seconds=1); p.save()
        with self.assertRaises(ValidationError): apply_workbook(self.ops, p.pk)
        self.assertEqual(self.dossier.revisions.count(), 2)

    def test_invalid_workbooks_never_persist_a_preview(self):
        self.client.force_login(self.ops)
        url = reverse('erp:cdc-workbook', args=[self.dossier.pk])
        self.assertEqual(self.client.post(url, {'mode': 'merge'}).status_code, 400)
        response = self.client.post(url, {'expected_version': self.dossier.version, 'file': self.upload({'F7': '=1+2'}),
            'mode': 'merge', 'reason': 'Test formule'})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(CdcWorkbookPreview.objects.exists())
        with self.assertRaises(ValidationError): self.preview(reason='')
        f = self.upload(); f.size = 21*1024*1024
        with self.assertRaises(ValidationError): self.preview(upload=f)
        other = create_dossier(self.ops, family='equipment', reference='82/SME/SDFM/SG/ESSBO/2026', title='Autre')
        with self.assertRaises(DocumentError): self.preview(upload=self.upload(family=other))

    def test_price_import_is_explicit_and_preserves_tax_and_catalog_link(self):
        self.article.specifications = 'Spécifications du catalogue'; self.article.save()
        save_cdc_item(self.ops, self.item.lot_id, expected=self.dossier.version, pk=self.item.pk,
            article=self.article, purchase_unit=self.unit, values={'quantity': 2, 'estimated_price': 25,
                'tax_rate': 9, 'price_source': 'Devis', 'currency': 'DZD'})
        self.dossier.refresh_from_db(); self.item.refresh_from_db()
        priced_revision = self.dossier.revisions.first()
        snapshot = self.item.article_snapshot
        p = self.preview(upload=self.upload({'F7': 8, 'H7': 42}), import_prices=True)
        self.client.force_login(self.ops)
        self.assertContains(self.client.get(reverse('erp:cdc-workbook-preview', args=[p.pk])), '42.00')
        apply_workbook(self.ops, p.pk)
        self.item.refresh_from_db()
        self.assertEqual(self.item.estimated_price, 42); self.assertEqual(self.item.tax_rate, 9)
        self.assertEqual(self.item.article_id, self.article.pk); self.assertEqual(self.item.article_snapshot, snapshot)
        self.dossier.refresh_from_db()
        p = self.preview(upload=self.upload({'D7': 'autre unité'}))
        with self.assertRaises(ValidationError): apply_workbook(self.ops, p.pk)
        self.assertEqual(self.dossier.revisions.count(), 3)
        restore_revision(self.ops, priced_revision.pk, expected=self.dossier.version, reason='Reprise du devis initial')
        self.item.refresh_from_db()
        self.assertEqual((self.item.quantity, self.item.estimated_price, self.item.tax_rate), (2, 25, 9))
        self.assertEqual(self.item.article_snapshot, snapshot)
        self.assertEqual(self.item.purchase_unit_id, self.unit.pk)

    def test_financial_permission_revocation_blocks_pending_preview(self):
        work = self.dossier.work
        work.allow_costs = True; work.save()
        p = self.preview(self.operator, import_prices=True, upload=self.upload({'H7': 15}))
        work.allow_costs = False; work.save()
        with self.assertRaises(PermissionDenied): apply_workbook(self.operator, p.pk)
        self.client.force_login(self.operator)
        self.assertEqual(self.client.get(reverse('erp:cdc-workbook-preview', args=[p.pk])).status_code, 403)
        self.assertEqual(self.dossier.revisions.count(), 1)

    def test_foreign_currency_not_converted_and_technical_import_keeps_prices(self):
        self.item.estimated_price = 90; self.item.price_source = 'Devis EUR'; self.item.currency = 'EUR'; self.item.save()
        with self.assertRaises(ValidationError): export_workbook(self.ops, self.dossier, prices=True)
        self.client.force_login(self.ops)
        self.assertContains(self.client.get(reverse('erp:cdc-workbook-download', args=[self.dossier.pk])+'?prices=1'), 'Export Excel impossible', status_code=400)
        p = self.preview(upload=self.upload({'H7': 120}), import_prices=True)
        with self.assertRaises(ValidationError): apply_workbook(self.ops, p.pk)
        p = self.preview(upload=self.upload({'F7': 7, 'H7': 120}))
        apply_workbook(self.ops, p.pk)
        self.item.refresh_from_db(); self.assertEqual(self.item.estimated_price, 90); self.assertEqual(self.item.currency, 'EUR')

    def test_merge_keeps_absent_rows_replace_retires_and_adds_without_deleting_history(self):
        changes = {f'{c}8': None for c in 'ABCDEFGHI'}
        changes.update({'A100': 94, 'B100': 'Nouvel appareil', 'C100': 'Spécifications vérifiables', 'D100': 'Pièce', 'F100': 2})
        p = self.preview(upload=self.upload(changes), mode='replace')
        removed = self.item.lot.items.get(position=2)
        apply_workbook(self.ops, p.pk)
        removed.refresh_from_db(); self.assertFalse(removed.active)
        self.assertTrue(CdcItem.objects.filter(lot=self.item.lot, designation='Nouvel appareil', active=True).exists())
        self.assertEqual(self.dossier.revisions.count(), 2)
        self.assertTrue(self.initial.data['lot_catalog']['lots'][0]['items'])

    def test_reassignment_and_submission_revoke_pending_import(self):
        p = self.preview(self.operator)
        work = self.dossier.work; work.assignee = self.second; work.save()
        with self.assertRaises(PermissionDenied): apply_workbook(self.operator, p.pk)
        work.status = 'SUBMITTED'; work.save()
        with self.assertRaises(PermissionDenied): self.preview()

    def test_lot_reuse_retirement_restoration_and_revision_recovery(self):
        source = self.item.lot
        original = document_data(self.dossier)
        lot = add_lot(self.ops, self.dossier.pk, expected=self.dossier.version, name='Lot complémentaire',
            name_ar='مجموعة إضافية', source=source, reason='Réutilisation technique')
        self.assertEqual(lot.items.count(), source.items.count())
        self.assertTrue(all(i.source_key.startswith('new-') for i in lot.items.all()))
        self.dossier.refresh_from_db()
        set_lot_active(self.ops, lot.pk, expected=self.dossier.version, active=False, reason='Retrait motivé')
        self.dossier.refresh_from_db()
        self.assertEqual(document_data(self.dossier), original)
        restore_revision(self.ops, self.initial.pk, expected=self.dossier.version, reason='Reprise initiale')
        self.dossier.refresh_from_db()
        self.assertEqual(document_data(self.dossier), original)
        self.assertEqual(self.dossier.revisions.count(), 4)
        lot.refresh_from_db()
        self.assertEqual(lot.position, 3)
        with self.assertRaises(PermissionDenied): restore_revision(self.operator, self.initial.pk, expected=self.dossier.version, reason='Non autorisé')
        with self.assertRaises(ValidationError): restore_revision(self.ops, self.initial.pk, expected=self.dossier.version, reason='')
        with self.assertRaises(ValidationError): set_lot_active(self.ops, lot.pk, expected=self.dossier.version, active=True, reason='')

    def test_lot_forms_and_restore_are_reachable_and_reject_invalid_operations(self):
        self.client.force_login(self.ops)
        add = reverse('erp:cdc-lot-create', args=[self.dossier.pk])
        self.assertEqual(self.client.get(add).status_code, 200)
        self.assertEqual(self.client.post(add, {}).status_code, 400)
        self.assertEqual(self.client.post(add, {'expected_version': 999, 'name': 'Périmé', 'reason': 'Conflit'}).status_code, 400)
        response = self.client.post(add, {'expected_version': self.dossier.version, 'name': 'Nouveau lot',
            'name_ar': '', 'source': '', 'reason': 'Besoins additionnels'})
        self.assertEqual(response.status_code, 302)
        lot = self.dossier.lots.get(name='Nouveau lot'); self.dossier.refresh_from_db()
        toggle = reverse('erp:cdc-lot-toggle', args=[lot.pk])
        self.assertEqual(self.client.get(toggle).status_code, 200)
        self.assertEqual(self.client.post(toggle, {'expected_version': 999, 'confirm': 'on', 'reason': 'Périmé'}).status_code, 400)
        self.assertEqual(self.client.post(toggle, {'expected_version': self.dossier.version, 'confirm': 'on', 'reason': 'Retrait'}).status_code, 302)
        self.dossier.refresh_from_db()
        restore = reverse('erp:cdc-revision-restore', args=[self.initial.pk])
        self.assertEqual(self.client.get(restore).status_code, 200)
        self.assertEqual(self.client.post(restore, {'expected_version': 999, 'confirm': 'on', 'reason': 'Périmé'}).status_code, 400)
        self.assertEqual(self.client.post(restore, {'expected_version': self.dossier.version, 'confirm': 'on', 'reason': 'Reprise'}).status_code, 302)

    def test_workbook_exports_blank_and_prices_and_all_families_import(self):
        self.client.force_login(self.ops)
        for family in ('equipment', 'reagents', 'works'):
            with self.subTest(family=family):
                d = create_dossier(self.ops, family=family, reference=f'{90+len(family)}/SME/SDFM/SG/ESSBO/2026', title=family)
                self.assertEqual(self.client.get(reverse('erp:cdc-workbook-download', args=[d.pk])+'?blank=1').status_code, 200)
                export_workbook(self.ops, d, prices=True)
                p = preview_workbook(self.ops, d.pk, expected=d.version, upload=self.upload(family=d), mode='merge', reason='Besoin')
                apply_workbook(self.ops, p.pk)
                self.assertEqual(d.lots.first().items.first().quantity, 12)
        works = d
        works.refresh_from_db()
        with self.assertRaises(ValidationError): add_lot(self.ops, works.pk, expected=works.version, name='Interdit')
        with self.assertRaises(ValidationError): add_lot(self.ops, self.dossier.pk, expected=self.dossier.version, name='Autre', source=works.lots.first())

    def test_stock_search_includes_supplier_manufacturer_and_location_without_scope_leaks(self):
        container, _, _ = self.receive('SEARCH')
        self.party.name = 'Fournisseur recherche unique'; self.party.save()
        self.freezer.name = 'Congélateur recherche unique'; self.freezer.save()
        self.client.force_login(self.ops)
        for q in ('Fournisseur recherche unique', 'Congélateur recherche unique', container.lot.manufacturer_lot):
            response = self.client.get(reverse('erp:stock-list'), {'q': q})
            self.assertContains(response, container.code)
        self.client.force_login(self.second)
        self.assertEqual(self.client.get(reverse('erp:stock-list'), {'q': self.party.name}).status_code, 403)

    def test_imported_quantities_reach_real_docx_for_each_family(self):
        from erp.cdc.catalog import generate_document
        from erp.cdc.docengine import Document
        for family, number in (('equipment', 201), ('reagents', 202), ('works', 203)):
            d = create_dossier(self.ops, family=family, reference=f'{number}/SME/SDFM/SG/ESSBO/2026', title=family)
            p = preview_workbook(self.ops, d.pk, expected=d.version, upload=self.upload({'F7': 123456}, family=d),
                mode='merge', reason='Quantité du besoin annuel')
            revision = apply_workbook(self.ops, p.pk)
            output, report = generate_document(revision.data)
            text = ' '.join(block['text'] for block in Document(output).source_index)
            self.assertIn('123456', text.replace(' ', ''))
            self.assertEqual(report['lot_catalog']['status'], 'GENERATED')

    def test_restore_handles_plan_seed_revision_whose_live_rows_were_replaced(self):
        CdcItem.objects.filter(lot__dossier=self.dossier).delete()
        self.dossier.lots.all().delete()
        restore_revision(self.ops, self.initial.pk, expected=self.dossier.version, reason='Reprise du modèle initial')
        self.dossier.refresh_from_db()
        self.assertEqual(document_data(self.dossier), self.initial.data)

    def test_lot_binding_is_editable_and_removed_lot_is_excluded_from_totals(self):
        from erp.services.cdc import estimate_totals, save_cdc_lot
        lot = self.item.lot
        save_cdc_lot(self.ops, lot.pk, expected=self.dossier.version, name=lot.name, name_ar=lot.name_ar,
            source_slot=1, reason='Correspondance documentaire confirmée')
        self.dossier.refresh_from_db()
        self.item.estimated_price = 15; self.item.tax_rate = 0; self.item.price_source = 'Devis'; self.item.save()
        self.assertIn('DZD', estimate_totals(self.ops, self.dossier)['currencies'])
        set_lot_active(self.ops, lot.pk, expected=self.dossier.version, active=False, reason='Retrait')
        self.assertNotIn('DZD', estimate_totals(self.ops, self.dossier)['currencies'])
        self.dossier.refresh_from_db()
        set_lot_active(self.ops, lot.pk, expected=self.dossier.version, active=True, reason='Réintégration')
        self.assertIn('DZD', estimate_totals(self.ops, self.dossier)['currencies'])

    def test_item_search_duplicate_move_and_recovery_preserve_history(self):
        from erp.services.cdc_exchange import arrange_item
        self.client.force_login(self.ops)
        url = reverse('erp:cdc-lot', args=[self.item.lot_id])
        self.assertEqual(self.client.get(url, {'q': 'introuvable-unique'}).context['page'].paginator.count, 0)
        self.assertGreater(self.client.get(url, {'q': self.item.designation[:12]}).context['page'].paginator.count, 0)
        self.item.estimated_price = 20; self.item.tax_rate = 9; self.item.price_source = 'Devis'; self.item.save()
        original_id = self.item.pk
        target = self.dossier.lots.last()
        duplicate = arrange_item(self.operator, self.item.pk, expected=self.dossier.version,
            destination=target, action='duplicate', position=1, reason='Besoin supplémentaire')
        self.assertIsNone(duplicate.estimated_price)
        self.assertIsNone(duplicate.tax_rate)
        self.assertNotEqual(duplicate.source_key, self.item.source_key)
        self.item.refresh_from_db(); self.assertTrue(self.item.active)
        self.dossier.refresh_from_db()
        moved = arrange_item(self.operator, self.item.pk, expected=self.dossier.version,
            destination=target, action='move', position=2, reason='Répartition des besoins')
        self.item.refresh_from_db(); self.assertFalse(self.item.active)
        self.assertEqual(moved.estimated_price, 20)
        self.assertEqual(target.items.filter(active=True).order_by('position')[1], moved)
        self.dossier.refresh_from_db()
        restore_revision(self.ops, self.initial.pk, expected=self.dossier.version, reason='Annulation des réorganisations')
        self.item.refresh_from_db(); moved.refresh_from_db(); duplicate.refresh_from_db()
        self.assertEqual(self.item.pk, original_id)
        self.assertTrue(self.item.active)
        self.assertFalse(moved.active); self.assertFalse(duplicate.active)
        self.assertEqual(self.dossier.revisions.count(), 4)

    def test_item_arrangement_http_and_invalid_destinations_are_atomic(self):
        from erp.services.cdc_exchange import arrange_item
        self.client.force_login(self.operator)
        url = reverse('erp:cdc-item-arrange', args=[self.item.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.post(url, {}).status_code, 400)
        values = {'expected_version': 999, 'action': 'move', 'destination': self.item.lot_id,
            'position': 2, 'reason': 'Réordonner', 'confirm': 'on'}
        self.assertEqual(self.client.post(url, values).status_code, 400)
        values['expected_version'] = self.dossier.version
        self.assertEqual(self.client.post(url, values).status_code, 302)
        self.item.refresh_from_db(); self.assertFalse(self.item.active)
        self.dossier.refresh_from_db()
        with self.assertRaises(ValidationError):
            arrange_item(self.ops, self.item.pk, expected=self.dossier.version, destination=self.item.lot,
                action='move', position=1, reason='Article retiré')
        self.assertEqual(self.client.get(url).status_code, 404)
        current = self.item.lot.items.filter(active=True).first()
        for action, reason, position in [('invalid', 'Motif', 1), ('move', '', 1), ('move', 'Motif', 9999)]:
            with self.assertRaises(ValidationError):
                arrange_item(self.ops, current.pk, expected=self.dossier.version, destination=current.lot,
                    action=action, position=position, reason=reason)
        other = create_dossier(self.ops, family='equipment', reference='650/SME/SDFM/SG/ESSBO/2026', title='Autre périmètre')
        with self.assertRaises(ValidationError):
            arrange_item(self.ops, current.pk, expected=self.dossier.version, destination=other.lots.first(),
                action='move', position=1, reason='Lot étranger')
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse('erp:cdc-item-arrange', args=[current.pk])).status_code, 404)
        self.assertEqual(self.dossier.revisions.count(), 2)

    def test_unassigned_dossiers_remain_visible_across_planning_and_stock_workflows(self):
        import uuid
        from django.utils.translation import override
        from erp.services.inventory import create_inventory
        from erp.services.procurement import create_plan
        from erp.services.planning import create_activity
        work = self.dossier.work; work.assignee = None; work.save()
        self.receive('UNASSIGNED')
        campaign = create_inventory(self.ops, title='Inventaire sans responsable')
        plan = create_plan(self.ops, reference='UNASSIGNED', year=timezone.localdate().year, title='Plan sans responsable')
        start = timezone.now() - timedelta(minutes=5)
        schedule = create_activity(self.ops, key=uuid.uuid4(), kind='CONTROL', title='Contrôle sans responsable',
            assignee=None, starts_at=start, ends_at=start+timedelta(hours=1))
        self.client.force_login(self.ops)
        routes = [('erp:planning', []), ('erp:work-list', []), ('erp:work-detail', [work.pk]),
            ('erp:activity-detail', [work.pk]), ('erp:activity-detail', [schedule.work_id]),
            ('erp:cdc-list', []), ('erp:inventory-list', []), ('erp:inventory-detail', [campaign.pk]),
            ('erp:procurement-list', []), ('erp:procurement-detail', [plan.pk])]
        for language in ('fr', 'en', 'ar'):
            with override(language):
                for name, args in routes:
                    with self.subTest(language=language, page=name):
                        self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 200)
        work.status = 'SUBMITTED'; work.submitted_at = timezone.now(); work.save()
        self.assertContains(self.client.get(reverse('erp:planning')), 'Non affecté')
