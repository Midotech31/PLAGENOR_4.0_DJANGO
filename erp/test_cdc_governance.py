from decimal import Decimal
import io

from openpyxl import load_workbook

from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse

from erp.cdc.catalog import document
from erp.cdc.schedule_adapter import managed_ids
from erp.models import (CdcClause, CdcClauseSelection, CdcClauseVersion, CdcCriterion,
                        CdcDossier, CdcItem, CdcRequirement, CdcReviewDecision, WorkItem)
from erp.services.cdc import (create_dossier, document_data, duplicate_dossier,
    edit_cdc_paragraph, save_cdc_item)
from erp.services.cdc_exchange import add_lot, arrange_item, restore_revision
from erp.services.cdc_governance import (apply_clause_selections, criteria_findings,
    governance_snapshot, publish_clause, review_revision, review_summary, save_criterion,
    save_requirement, select_clause)
from erp.services.work import create_work
from erp.test_operations import OperationFixtures
from notifications.models import Notification


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],
    SECURE_SSL_REDIRECT=False,
    STORAGES={'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
              'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class CdcNativeGovernanceTests(OperationFixtures, TestCase):
    def setUp(self):
        self.dossier = create_dossier(self.ops, family='equipment',
            reference='81/SME/SDFM/SG/ESSBO/2026', title='CDC gouvernance',
            assignee=self.operator, allow_costs=True)
        managed, _ = managed_ids(self.dossier.family)
        self.block = next(row for row in document(self.dossier.family).source_index
            if row['text'].strip() and not row['guard'] and row['id'] not in managed)

    def refresh(self):
        self.dossier.refresh_from_db()
        self.dossier.work.refresh_from_db()
        return self.dossier

    def criterion_values(self, **changes):
        values = {
            'lot': self.dossier.lots.first(), 'code': 'TECH-01',
            'category': CdcCriterion.Category.TECHNICAL, 'title': 'Conformité technique',
            'description': 'Évaluer les performances annoncées',
            'expected_evidence': 'Fiche technique vérifiable',
            'min_score': Decimal('0'), 'max_score': Decimal('20'),
            'weight': Decimal('100'), 'threshold': Decimal('10'),
            'formula': '', 'rounding_rule': 'Deux décimales', 'eliminatory': False,
            'source': 'Cahier des charges institutionnel', 'justification': 'Critère requis',
            'position': 1, 'active': True,
        }
        values.update(changes)
        return values


    def requirement_values(self, **changes):
        values = {
            'item': self.dossier.lots.first().items.filter(active=True).first(),
            'code': 'REQ-01', 'kind': CdcRequirement.Kind.MANDATORY,
            'statement': 'Performance minimale vérifiable',
            'evidence': 'Fiche technique du fabricant',
            'verification': 'Contrôle documentaire à la réception',
            'justification': 'Besoin technique de la plateforme',
            'position': 1, 'active': True,
        }
        values.update(changes)
        return values

    def test_clause_library_versions_are_canonical_and_snapshotted(self):
        with self.assertRaisesRegex(ValidationError, 'titre, la source'):
            publish_clause(self.ops, self.dossier, expected=self.dossier.version,
                paragraph_id=self.block['id'], title='', body=self.block['text'],
                source='', reason='')
        with self.assertRaisesRegex(ValidationError, 'inconnu'):
            publish_clause(self.ops, self.dossier, expected=self.dossier.version,
                paragraph_id='missing', title='Clause', body='Texte', source='Source', reason='Motif')
        managed, _ = managed_ids(self.dossier.family)
        protected = next(row for row in document(self.dossier.family).source_index
            if row['guard'] or row['id'] in managed)
        with self.assertRaisesRegex(ValidationError, 'donnée structurée'):
            publish_clause(self.ops, self.dossier, expected=self.dossier.version,
                paragraph_id=protected['id'], title='Clause', body=protected['text'],
                source='Source', reason='Motif')
        with self.assertRaisesRegex(ValidationError, 'vide'):
            publish_clause(self.ops, self.dossier, expected=self.dossier.version,
                paragraph_id=self.block['id'], title='Clause', body=' ', source='Source', reason='Motif')
        with self.assertRaises(PermissionDenied):
            publish_clause(self.operator, self.dossier, expected=self.dossier.version,
                paragraph_id=self.block['id'], title='Clause', body=self.block['text'],
                source='Source', reason='Motif')

        clause, version1, revision1 = publish_clause(self.ops, self.dossier,
            expected=self.dossier.version, paragraph_id=self.block['id'], title='Clause qualité',
            category='Qualité', body='Texte canonique version 1', source='Référence ESSBO',
            mandatory=True, reason='Création de la clause commune')
        self.refresh()
        self.assertEqual(clause.current_version_id, version1.pk)
        self.assertEqual(document_data(self.dossier)['paragraphs'][self.block['id']], 'Texte canonique version 1')
        self.assertEqual(revision1.governance['clauses'][0]['sha256'], version1.sha256)
        with self.assertRaisesRegex(ValidationError, 'bibliothèque'):
            edit_cdc_paragraph(self.ops, self.dossier.pk, expected=self.dossier.version,
                paragraph_id=self.block['id'], value='Texte concurrent', reason='Interdit')

        clause2, version2, revision2 = publish_clause(self.ops, self.dossier,
            expected=self.dossier.version, paragraph_id=self.block['id'], title='Clause qualité',
            category='Qualité', body='Texte canonique version 2', source='Référence ESSBO v2',
            mandatory=True, reason='Mise à jour validée')
        self.assertEqual(clause.pk, clause2.pk)
        self.assertEqual(version2.number, 2)
        self.assertEqual(revision1.data['paragraphs'][self.block['id']], 'Texte canonique version 1')
        self.assertEqual(revision2.data['paragraphs'][self.block['id']], 'Texte canonique version 2')

        other = CdcClause.objects.create(family='works', code='OTHER', paragraph_id='other',
            title='Autre', created_by=self.ops)
        other_version = CdcClauseVersion.objects.create(clause=other, number=1, body='x',
            source='x', actor=self.ops, sha256='0' * 64)
        with self.assertRaisesRegex(ValidationError, 'ne peut pas'):
            select_clause(self.ops, self.dossier, expected=self.dossier.version,
                clause=other, version=other_version, reason='Mauvaise famille')
        with self.assertRaisesRegex(ValidationError, 'Justifiez'):
            select_clause(self.ops, self.dossier, expected=self.dossier.version,
                clause=clause, version=version1, reason='')
        self.refresh()
        restored = select_clause(self.ops, self.dossier, expected=self.dossier.version,
            clause=clause, version=version1, reason='Retour à la version approuvée')
        self.assertEqual(restored.data['paragraphs'][self.block['id']], 'Texte canonique version 1')

    def test_clause_selection_integrity_guard_rejects_wrong_version(self):
        clause, version, _ = publish_clause(self.ops, self.dossier, expected=self.dossier.version,
            paragraph_id=self.block['id'], title='Clause', body='Version', source='Source', reason='Motif')
        second = CdcClause.objects.create(family='equipment', code='SECOND', paragraph_id='second',
            title='Seconde', created_by=self.ops)
        wrong = CdcClauseVersion.objects.create(clause=second, number=1, body='Autre',
            source='Source', actor=self.ops, sha256='1' * 64)
        selection = CdcClauseSelection.objects.get(dossier=self.dossier, clause=clause)
        selection.selected_version = wrong
        selection.save()
        with self.assertRaisesRegex(ValidationError, 'liaison incohérente'):
            apply_clause_selections(self.dossier, self.dossier.data.copy())
        selection.selected_version = version
        selection.save()

    def test_criteria_are_versioned_checked_and_do_not_duplicate_catalog(self):
        with self.assertRaisesRegex(ValidationError, 'Justifiez'):
            save_criterion(self.ops, self.dossier, expected=self.dossier.version,
                values=self.criterion_values(), reason='')
        with self.assertRaisesRegex(ValidationError, 'note maximale ou une formule'):
            save_criterion(self.ops, self.dossier, expected=self.dossier.version,
                values=self.criterion_values(max_score=None, formula='', weight=Decimal('50')), reason='Test')
        with self.assertRaisesRegex(ValidationError, 'seuil'):
            save_criterion(self.ops, self.dossier, expected=self.dossier.version,
                values=self.criterion_values(max_score=Decimal('10'), threshold=Decimal('11')), reason='Test')
        other = create_dossier(self.ops, family='equipment',
            reference='82/SME/SDFM/SG/ESSBO/2026', title='Autre CDC')
        with self.assertRaisesRegex(ValidationError, 'lot'):
            save_criterion(self.ops, self.dossier, expected=self.dossier.version,
                values=self.criterion_values(lot=other.lots.first()), reason='Test')

        criterion, revision = save_criterion(self.operator, self.dossier,
            expected=self.dossier.version, values=self.criterion_values(), reason='Grille technique')
        self.refresh()
        self.assertEqual(revision.governance['criteria'][0]['code'], 'TECH-01')
        self.assertEqual(criterion.lot.dossier_id, self.dossier.pk)
        self.assertEqual(criteria_findings(self.dossier), [])

        _, revision2 = save_criterion(self.operator, self.dossier, expected=self.dossier.version,
            pk=criterion.pk, values=self.criterion_values(weight=Decimal('60')),
            reason='Rééquilibrage')
        self.refresh()
        findings = criteria_findings(self.dossier)
        self.assertTrue(any(row['id'] == 'CRITERIA_WEIGHT_TOTAL' for row in findings))
        self.assertEqual(revision.governance['criteria'][0]['weight'], '100.00')
        self.assertEqual(revision2.governance['criteria'][0]['weight'], '60.00')

        criterion.max_score = Decimal('10')
        criterion.formula = ''
        criterion.source = ''
        criterion.threshold = Decimal('999')
        criterion.save()
        ids = {row['id'] for row in criteria_findings(self.dossier)}
        self.assertTrue({'CRITERION_THRESHOLD', 'CRITERION_SOURCE'} <= ids)
        criterion.max_score = None
        criterion.save(update_fields=['max_score'])
        self.assertTrue(any(row['id'] == 'CRITERION_METHOD' for row in criteria_findings(self.dossier)))

    def test_review_sequence_permissions_notifications_and_correction_loop(self):
        revision = self.dossier.revisions.get(number=self.dossier.revision_number)
        with self.assertRaisesRegex(ValidationError, 'soumis'):
            review_revision(self.operator, revision, stage='TECHNICAL', decision='APPROVED', comment='OK')
        self.dossier.work.status = WorkItem.Status.SUBMITTED
        self.dossier.work.save(update_fields=['status'])
        with self.assertRaisesRegex(ValidationError, 'Étape'):
            review_revision(self.ops, revision, stage='UNKNOWN', decision='APPROVED', comment='OK')
        with self.assertRaisesRegex(ValidationError, 'Décision'):
            review_revision(self.operator, revision, stage='TECHNICAL', decision='UNKNOWN', comment='OK')
        with self.assertRaisesRegex(ValidationError, 'compte rendu'):
            review_revision(self.operator, revision, stage='TECHNICAL', decision='APPROVED', comment='')
        with self.assertRaisesRegex(ValidationError, 'étapes précédentes'):
            review_revision(self.ops, revision, stage='FINANCIAL', decision='APPROVED', comment='Trop tôt')
        review_revision(self.operator, revision, stage='TECHNICAL', decision='APPROVED', comment='Technique conforme')
        with self.assertRaisesRegex(ValidationError, 'déjà'):
            review_revision(self.operator, revision, stage='TECHNICAL', decision='APPROVED', comment='Bis')
        with self.assertRaises(PermissionDenied):
            review_revision(self.operator, revision, stage='ADMIN', decision='APPROVED', comment='Interdit')
        review_revision(self.ops, revision, stage='ADMIN', decision='APPROVED', comment='Administratif conforme')
        result = review_revision(self.ops, revision, stage='FINANCIAL', decision='CHANGES', comment='Corriger le financement')
        self.dossier.work.refresh_from_db()
        self.assertEqual(result.decision, 'CHANGES')
        self.assertEqual(self.dossier.work.status, WorkItem.Status.CHANGES_REQUESTED)
        self.assertTrue(Notification.objects.filter(user=self.operator, link_url=reverse('erp:cdc-detail', args=[self.dossier.pk])).exists())
        self.assertEqual(len(review_summary(self.dossier)), 3)
        empty_work = create_work(self.ops, kind=WorkItem.Kind.CDC, title='Résumé vide')
        empty = CdcDossier.objects.create(work=empty_work, family='works',
            reference='83/SME/SDFM/SG/ESSBO/2026', data={})
        self.assertEqual(review_summary(empty), [])

        stale = revision
        self.dossier.work.status = WorkItem.Status.IN_PROGRESS
        self.dossier.work.save(update_fields=['status'])
        criterion, new_revision = save_criterion(self.operator, self.dossier,
            expected=self.dossier.version, pk=criterion.pk, values=self.criterion_values(), reason='Nouvelle révision')
        self.assertNotEqual(stale.number, new_revision.number)
        self.dossier.work.status = WorkItem.Status.SUBMITTED
        self.dossier.work.save(update_fields=['status'])
        with self.assertRaisesRegex(ValidationError, 'révision courante'):
            review_revision(self.ops, stale, stage='TECHNICAL', decision='APPROVED', comment='Ancienne')

    def test_supplier_is_canonical_cost_data_and_is_not_erased_by_technical_edit(self):
        lot = self.dossier.lots.first()
        item = lot.items.first()
        save_cdc_item(self.ops, lot.pk, pk=item.pk, expected=self.dossier.version,
            values={'estimated_price': Decimal('100'), 'tax_rate': Decimal('19'),
                    'price_source': 'Devis fournisseur', 'currency': 'DZD'},
            supplier=self.party, supplier_provided=True, reason='Estimation')
        self.refresh()
        item.refresh_from_db()
        self.assertEqual(item.supplier_id, self.party.pk)
        with self.assertRaises(PermissionDenied):
            save_cdc_item(self.operator, lot.pk, pk=item.pk, expected=self.dossier.version,
                values={'quantity': item.quantity}, supplier=self.party, supplier_provided=True,
                reason='Tentative financière')
        save_cdc_item(self.operator, lot.pk, pk=item.pk, expected=self.dossier.version,
            values={'quantity': item.quantity}, reason='Correction technique')
        item.refresh_from_db()
        self.assertEqual(item.supplier_id, self.party.pk)

    def test_native_http_surfaces_persist_governance(self):
        self.client.force_login(self.ops)
        criteria_url = reverse('erp:cdc-criteria', args=[self.dossier.pk])
        self.assertEqual(self.client.get(criteria_url).status_code, 200)
        create_url = reverse('erp:cdc-criterion-new', args=[self.dossier.pk])
        self.assertEqual(self.client.get(create_url).status_code, 200)
        payload = {key: (value.pk if hasattr(value, 'pk') else value) for key, value in self.criterion_values().items()}
        payload.update(expected_version=self.dossier.version, reason='Création HTTP')
        response = self.client.post(create_url, payload)
        self.assertEqual(response.status_code, 302, response.context['form'].errors if response.context else '')
        self.refresh()
        criterion = self.dossier.criteria.get(code='TECH-01')
        edit_url = reverse('erp:cdc-criterion-edit', args=[self.dossier.pk, criterion.pk])
        bad = dict(payload); bad['expected_version'] = self.dossier.version + 10
        self.assertEqual(self.client.post(edit_url, bad).status_code, 400)

        publish_url = reverse('erp:cdc-clause-publish', args=[self.dossier.pk])
        self.assertEqual(self.client.get(publish_url + '?paragraph_id=' + self.block['id']).status_code, 200)
        clause_payload = {'expected_version': self.dossier.version, 'paragraph_id': self.block['id'],
            'title': 'Clause HTTP', 'category': 'Qualité', 'body': self.block['text'],
            'source': 'Référence ESSBO', 'mandatory': 'on', 'reason': 'Version canonique'}
        self.assertEqual(self.client.post(publish_url, clause_payload).status_code, 302)
        self.refresh()
        self.assertTrue(self.dossier.clause_selections.exists())
        clauses_page = self.client.get(reverse('erp:cdc-clauses', args=[self.dossier.pk]))
        self.assertContains(clauses_page, 'bibliothèque versionnée')
        selection = self.dossier.clause_selections.select_related('clause', 'selected_version').get()
        select_url = reverse('erp:cdc-clause-select', args=[self.dossier.pk, selection.clause_id])
        self.assertEqual(self.client.get(select_url).status_code, 200)
        self.assertEqual(self.client.post(select_url, {'expected_version': self.dossier.version,
            'version': selection.selected_version_id, 'reason': 'Version confirmée'}).status_code, 302)
        self.refresh()

        export = self.client.get(reverse('erp:cdc-criteria-export', args=[self.dossier.pk]))
        self.assertEqual(export.status_code, 200)
        payload = b''.join(export.streaming_content)
        book = load_workbook(io.BytesIO(payload), data_only=True)
        self.assertEqual(book['Critères']['A2'].value, 'TECH-01')
        self.assertEqual(book['Traçabilité']['B1'].value, self.dossier.reference)

        self.dossier.work.status = WorkItem.Status.SUBMITTED
        self.dossier.work.save(update_fields=['status'])
        review_url = reverse('erp:cdc-review', args=[self.dossier.pk])
        self.assertEqual(self.client.get(review_url).status_code, 200)
        response = self.client.post(review_url, {'stage': 'TECHNICAL', 'decision': 'APPROVED',
            'comment': 'Revue HTTP'})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(CdcReviewDecision.objects.filter(revision__dossier=self.dossier, stage='TECHNICAL').exists())


    def test_duplication_and_revision_restore_preserve_governance_without_duplicate_references(self):
        lot = self.dossier.lots.first()
        item = lot.items.first()
        save_cdc_item(self.ops, lot.pk, pk=item.pk, expected=self.dossier.version,
            values={'estimated_price': Decimal('25'), 'tax_rate': Decimal('19'),
                    'price_source': 'Devis A', 'currency': 'DZD'},
            supplier=self.party, supplier_provided=True, reason='Prix et fournisseur')
        self.refresh()
        clause, version1, _ = publish_clause(self.ops, self.dossier, expected=self.dossier.version,
            paragraph_id=self.block['id'], title='Clause de reprise', body='Texte historique',
            source='Source validée', reason='Clause initiale')
        self.refresh()
        criterion, source_revision = save_criterion(self.ops, self.dossier, expected=self.dossier.version,
            values=self.criterion_values(), reason='Grille initiale')
        self.refresh()

        duplicate = duplicate_dossier(self.ops, self.dossier.pk, expected=self.dossier.version,
            reference='84/SME/SDFM/SG/ESSBO/2026', title='Copie gouvernée',
            assignee=self.second, allow_costs=True, copy_estimates=False,
            reason='Nouvelle procédure fondée sur le modèle validé')
        copied_selection = duplicate.clause_selections.get()
        self.assertEqual(copied_selection.clause_id, clause.pk)
        self.assertEqual(copied_selection.selected_version_id, version1.pk)
        self.assertEqual(duplicate.criteria.get(code='TECH-01').weight, Decimal('100'))
        self.assertIsNone(duplicate.lots.filter(active=True).first().items.filter(active=True).first().supplier_id)

        _, version2, _ = publish_clause(self.ops, self.dossier, expected=self.dossier.version,
            paragraph_id=self.block['id'], title='Clause de reprise', body='Texte modifié',
            source='Source validée v2', reason='Nouvelle version')
        self.refresh()
        save_criterion(self.ops, self.dossier, expected=self.dossier.version, pk=criterion.pk,
            values=self.criterion_values(weight=Decimal('60')), reason='Modification provisoire')
        self.refresh()
        item.refresh_from_db()
        save_cdc_item(self.ops, item.lot_id, pk=item.pk, expected=self.dossier.version,
            values={}, supplier=None, supplier_provided=True, reason='Fournisseur retiré provisoirement')
        self.refresh()
        self.assertEqual(self.dossier.clause_selections.get().selected_version_id, version2.pk)
        self.assertIsNone(CdcItem.objects.get(pk=item.pk).supplier_id)

        restored = restore_revision(self.ops, source_revision.pk, expected=self.dossier.version,
            reason='Retour contrôlé à la version antérieure')
        self.refresh()
        self.assertEqual(restored.data['paragraphs'][self.block['id']], 'Texte historique')
        self.assertEqual(self.dossier.clause_selections.get().selected_version_id, version1.pk)
        self.assertEqual(self.dossier.criteria.get(code='TECH-01').weight, Decimal('100'))
        self.assertEqual(CdcItem.objects.get(pk=item.pk).supplier_id, self.party.pk)


    def test_structured_requirements_validate_generate_snapshot_and_score_linkage(self):
        with self.assertRaisesRegex(ValidationError, 'Justifiez'):
            save_requirement(self.ops, self.dossier, expected=self.dossier.version,
                values=self.requirement_values(), reason='')
        with self.assertRaisesRegex(ValidationError, 'preuve attendue'):
            save_requirement(self.ops, self.dossier, expected=self.dossier.version,
                values=self.requirement_values(evidence=''), reason='Test incomplet')
        with self.assertRaisesRegex(ValidationError, 'éliminatoire'):
            save_requirement(self.ops, self.dossier, expected=self.dossier.version,
                values=self.requirement_values(kind=CdcRequirement.Kind.ELIMINATORY,
                    justification=''), reason='Test éliminatoire')
        other = create_dossier(self.ops, family='equipment',
            reference='85/SME/SDFM/SG/ESSBO/2026', title='Autre CDC')
        with self.assertRaisesRegex(ValidationError, 'article actif'):
            save_requirement(self.ops, self.dossier, expected=self.dossier.version,
                values=self.requirement_values(item=other.lots.first().items.first()),
                reason='Mauvais dossier')

        requirement, revision = save_requirement(self.operator, self.dossier,
            expected=self.dossier.version, values=self.requirement_values(),
            reason='Exigence vérifiable')
        self.refresh()
        self.assertEqual(revision.governance['requirements'][0]['code'], 'REQ-01')
        self.assertIn('Performance minimale vérifiable',
            document_data(self.dossier)['lot_catalog']['lots'][0]['items'][0]['specifications'])
        self.assertFalse(any(row['id'].startswith('REQUIREMENT_') for row in criteria_findings(self.dossier)))

        scored_values = self.requirement_values(code='REQ-SCORE', kind=CdcRequirement.Kind.SCORED,
            statement='Performance notée')
        scored, _ = save_requirement(self.operator, self.dossier, expected=self.dossier.version,
            values=scored_values, reason='Exigence notée')
        self.refresh()
        self.assertTrue(any(row['id'] == 'REQUIREMENT_SCORING' for row in criteria_findings(self.dossier)))
        criterion_values = self.criterion_values(requirement=scored)
        criterion, criterion_revision = save_criterion(self.operator, self.dossier,
            expected=self.dossier.version, values=criterion_values, reason='Lien de notation')
        self.refresh()
        self.assertFalse(any(row['id'] == 'REQUIREMENT_SCORING' for row in criteria_findings(self.dossier)))
        self.assertEqual(criterion_revision.governance['criteria'][0]['requirement_code'], 'REQ-SCORE')

        updated, _ = save_requirement(self.operator, self.dossier, expected=self.dossier.version,
            pk=requirement.pk, values=self.requirement_values(statement='Performance minimale révisée'),
            reason='Clarification technique')
        self.assertEqual(updated.version, 2)

    def test_criteria_scope_and_requirement_scope_are_strict(self):
        requirement, _ = save_requirement(self.operator, self.dossier,
            expected=self.dossier.version, values=self.requirement_values(), reason='Exigence locale')
        self.refresh()
        _, _ = save_criterion(self.operator, self.dossier, expected=self.dossier.version,
            values=self.criterion_values(requirement=requirement), reason='Critère local')
        self.refresh()
        global_values = self.criterion_values(lot=None, requirement=None, code='GLOBAL-01',
            title='Critère global', weight=Decimal('100'))
        _, _ = save_criterion(self.operator, self.dossier, expected=self.dossier.version,
            values=global_values, reason='Critère global')
        self.refresh()
        self.assertTrue(any(row['id'] == 'CRITERIA_SCOPE_MIXED' for row in criteria_findings(self.dossier)))

        second = create_dossier(self.ops, family='equipment',
            reference='86/SME/SDFM/SG/ESSBO/2026', title='CDC hors périmètre')
        foreign_req, _ = save_requirement(self.ops, second, expected=second.version,
            values={**self.requirement_values(), 'item': second.lots.first().items.first(),
                    'code': 'FOREIGN'}, reason='Exigence externe')
        self.refresh()
        with self.assertRaisesRegex(ValidationError, 'appartenir à ce cahier'):
            save_criterion(self.operator, self.dossier, expected=self.dossier.version,
                values=self.criterion_values(code='BAD-REQ', requirement=foreign_req),
                reason='Lien invalide')

    def test_requirements_follow_lot_reuse_item_duplication_and_controlled_dossier_duplication(self):
        requirement, _ = save_requirement(self.operator, self.dossier,
            expected=self.dossier.version, values=self.requirement_values(), reason='Exigence source')
        self.refresh()
        source_lot = self.dossier.lots.first()
        reused = add_lot(self.operator, self.dossier.pk, expected=self.dossier.version,
            name='Lot réutilisé', source=source_lot, reason='Réutilisation technique')
        self.refresh()
        reused_item = reused.items.filter(active=True).first()
        self.assertEqual(reused_item.requirements.get().statement, requirement.statement)

        duplicated_item = arrange_item(self.operator, source_lot.items.filter(active=True).first().pk,
            expected=self.dossier.version, destination=reused, action='duplicate',
            position=reused.items.filter(active=True).count() + 1, reason='Duplication contrôlée')
        self.refresh()
        self.assertEqual(duplicated_item.requirements.get().code, 'REQ-01')

        duplicate = duplicate_dossier(self.ops, self.dossier.pk, expected=self.dossier.version,
            reference='87/SME/SDFM/SG/ESSBO/2026', title='CDC copie exigences',
            assignee=self.second, reason='Nouvelle procédure')
        copied = CdcRequirement.objects.filter(item__lot__dossier=duplicate, code='REQ-01')
        self.assertGreaterEqual(copied.count(), 1)
        self.assertNotEqual(copied.first().item_id, requirement.item_id)

    def test_requirement_revision_restore_restores_text_and_criterion_link(self):
        requirement, source_revision = save_requirement(self.operator, self.dossier,
            expected=self.dossier.version, values=self.requirement_values(), reason='Exigence initiale')
        self.refresh()
        criterion, source_revision = save_criterion(self.operator, self.dossier,
            expected=self.dossier.version,
            values=self.criterion_values(requirement=requirement), reason='Critère lié')
        self.refresh()
        original_revision = source_revision
        save_requirement(self.operator, self.dossier, expected=self.dossier.version,
            pk=requirement.pk, values=self.requirement_values(statement='Texte provisoire'),
            reason='Modification provisoire')
        self.refresh()
        restored = restore_revision(self.ops, original_revision.pk, expected=self.dossier.version,
            reason='Restauration gouvernance complète')
        self.refresh()
        restored_requirement = CdcRequirement.objects.get(item=requirement.item, code='REQ-01')
        self.assertEqual(restored_requirement.statement, 'Performance minimale vérifiable')
        self.assertEqual(self.dossier.criteria.get(code='TECH-01').requirement_id,
            restored_requirement.pk)
        self.assertEqual(restored.governance['requirements'][0]['statement'],
            'Performance minimale vérifiable')

    def test_requirement_http_and_export_surfaces_are_native(self):
        self.client.force_login(self.operator)
        list_url = reverse('erp:cdc-requirements', args=[self.dossier.pk])
        self.assertEqual(self.client.get(list_url).status_code, 200)
        create_url = reverse('erp:cdc-requirement-new', args=[self.dossier.pk])
        self.assertEqual(self.client.get(create_url).status_code, 200)
        payload = {key: (value.pk if hasattr(value, 'pk') else value)
            for key, value in self.requirement_values().items()}
        payload.update(expected_version=self.dossier.version, reason='Création HTTP')
        response = self.client.post(create_url, payload)
        self.assertEqual(response.status_code, 302)
        self.refresh()
        requirement = CdcRequirement.objects.get(item__lot__dossier=self.dossier, code='REQ-01')
        edit_url = reverse('erp:cdc-requirement-edit', args=[self.dossier.pk, requirement.pk])
        self.assertEqual(self.client.get(edit_url).status_code, 200)
        export = self.client.get(reverse('erp:cdc-criteria-export', args=[self.dossier.pk]))
        payload_xlsx = b''.join(export.streaming_content)
        book = load_workbook(io.BytesIO(payload_xlsx), data_only=True)
        self.assertEqual(book['Exigences']['C2'].value, 'REQ-01')
        self.assertEqual(book['Exigences']['E2'].value, 'Performance minimale vérifiable')
