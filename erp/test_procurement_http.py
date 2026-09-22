from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import uuid

from django.core.exceptions import ValidationError
from django.test import TestCase,override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from erp.models import ProcurementPlan,ProcurementLine,PurchaseOrder,PurchaseOrderLine
from erp.services.procurement import create_plan,add_plan_article,decide_plan_line,submit_plan,approve_plan
from erp.services.work import delegate_work
from erp.test_operations import OperationFixtures


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],SECURE_SSL_REDIRECT=False,
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
        'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ProcurementHTTPTests(OperationFixtures,TestCase):
    def setUp(self):
        self.plan=create_plan(self.ops,reference='PLAN-HTTP',year=timezone.localdate().year+1,
            title='Plan HTTP',assignee=self.operator,allow_costs=True)
        self.line=add_plan_article(self.ops,self.plan.pk,expected=self.plan.version,article=self.article,lot_name='Consommables')
        self.plan.refresh_from_db()

    def valid_decision(self):
        return {'expected_version':self.plan.version,'retained_quantity':'10','included':'on',
            'lot_name':'Consommables','priority':'HIGH','estimated_price':'12.34','tax_rate':'19',
            'currency':'DZD','supplier':self.party.pk,'price_source':'Devis contrôlé',
            'decision_reason':'Quantité vérifiée sur la base de l’activité prévue'}

    def approve(self):
        data=self.valid_decision()
        data.pop('expected_version')
        data['included']=True
        data['supplier']=self.party
        decide_plan_line(self.operator,self.line.pk,expected=self.plan.version,values=data)
        self.plan.refresh_from_db()
        submit_plan(self.operator,self.plan.pk,expected=self.plan.version,reason='Préparation terminée')
        self.plan.refresh_from_db()
        approve_plan(self.ops,self.plan.pk,expected=self.plan.version,reason='Plan examiné')
        self.plan.refresh_from_db()

    def test_member_decides_submits_and_admin_approves_through_native_forms(self):
        self.client.force_login(self.operator)
        routes=[('erp:procurement-list',[]),('erp:procurement-detail',[self.plan.pk]),
            ('erp:procurement-article',[self.plan.pk]),('erp:procurement-line',[self.line.pk]),
            ('erp:forecast-detail',[self.line.pk])]
        for name,args in routes:
            with self.subTest(route=name):
                self.assertEqual(self.client.get(reverse(name,args=args)).status_code,200)
        decision=self.client.post(reverse('erp:procurement-line',args=[self.line.pk]),self.valid_decision())
        self.assertEqual(decision.status_code,302,decision.context['form'].errors if decision.context else '')
        self.line.refresh_from_db()
        self.assertEqual(self.line.retained_quantity,10)
        self.plan.refresh_from_db()
        response=self.client.post(reverse('erp:procurement-detail',args=[self.plan.pk]),
            {'action':'submit','expected_version':self.plan.version,'reason':'Plan prêt pour revue'})
        self.assertEqual(response.status_code,302)
        response=self.client.post(reverse('erp:procurement-detail',args=[self.plan.pk]),
            {'action':'approve','expected_version':self.plan.version,'reason':'Auto-validation'})
        self.assertEqual(response.status_code,403)
        self.client.force_login(self.ops)
        response=self.client.post(reverse('erp:procurement-detail',args=[self.plan.pk]),
            {'action':'approve','expected_version':self.plan.version,'reason':'Dossier et budget vérifiés'})
        self.assertEqual(response.status_code,302)
        self.plan.refresh_from_db()
        self.assertIsNotNone(self.plan.approved_revision_id)
        self.assertContains(self.client.get(reverse('erp:procurement-detail',args=[self.plan.pk])),'Préparer le cahier des charges')

    def test_forecast_compute_and_stale_form_are_controlled(self):
        self.client.force_login(self.operator)
        url=reverse('erp:forecast-detail',args=[self.line.pk])
        data={'expected_version':self.plan.version}
        response=self.client.post(url,data)
        self.assertEqual(response.status_code,302)
        rendered=self.client.get(url)
        self.assertContains(rendered,'Historique inférieur à douze mois')
        self.assertEqual(self.client.post(url,data).status_code,400)
        self.assertEqual(self.plan.forecasts.count(),1)

    def test_export_uses_formula_cells_and_never_exposes_costs_without_permission(self):
        self.approve()
        self.client.force_login(self.ops)
        response=self.client.get(reverse('erp:procurement-export',args=[self.plan.pk]))
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(b''.join(response.streaming_content)),data_only=False)
        sheet=book['Données']
        self.assertEqual(sheet['F5'].value,10)
        self.assertEqual(sheet['J5'].value,12.34)
        self.assertEqual(sheet['M5'].value,'=ROUND(F5*J5,2)')
        self.assertEqual(sheet['O5'].value,'=M5+N5')
        self.assertEqual(book['Traçabilité']['B3'].value,self.plan.approved_revision.sha256)
        self.plan.work.allow_costs=False
        self.plan.work.save(update_fields=['allow_costs'])
        self.client.force_login(self.operator)
        response=self.client.get(reverse('erp:procurement-export',args=[self.plan.pk]))
        book=load_workbook(BytesIO(b''.join(response.streaming_content)),data_only=False)
        self.assertEqual(book['Données'].max_column,9)
        cells=[cell.value for row in book['Données'] for cell in row]
        self.assertNotIn('Prix HT',cells)
        self.assertNotIn(12.34,cells)
        self.assertNotIn('Devis contrôlé',cells)
        from erp.services.report_exports import literal
        self.assertEqual(literal('=HYPERLINK(1)'), chr(39)+'=HYPERLINK(1)')
        self.assertEqual(literal('  @SUM(A1)'),chr(39)+'  @SUM(A1)')

    def test_order_creation_line_confirmation_partial_receipt_and_completion(self):
        self.approve()
        self.client.force_login(self.ops)
        url=reverse('erp:order-create',args=[self.plan.pk])
        self.assertEqual(self.client.get(url).status_code,200)
        response=self.client.post(url,{'expected_version':self.plan.version,'reference':'PO-HTTP','supplier':self.party.pk,
            'ordered_on':timezone.localdate().isoformat(),'expected_on':(timezone.localdate()+timedelta(days=14)).isoformat()})
        self.assertEqual(response.status_code,302,response.context['form'].errors if response.context else '')
        order=PurchaseOrder.objects.get(reference='PO-HTTP')
        url=reverse('erp:order-line-new',args=[order.pk])
        self.assertEqual(self.client.get(url).status_code,200)
        response=self.client.post(url,{'expected_version':order.version,'plan_line':self.line.pk,'quantity':'10',
            'unit_price':'12.34','tax_rate':'19','currency':'DZD'})
        self.assertEqual(response.status_code,302,response.context['form'].errors if response.context else '')
        line=order.lines.get()
        order.refresh_from_db()
        response=self.client.post(reverse('erp:order-detail',args=[order.pk]),
            {'action':'confirm','expected_version':order.version,'reason':'Commande institutionnelle signée'})
        self.assertEqual(response.status_code,302)
        order.refresh_from_db()
        url=reverse('erp:order-receive',args=[line.pk])
        self.assertEqual(self.client.get(url).status_code,200)
        payload={'expected_version':order.version,'key':str(uuid.uuid4()),'amount':'10','location':self.freezer.pk,
            'lot_code':'LOT-HTTP','manufacturer_lot':'MFG-HTTP','container_code':'CONT-HTTP',
            'received_on':timezone.localdate().isoformat(),'condition':'Intact','cold_chain_ok':'unknown'}
        response=self.client.post(url,payload)
        self.assertEqual(response.status_code,302,response.context['form'].errors if response.context else '')
        self.assertEqual(self.client.post(url,payload).status_code,302)
        order.refresh_from_db()
        self.assertEqual(order.status,'RECEIVED')
        self.assertEqual(line.deliveries.count(),1)
        self.assertContains(self.client.get(reverse('erp:order-detail',args=[order.pk])),'CONT-HTTP')

    def test_scope_revocation_cost_forgery_and_unassigned_access(self):
        self.plan.work.allow_costs=False
        self.plan.work.save(update_fields=['allow_costs'])
        self.client.force_login(self.operator)
        url=reverse('erp:procurement-line',args=[self.line.pk])
        self.assertNotIn('estimated_price',self.client.get(url).context['form'].fields)
        response=self.client.post(url,self.valid_decision())
        self.assertEqual(response.status_code,302)
        self.line.refresh_from_db()
        self.assertIsNone(self.line.estimated_price)
        self.plan.work.refresh_from_db()
        delegate_work(self.ops,self.plan.work_id,expected=self.plan.work.version,assignee=self.second,reason='Réaffectation contrôlée')
        self.assertEqual(self.client.get(url).status_code,404)
        self.assertEqual(self.client.get(reverse('erp:procurement-export',args=[self.plan.pk])).status_code,404)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse('erp:procurement-list')).status_code,403)
        self.assertEqual(self.client.get(reverse('erp:procurement-create')).status_code,403)

    def test_admin_create_and_add_article_forms_reject_duplicates(self):
        self.client.force_login(self.ops)
        response=self.client.get(reverse('erp:procurement-create'))
        self.assertEqual(response.status_code,200)
        response=self.client.post(reverse('erp:procurement-create'),{'title':'Nouveau plan','reference':'HTTP-CREATE',
            'year':timezone.localdate().year+1,'assignee':self.second.pk,'allow_costs':'on'})
        self.assertEqual(response.status_code,302,response.context['form'].errors if response.context else '')
        plan=ProcurementPlan.objects.get(reference='HTTP-CREATE')
        self.assertEqual(plan.work.assignee_id,self.second.pk)
        url=reverse('erp:procurement-article',args=[plan.pk])
        data={'expected_version':plan.version,'article':self.article.pk,'lot_name':'Fournitures'}
        self.assertEqual(self.client.post(url,data).status_code,302)
        plan.refresh_from_db()
        data['expected_version']=plan.version
        self.assertEqual(self.client.post(url,data).status_code,400)
        self.assertEqual(plan.lines.count(),1)
