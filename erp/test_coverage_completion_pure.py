import copy
import io
import json
import sys
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import Mock, patch

from django.core.exceptions import PermissionDenied, ValidationError
from django.test import SimpleTestCase, TestCase
from xml.etree import ElementTree as ET

from erp import table_probe
from erp.cdc import catalog, consultation, lot_workbook, procurement
from erp.cdc.docengine import DocumentError
from erp.services import bulk_imports
from erp.test_operations import OperationFixtures


def workbook_bytes(sheets=("Data",), rows=(("a", "b"), ("1", "2"))):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheets[0]
    for row in rows:
        ws.append(list(row))
    for title in sheets[1:]:
        wb.create_sheet(title)
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def rewrite_zip(payload, changes=(), additions=()):
    src = io.BytesIO(payload)
    out = io.BytesIO()
    mapping = dict(changes)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = mapping.get(info.filename, zin.read(info.filename))
            zout.writestr(info.filename, data)
        for name, data in additions:
            zout.writestr(name, data)
    return out.getvalue()


class TableProbeCoverageTests(SimpleTestCase):
    def helpers(self):
        import runpy
        return runpy.run_path(str(table_probe.Path(table_probe.__file__).resolve().parents[1] / "core/document_probe.py"))

    def test_normalize_csv_and_matrix_validation(self):
        self.assertEqual(table_probe.normalize(date(2026, 9, 23)), "2026-09-23")
        self.assertTrue(table_probe.normalize(True) == "true")
        self.assertEqual(table_probe.normalize(None), "")
        with self.assertRaisesRegex(ValueError, "CELL"):
            table_probe.normalize("x" * 10001)
        with self.assertRaisesRegex(ValueError, "CELL"):
            table_probe.normalize("a\x01b")
        self.assertEqual(table_probe.csv_rows(b"a;b\n1;2\n"), [["a", "b"], ["1", "2"]])
        with patch.object(table_probe, "MAX_COLUMNS", 1):
            with self.assertRaisesRegex(ValueError, "SIZE"):
                table_probe.csv_rows(b"a,b\n1,2\n")
        h = self.helpers()
        self.assertEqual(table_probe.read_matrix(b"a,b\n1,2\n\n", ".csv", h), [["a", "b"], ["1", "2"]])
        for payload, ext, code in [(b"", ".csv", "SIZE"), (b"a\n", ".csv", "EMPTY"), (b"a,b\n1,2", ".txt", "EXTENSION")]:
            with self.subTest(ext=ext, code=code), self.assertRaisesRegex(ValueError, code):
                table_probe.read_matrix(payload, ext, h)
        with patch.object(table_probe, "MAX_TEXT", 2):
            with self.assertRaisesRegex(ValueError, "SIZE"):
                table_probe.read_matrix(b"a,b\n1,2\n", ".csv", h)

    def test_xlsx_happy_path_sheet_and_size_guards(self):
        h = self.helpers()
        payload = workbook_bytes()
        rows = table_probe.xlsx_rows(payload, h)
        self.assertEqual(rows[0][:2], ["a", "b"])
        with self.assertRaisesRegex(ValueError, "SHEETS"):
            table_probe.xlsx_rows(workbook_bytes(("Data", "Other")), h)
        wide = workbook_bytes(rows=(tuple(str(i) for i in range(41)), tuple("x" for _ in range(41))))
        with self.assertRaisesRegex(ValueError, "SIZE"):
            table_probe.xlsx_rows(wide, h)
        tall = workbook_bytes(rows=tuple(("x",) for _ in range(502)))
        with self.assertRaisesRegex(ValueError, "SIZE"):
            table_probe.xlsx_rows(tall, h)

    def test_xlsx_zip_security_guards(self):
        h = self.helpers()
        payload = workbook_bytes()
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            sheet_name = next(n for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
            sheet = z.read(sheet_name)
            rel_name = "xl/_rels/workbook.xml.rels"
            rel = z.read(rel_name)
        formula = sheet.replace(b"</c>", b"<f>1+1</f></c>", 1)
        with self.assertRaisesRegex(ValueError, "FORMULA"):
            table_probe.xlsx_rows(rewrite_zip(payload, ((sheet_name, formula),)), h)
        active = rewrite_zip(payload, additions=(("xl/vbaProject.bin", b"x"),))
        with self.assertRaisesRegex(ValueError, "Active"):
            table_probe.xlsx_rows(active, h)
        external = rel.replace(b"</Relationships>", b'<Relationship Id="rExt" TargetMode="External" Target="https://example.test" Type="x"/></Relationships>')
        with self.assertRaisesRegex(ValueError, "EXTERNAL"):
            table_probe.xlsx_rows(rewrite_zip(payload, ((rel_name, external),)), h)
        duplicate = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(payload)) as zin, zipfile.ZipFile(duplicate, "w") as zout:
            for info in zin.infolist():
                zout.writestr(info.filename, zin.read(info.filename))
            zout.writestr(zin.infolist()[0].filename, b"duplicate")
        with self.assertRaisesRegex(ValueError, "STRUCTURE"):
            table_probe.xlsx_rows(duplicate.getvalue(), h)

    def test_main_success_and_error_mapping(self):
        helpers = {"limit_resources": Mock(), "_entry_key": lambda info: info.filename, "_xml": lambda data: ET.fromstring(data)}
        stdin = Mock(buffer=io.BytesIO(b"a,b\n1,2\n"))
        stdout = Mock(buffer=io.BytesIO())
        with patch.object(table_probe.runpy, "run_path", return_value=helpers), patch.object(table_probe.sys, "stdin", stdin), patch.object(table_probe.sys, "stdout", stdout), patch.object(table_probe.sys, "argv", ["probe", ".csv"]):
            self.assertEqual(table_probe.main(), 0)
            self.assertEqual(json.loads(stdout.buffer.getvalue()), {"rows": [["a", "b"], ["1", "2"]]})
            helpers["limit_resources"].assert_called_once()
        stdin = Mock(buffer=io.BytesIO(b"bad"))
        stdout = Mock(buffer=io.BytesIO())
        with patch.object(table_probe.runpy, "run_path", return_value=helpers), patch.object(table_probe.sys, "stdin", stdin), patch.object(table_probe.sys, "stdout", stdout), patch.object(table_probe.sys, "argv", ["probe", ".bad"]):
            self.assertEqual(table_probe.main(), 1)
            self.assertEqual(json.loads(stdout.buffer.getvalue()), {"error": "EXTENSION"})
        stdout = Mock(buffer=io.BytesIO())
        with patch.object(table_probe.runpy, "run_path", side_effect=RuntimeError("surprise")), patch.object(table_probe.sys, "stdout", stdout):
            self.assertEqual(table_probe.main(), 1)
            self.assertEqual(json.loads(stdout.buffer.getvalue()), {"error": "INVALID"})


class ConsultationCoverageTests(SimpleTestCase):
    def valid(self, family="equipment"):
        return consultation.initial_consultation(family)

    def test_initial_upgrade_and_value_guardrails(self):
        with self.assertRaises(DocumentError):
            consultation.initial_consultation("unknown")
        value = self.valid()
        self.assertEqual(consultation.consultation_value({"consultation": value}, "equipment"), value)
        self.assertEqual(consultation.consultation_value({}, "equipment")["schema"], consultation.SCHEMA)
        legacy = copy.deepcopy(value)
        legacy["schema"] = 1
        legacy.pop("withdrawal_fr")
        legacy.pop("withdrawal_ar")
        legacy.pop("confirmed")
        upgraded = consultation._upgrade(legacy, "equipment")
        self.assertEqual(upgraded["schema"], consultation.SCHEMA)
        self.assertFalse(upgraded["confirmed"])
        for bad in (None, [], {"schema": 999}):
            with self.subTest(bad=bad), self.assertRaises(DocumentError):
                consultation._upgrade(bad, "equipment")

    def test_validation_rejects_each_semantic_family(self):
        base = self.valid()
        cases = []
        bad = copy.deepcopy(base); bad["extra"] = 1; cases.append(bad)
        bad = copy.deepcopy(base); bad["object_fr"] = ""; cases.append(bad)
        bad = copy.deepcopy(base); bad["budget_year"] = 1999; cases.append(bad)
        bad = copy.deepcopy(base); bad["preparation_days"] = 0; cases.append(bad)
        bad = copy.deepcopy(base); bad["validity_months"] = 61; cases.append(bad)
        bad = copy.deepcopy(base); bad["confirmed"] = 1; cases.append(bad)
        bad = copy.deepcopy(base); bad["deposit_time"] = "25:00"; cases.append(bad)
        bad = copy.deepcopy(base); bad["preparation_fr"] = "vingt jours"; cases.append(bad)
        bad = copy.deepcopy(base); bad["preparation_ar"] = "عشرون"; cases.append(bad)
        bad = copy.deepcopy(base); bad["validity_fr"] = "deux mois"; cases.append(bad)
        bad = copy.deepcopy(base); bad["validity_ar"] = "شهران"; cases.append(bad)
        for item in cases:
            with self.subTest(item=item), self.assertRaises(DocumentError):
                consultation.validate_consultation(item, "equipment")
        reagent = self.valid("reagents")
        reagent["confirmed"] = True
        reagent["withdrawal_ar"] = "sans compte commun"
        with self.assertRaises(DocumentError):
            consultation.validate_consultation(reagent, "reagents")
        reagent["withdrawal_ar"] = consultation.DEFAULTS["reagents"]["withdrawal_fr"]
        self.assertTrue(consultation.validate_consultation(reagent, "reagents")["confirmed"])

    def test_consultation_edits_short_circuits_conflict_and_missing_binding(self):
        spans = {"p": [{"segment": 0, "before": "wrong", "after": "x"}]}
        data = {"family": "equipment"}
        result, report = consultation.consultation_edits(data, spans, Mock())
        self.assertEqual(report["status"], "NOT_REQUESTED")
        data["consultation"] = self.valid()
        result, report = consultation.consultation_edits(data, {}, Mock())
        self.assertEqual(report["status"], "UNCHANGED")

        changed = self.valid()
        changed["object_fr"] = "Objet de couverture"
        data["consultation"] = changed
        nodes = [Mock(characters=consultation.DEFAULTS["equipment"]["object_fr"])]
        block = {"id": "p"}
        doc = Mock(source_index=[block])
        doc.editable_segments.return_value = [nodes]
        with patch.object(consultation, "_replace_object", return_value="Objet de couverture"):
            result, report = consultation.consultation_edits(data, {}, doc)
        self.assertEqual(report["status"], "GENERATED")
        self.assertEqual(report["fields"]["object_fr"], 1)

        with patch.object(consultation, "_replace_object", return_value="Objet de couverture"):
            with self.assertRaises(DocumentError):
                consultation.consultation_edits(data, spans, doc)

        empty_doc = Mock(source_index=[])
        with self.assertRaises(DocumentError):
            consultation.consultation_edits(data, {}, empty_doc)

    def test_findings_missing_confirmation_and_calendar_order(self):
        self.assertEqual(consultation.consultation_findings({"family": "equipment"})[0]["id"], "CONSULTATION_REQUIRED")
        data = {"family": "equipment", "consultation": self.valid()}
        ids = {x["id"] for x in consultation.consultation_findings(data)}
        self.assertIn("CONSULTATION_CONFIRMATION", ids)
        data["consultation"]["deposit_time"] = "13:00"
        data["consultation"]["opening_time"] = "12:00"
        ids = {x["id"] for x in consultation.consultation_findings(data)}
        self.assertIn("CALENDAR_ORDER", ids)


class CatalogCoverageTests(SimpleTestCase):
    def fresh(self, family="equipment"):
        return catalog.initial_data(family)

    def test_validate_data_metadata_and_structure_guardrails(self):
        with self.assertRaises(DocumentError):
            catalog.validate_data({}, "unknown")
        base = self.fresh()
        mutations = []
        x=copy.deepcopy(base); x["format"]=2; mutations.append(x)
        x=copy.deepcopy(base); x["unexpected"]=1; mutations.append(x)
        x=copy.deepcopy(base); x["reference"]=""; mutations.append(x)
        x=copy.deepcopy(base); x["paragraphs"]="bad"; mutations.append(x)
        x=copy.deepcopy(base); x["rows"]={"x":"bad"}; mutations.append(x)
        x=copy.deepcopy(base); x["omitted"]="bad"; mutations.append(x)
        x=copy.deepcopy(base); x["review_acknowledged"]="yes"; mutations.append(x)
        for item in mutations:
            with self.subTest(keys=item.keys()), self.assertRaises(DocumentError):
                catalog.validate_data(item, "equipment")

        profile = catalog.profile("equipment")
        editable = next(p for p in profile["paragraphs"] if not p["guard"])
        bad = copy.deepcopy(base); bad["paragraphs"]={editable["id"]:"a\nb"}
        with self.assertRaises(DocumentError):
            catalog.validate_data(bad, "equipment")
        bad = copy.deepcopy(base); bad["omitted"]=[editable["id"]]
        with self.assertRaises(DocumentError):
            catalog.validate_data(bad, "equipment")

    def test_controls_preview_anchor_and_reference_replacement(self):
        base = self.fresh()
        original = catalog.profile("equipment")["reference"]
        canonical = "17/SME/SDFM/SG/ESSBO/2026"
        data = copy.deepcopy(base); data["reference"] = canonical
        self.assertIn("17", catalog.reference_replacement(original, data))
        with patch.object(catalog, "generation_edits", return_value=({}, {"p":[{"segment":0,"before":"X","after":"Y"}]})):
            node=Mock(characters="Z")
            paragraph=Mock()
            doc=Mock()
            doc.paragraphs={"p":("part",paragraph)}
            doc.editable_segments.return_value=[[node]]
            with patch.object(catalog, "document", return_value=doc), \
                 patch("erp.cdc.docengine.own_text_nodes", return_value=[node]):
                with self.assertRaises(DocumentError):
                    catalog.effective_edits(data)

        with patch.object(catalog, "validate_data", return_value=data), \
             patch.object(catalog, "profile", return_value={"issues":[{"id":"I1","summary":"warn"}],"reference":original}), \
             patch.object(catalog, "resolved_issue", return_value=False), \
             patch("erp.cdc.schedule_adapter.binding_findings", return_value=[]), \
             patch.object(catalog, "consultation_findings", return_value=[]), \
             patch("erp.cdc.common_data.manual_conflicts", return_value=[]), \
             patch.object(catalog, "effective_edits", return_value={"p":"{{ unresolved }}"}):
            findings = catalog.controls(data)
        ids={x["id"] for x in findings}
        self.assertIn("I1",ids); self.assertIn("UNRESOLVED_MARKER",ids); self.assertIn("AR_MAPPING_REQUIRED",ids)


class LotWorkbookCoverageTests(SimpleTestCase):
    def xml_cell_sheet(self, cells):
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        root=ET.Element("{%s}worksheet"%ns)
        data=ET.SubElement(root,"{%s}sheetData"%ns)
        row=ET.SubElement(data,"{%s}row"%ns,{"r":"1"})
        for attrs, children in cells:
            c=ET.SubElement(row,"{%s}c"%ns,attrs)
            for tag,text in children:
                child=ET.SubElement(c,"{%s}%s"%(ns,tag))
                child.text=text
        return ET.tostring(root)

    def archive(self, name, data):
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,"w") as z:z.writestr(name,data)
        return zipfile.ZipFile(io.BytesIO(buf.getvalue()))

    def test_xml_and_read_cells_guardrails(self):
        for payload in (b"<!DOCTYPE x><x/>", b"<!ENTITY e 'x'><x/>", b"<x>"):
            with self.assertRaises(DocumentError):
                lot_workbook.xml(payload)
        self.assertEqual(lot_workbook.xml(b"<x><y>1</y></x>").find("y").text,"1")

        path="xl/worksheets/sheet1.xml"
        cases=[
          ([({"r":"A1"},[("f","1+1"),("v","2")])], "formules"),
           ([({"r":"J1"},[("v","1")])], "Adresse"),
           ([({"r":"A1","t":"s"},[("v","-1")])], "Chaîne"),
           ([({"r":"A1","t":"e"},[("v","#N/A")])], "texte ou nombre"),
           ([({"r":"A1"},[("v","x"*32768)])], "limite Excel"),
        ]
        for cells, message in cases:
            z=self.archive(path,self.xml_cell_sheet(cells))
            with z, self.assertRaisesRegex(DocumentError,message):
                lot_workbook.read_cells(z,path,[])
        z=self.archive(path,self.xml_cell_sheet([({"r":"A1","t":"inlineStr"},[("is","")])]))
        with z:
            self.assertIn("A1",lot_workbook.read_cells(z,path,[]))

    def test_parse_workbook_front_door_guards(self):
        data=catalog.initial_data("equipment")
        for mode in ("bad",None):
            with self.subTest(mode=mode), self.assertRaises(DocumentError):
                lot_workbook.parse_workbook(b"x","x.xlsx",data,uuid.uuid4(),1,lambda x:x,mode=mode)
        with self.assertRaises(DocumentError):
            lot_workbook.parse_workbook(b"x","x.txt",data,uuid.uuid4(),1,lambda x:x)
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,"w") as z:z.writestr("xl/vbaProject.bin",b"x")
        with self.assertRaises(DocumentError):
            lot_workbook.parse_workbook(buf.getvalue(),"x.xlsx",data,uuid.uuid4(),1,lambda x:x)
        plain = workbook_bytes()
        with self.assertRaises(DocumentError):
            lot_workbook.parse_workbook(plain,"x.xlsx",data,uuid.uuid4(),1,lambda x:x)


class BulkImportPrimitiveCoverageTests(OperationFixtures, TestCase):
    def test_scalar_parsers_and_reference_errors(self):
        self.assertEqual(bulk_imports._decimal("1 234,5"),Decimal("1234.5"))
        for bad in ("", "x", "NaN", "-1", "9"*70):
            with self.subTest(bad=bad), self.assertRaises(ValidationError):
                bulk_imports._decimal(bad)
        self.assertEqual(bulk_imports._decimal("-1",signed=True),Decimal("-1"))
        self.assertEqual(bulk_imports._integer("2"),2)
        with self.assertRaises(ValidationError): bulk_imports._integer("2.5")
        self.assertEqual(bulk_imports._date("2026-09-23"),date(2026,9,23))
        self.assertEqual(bulk_imports._date("2026-09-23T00:00:00"),date(2026,9,23))
        for bad in ("bad","2026-09-23T01:00:00","2026-09-23T00:00:00+01:00"):
            with self.subTest(bad=bad), self.assertRaises(ValidationError): bulk_imports._date(bad)
        self.assertTrue(bulk_imports._boolean("oui")); self.assertFalse(bulk_imports._boolean("no"))
        self.assertIsNone(bulk_imports._boolean("unknown"))
        with self.assertRaises(ValidationError): bulk_imports._boolean("maybe")
        self.assertIsNotNone(bulk_imports._moment("2026-09-23T12:00:00"))
        with self.assertRaises(ValidationError): bulk_imports._moment("2026-09-23")
        with self.assertRaises(ValidationError): bulk_imports._criticality("not-a-level")
        with self.assertRaises(ValidationError): bulk_imports._get(type(self.article), "missing")
        self.assertIsNone(bulk_imports._get(type(self.article),"",optional=True))

    def test_require_kind_denials_and_unknown_domain(self):
        outsider=self.outsider
        with self.assertRaises(PermissionDenied): bulk_imports.require_kind(outsider,"CATALOG")
        with self.assertRaises(ValidationError): bulk_imports.require_kind(self.ops,"PLAN")
        with self.assertRaises(ValidationError): bulk_imports.require_kind(self.ops,"UNKNOWN")
