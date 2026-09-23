from datetime import datetime,timedelta
from decimal import Decimal
from io import BytesIO,StringIO
import csv
import json
from pathlib import Path
import runpy
import uuid
import zipfile

from django.core.exceptions import PermissionDenied,ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import SimpleTestCase,TestCase,override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook,load_workbook

from erp import table_probe
from erp.models import (Article,AuditEvent,BiologicalSample,Capability,ImportBatch,Location,PriceObservation,
    StockContainer,StockMovement,TemperatureReading)
from erp.services.bulk_imports import apply_import,cancel_import,preview_import,require_batch
from erp.services.catalog import save_article
from erp.services.common import Conflict
from erp.services.procurement import create_plan
from erp.services.stock import reconcile_stock
from erp.services.table_intake import import_template,parse_table
from erp.test_operations import OperationFixtures


def csv_file(rows):
    stream=StringIO(newline='')
    writer=csv.DictWriter(stream,fieldnames=list(rows[0]),delimiter=';')
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode('utf-8')


def xlsx_file(rows):
    book=Workbook()
    sheet=book.active
    sheet.title='Données'
    for row in rows:
        sheet.append(row)
    buffer=BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],SECURE_SSL_REDIRECT=False,
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
        'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ImportIntegrationTests(OperationFixtures,TestCase):
    def preview(self,kind,rows,user=None,**kwargs):
        return preview_import(user or self.ops,key=uuid.uuid4(),kind=kind,filename='donnees.csv',
            data=csv_file(rows),reason='Reprise documentée de recette',**kwargs)

    def receipt_row(self,code='IMPORT-CONT'):
        return {'article_code':self.article.code,'location_code':self.freezer.code,'lot_code':'IMPORT-LOT',
            'manufacturer_lot':'IMPORT-MANUF','container_code':code,'amount':'10','unit_code':self.unit.code,
            'received_on':timezone.localdate().isoformat(),'condition':'Intact'}

    def test_catalogue_preview_is_nonpersistent_and_apply_is_idempotent(self):
        rows=[{'code':'IMPORTED','name':'Article importé','category_code':self.category.code,'base_unit_code':self.unit.code,
            'safety_stock':'12','minimum_order_quantity':'2','lead_time_days':'30','criticality':'Haute'}]
        count=Article.objects.count()
        batch=self.preview('CATALOG',rows)
        self.assertTrue(batch.report['valid'],batch.report)
        self.assertEqual(Article.objects.count(),count)
        result=apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        item=Article.objects.get(code='IMPORTED')
        self.assertEqual((item.safety_stock,item.minimum_order_quantity,item.lead_time_days),(12,2,30))
        self.assertEqual(item.criticality,'HIGH')
        version=result.version
        second=apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        self.assertEqual(second.version,version)
        self.assertEqual(Article.objects.count(),count+1)

    def test_initial_stock_uses_ledger_and_can_create_explicit_catalogue_reference(self):
        row=self.receipt_row()
        row.update(article_code='NEW-INITIAL',new_name='Référence nouvelle',new_category_code=self.category.code,
            new_base_unit_code=self.unit.code)
        batch=self.preview('INITIAL',[row])
        self.assertTrue(batch.report['valid'],batch.report)
        self.assertFalse(StockMovement.objects.exists())
        self.assertFalse(Article.objects.filter(code='NEW-INITIAL').exists())
        apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        movement=StockMovement.objects.get(kind='INITIAL')
        self.assertEqual(movement.receipt.container.quantity,10)
        self.assertEqual(movement.receipt.container.lot.article.code,'NEW-INITIAL')
        self.assertEqual(reconcile_stock(self.ops),[])

    def test_invalid_row_prevents_all_partial_business_writes(self):
        rows=[self.receipt_row('FIRST-VALID'),{**self.receipt_row('SECOND-BAD'),'amount':'-1'}]
        batch=self.preview('RECEIPTS',rows)
        self.assertFalse(batch.report['valid'])
        self.assertTrue(batch.report['preview'][0]['valid'])
        self.assertFalse(batch.report['preview'][1]['valid'])
        with self.assertRaises(ValidationError):
            apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        self.assertFalse(StockContainer.objects.exists())
        self.assertFalse(StockMovement.objects.exists())

    def test_duplicate_container_rows_cannot_be_confirmed(self):
        batch=self.preview('INITIAL',[self.receipt_row('SAME-CONT'),self.receipt_row('SAME-CONT')])
        self.assertFalse(batch.report['valid'])
        self.assertTrue(batch.report['preview'][0]['valid'])
        self.assertFalse(batch.report['preview'][1]['valid'])
        with self.assertRaises(ValidationError):
            apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        self.assertFalse(StockContainer.objects.filter(code='SAME-CONT').exists())

    def test_storage_import_enforces_delegated_tree_and_position_limit(self):
        self.grant(Capability.EDIT_STORAGE,location=self.freezer)
        rows=[{'code':'SCOPED-RACK','name':'Rack délégué','kind_code':self.kind.code,
               'parent_code':self.freezer.code},
              {'code':'SCOPED-BOX','name':'Boîte déléguée','kind_code':self.kind.code,
               'parent_code':'SCOPED-RACK'}]
        batch=self.preview('LOCATIONS',rows,user=self.operator)
        self.assertTrue(batch.report['valid'],batch.report)
        require_batch(self.operator,batch)
        apply_import(self.operator,batch.pk,expected=batch.version,confirmed=True)
        self.assertEqual(Location.objects.get(code='SCOPED-BOX').parent.code,'SCOPED-RACK')
        outside=self.preview('LOCATIONS',[{'code':'OUTSIDE','name':'Hors périmètre',
            'kind_code':self.kind.code,'parent_code':self.lab.code}],user=self.operator)
        self.assertFalse(outside.report['valid'])
        self.assertFalse(Location.objects.filter(code='OUTSIDE').exists())
        with self.assertRaises(ValidationError):
            self.preview('LOCATIONS',[{'code':'TOO-MANY','name':'Boîte surdimensionnée',
                'kind_code':self.storage_kind.code,'parent_code':self.freezer.code,
                'grid_rows':'201','grid_columns':'100'}])
        self.assertFalse(Location.objects.filter(code='TOO-MANY').exists())

    def test_unknown_sample_source_never_creates_a_biobank_record(self):
        from accounts.models import MemberProfile
        from core.models import Request,Service
        member,_=MemberProfile.objects.get_or_create(user=self.operator)
        service=Service.objects.create(code='SOURCE-CHECK',name='Service échantillons')
        req=Request.objects.create(requester=self.outsider,assigned_to=member,service=service,
            title='Échantillons déclarés',status='IN_PROGRESS',sample_table=[{'sample_code':'KNOWN'}])
        base={'code':'UNKNOWN-SOURCE','location_code':self.freezer.code,'amount':'5',
            'unit_code':self.ul.code,'received_on':timezone.localdate().isoformat(),
            'request_reference':str(req.pk),'source_code':'MISSING'}
        wrong_sample=self.preview('SAMPLES',[base])
        self.assertFalse(wrong_sample.report['valid'])
        self.assertIn('introuvable',wrong_sample.report['preview'][0]['message'].lower())
        unknown_request=self.preview('SAMPLES',[{**base,'request_reference':str(uuid.uuid4())}])
        self.assertFalse(unknown_request.report['valid'])
        unlinked=self.preview('SAMPLES',[{**base,'request_reference':''}])
        self.assertFalse(unlinked.report['valid'])
        missing_position=self.preview('SAMPLES',[{**base,'request_reference':'',
            'source_code':'','position_row':'1','position_column':'1'}])
        self.assertFalse(missing_position.report['valid'])
        self.assertFalse(BiologicalSample.objects.filter(code='UNKNOWN-SOURCE').exists())

    def test_preview_rejects_reference_creation_races_at_confirmation(self):
        catalog={'code':'NEW-RACE','name':'Article prévu','category_code':self.category.code,
            'base_unit_code':self.unit.code}
        article_batch=self.preview('CATALOG',[catalog])
        self.assertTrue(article_batch.report['valid'])
        save_article(self.ops,{'code':'NEW-RACE','name':'Créé par une autre opération',
            'category':self.category,'base_unit':self.unit})
        with self.assertRaises(Conflict):
            apply_import(self.ops,article_batch.pk,expected=article_batch.version,confirmed=True)
        self.assertEqual(Article.objects.get(code='NEW-RACE').name,'Créé par une autre opération')

        row=self.receipt_row('RACE-CONTAINER')
        row.update(article_code='RACE-ARTICLE',new_name='Article prévu',
            new_category_code=self.category.code,new_base_unit_code=self.unit.code)
        stock_batch=self.preview('INITIAL',[row])
        self.assertTrue(stock_batch.report['valid'])
        save_article(self.ops,{'code':'RACE-ARTICLE','name':'Créé hors import',
            'category':self.category,'base_unit':self.unit})
        with self.assertRaises(Conflict):
            apply_import(self.ops,stock_batch.pk,expected=stock_batch.version,confirmed=True)
        self.assertFalse(StockContainer.objects.filter(code='RACE-CONTAINER').exists())

    def test_imported_receipt_preserves_dated_supplier_and_price_evidence(self):
        today=timezone.localdate()
        row=self.receipt_row('DATED-CONTAINER')
        row.update(expires_on=(today+timedelta(days=90)).isoformat(),
            manufactured_on=(today-timedelta(days=10)).isoformat(),
            supplier_code=self.party.code,unit_price_base='12,50',cold_chain_ok='oui')
        batch=self.preview('RECEIPTS',[row])
        self.assertTrue(batch.report['valid'],batch.report)
        apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        receipt=StockMovement.objects.get(kind='RECEIPT').receipt
        self.assertEqual(receipt.supplier,self.party)
        self.assertEqual(receipt.unit_price,Decimal('12.50'))
        self.assertEqual(receipt.container.lot.expires_on,today+timedelta(days=90))
        self.assertEqual(receipt.container.lot.manufactured_on,today-timedelta(days=10))

    def test_catalogue_change_after_preview_aborts_receipt_import(self):
        batch=self.preview('INITIAL',[self.receipt_row()])
        self.assertTrue(batch.report['valid'])
        save_article(self.ops,{'name':'Désignation modifiée'},pk=self.article.pk,expected=self.article.version)
        with self.assertRaises(Conflict):
            apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        self.assertFalse(StockMovement.objects.exists())

    def test_storage_tree_import_creates_parent_before_children_and_positions(self):
        rows=[{'code':'IMPORT-FREEZER','name':'Congélateur importé','kind_code':self.storage_kind.code,
            'parent_code':self.lab.code,'grid_rows':'','grid_columns':'','temperature_min':'-90','temperature_max':'-70'},
            {'code':'IMPORT-BOX','name':'Boîte importée','kind_code':self.storage_kind.code,
            'parent_code':'IMPORT-FREEZER','grid_rows':'2','grid_columns':'3','temperature_min':'','temperature_max':''}]
        batch=self.preview('LOCATIONS',rows)
        self.assertTrue(batch.report['valid'],batch.report)
        self.assertFalse(Location.objects.filter(code='IMPORT-BOX').exists())
        apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        box=Location.objects.get(code='IMPORT-BOX')
        self.assertEqual(box.parent.code,'IMPORT-FREEZER')
        self.assertEqual(box.positions.count(),6)

    def test_temperature_duplicates_are_recognized_without_duplicate_measurement(self):
        moment=(timezone.now()-timedelta(minutes=5)).replace(microsecond=0)
        row={'location_code':self.freezer.code,'measured_at':moment.isoformat(),'value':'-65','comment':'Relevé importé'}
        first=self.preview('TEMPERATURE',[row])
        self.assertTrue(first.report['valid'],first.report)
        apply_import(self.ops,first.pk,expected=first.version,confirmed=True)
        self.assertEqual(TemperatureReading.objects.count(),1)
        reading=TemperatureReading.objects.get()
        self.assertTrue(reading.out_of_range)
        self.assertEqual(reading.source,'IMPORT')
        second=self.preview('TEMPERATURE',[row])
        apply_import(self.ops,second.pk,expected=second.version,confirmed=True)
        self.assertEqual(TemperatureReading.objects.count(),1)

    def test_price_import_uses_documented_unit_and_does_not_duplicate_observation(self):
        row={'article_code':self.article.code,'unit_code':self.unit.code,'price':'12,34','currency':'DZD',
            'observed_on':timezone.localdate().isoformat(),'source':'Devis reçu','supplier_code':self.party.code}
        batch=self.preview('PRICES',[row])
        self.assertTrue(batch.report['valid'],batch.report)
        self.assertFalse(PriceObservation.objects.exists())
        apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        price=PriceObservation.objects.get()
        self.assertEqual(price.amount,Decimal('12.34'))
        next_batch=self.preview('PRICES',[row])
        apply_import(self.ops,next_batch.pk,expected=next_batch.version,confirmed=True)
        self.assertEqual(PriceObservation.objects.count(),1)

    def test_plan_import_has_one_combined_revision_and_does_not_auto_approve(self):
        plan=create_plan(self.ops,reference='IMPORT-PLAN',year=timezone.localdate().year+1,title='Plan importé',assignee=self.operator,allow_costs=True)
        rows=[{'article_code':self.article.code,'lot_name':'Réactifs','retained_quantity':'5','decision_reason':'Besoin vérifié',
            'estimated_price':'20','tax_rate':'19','currency':'DZD','source':'Devis fournisseur'},
            {'article_code':self.liquid.code,'lot_name':'Produits','retained_quantity':'10','decision_reason':'Consommation prévue',
            'estimated_price':'10','tax_rate':'0','currency':'DZD','source':'Document de prix'}]
        baseline=plan.revision_number
        batch=self.preview('PLAN',rows,user=self.operator,plan=plan)
        self.assertTrue(batch.report['valid'],batch.report)
        self.assertEqual(plan.lines.count(),0)
        apply_import(self.operator,batch.pk,expected=batch.version,confirmed=True)
        plan.refresh_from_db()
        self.assertEqual(plan.lines.count(),2)
        self.assertEqual(plan.revision_number,baseline+1)
        self.assertIsNone(plan.approved_revision_id)
        self.assertEqual(plan.work.status,'ASSIGNED')

    def test_sample_import_links_existing_source_instead_of_unrelated_duplicate(self):
        from accounts.models import MemberProfile
        from core.models import Request,Service
        member,_=MemberProfile.objects.get_or_create(user=self.operator)
        service=Service.objects.create(code='IMPORT-SERVICE',name='Service de recette')
        req=Request.objects.create(requester=self.outsider,assigned_to=member,service=service,
            title='Demande importée',display_id='PLAGENOR-IMPORT-1',status='IN_PROGRESS',sample_table=[{'sample_code':'DECLARED-1','sample_type':'ADN'}])
        row={'code':'IMPORTED-SAMPLE','location_code':self.freezer.code,'amount':'25','unit_code':self.ul.code,
            'received_on':timezone.localdate().isoformat(),'request_reference':req.display_id,'source_code':'DECLARED-1'}
        batch=self.preview('SAMPLES',[row])
        self.assertTrue(batch.report['valid'],batch.report)
        apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        sample=BiologicalSample.objects.get(code='IMPORTED-SAMPLE')
        self.assertEqual(sample.origin_request_id,req.pk)
        self.assertEqual(sample.source_snapshot['row']['sample_code'],'DECLARED-1')
        self.assertEqual(sample.events.count(),1)

    def test_import_cancellation_expiration_and_confirmation_are_enforced(self):
        batch=self.preview('INITIAL',[self.receipt_row()])
        with self.assertRaises(ValidationError):
            apply_import(self.ops,batch.pk,expected=batch.version,confirmed=False)
        ImportBatch.objects.filter(pk=batch.pk).update(expires_at=timezone.now()-timedelta(seconds=1))
        with self.assertRaises(ValidationError):
            apply_import(self.ops,batch.pk,expected=batch.version,confirmed=True)
        other=self.preview('INITIAL',[self.receipt_row('CANCELLED-CONT')])
        cancel_import(self.ops,other.pk,expected=other.version,reason='Fichier à reprendre')
        other.refresh_from_db()
        self.assertEqual(other.status,'CANCELLED')
        with self.assertRaises(ValidationError):
            apply_import(self.ops,other.pk,expected=other.version,confirmed=True)
        self.assertFalse(StockMovement.objects.exists())

    def test_scoped_import_permissions_and_revocation_are_checked_server_side(self):
        self.grant(Capability.RECEIVE_STOCK,location=self.freezer,category=self.category)
        batch=self.preview('INITIAL',[self.receipt_row()],user=self.operator)
        self.assertTrue(batch.report['valid'],batch.report)
        with self.assertRaises(PermissionDenied):
            apply_import(self.second,batch.pk,expected=batch.version,confirmed=True)
        from erp.models import AccessGrant
        AccessGrant.objects.filter(user=self.operator).update(active=False)
        with self.assertRaises(PermissionDenied):
            require_batch(self.operator,batch)
        self.assertFalse(StockMovement.objects.exists())
        with self.assertRaises(PermissionDenied):
            self.preview('INITIAL',[self.receipt_row()],user=self.outsider)

    def test_same_preview_key_replays_exactly_and_rejects_changed_payload(self):
        key=uuid.uuid4()
        data=csv_file([self.receipt_row()])
        kwargs={'key':key,'kind':'INITIAL','filename':'stock.csv','data':data,'reason':'Stock vérifié'}
        first=preview_import(self.ops,**kwargs)
        second=preview_import(self.ops,**kwargs)
        self.assertEqual(first.pk,second.pk)
        self.assertEqual(ImportBatch.objects.count(),1)
        with self.assertRaises(Conflict):
            preview_import(self.ops,**{**kwargs,'reason':'Autre opération'})

    def test_native_import_upload_preview_report_and_confirmation_pages(self):
        self.client.force_login(self.ops)
        self.assertEqual(self.client.get(reverse('erp:imports')).status_code,200)
        self.assertEqual(self.client.get(reverse('erp:import-template',args=['CATALOG'])).status_code,200)
        data=csv_file([self.receipt_row()])
        response=self.client.post(reverse('erp:imports'),{'key':str(uuid.uuid4()),'kind':'INITIAL',
            'file':SimpleUploadedFile('stock.csv',data,content_type='text/csv'),'reason':'Reprise du stock réel'})
        self.assertEqual(response.status_code,302,response.context['form'].errors if response.context else '')
        batch=ImportBatch.objects.get()
        self.assertContains(self.client.get(reverse('erp:import-detail',args=[batch.pk])),'Aucune donnée métier')
        report=self.client.get(reverse('erp:import-report',args=[batch.pk]))
        book=load_workbook(BytesIO(b''.join(report.streaming_content)),data_only=False)
        self.assertEqual(book['Données']['A5'].value,2)
        result=self.client.post(reverse('erp:import-detail',args=[batch.pk]),{'expected_version':batch.version,'confirmed':'on'})
        self.assertEqual(result.status_code,302)
        self.assertEqual(StockContainer.objects.get().quantity,10)
        self.assertEqual(self.client.post(reverse('erp:import-detail',args=[batch.pk]),{'expected_version':batch.version,'confirmed':'on'}).status_code,302)


    def test_http_import_failures_cancellation_and_access_boundaries(self):
        self.client.force_login(self.ops)
        invalid = SimpleUploadedFile('broken.csv', b'wrong;header\nA;B\n')
        response = self.client.post(reverse('erp:imports'), {'key': str(uuid.uuid4()),
            'kind': 'CATALOG', 'reason': 'Fichier invalide', 'file': invalid})
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.context['form'].errors)
        self.assertFalse(ImportBatch.objects.exists())
        batch = self.preview('RECEIPTS', [self.receipt_row()])
        detail = reverse('erp:import-detail', args=[batch.pk])
        cancel = reverse('erp:import-cancel', args=[batch.pk])
        self.assertEqual(self.client.get(cancel).status_code, 200)
        stale = self.client.post(detail, {'expected_version': batch.version + 1, 'confirmed': 'on'})
        self.assertEqual(stale.status_code, 400)
        self.assertFalse(StockContainer.objects.exists())
        response = self.client.post(cancel, {'expected_version': batch.version + 1, 'reason': 'Abandon'})
        self.assertEqual(response.status_code, 400)
        response = self.client.post(cancel, {'expected_version': batch.version, 'reason': 'Reprise à corriger'})
        self.assertEqual(response.status_code, 302)
        batch.refresh_from_db()
        self.assertEqual(batch.status, 'CANCELLED')
        response = self.client.post(cancel, {'expected_version': batch.version, 'reason': 'Déjà abandonné'})
        self.assertEqual(response.status_code, 400)
        response = self.client.post(detail, {'expected_version': batch.version, 'confirmed': 'on'})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(StockContainer.objects.exists())
        self.assertEqual(self.client.get(reverse('erp:import-template', args=['UNKNOWN'])).status_code, 404)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse('erp:imports')).status_code, 403)
        self.assertEqual(self.client.get(reverse('erp:import-template', args=['CATALOG'])).status_code, 403)


class TableIntakeTests(SimpleTestCase):
    def setUp(self):
        self.helpers=runpy.run_path(str(Path(table_probe.__file__).resolve().parents[1]/'core/document_probe.py'))

    def test_empty_template_is_not_imported_as_fake_data(self):
        data=import_template('CATALOG')
        with self.assertRaises(ValidationError):
            parse_table('CATALOG','catalogue.xlsx',data)
        book=load_workbook(BytesIO(data))
        sheet=book['Données']
        sheet.append(['CODE','Article','REAGENTS','PIECE'])
        buffer=BytesIO()
        book.save(buffer)
        rows=parse_table('CATALOG','catalogue.xlsx',buffer.getvalue())
        self.assertEqual(rows[0]['data']['code'],'CODE')

    def test_formula_macros_external_content_and_unknown_headers_are_rejected(self):
        data=xlsx_file([['code','name','category_code','base_unit_code'],['CODE','=SUM(1,2)','REAGENTS','PIECE']])
        with self.assertRaises(ValueError):
            table_probe.read_matrix(data,'.xlsx',self.helpers)
        with self.assertRaises(ValidationError):
            parse_table('CATALOG','catalogue.csv',b'code;unknown\nX;Y\n')
        with self.assertRaises(ValidationError):
            parse_table('CATALOG','old.xls',b'old binary format')
        valid=xlsx_file([['code'],['A']])
        with zipfile.ZipFile(BytesIO(valid)) as source:
            output=BytesIO()
            with zipfile.ZipFile(output,'w') as target:
                for info in source.infolist():
                    target.writestr(info,source.read(info.filename))
                target.writestr('xl/vbaProject.bin',b'active content')
        with self.assertRaises(ValueError):
            table_probe.read_matrix(output.getvalue(),'.xlsx',self.helpers)

    def test_utf16_dtd_and_duplicate_zip_parts_are_not_accepted(self):
        data=xlsx_file([['code'],['A']])
        output=BytesIO()
        with zipfile.ZipFile(BytesIO(data)) as source,zipfile.ZipFile(output,'w') as target:
            for info in source.infolist():
                raw=source.read(info.filename)
                if info.filename=='xl/worksheets/sheet1.xml':
                    raw='<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE a [<!ENTITY value "blocked">]><a>&value;</a>'.encode('utf-16')
                target.writestr(info,raw)
        with self.assertRaises(ValueError):
            table_probe.read_matrix(output.getvalue(),'.xlsx',self.helpers)

    def test_dates_booleans_row_limits_and_control_characters_are_checked(self):
        self.assertEqual(table_probe.normalize(True),'true')
        self.assertEqual(table_probe.normalize(False),'false')
        self.assertEqual(table_probe.normalize(None),'')
        self.assertEqual(table_probe.normalize(datetime(2026,9,21,10,30)),'2026-09-21T10:30:00')
        with self.assertRaises(ValueError):
            table_probe.normalize(chr(0))
        with self.assertRaises(ValueError):
            table_probe.normalize('x'*10001)
        with self.assertRaises(ValueError):
            table_probe.read_matrix(b'a;b\n'+b'1;2\n'*501,'.csv',self.helpers)
        with self.assertRaises(ValueError):
            table_probe.read_matrix(b'','.csv',self.helpers)
