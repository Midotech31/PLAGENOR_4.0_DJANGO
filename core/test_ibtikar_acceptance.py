from plagenor.test_support import close_response
import io
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.signing import TimestampSigner
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from docx import Document

from accounts.models import User, MemberProfile
from core.exceptions import PricingConfigurationError, InvalidTransitionError
from core.ibtikar.forms import RetainedField, SchemaForm, make_sample_formset
from core.ibtikar.schema import get_schema, schema_for_service, matches, display_value
from core.ibtikar.legacy import mapped_values, legacy_initial
from core.ibtikar.pricing import amount, resolve_schema_cost
from core.ibtikar.services import save_submission, save_staff, record_code
from core.ibtikar.models import IbtikarSubmission
from core.models import Service, ServiceFormField, ServicePricing, Request
from core import test_ibtikar_forms as contracts
from core.test_ibtikar_forms import STATIC, fixture, posted, document_text, upload_image, upload_document


class EdgeSchemaContracts(SimpleTestCase):
    def test_inactive_payload_still_has_a_length_limit(self):
        with self.assertRaises(ValidationError):
            RetainedField().clean('x' * 20001)

    def test_boolean_and_multichoice_display_are_not_defaults(self):
        self.assertEqual(display_value({}, False, 'en'), 'No')
        self.assertEqual(display_value({}, True, 'ar'), 'نعم')
        self.assertEqual(display_value({}, ['unmapped', 'raw'], 'en'), 'unmapped ; raw')
        self.assertTrue(matches({'all': [{'field':'a','in':['yes']},{'field':'b','in':['yes']}]}, {'a':'yes','b':'yes'}))

    def test_missing_service_has_no_invented_reference(self):
        schema = schema_for_service(None)
        self.assertEqual(schema['service_code'], '')
        self.assertEqual(schema['parameters'], [])
        self.assertEqual(schema['samples'], [])

    def test_invalid_upload_fails_form_validation(self):
        schema = get_schema('EGTP-IMT')
        bad = upload_document(); bad.content_type = 'text/html'
        form = SchemaForm({}, {'biosafety_declaration':bad}, specs=schema['attachments'], samples=[{'risk_status':'clinical'}])
        self.assertFalse(form.is_valid())
        self.assertIn('biosafety_declaration', form.errors)

    def test_duplicate_samples_and_incompatible_units_are_rejected(self):
        schema,a,p,rows = fixture('EGTP-GDE',2)
        rows[1]['sample_code'] = rows[0]['sample_code']
        rows[0]['quantity_unit'] = 'g'
        forms = make_sample_formset(schema, posted(a,p,rows), parameters=p)
        self.assertFalse(forms.is_valid())
        self.assertIn('quantity_unit', forms.errors[0])
        self.assertIn('__all__', forms.errors[1])
        schema,a,p,rows = fixture('EGTP-IMT')
        rows[0].update(origin='clinical', risk_status='standard')
        forms = make_sample_formset(schema,posted(a,p,rows),parameters=p)
        self.assertFalse(forms.is_valid())
        self.assertIn('risk_status',forms.errors[0])

    def test_legacy_unknown_choice_is_not_replaced_with_first_option(self):
        schema = get_schema('EGTP-CAN')
        self.assertNotIn('nucleic_acid_type',mapped_values(schema['samples'],{'nucleic_acid_type':'unspecified historical'}))
        req=SimpleNamespace(service_params={},requester_data={},title='Old',guest_name='Guest',guest_email='guest@example.test',guest_phone='0123',sample_table=[],pricing={})
        old=legacy_initial(req,schema)
        self.assertEqual(old['applicant']['email'],'guest@example.test')
        self.assertEqual(old['applicant']['full_name'],'Guest')

    def test_nonfinite_tariff_and_non_numeric_tariff_fail_closed(self):
        for value in ('not-money','NaN','Infinity','-1'):
            with self.subTest(value=value), self.assertRaises(PricingConfigurationError):
                amount(value)


@override_settings(STORAGES=STATIC, EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend', PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], RATE_LIMIT_BACKEND='cache', DOCUMENT_PDF_ENABLED=False)
class AcceptanceFlows(TestCase):
    setUpTestData = contracts.PersistenceContractTests.__dict__['setUpTestData']
    setUp = contracts.PersistenceContractTests.setUp
    tearDown = contracts.PersistenceContractTests.tearDown
    submit = contracts.PersistenceContractTests.submit

    def test_custom_service_and_admin_fields_preserve_explicit_options(self):
        service=Service.objects.create(code='CUSTOM-IBK',name='Custom',name_fr='Sur mesure',name_en='Custom',name_ar='خاصة')
        ServiceFormField.objects.create(service=service,name='specific_choice',label='Choice',field_type='enum',options=['A','B'])
        ServiceFormField.objects.create(service=service,name='sample_ready',label='Ready',field_category='sample_column',field_type='boolean',conditional_logic=[{'trigger_field':'sample_type','trigger_value':'water'}])
        schema=schema_for_service(service)
        self.assertEqual(schema['source']['kind'],'administrator-defined')
        self.assertIn('specific_choice',[x['name'] for x in schema['parameters']])
        self.assertEqual(next(x for x in schema['samples'] if x['name']=='sample_ready')['options'][1]['value'],'no')
        pcr=Service.objects.get(code='EGTP-PCR')
        ServiceFormField.objects.create(service=pcr,name='pcr_kit',label='Approved kits',field_type='enum',options=['Approved A'])
        self.assertEqual(next(x for x in schema_for_service(pcr)['parameters'] if x['name']=='pcr_kit')['options'][0]['value'],'Approved A')

    def test_legacy_import_payload_limit_and_guest_query_routing(self):
        for query,target in [({'channel':'IBTIKAR'},reverse('ibtikar:index')),({'channel':'IBTIKAR','service':'EGTP-CAN'},reverse('ibtikar:new',args=['EGTP-CAN']))]:
            self.assertRedirects(self.client.get(reverse('guest_submit'),query),target,fetch_redirect_response=False)
        service=Service.objects.get(code='EGTP-CAN')
        response=self.client.post(reverse('dashboard:requester_create'),{'service_id':str(service.pk),'param_description':'x'*500001})
        self.assertEqual(response.status_code,400)
        self.assertEqual(IbtikarSubmission.objects.count(),0)
        response=self.client.post(reverse('guest_submit'),{'guest_name':'Guest','guest_email':'g@example.test','service_id':'bad-uuid'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(IbtikarSubmission.objects.count(),0)

    def test_guest_conversion_opens_the_ibtikar_index(self):
        self.client.logout();obj=self.submit()
        token=TimestampSigner(salt='guest-conversion').sign(obj.request.guest_email)
        response=self.client.post(reverse('accounts:convert_guest_verify',args=[token]),{'password':'Fixture-conversion-only!2026'})
        self.assertRedirects(response,reverse('ibtikar:index'),fetch_redirect_response=False)
        obj.request.refresh_from_db();self.assertIsNotNone(obj.request.requester_id)
        self.assertContains(self.client.get(reverse('ibtikar:index')),obj.request.display_id)

    def test_staff_assignment_permissions_and_closed_requests(self):
        from core.ibtikar.views import may_read
        obj=self.submit();member=User.objects.create_user(username='assigned-reader',role='MEMBER')
        profile=MemberProfile.objects.get(user=member)
        self.assertFalse(may_read(member,obj.request))
        obj.request.informed_members.add(profile);self.assertTrue(may_read(member,obj.request))
        obj.request.assigned_to=profile;obj.request.save(update_fields=['assigned_to'])
        self.client.force_login(member)
        response=self.client.get(reverse('ibtikar:staff',args=[obj.request_id]));self.assertEqual(response.status_code,200)
        self.assertNotIn('validated_price',response.context['form'].fields)
        save_staff(obj.pk,{'condition':'Intact'},member,obj.revision)
        obj.refresh_from_db();self.assertEqual(obj.staff['condition'],'Intact')
        with self.assertRaises(ValidationError):save_staff(obj.pk,{'validated_price':'100'},member,obj.revision)
        self.assertEqual(self.client.get(reverse('ibtikar:new',args=['EGTP-CAN'])).status_code,403)
        obj.request.status='COMPLETED';obj.request.save(update_fields=['status'])
        self.assertEqual(self.client.get(reverse('ibtikar:edit',args=[obj.request_id])).status_code,403)
        with self.assertRaises(ValidationError):save_staff(obj.pk,{},self.ops,obj.revision)

    def test_staff_conflicting_invalid_and_overbudget_values_are_not_persisted(self):
        obj=self.submit('EGTP-GDE')
        for values,actor,revision in [({},self.ops,0),({},self.requester,1),({'validated_price':'10','price_justification':'short'},self.ops,1),({'validated_price':'NaN','price_justification':'Explicit pricing reason'},self.ops,1),({'validated_price':'200001','price_justification':'Explicit pricing reason'},self.ops,1)]:
            with self.subTest(values=values),self.assertRaises(ValidationError):save_staff(obj.pk,values,actor,revision)
        save_staff(obj.pk,{'validated_price':'100','price_justification':'Explicit pricing reason'},self.ops,1)
        with self.assertRaises(ValidationError):save_staff(obj.pk,{},self.ops,2)
        self.client.force_login(self.ops)
        response=self.client.post(reverse('ibtikar:staff',args=[obj.request_id]),{'revision':'bad'})
        self.assertEqual(response.status_code,400)
        obj.refresh_from_db();self.assertEqual(Decimal(obj.estimate['total']),100)

    def test_revision_errors_and_draft_rules_preserve_saved_values(self):
        obj=self.submit();url=reverse('ibtikar:edit',args=[obj.request_id])
        response=self.client.post(url,posted(obj.applicant,obj.parameters,obj.samples,revision='invalid'))
        self.assertEqual(response.status_code,400)
        response=self.client.post(url,posted(obj.applicant,obj.parameters,obj.samples,revision=obj.revision,action='draft'))
        self.assertEqual(response.status_code,400)
        schema,a,p,rows=fixture('EGTP-CAN')
        self.assertEqual(self.client.post(reverse('ibtikar:new',args=['EGTP-CAN']),posted(a,p,[])).status_code,400)
        obj.request.status='ASSIGNED';obj.request.save(update_fields=['status'])
        with self.assertRaises(ValidationError):save_submission(service=obj.request.service,schema=obj.schema,applicant=obj.applicant,parameters=obj.parameters,samples=obj.samples,files={},actor=self.requester,req=obj.request,revision=obj.revision)
        for balance in ('-1','200001'):
            a['declared_balance']=balance
            with self.assertRaises(ValidationError):save_submission(service=obj.request.service,schema=obj.schema,applicant=a,parameters=p,samples=rows,files={},actor=self.requester)
        a['declared_balance']='1'
        with self.assertRaises(ValidationError):save_submission(service=obj.request.service,schema=obj.schema,applicant=a,parameters=p,samples=rows,files={},actor=self.requester)

    def test_attachment_rollbacks_leave_no_orphan_and_guest_download_works(self):
        schema,a,p,rows=fixture('EGTP-CAN');service=Service.objects.get(code='EGTP-CAN')
        with patch('core.ibtikar.models.IbtikarRevision.objects.create',side_effect=RuntimeError('rollback')):
            with self.assertRaises(RuntimeError):save_submission(service=service,schema=schema,applicant=a,parameters=p,samples=rows,files={'applicant_signature':upload_image()},actor=self.requester)
        self.assertEqual(IbtikarSubmission.objects.count(),0)
        self.assertFalse(any(x.is_file() for x in Path(self.folder.name).rglob('*')))
        self.client.logout();data=posted(a,p,rows);data['attachments-applicant_signature']=upload_image()
        self.assertEqual(self.client.post(reverse('ibtikar:new',args=['EGTP-CAN']),data).status_code,302)
        obj=IbtikarSubmission.objects.latest('pk');file=obj.attachments.get()
        edit=self.client.get(reverse('ibtikar:guest_edit',args=[obj.request.guest_token]));self.assertEqual(edit.status_code,200)
        response=self.client.get(reverse('ibtikar:attachment',args=[file.pk]),{'access':str(obj.request.guest_token)})
        self.assertEqual(response.status_code,200);self.assertTrue(b''.join(response.streaming_content));close_response(response)
        file.file.delete(save=False)
        self.assertEqual(self.client.get(reverse('ibtikar:attachment',args=[file.pk]),{'access':str(obj.request.guest_token)}).status_code,404)

    def test_reference_stage_and_owner_are_checked(self):
        obj=self.submit();req=obj.request
        with self.assertRaises(ValidationError):record_code(req,'REFERENCE',self.requester)
        self.assertEqual(self.client.post(reverse('ibtikar:code',args=[req.pk]),{'ibtikar_code':''}).status_code,302)
        req.status='IBTIKAR_SUBMISSION_PENDING';req.save(update_fields=['status'])
        for code,actor in [('x'*51,self.requester),('REF',None),('REF',self.ops)]:
            with self.assertRaises(ValidationError):record_code(req,code,actor)
        response=self.client.post(reverse('ibtikar:code',args=[req.pk]),{'ibtikar_code':'SAVED-REF'})
        self.assertEqual(response.status_code,302);req.refresh_from_db();self.assertEqual(req.ibtikar_external_code,'SAVED-REF')
        self.client.force_login(self.ops)
        self.assertEqual(self.client.post(reverse('ibtikar:code',args=[req.pk]),{'ibtikar_code':'ADMIN'}).status_code,404)
        req.submitted_as_guest=True;req.guest_token=__import__('uuid').uuid4();req.status='SUBMITTED';req.save()
        response=self.client.post(reverse('guest_ibtikar_code',args=[req.guest_token]),{'ibtikar_code':'EARLY'})
        self.assertEqual(response.status_code,302);req.refresh_from_db();self.assertNotEqual(req.ibtikar_external_code,'EARLY')

    def test_estimate_current_snapshot_and_invalid_values(self):
        obj=self.submit();self.client.force_login(self.ops);url=reverse('ibtikar:estimate',args=['EGTP-CAN'])+'?request='+str(obj.request_id)
        data=posted(obj.applicant,obj.parameters,obj.samples)
        self.assertIsNotNone(self.client.post(url,data).json()['total'])
        invalid=deepcopy(data);invalid['parameters-qc_methods']=['bad']
        self.assertEqual(self.client.post(url,invalid).json()['status'],'incomplete')
        with patch('core.pricing.resolve_cost',side_effect=PricingConfigurationError('missing tariff')):
            self.assertEqual(self.client.post(url,data).json()['status'],'pending')
        self.assertEqual(self.client.post(reverse('ibtikar:estimate',args=['EGTP-PCR'])+'?request='+str(obj.request_id),data).status_code,404)
        self.assertEqual(self.client.post(reverse('ibtikar:estimate',args=['EGTP-CAN'])+'?request=not-uuid',data).status_code,404)
        self.assertEqual(self.client.post(reverse('dashboard:admin_adjust_cost',args=[obj.request_id]),{}).status_code,302)

    def test_pending_price_and_technical_review_block_operational_advance(self):
        from core.workflow import transition
        obj=self.submit('EGTP-Lyoph');req=obj.request
        req.status='VALIDATION_PEDAGOGIQUE';req.save(update_fields=['status'])
        with self.assertRaises(InvalidTransitionError):transition(req,'VALIDATION_FINANCE',self.ops)
        save_staff(obj.pk,{'validated_price':'1000','price_justification':'Validated scope and cost'},self.ops,obj.revision)
        req.refresh_from_db();req.status='SAMPLE_RECEIVED';req.save(update_fields=['status'])
        with self.assertRaises(InvalidTransitionError):transition(req,'ANALYSIS_STARTED',self.ops)

    def test_document_attachments_signature_staff_and_history(self):
        from documents.generators import generate_ibtikar_form,_render_sample_table,_render_service_params
        schema,a,p,rows=fixture('EGTP-IMT');rows[0].update(risk_status='clinical',maldi_target='disposable')
        data=posted(a,p,rows);data['attachments-biosafety_declaration']=upload_document();data['attachments-applicant_signature']=upload_image()
        self.assertEqual(self.client.post(reverse('ibtikar:new',args=['EGTP-IMT']),data).status_code,302)
        obj=IbtikarSubmission.objects.latest('pk');save_staff(obj.pk,{'condition':'Réception intacte'},self.ops,obj.revision)
        doc=Document(generate_ibtikar_form(obj.request));self.assertIn('Réception intacte',document_text(doc));self.assertTrue(doc.inline_shapes)
        self.assertTrue(doc.sections[0].header._element.xpath('.//w:drawing'))
        req=Request.objects.create(display_id='LEGACY-UNASSIGNED',channel='IBTIKAR',requester=self.requester,service=None,title='Dossier historique',service_params={'unmapped':'Retained'},sample_table=[{'unmapped':'Last historical sample'}])
        text=document_text(Document(generate_ibtikar_form(req)));self.assertIn('Last historical sample',text);self.assertIn('Retained',text)
        doc=Document();_render_sample_table(doc,[]);_render_service_params(doc,[]);_render_service_params(doc,{'empty':''});_render_service_params(doc,{'custom':'Value'})
        self.assertIn('Value',document_text(doc))
        self.assertEqual(self.client.get(reverse('ibtikar:detail',args=[req.pk])).status_code,200)
        self.assertEqual(self.client.get(reverse('ibtikar:edit',args=[req.pk])).status_code,404)

    def test_unpriced_options_never_become_free_or_silent_defaults(self):
        cases=[('EGTP-SeqS',{'submitted_type':'bigdye'},{}),('EGTP-SeqS',{'submitted_type':''},{}),('EGTP-SeqS',{'submitted_type':'unpurified_pcr'},{}),('EGTP-SeqS',{'sequencing_mode':'unknown'},{}),('EGTP-Seq02',{'extraction_method':'commercial','pcr_kit':'Other','qc_methods':['gel']},{}),('EGTP-PCR',{'pcr_kit':'Other','qc_methods':['fluorimetry'],'product_recovery':'yes','product_volume_ul':'30'},{}),('EGTP-CAN',{'qc_methods':['spectrophotometry']},{}),('EGTP-PS',{'purification_method':'other','delivery_format':'solution','concentration_um':'20'},{}),('EGTP-IMT',{}, {'fresh_culture':'no','preparation':['purification','replating'],'analysis_mode':'unknown'})]
        for code,updates,row_updates in cases:
            with self.subTest(code=code,updates=updates):
                schema,a,p,rows=fixture(code);p.update(updates);rows[0].update(row_updates)
                result=resolve_schema_cost(Service.objects.get(code=code),'IBTIKAR',rows,p)
                if code=='EGTP-SeqS' and updates.get('submitted_type')=='unpurified_pcr':self.assertIsNotNone(result['total'])
                else:self.assertIsNone(result['total']);self.assertTrue(result['reasons'])
        service=Service.objects.get(code='EGTP-CAN');schema,a,p,rows=fixture(service.code)
        self.assertIsNone(resolve_schema_cost(service,'IBTIKAR',[],p)['total'])
        with self.assertRaises(PricingConfigurationError):resolve_schema_cost(service,'GENOCLAB',rows,p)
        for config in (['invalid'],{'ibtikar':['invalid']}):
            service.pricing_data=config
            with self.assertRaises(PricingConfigurationError):resolve_schema_cost(service,'IBTIKAR',rows,p)
        service.pricing_data={'base_price':1};self.assertIn('legacy_pricing_configuration_requires_review',resolve_schema_cost(service,'IBTIKAR',rows,p)['reasons'])
        service.pricing_data={};self.assertIn('urgency_price_requires_review',resolve_schema_cost(service,'IBTIKAR',rows,p,'Express')['reasons'])
        ServicePricing.objects.create(service=service,channel='IBTIKAR',name='Option to review',pricing_type='BASE',amount=100)
        self.assertIn('configured_tariffs_require_option_review',resolve_schema_cost(service,'IBTIKAR',rows,p)['reasons'])
        ServiceFormField.objects.create(service=service,name='extra_assay',label='Additional assay',field_type='string',affects_pricing=True)
        self.assertIn('administrator_option_tariff_requires_review',resolve_schema_cost(service,'IBTIKAR',rows,{**p,'extra_assay':'Requested'})['reasons'])

    def test_field_dependencies_reject_cycles_and_accept_reverse_order(self):
        from core.ibtikar.schema import active_names
        fields=[{'name':'child','when':{'field':'parent','in':['yes']}},{'name':'parent'}]
        self.assertEqual(active_names(fields,{'parent':'yes'}),{'parent','child'})
        with self.assertRaises(ValueError):
            active_names([{'name':'cycle','when':{'field':'cycle','in':['yes']}}],{'cycle':'yes'})

    def test_invalid_sample_value_is_reported_without_rule_evaluation(self):
        schema,a,p,rows=fixture('EGTP-Lyoph');rows[0]['quantity']='not-a-number'
        forms=make_sample_formset(schema,posted(a,p,rows),parameters=p)
        self.assertFalse(forms.is_valid());self.assertIn('quantity',forms.errors[0])

    def test_admin_field_does_not_replace_a_canonical_business_enum(self):
        service=Service.objects.get(code='EGTP-CAN')
        ServiceFormField.objects.create(service=service,name='qc_methods',label='Legacy',field_type='enum',options=['invalid replacement'])
        spec=next(x for x in schema_for_service(service)['parameters'] if x['name']=='qc_methods')
        self.assertIn('gel',[x['value'] for x in spec['options']])

    def test_empty_stored_schema_is_not_rendered_or_priced(self):
        obj=self.submit();IbtikarSubmission.objects.filter(pk=obj.pk).update(schema={})
        self.assertEqual(self.client.get(reverse('ibtikar:edit',args=[obj.request_id])).status_code,404)
        self.client.force_login(self.ops)
        response=self.client.post(reverse('ibtikar:estimate',args=['EGTP-CAN'])+'?request='+str(obj.request_id),{})
        self.assertEqual(response.status_code,404)

    def test_inactive_risk_attachment_is_preserved_but_not_printed(self):
        from documents.generators import generate_ibtikar_form
        schema,a,p,rows=fixture('EGTP-IMT');rows[0].update(risk_status='clinical',maldi_target='disposable')
        data=posted(a,p,rows);data['attachments-biosafety_declaration']=upload_document()
        response=self.client.post(reverse('ibtikar:new',args=['EGTP-IMT']),data)
        self.assertEqual(response.status_code,302)
        obj=IbtikarSubmission.objects.latest('pk');saved_file=obj.attachments.get()
        rows[0].update(risk_status='standard',maldi_target='reusable')
        response=self.client.post(reverse('ibtikar:edit',args=[obj.request_id]),posted(a,p,rows,revision=obj.revision))
        self.assertEqual(response.status_code,302);obj.refresh_from_db()
        self.assertTrue(obj.attachments.filter(pk=saved_file.pk,active=True).exists())
        rendered=document_text(Document(generate_ibtikar_form(obj.request)))
        self.assertNotIn(saved_file.original_name,rendered)

    def test_two_hundred_rows_survive_the_http_parser_and_draft_save(self):
        schema,a,p,rows=fixture('EGTP-CAN',200)
        data=posted(a,p,rows,action='draft')
        self.assertGreater(len(data),1000)
        response=self.client.post(reverse('ibtikar:new',args=['EGTP-CAN']),data)
        self.assertEqual(response.status_code,302)
        obj=IbtikarSubmission.objects.latest('pk')
        self.assertEqual(len(obj.samples),200)
        self.assertEqual(obj.samples[-1]['sample_code'],'S200')
        self.assertEqual(obj.request.status,'DRAFT')
        self.assertIsNone(obj.estimate['total'])

    def test_reference_checklist_and_instructions_are_service_specific(self):
        from core.ibtikar.schema import projection
        expected={'EGTP-PCR':'primers','EGTP-SeqS':'dna_qc','EGTP-PS':'sequence_format','EGTP-Illumina-Microbial-WGS':'medium_quantity'}
        for code,key in expected.items():
            schema,a,p,rows=fixture(code)
            field=next(f for f in schema['staff'] if f['name']=='checklist')
            self.assertIn(key,[o['value'] for o in field['options']])
            output=projection(schema,a,p,rows,print_blank_staff=True)
            checklist=next(r for r in output['staff'] if r['name']=='checklist')
            self.assertTrue(checklist['all_options'])
            self.assertTrue(all(not o['selected'] for o in checklist['options']))
        self.assertTrue(any('15 µL' in n['fr'] for n in get_schema('EGTP-CAN')['notices']))
        self.assertTrue(any('FASTQ' in n['fr'] for n in get_schema('EGTP-Illumina-Microbial-WGS')['notices']))

    def test_internal_tariff_does_not_leak_through_staff_summary_or_form(self):
        from documents.generators import generate_ibtikar_form
        obj=self.submit('EGTP-GDE')
        save_staff(obj.pk,{'validated_price':'1234.56','price_justification':'INTERNAL-TARIFF-TEST-ONLY'},self.ops,obj.revision)
        obj.refresh_from_db()
        response=self.client.get(reverse('ibtikar:detail',args=[obj.request_id]))
        self.assertNotContains(response,'INTERNAL-TARIFF-TEST-ONLY')
        rendered=document_text(Document(generate_ibtikar_form(obj.request)))
        self.assertNotIn('INTERNAL-TARIFF-TEST-ONLY',rendered)
        self.assertNotIn('1234.56',rendered)
        self.client.force_login(self.ops)
        self.assertContains(self.client.get(reverse('ibtikar:detail',args=[obj.request_id])),'INTERNAL-TARIFF-TEST-ONLY')

    def test_nested_dependencies_follow_parent_applicability(self):
        from core.ibtikar.schema import active_names
        fields=[{'name':'technique'},{'name':'other','when':{'field':'technique','in':['other']}},
                {'name':'details','when':{'all':[{'field':'technique','in':['other']},
                    {'any':[{'field':'other','in':['specific']},{'field':'technique','in':['special']}]}]}}]
        self.assertEqual(active_names(fields,{'technique':'other','other':'specific'}),{'technique','other','details'})
        self.assertEqual(active_names(fields,{'technique':'standard','other':'specific'}),{'technique'})

    def test_schema_editor_uses_its_own_validation_and_submit_actions(self):
        response=self.client.get(reverse('ibtikar:new',args=['EGTP-CAN']))
        self.assertContains(response,'class="no-validation" id="ibk-editor"')
        self.assertContains(response,'value="draft" formnovalidate')
        self.assertContains(response,'name="action" value="submit"')
