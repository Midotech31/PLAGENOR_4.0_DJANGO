import io
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import Http404
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse

from accounts.models import User
from core.models import Request
from core.uploads import validate_upload
from dashboard.views.report import protected_report_media, serve_media
from documents import genoclab_layout as layout
from plagenor.test_documents import valid_pdf_bytes


@override_settings(STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ReleaseBoundaryContracts(TestCase):
    def test_ordinary_and_report_storage_errors_are_denied_before_open(self):
        user = User.objects.create_user('boundary-owner', role='CLIENT')
        Request.objects.create(requester=user, channel='GENOCLAB', order_file='orders/one.pdf', report_file='reports/one.pdf')
        request = RequestFactory().get('/'); request.user=user
        with self.assertRaises(Http404): protected_report_media(request, '../orders/one.pdf')
        with patch('dashboard.views.report.validate_storage_path', side_effect=ValidationError('invalid')), patch('dashboard.views.report.default_storage') as storage:
            with self.assertRaises(Http404): serve_media(request, 'orders/one.pdf')
            with self.assertRaises(Http404): protected_report_media(request, 'one.pdf')
            storage.open.assert_not_called()

    def test_report_wrong_local_target_or_os_error_is_denied(self):
        user = User.objects.create_user('boundary-staff', role='MEMBER')
        req = RequestFactory().get('/'); req.user=user
        Request.objects.create(channel='GENOCLAB', report_file='reports/one.pdf')
        with patch('dashboard.views.report.validate_storage_path', side_effect=OSError('unavailable')):
            with self.assertRaises(Http404): protected_report_media(req, 'one.pdf')
            with self.assertRaises(Http404): serve_media(req, 'documents/one.pdf')

    def test_required_service_and_role_do_not_reach_business_write(self):
        response=self.client.post(reverse('guest_submit'), {'guest_name':'Test Guest','guest_email':'guest@example.test'})
        self.assertEqual(response.status_code,200)
        self.assertFalse(Request.objects.exists())
        member=User.objects.create_user('blocked-member', role='MEMBER')
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse('dashboard:requester')).status_code,403)

    def test_declared_file_size_cannot_bypass_bounded_actual_read(self):
        for data in (valid_pdf_bytes(), b''):
            file=SimpleUploadedFile('input.pdf',data,'application/pdf'); file.size=1
            with self.assertRaises(ValidationError): validate_upload(file,'report',max_bytes=32)
            self.assertEqual(file.tell(),0)

    def test_custom_document_notices_and_blank_legacy_values(self):
        self.assertEqual(layout._money('Infinity'),'Infinity')
        self.assertEqual(layout._money_int('Infinity'),'Infinity')
        doc=Document()
        identity={'values':{'genoclab_footer_legal':'Arrêtée la présente facture à la somme de {amount_words}.','genoclab_issuer_address1':'Address'},'commercial_terms':'Custom terms'}
        layout.add_genoclab_footer(doc,total_amount=10,identity=identity,document_kind='quote')
        self.assertIn('Arrêté le présent devis',' '.join(p.text for p in doc.paragraphs))
        self.assertIn('Custom terms',' '.join(p.text for p in doc.paragraphs))
        blank=Document(); layout.add_genoclab_footer(blank,identity={'values':{}})
        self.assertTrue(blank.sections[0].footer.tables)
        cell=Document().add_table(rows=1,cols=1).cell(0,0)
        layout._cell_border(cell,'top','000000',4)
        layout._cell_border(cell,'top','000000',6)
        self.assertIn('w:sz="6"',cell._tc.xml)
        self.assertEqual(layout._quantity('3.50'),'3.5')
