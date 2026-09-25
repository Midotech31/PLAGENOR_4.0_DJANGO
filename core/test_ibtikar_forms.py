from plagenor.test_documents import valid_pdf_bytes
from plagenor.test_support import close_response
import io
import json
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
import tempfile
import uuid
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone, translation
from docx import Document
from PIL import Image

from accounts.models import User, MemberProfile
from core.models import Service, ServicePricing, ServiceFormField, Request, FinancialVisibility
from core.ibtikar.forms import SchemaForm, SequenceField, make_sample_formset, cleaned_samples
from core.ibtikar.models import IbtikarSubmission, IbtikarRevision, IbtikarAttachment
from core.ibtikar.schema import definitions, get_schema, schema_for_service, active_data, active_names, projection, schema_digest, display_value
from core.ibtikar.services import save_submission, save_staff, record_code
from core.ibtikar.legacy import legacy_initial
from core.pricing import resolve_cost
from core.registry import get_service_def


STATIC = {'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
          'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}}


def upload_image():
    buffer = io.BytesIO()
    Image.new('RGB', (80, 20), 'white').save(buffer, 'PNG')
    return SimpleUploadedFile('signature.png', buffer.getvalue(), content_type='image/png')


def upload_document():
    return SimpleUploadedFile('declaration.pdf', valid_pdf_bytes(), content_type='application/pdf')


def raw_group(specs):
    values = {}
    for f in specs:
        kind = f['type']
        if kind == 'computed':
            continue
        if kind in ('choice', 'multi'):
            value = f['options'][0]['value']
            values[f['name']] = [value] if kind == 'multi' else value
        elif kind in ('integer', 'decimal'):
            values[f['name']] = '10'
        elif kind == 'consent':
            values[f['name']] = True
        elif kind == 'email':
            values[f['name']] = 'fixture@example.test'
        elif kind == 'date':
            values[f['name']] = '2026-09-15'
        elif kind == 'sequence':
            values[f['name']] = 'ACGTACGTACGTACGTACGTACGT'
        elif kind not in ('file', 'image'):
            values[f['name']] = 'Valeur témoin ' + f['name']
    return values


def fixture(code, count=1):
    schema = get_schema(code)
    applicant = raw_group(schema['applicant'])
    applicant.update(full_name='Demandeur Fictif', institution='Institution de test',
                     phone='0555000000', status='Doctorant', analysis_frame='phd',
                     project_title='Projet scientifique de contrôle', supervisor='Encadrant Fictif',
                     ibtikar_id='IDGRSTD12345', declared_balance='200000')
    parameters = raw_group(schema['parameters'])
    parameters.update({k: v for k, v in {
        'qc_methods': ['spectrophotometry'], 'gel_percentage': '2', 'size_marker': '1 kb',
        'submitted_type': 'purified_pcr', 'sequencing_mode': 'forward', 'pcr_kit': 'DreamTaq',
        'sequencing_depth': 'standard', 'duration_units_24h': '1', 'product_volume_ul': '25',
    }.items() if k in parameters})
    if code == 'EGTP-CAN': parameters['qc_methods'] = ['fluorimetry', 'gel']
    if code == 'EGTP-PCR': parameters.update(pcr_kit='Standard Taq', qc_methods=['none'])
    rows = []
    for index in range(count):
        row = raw_group(schema['samples'])
        row.update({k: v for k, v in {'sample_code': f'S{index+1:03d}', 'amplicon_bp': '500',
                    'risk_status': 'standard', 'sample_type': 'culture', 'quantity': '10',
                    'quantity_unit': 'mL', 'container_volume_ml': '100', 'fill_volume_ml': '25',
                    'pair_code': f'P{index//2+1}', 'primer_name': f'PRIMER{index+1}',
                    'direction': 'forward' if index % 2 == 0 else 'reverse',
                    'maldi_target': 'reusable', 'analysis_mode': 'single',
                    'biological_origin': 'microbial'}.items() if k in row})
        for f in schema['samples']:
            if f['type'] == 'choice' and row.get(f['name']) not in [o['value'] for o in f['options']]:
                row[f['name']] = f['options'][0]['value']
        rows.append(row)
    return schema, applicant, parameters, rows


def posted(applicant, parameters, rows, action='submit', revision=0):
    data = {'action': action, 'revision': str(revision), 'samples-TOTAL_FORMS': str(len(rows)),
            'samples-INITIAL_FORMS': '0', 'samples-MIN_NUM_FORMS': '0', 'samples-MAX_NUM_FORMS': '200'}
    data.update({'applicant-' + k: v for k, v in applicant.items()})
    data.update({'parameters-' + k: v for k, v in parameters.items()})
    for index, row in enumerate(rows): data.update({f'samples-{index}-{k}': v for k, v in row.items()})
    return data


def activate(rule, values, parameters=None, rows=None):
    if not rule: return
    if 'all' in rule:
        for child in rule['all']: activate(child, values, parameters, rows)
    elif 'any' in rule:
        activate(rule['any'][0], values, parameters, rows)
    elif 'any_row' in rule:
        activate(rule['any_row'], rows[0], parameters, rows)
    else:
        target = parameters if rule.get('scope') == 'parameters' else values
        key = rule['field']; value = rule['in'][0]
        target[key] = [value] if isinstance(target.get(key), list) else value


def document_text(doc):
    return '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)


class SchemaContractTests(SimpleTestCase):
    def test_source_inventory_and_stable_labels(self):
        self.assertEqual(len(definitions()['services']), 10)
        for code, schema in definitions()['services'].items():
            for group in ('applicant', 'parameters', 'samples', 'staff', 'attachments'):
                names = [f['name'] for f in schema[group]]
                self.assertEqual(len(names), len(set(names)), (code, group))
                for f in schema[group]:
                    self.assertTrue(all(f['label'].get(lang) for lang in ('fr', 'en', 'ar')))
                    self.assertNotIn('default', f)
            self.assertEqual(len(schema_digest(schema)), 64)

    def test_inactive_values_are_retained_but_not_projected(self):
        schema, applicant, params, rows = fixture('EGTP-CAN')
        params.update(qc_methods=['spectrophotometry'], gel_percentage='3', size_marker='100 pb')
        form = SchemaForm(params, specs=schema['parameters'])
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['gel_percentage'], '3')
        for lang in ('fr', 'en', 'ar'):
            result = projection(schema, applicant, form.cleaned_data, rows, language=lang)
            self.assertNotIn('gel_percentage', [x['name'] for x in result['parameters']])
            self.assertNotIn('size_marker', [x['name'] for x in result['parameters']])

    def test_sequences_are_normalized_without_invented_tm(self):
        self.assertEqual(SequenceField().clean('ac gt\nTA'), 'ACGTTA')
        with self.assertRaises(ValidationError): SequenceField().clean('ACGT-E')
        schema, applicant, params, rows = fixture('EGTP-PS')
        rows[0].update(sequence='ACGTN', tm=None)
        out = projection(schema, applicant, params, rows)['samples'][0]
        values = {x['name']: x['value'] for x in out}
        self.assertEqual(values['length_nt'], 5)
        self.assertIsNone(values['gc_percent'])
        self.assertNotIn('tm', values)

    def test_multi_selection_and_none_exclusivity(self):
        schema = get_schema('EGTP-PCR'); data = raw_group(schema['parameters'])
        data['qc_methods'] = ['spectrophotometry', 'gel']
        form = SchemaForm(data, specs=schema['parameters'])
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['qc_methods'], data['qc_methods'])
        data['qc_methods'] = ['none', 'gel']
        self.assertFalse(SchemaForm(data, specs=schema['parameters']).is_valid())

    def test_conditional_risk_document_requires_file_not_a_checkbox(self):
        schema = get_schema('EGTP-IMT')
        form = SchemaForm({}, specs=schema['attachments'], samples=[{'risk_status': 'clinical'}])
        self.assertFalse(form.is_valid()); self.assertIn('biosafety_declaration', form.errors)
        form = SchemaForm({}, {'biosafety_declaration': upload_document()}, specs=schema['attachments'], samples=[{'risk_status': 'clinical'}])
        self.assertTrue(form.is_valid(), form.errors)

    def test_primer_rows_not_artificially_limited_to_paper_grid(self):
        schema, a, p, rows = fixture('EGTP-PS', 24)
        fs = make_sample_formset(schema, posted(a, p, rows), parameters=p)
        self.assertTrue(fs.is_valid(), fs.errors)
        self.assertEqual(len(cleaned_samples(fs)), 24)

    def test_sample_rules_and_technical_limit(self):
        for code, updates, field in [('EGTP-IMT', {'risk_status':'pathogenic','maldi_target':'reusable'}, 'maldi_target'),
            ('EGTP-Lyoph', {'container_volume_ml':'100','fill_volume_ml':'51'}, 'fill_volume_ml'),
            ('EGTP-GDE', {'sample_type':'culture','quantity':'9','quantity_unit':'mL'}, 'quantity')]:
            schema,a,p,rows=fixture(code);rows[0].update(updates)
            fs=make_sample_formset(schema,posted(a,p,rows),parameters=p)
            self.assertFalse(fs.is_valid(), code);self.assertIn(field, fs.errors[0])
        schema,a,p,rows=fixture('EGTP-CAN',201)
        fs=make_sample_formset(schema,posted(a,p,rows),parameters=p)
        self.assertFalse(fs.is_valid());self.assertTrue(fs.non_form_errors())


def field_contract(code, group, name):
    def test(self):
        schema=get_schema(code); specs=schema[group]; spec=next(f for f in specs if f['name']==name)
        values=raw_group(specs);params=raw_group(schema['parameters']);rows=[raw_group(schema['samples'])]
        activate(spec.get('when'),values,params,rows)
        if code == 'EGTP-IMT' and name == 'preparation_other': values['fresh_culture'] = 'no'
        form=SchemaForm(values,specs=specs,parameters=params,samples=rows,require_complete=False)
        self.assertTrue(form.is_valid(),form.errors)
        if spec['type'] in ('choice','multi'):
            for option in spec['options']:
                with self.subTest(option=option['value']):
                    values[name]=[option['value']] if spec['type']=='multi' else option['value']
                    f=SchemaForm(values,specs=specs,parameters=params,samples=rows,require_complete=False)
                    self.assertTrue(f.is_valid(),f.errors)
                    self.assertEqual(f.cleaned_data[name],values[name])
            values[name]=['__invalid__'] if spec['type']=='multi' else '__invalid__'
            self.assertFalse(SchemaForm(values,specs=specs,parameters=params,samples=rows,require_complete=False).is_valid())
        if spec['required']:
            values=raw_group(specs);activate(spec.get('when'),values,params,rows);values.pop(name,None)
            if code == 'EGTP-IMT' and name == 'preparation_other': values['fresh_culture'] = 'no'
            f=SchemaForm(values,specs=specs,parameters=params,samples=rows)
            f.is_valid();self.assertIn(name,f.errors,(code,group,name))
    return test

for code,schema in definitions()['services'].items():
    for group in ('applicant','parameters','samples','staff'):
        for spec in schema[group]:
            if spec['type']!='computed':
                setattr(SchemaContractTests,'test_'+code.replace('-','_')+'_'+group+'_'+spec['name'],field_contract(code,group,spec['name']))


@override_settings(STORAGES=STATIC, EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
                   PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],
                   RATE_LIMIT_BACKEND='cache', DOCUMENT_PDF_ENABLED=False)
class PersistenceContractTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('seed_services', stdout=io.StringIO())
        cls.requester=User.objects.create_user(username='ibk-test-owner',email='owner@example.test',role='REQUESTER',password='Test-owner-only!2026')
        cls.ops=User.objects.create_user(username='ibk-test-ops',email='ops@example.test',role='PLATFORM_ADMIN',password='Test-ops-only!2026')
        cls.client_user=User.objects.create_user(username='ibk-test-client',email='client@example.test',role='CLIENT',password='Test-client-only!2026')

    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        self.folder=tempfile.TemporaryDirectory(prefix='plagenor-forms-test-')
        self.storage_override=override_settings(MEDIA_ROOT=self.folder.name)
        self.storage_override.enable()
        self.client.force_login(self.requester)

    def tearDown(self):
        self.storage_override.disable();self.folder.cleanup()

    def submit(self, code='EGTP-CAN', count=1, **changes):
        schema,a,p,rows=fixture(code,count)
        p.update(changes)
        service=Service.objects.get(code=code)
        response=self.client.post(reverse('ibtikar:new',args=[code]),posted(a,p,rows))
        self.assertEqual(response.status_code,302,(code,response.content[:1800]))
        return IbtikarSubmission.objects.select_related('request').latest('pk')

    def test_persistence_edit_and_document_agree(self):
        obj=self.submit(); req=obj.request
        detail=self.client.get(reverse('ibtikar:detail',args=[req.pk]));self.assertEqual(detail.status_code,200)
        self.assertContains(detail,'1 kb')
        edit=self.client.get(reverse('ibtikar:edit',args=[req.pk]));self.assertEqual(edit.status_code,200)
        self.assertEqual(edit.context['parameter_form'].initial,obj.parameters)
        params=deepcopy(obj.parameters);params['qc_methods']=['spectrophotometry'];params['size_marker']='Retained marker'
        data=posted(obj.applicant,params,obj.samples,revision=obj.revision)
        result=self.client.post(reverse('ibtikar:edit',args=[req.pk]),data);self.assertEqual(result.status_code,302)
        obj.refresh_from_db();self.assertEqual(obj.parameters['size_marker'],'Retained marker')
        self.assertEqual(obj.revisions.count(),2)
        from documents.generators import generate_ibtikar_form, build_field_map
        req.refresh_from_db()
        doc=Document(generate_ibtikar_form(req));text=document_text(doc)
        self.assertNotIn('Retained marker',text);self.assertNotIn('100 pb',text)
        self.assertEqual(build_field_map(req)['FULL_NAME'],obj.applicant['full_name'])
        self.assertContains(self.client.get(reverse('dashboard:requester_request_detail',args=[req.pk])),obj.applicant['project_title'])

    def test_all_services_round_trip_and_source_projection(self):
        from documents.generators import generate_ibtikar_form
        for code in definitions()['services']:
            with self.subTest(service=code):
                obj=self.submit(code,2 if code=='EGTP-PS' else 1)
                obj.refresh_from_db()
                self.assertEqual(obj.request.service_params['_ibtikar_schema'],'1.0.0')
                self.assertEqual(self.client.get(reverse('ibtikar:detail',args=[obj.request_id])).status_code,200)
                self.assertEqual(self.client.get(reverse('ibtikar:edit',args=[obj.request_id])).status_code,200)
                for lang in ('fr','en','ar'):
                    with translation.override(lang):
                        document=Document(generate_ibtikar_form(obj.request));text=document_text(document)
                        self.assertIn(obj.applicant['full_name'],text)
                        self.assertIn(obj.applicant['project_title'],text)
                        for row in obj.samples:
                            self.assertIn(row.get('sample_code') or row.get('primer_name'),text)
                        if code=='EGTP-PSM':self.assertNotIn('Type de prestation : Lyophilisation',text)

    def test_guest_creation_document_ownership_and_reference(self):
        self.client.logout()
        obj=self.submit();req=obj.request
        self.assertTrue(req.submitted_as_guest);self.assertIsNotNone(req.guest_token)
        self.assertEqual(self.client.get(reverse('ibtikar:detail',args=[req.pk])).status_code,404)
        self.assertEqual(self.client.get(reverse('ibtikar:guest_detail',args=[req.guest_token])).status_code,200)
        response=self.client.get(reverse('documents:guest_ibtikar_form',args=[req.guest_token]))
        self.assertEqual(response.status_code,200);b''.join(response.streaming_content);close_response(response)
        req.status='IBTIKAR_SUBMISSION_PENDING';req.save(update_fields=['status'])
        r=self.client.post(reverse('ibtikar:guest_code',args=[req.guest_token]),{'ibtikar_code':'IBK-TEST-REFERENCE'})
        self.assertEqual(r.status_code,302);req.refresh_from_db();self.assertEqual(req.status,'IBTIKAR_CODE_SUBMITTED')

    def test_draft_has_no_invented_required_values_and_can_be_completed(self):
        schema,a,p,rows=fixture('EGTP-CAN');data=posted({}, {}, [], action='draft')
        r=self.client.post(reverse('ibtikar:new',args=['EGTP-CAN']),data);self.assertEqual(r.status_code,302)
        obj=IbtikarSubmission.objects.latest('pk');self.assertIsNone(obj.submitted_at);self.assertEqual(obj.request.status,'DRAFT')
        self.assertFalse(obj.applicant.get('full_name'));self.assertIsNone(obj.estimate['total'])
        r=self.client.post(reverse('ibtikar:edit',args=[obj.request_id]),posted(a,p,rows,revision=1))
        self.assertEqual(r.status_code,302);obj.refresh_from_db();self.assertIsNotNone(obj.submitted_at)

    def test_guest_empty_draft_can_be_saved_then_completed(self):
        self.client.logout()
        response = self.client.post(reverse('ibtikar:new', args=['EGTP-CAN']), posted({}, {}, [], action='draft'))
        self.assertEqual(response.status_code, 302)
        obj = IbtikarSubmission.objects.latest('pk')
        self.assertEqual(obj.request.status, 'DRAFT')
        self.assertIsNone(obj.request.requester_id)
        schema, applicant, params, rows = fixture('EGTP-CAN')
        response = self.client.post(reverse('ibtikar:guest_edit', args=[obj.request.guest_token]), posted(applicant, params, rows, revision=obj.revision))
        self.assertEqual(response.status_code, 302)
        obj.refresh_from_db()
        self.assertIsNotNone(obj.submitted_at)
        obj.request.refresh_from_db()
        self.assertEqual(obj.request.guest_email, applicant['email'])
        self.assertEqual(obj.request.guest_name, applicant['full_name'])
        self.assertEqual(obj.request.guest_phone, applicant['phone'])

    def test_invalid_input_does_not_create_or_erase_a_request(self):
        schema,a,p,rows=fixture('EGTP-CAN');a['email']='invalid'
        r=self.client.post(reverse('ibtikar:new',args=['EGTP-CAN']),posted(a,p,rows))
        self.assertEqual(r.status_code,400);self.assertFalse(IbtikarSubmission.objects.exists())
        obj=self.submit();before=deepcopy(obj.parameters)
        p=deepcopy(obj.parameters);p['qc_methods']=['invalid']
        r=self.client.post(reverse('ibtikar:edit',args=[obj.request_id]),posted(obj.applicant,p,obj.samples,revision=1))
        self.assertEqual(r.status_code,400);obj.refresh_from_db();self.assertEqual(obj.parameters,before)
        r=self.client.post(reverse('ibtikar:edit',args=[obj.request_id]),posted(obj.applicant,obj.parameters,obj.samples,revision=0))
        self.assertEqual(r.status_code,400);self.assertEqual(obj.revisions.count(),1)

    def test_schema_and_identity_are_snapshots(self):
        obj=self.submit(); before=deepcopy(obj.schema)
        self.requester.first_name='A different current identity';self.requester.save(update_fields=['first_name'])
        ServiceFormField.objects.create(service=obj.request.service,name='qc_methods',label='Changed field',field_type='enum',options=['Other'])
        response=self.client.get(reverse('ibtikar:edit',args=[obj.request_id]));self.assertEqual(response.status_code,200)
        self.assertEqual(response.context['schema'],before)
        obj.refresh_from_db();self.assertEqual(obj.applicant['full_name'],'Demandeur Fictif')

    def test_staff_permissions_validation_and_revision(self):
        obj=self.submit('EGTP-GDE');self.assertIsNone(obj.estimate['total'])
        self.assertEqual(self.client.get(reverse('ibtikar:staff',args=[obj.request_id])).status_code,403)
        self.client.force_login(self.ops)
        response=self.client.get(reverse('ibtikar:staff',args=[obj.request_id]));self.assertEqual(response.status_code,200)
        response=self.client.post(reverse('ibtikar:staff',args=[obj.request_id]),{'revision':str(obj.revision),'staff-validated_price':'1200','staff-price_justification':'Tarif validé pour la combinaison demandée.'})
        self.assertEqual(response.status_code,302);obj.refresh_from_db();self.assertEqual(Decimal(obj.estimate['total']),1200)
        self.assertEqual(obj.revisions.count(),2)
        req=obj.request;req.refresh_from_db();self.assertEqual(req.admin_validated_price,1200)
        self.client.force_login(self.requester)
        r=self.client.post(reverse('ibtikar:edit',args=[req.pk]),posted(obj.applicant,obj.parameters,obj.samples,revision=2))
        self.assertEqual(r.status_code,302);obj.refresh_from_db();self.assertIsNone(obj.estimate['total'])

    def test_legacy_values_preserved_without_silent_backfill(self):
        service=Service.objects.get(code='EGTP-CAN')
        req=Request.objects.create(display_id='IBK-LEGACY-ONLY',title='Ancien projet',channel='IBTIKAR',status='SUBMITTED',requester=self.requester,service=service,service_params={'qc_level':'Unknown historical'},sample_table=[{'sample_code':'Historical sample','unknown':'Preserved'}])
        original=deepcopy(req.service_params);initial=legacy_initial(req,get_schema(service.code))
        self.assertEqual(initial['legacy_data']['service_params'],original)
        self.assertNotIn('email',initial['applicant'])
        self.assertEqual(self.client.get(reverse('ibtikar:edit',args=[req.pk])).status_code,200)
        schema,a,p,rows=fixture('EGTP-CAN')
        r=self.client.post(reverse('ibtikar:edit',args=[req.pk]),posted(a,p,rows,revision=0));self.assertEqual(r.status_code,302)
        obj=IbtikarSubmission.objects.get(request=req);self.assertEqual(obj.legacy_data['sample_table'][0]['unknown'],'Preserved')

    def test_attachment_binding_and_private_delivery(self):
        schema,a,p,rows=fixture('EGTP-IMT');rows[0].update(risk_status='clinical',maldi_target='disposable')
        data=posted(a,p,rows);data['attachments-biosafety_declaration']=upload_document()
        r=self.client.post(reverse('ibtikar:new',args=['EGTP-IMT']),data);self.assertEqual(r.status_code,302)
        obj=IbtikarSubmission.objects.latest('pk');attachment=obj.attachments.get(active=True)
        self.assertEqual(len(attachment.sha256),64)
        response=self.client.get(reverse('ibtikar:attachment',args=[attachment.pk]));self.assertEqual(response.status_code,200)
        b''.join(response.streaming_content);close_response(response)
        self.client.force_login(self.client_user)
        self.assertEqual(self.client.get(reverse('ibtikar:attachment',args=[attachment.pk])).status_code,404)
        self.assertEqual(self.client.get('/media/'+attachment.file.name).status_code,404)

    def test_pricing_uses_selected_direction_and_pathogenic_rates(self):
        service=Service.objects.get(code='EGTP-SeqS');schema,a,p,rows=fixture(service.code)
        one=resolve_cost(service,'IBTIKAR',rows,{**p,'_ibtikar_schema':'1.0.0'})
        p['sequencing_mode']='both';two=resolve_cost(service,'IBTIKAR',rows,{**p,'_ibtikar_schema':'1.0.0'})
        self.assertGreater(two['total'],one['total'])
        self.assertEqual(projection(schema,a,p,rows)['read_count'],2)
        service=Service.objects.get(code='EGTP-IMT');schema,a,p,rows=fixture(service.code)
        for mode,expected in [('single',4000),('duplicate',7000),('triplicate',10000)]:
            rows[0].update(risk_status='pathogenic',analysis_mode=mode,maldi_target='disposable')
            result=resolve_cost(service,'IBTIKAR',rows,{'_ibtikar_schema':'1.0.0'})
            self.assertEqual(Decimal(result['known_subtotal']),expected)
            self.assertIsNone(result['total'])
            service.pricing_data={'ibtikar':{'disposable_target_price':'500'}};service.save(update_fields=['pricing_data'])
            result=resolve_cost(service,'IBTIKAR',rows,{'_ibtikar_schema':'1.0.0'})
            self.assertEqual(result['total'],expected+500)
            service.pricing_data={};service.save(update_fields=['pricing_data'])

    def test_manual_price_is_not_silently_zero_or_a_fabricated_option(self):
        for code in ('EGTP-GDE','EGTP-PSM','EGTP-Lyoph'):
            schema,a,p,rows=fixture(code);service=Service.objects.get(code=code)
            result=resolve_cost(service,'IBTIKAR',rows,{**p,'_ibtikar_schema':'1.0.0'})
            self.assertIsNone(result['total']);self.assertEqual(result['status'],'pending')
            ServicePricing.objects.create(service=service,channel='IBTIKAR',name='Forfait expressément validé',pricing_type='OVERRIDE',amount=2500)
            result=resolve_cost(service,'IBTIKAR',rows,{**p,'_ibtikar_schema':'1.0.0'})
            self.assertEqual(result['total'],2500)

    def test_estimate_visibility_and_scope_are_enforced(self):
        schema,a,p,rows=fixture('EGTP-CAN');url=reverse('ibtikar:estimate',args=['EGTP-CAN'])
        self.assertFalse(self.client.post(url,posted(a,p,rows)).json()['visible'])
        FinancialVisibility.objects.create(pk=1,show_estimates=True,valid_until=timezone.localdate())
        r=self.client.post(url,posted(a,p,rows));self.assertEqual(r.status_code,200);self.assertTrue(r.json()['visible'])
        Service.objects.filter(code='EGTP-CAN').update(channel_availability='GENOCLAB')
        self.assertEqual(self.client.post(url,posted(a,p,rows)).status_code,404)

    def test_invalid_legacy_service_identifier_is_not_a_500(self):
        r=self.client.post(reverse('dashboard:requester_create'),{'service_id':'invalid'})
        self.assertEqual(r.status_code,400)

    def test_many_samples_document_contains_last_and_first(self):
        from documents.generators import generate_ibtikar_form
        obj=self.submit('EGTP-CAN',12)
        text=document_text(Document(generate_ibtikar_form(obj.request)))
        self.assertIn('S001',text);self.assertIn('S012',text)
        self.assertEqual(len(obj.samples),12)


    def test_ibtikar_print_layout_highlights_requested_service(self):
        from documents.generators import generate_ibtikar_form
        obj = self.submit('EGTP-PCR')
        document = Document(generate_ibtikar_form(obj.request))
        flattened = document_text(document)
        self.assertIn('FICHE DE DEMANDE IBTIKAR', flattened)
        self.assertIn('Service demandé', flattened)
        self.assertIn('Amplification PCR', flattened)
        self.assertIn('1. INFORMATIONS GÉNÉRALES', flattened)
        self.assertIn('2. DEMANDEUR ET PROJET', flattened)
        title_tables = [table for table in document.tables
                        if 'FICHE DE DEMANDE IBTIKAR' in ' '.join(
                            cell.text for row in table.rows for cell in row.cells)]
        self.assertEqual(len(title_tables), 1)
        self.assertIn('Amplification PCR', ' '.join(
            cell.text for row in title_tables[0].rows for cell in row.cells))
