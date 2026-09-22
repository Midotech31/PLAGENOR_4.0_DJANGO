from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import translation
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn

from accounts.models import User, MemberProfile
from core.models import Request, Service
from core import test_ibtikar_forms as contracts
from core.ibtikar.models import IbtikarSubmission
from core.ibtikar.schema import definitions, reference_projection, project_group
from core.ibtikar.services import save_staff
from documents.generators import generate_ibtikar_form
from documents.views import _cached_doc_path
from plagenor.test_support import close_response


@override_settings(STORAGES=contracts.STATIC, EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
    PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], RATE_LIMIT_BACKEND='cache', DOCUMENT_PDF_ENABLED=False)
class HybridIbtikarTests(TestCase):
    setUpTestData = contracts.PersistenceContractTests.__dict__['setUpTestData']
    setUp = contracts.PersistenceContractTests.setUp
    tearDown = contracts.PersistenceContractTests.tearDown
    submit = contracts.PersistenceContractTests.submit

    def test_guest_channel_picker_does_not_create_a_request(self):
        self.client.logout()
        response = self.client.get(reverse('guest_submit'))
        self.assertContains(response, 'data-guest-channel-picker')
        self.assertContains(response, '<option value="IBTIKAR"')
        self.assertNotContains(response, 'ouvrir les formulaires dédiés')
        self.assertNotContains(response, 'name="guest_name"')
        response = self.client.get(reverse('guest_submit'), {'channel': 'IBTIKAR'})
        self.assertContains(response, 'value="IBTIKAR" selected')
        self.assertEqual(Request.objects.count(), 0)
        self.assertEqual(IbtikarSubmission.objects.count(), 0)

    def test_chosen_service_opens_the_canonical_form(self):
        service = Service.objects.get(code='EGTP-CAN')
        response = self.client.get(reverse('guest_submit'), {'channel': 'IBTIKAR', 'service': service.code})
        self.assertRedirects(response, reverse('ibtikar:new', args=[service.code]), fetch_redirect_response=False)
        response = self.client.post(reverse('guest_submit'), {'channel': 'IBTIKAR', 'service_id': str(service.pk)})
        self.assertRedirects(response, reverse('ibtikar:new', args=[service.code]), fetch_redirect_response=False)
        self.assertEqual(Request.objects.count(), 0)

    def test_commercial_service_cannot_be_selected_for_ibtikar(self):
        service = Service.objects.get(code='EGTP-CAN')
        service.channel_availability = 'GENOCLAB'
        service.save(update_fields=['channel_availability'])
        response = self.client.get(reverse('guest_submit'), {'channel': 'IBTIKAR', 'service': service.code})
        self.assertEqual(response.status_code, 404)
        response = self.client.post(reverse('guest_submit'), {'channel': 'IBTIKAR', 'service_id': str(service.pk)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Request.objects.count(), 0)

    def test_commercial_submission_keeps_contact_validation(self):
        service = Service.objects.get(code='EGTP-CAN')
        service.channel_availability = 'BOTH'
        service.save(update_fields=['channel_availability'])
        response = self.client.get(reverse('guest_submit'), {'channel': 'GENOCLAB', 'service': service.code})
        self.assertContains(response, 'name="guest_email"')
        self.assertNotContains(response, 'name="ibtikar_id"')
        response = self.client.post(reverse('guest_submit'), {'channel': 'GENOCLAB', 'service_id': str(service.pk)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Request.objects.count(), 0)

    def test_profile_picker_and_editor_share_authenticated_workspace(self):
        response = self.client.get(reverse('dashboard:requester'), {'tab': 'new'})
        self.assertContains(response, 'id="profile-ibtikar-service"')
        response = self.client.get(reverse('ibtikar:new', args=['EGTP-CAN']))
        self.assertTemplateUsed(response, 'base.html')
        self.assertNotContains(response, 'name="staff-received_date"')
        self.assertEqual(response.context['applicant_form'].initial['email'], self.requester.email)

    def test_operator_blanks_and_checkboxes_are_consistent_for_all_services(self):
        for code in definitions()['services']:
            form = self.submit(code, 2 if code == 'EGTP-PS' else 1)
            for language in ('fr', 'en', 'ar'):
                with self.subTest(code=code, language=language), translation.override(language):
                    projected = reference_projection(form, language)
                    document = Document(generate_ibtikar_form(form.request))
                    tables = [table for table in document.tables if len(table.columns) == 2]
                    staff_labels = {row['label']: row for row in projected['staff']}
                    staff_table = next(table for table in tables if table.rows[0].cells[0].text in staff_labels)
                    self.assertNotIn('Non renseigné', '\n'.join(cell.text for row in staff_table.rows for cell in row.cells))
                    for row in staff_table.rows:
                        spec = staff_labels[row.cells[0].text]
                        self.assertEqual(spec['display'], '')
                        self.assertNotIn('☑', row.cells[1].text)
                        self.assertEqual(row.height_rule, WD_ROW_HEIGHT_RULE.AT_LEAST)
                        self.assertGreaterEqual(row.height.cm, 0.84)
                        if not spec['options']:
                            self.assertEqual(row.cells[1].text, '')
                        else:
                            for choice in spec['options']:
                                self.assertIn(choice['label'], row.cells[1].text)
                    self.assertNotIn('validated_price', {row['name'] for row in projected['staff']})
                    self.assertNotIn('price_justification', {row['name'] for row in projected['staff']})

    def test_zero_and_false_staff_values_are_not_treated_as_missing(self):
        specs = [{'name': 'measurement', 'type': 'decimal', 'label': 'Mesure', 'document_blank': True},
                 {'name': 'confirmed', 'type': 'consent', 'label': 'Vérifié', 'document_blank': True}]
        rows = project_group(specs, {'measurement': 0, 'confirmed': False}, include_empty=True)
        self.assertEqual([row['display'] for row in rows], ['0', 'Non'])
        self.assertFalse(any(row['write_in'] for row in rows))

    def test_operator_revision_preserves_requester_data_and_refreshes_download(self):
        form = self.submit('EGTP-SeqS')
        original = deepcopy((form.applicant, form.parameters, form.samples))
        before = _cached_doc_path(form.request, 'IBTIKAR_FORM')
        download = reverse('documents:ibtikar_form', args=[form.request_id])
        first = self.client.get(download)
        first_text = contracts.document_text(Document(BytesIO(b''.join(first.streaming_content))))
        close_response(first)
        save_staff(form.pk, {'condition': 'Tubes intacts', 'received_quantity': 'Deux tubes de 20 µL'}, self.ops, form.revision)
        form.refresh_from_db()
        save_staff(form.pk, {'comment': 'Contrôle de réception effectué'}, self.ops, form.revision)
        form.refresh_from_db()
        self.assertEqual((form.applicant, form.parameters, form.samples), original)
        self.assertEqual(form.staff['condition'], 'Tubes intacts')
        self.assertEqual(form.staff['received_quantity'], 'Deux tubes de 20 µL')
        self.assertEqual(form.revisions.count(), 3)
        self.assertEqual(form.revisions.get(revision=1).data['staff'], {})
        form.request.refresh_from_db()
        self.assertNotEqual(before, _cached_doc_path(form.request, 'IBTIKAR_FORM'))
        second = self.client.get(download)
        second_text = contracts.document_text(Document(BytesIO(b''.join(second.streaming_content))))
        close_response(second)
        self.assertNotIn('Tubes intacts', first_text)
        self.assertIn('Tubes intacts', second_text)
        self.assertIn('Deux tubes de 20 µL', second_text)
        self.assertIn('Contrôle de réception effectué', second_text)
        detail = self.client.get(reverse('ibtikar:detail', args=[form.request_id]))
        self.assertContains(detail, 'Tubes intacts')
        self.assertTemplateUsed(detail, 'base.html')

    def test_assigned_operator_can_complete_only_the_operator_section(self):
        form = self.submit()
        original = deepcopy(form.applicant)
        operator = User.objects.create_user(username='hybrid-operator', role='MEMBER')
        form.request.assigned_to = MemberProfile.objects.get(user=operator)
        form.request.save(update_fields=['assigned_to'])
        url = reverse('ibtikar:staff', args=[form.request_id])
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_login(operator)
        self.assertTemplateUsed(self.client.get(url), 'base.html')
        response = self.client.post(url, {'revision': form.revision, 'staff-condition': 'Contrôle visuel réalisé'})
        self.assertRedirects(response, reverse('ibtikar:detail', args=[form.request_id]), fetch_redirect_response=False)
        form.refresh_from_db()
        self.assertEqual(form.applicant, original)
        self.assertEqual(form.staff['condition'], 'Contrôle visuel réalisé')
        self.assertEqual(form.staff['operator_id'], operator.pk)
        with self.assertRaises(ValidationError):
            save_staff(form.pk, {'comment': 'Révision périmée'}, operator, form.revision - 1)
        form.request.status = 'ARCHIVED'
        form.request.save(update_fields=['status'])
        with self.assertRaises(ValidationError):
            save_staff(form.pk, {'comment': 'Après archivage'}, operator, form.revision)

    def test_profile_changes_do_not_rewrite_an_existing_request_form(self):
        form = self.submit()
        self.requester.organization = 'Nouvelle organisation du compte'
        self.requester.save(update_fields=['organization'])
        response = self.client.get(reverse('ibtikar:edit', args=[form.request_id]))
        self.assertEqual(response.context['applicant_form'].initial, form.applicant)
        rendered = contracts.document_text(Document(generate_ibtikar_form(form.request)))
        self.assertIn(form.applicant['institution'], rendered)
        self.assertNotIn(self.requester.organization, rendered)
