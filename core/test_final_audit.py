"""Commercial publication, identity, tariff and document regressions."""
import tempfile
from datetime import timedelta
from decimal import Decimal
from pathlib import Path
from django.test import TestCase, override_settings
from django.utils import timezone
from django.urls import reverse
from django.core.exceptions import PermissionDenied, ValidationError
from docx import Document
from accounts.models import User
from core.models import Request, Service, Invoice, FinancialVisibility, ServicePricing
from core.commercial import assign_billing_channel, document_identity, ESSBO_NAME
from core.pricing import resolve_cost
from core.exceptions import PricingConfigurationError


@override_settings(STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}}, PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], PRIVILEGED_MFA_ENFORCEMENT=False)
class FinalFinancialTests(TestCase):
    def setUp(self):
        self.ops = User.objects.create_user(username='ops', role='PLATFORM_ADMIN')
        self.customer = User.objects.create_user(username='customer', role='CLIENT')
        self.service = Service.objects.create(code='AUDIT', name='Audit', genoclab_price=100)
        self.req = Request.objects.create(display_id='GCL-AUDIT', channel='GENOCLAB',
            status='REQUEST_CREATED', service=self.service, requester=self.customer)
        self.client.force_login(self.ops)
        self.quote = {'item_label_0':'Prestation', 'item_unit_price_0':'100',
                      'item_quantity_0':'2', 'vat_rate':'19', 'action':'save'}
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.override = override_settings(MEDIA_ROOT=self.tmp.name)
        self.override.enable()
        self.addCleanup(self.override.disable)

    def test_assignment_exclusive_and_locked(self):
        with self.assertRaises(PermissionDenied):
            assign_billing_channel(self.req.pk, 'OHB', self.customer)
        superadmin = User.objects.create_user(username='super', role='SUPER_ADMIN')
        with self.assertRaises(PermissionDenied):
            assign_billing_channel(self.req.pk, 'OHB', superadmin)
        assign_billing_channel(self.req.pk, 'OHB', self.ops)
        self.req.refresh_from_db()
        self.assertEqual(self.req.billing_channel, 'OHB')
        self.req.status = 'QUOTE_SENT'
        self.req.save()
        with self.assertRaises(ValidationError):
            assign_billing_channel(self.req.pk, 'GENOCLAB', self.ops)

    def test_ohb_quote_invoice_and_identity_snapshot(self):
        assign_billing_channel(self.req.pk, 'OHB', self.ops)
        response = self.client.post(reverse('dashboard:admin_prepare_quote', args=[self.req.pk]), self.quote)
        self.assertEqual(response.status_code, 302)
        self.req.refresh_from_db()
        self.assertEqual(self.req.quote_detail['vat_rate'], 0)
        self.assertEqual(self.req.quote_amount, 200)
        self.req.status = 'ORDER_UPLOADED'
        self.req.save()
        from dashboard.views.admin_ops import generate_invoice
        from django.test import RequestFactory
        from django.contrib.messages.storage.fallback import FallbackStorage
        request = RequestFactory().post('/')
        request.user = self.ops
        request.session = {}
        request._messages = FallbackStorage(request)
        generate_invoice(request, self.req.pk)
        invoice = Invoice.objects.get(request=self.req)
        self.assertTrue(invoice.invoice_number.startswith('OHB-INV'))
        self.assertEqual(invoice.vat_rate, 0)
        from documents.generators import generate_invoice_document
        path = generate_invoice_document(invoice)
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn(ESSBO_NAME, text)
        self.assertIn('non assujetti', text)
        self.assertNotIn('19 %', text)
        self.assertEqual(invoice.document_snapshot['client_name'], 'customer')

    def test_sent_quote_cannot_be_repriced(self):
        self.req.status = 'QUOTE_SENT'
        self.req.quote_amount = 999
        self.req.save()
        r = self.client.post(reverse('dashboard:admin_prepare_quote', args=[self.req.pk]), self.quote)
        self.assertEqual(r.status_code, 403)
        self.req.refresh_from_db()
        self.assertEqual(self.req.quote_amount, 999)

    def test_quote_number_and_date_stable_on_download(self):
        self.client.post(reverse('dashboard:admin_prepare_quote', args=[self.req.pk]), self.quote)
        self.req.refresh_from_db()
        from documents.generators import generate_quote
        generate_quote(self.req)
        number, date = self.req.quote_number, self.req.quote_date
        generate_quote(self.req)
        self.assertEqual(self.req.quote_number, number)
        self.assertEqual(self.req.quote_date, date)

    def test_visibility_requires_explicit_unexpired_publication(self):
        self.assertFalse(FinancialVisibility.estimates_visible())
        p = FinancialVisibility.objects.create(pk=1, show_estimates=True)
        self.assertFalse(FinancialVisibility.estimates_visible())
        p.valid_until = timezone.localdate()
        p.save()
        self.assertTrue(FinancialVisibility.estimates_visible())
        p.valid_until -= timedelta(days=1)
        p.save()
        self.assertFalse(FinancialVisibility.estimates_visible())

    def test_expired_tariff_cannot_fall_back(self):
        ServicePricing.objects.create(service=self.service, name='Past', channel='GENOCLAB', amount=400,
            valid_until=timezone.localdate()-timedelta(days=1))
        with self.assertRaises(PricingConfigurationError):
            resolve_cost(self.service, 'GENOCLAB', sample_table=[{'id':1}])

    def test_zero_vat_document_and_cent_carry(self):
        from documents.generators import generate_invoice_document
        from documents.genoclab_layout import amount_in_words_fr
        invoice = Invoice.objects.create(invoice_number='ZERO', request=self.req, vat_rate=0,
            line_items=[{'description':'Test', 'quantity':1, 'unit_price':100,'total':100}],
            subtotal_ht=100, total_ttc=100)
        doc = Document(generate_invoice_document(invoice))
        text = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn('TVA (0 %)', text)
        self.assertNotIn('119', text)
        self.assertEqual(amount_in_words_fr('1.995'), 'deux')

    def test_visibility_ui_authorization_expiry_and_audit(self):
        url = reverse('dashboard:financial_visibility')
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.post(url, {'show_estimates':'on','valid_until':'invalid'})
        self.assertFalse(FinancialVisibility.estimates_visible())
        self.client.post(url, {'show_estimates':'on','valid_until':timezone.localdate().isoformat()})
        self.assertTrue(FinancialVisibility.estimates_visible())
        self.service.genoclab_price = 333
        self.service.save()
        self.assertFalse(FinancialVisibility.estimates_visible())
        self.client.force_login(self.customer)
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_public_fragment_has_no_prices_when_hidden(self):
        from unittest.mock import patch
        url = reverse('dashboard:service_form_fragment', args=[self.service.code])
        definition={'parameters':[{'name':'mode','label':'Mode','type':'enum', 'options':['Double'],
            'option_pricing':{'Double':987654}}], 'pricing':{'base_price':987654}}
        self.client.logout()
        with patch('dashboard.views.service_form_api.get_service_def', return_value=definition):
            r = self.client.get(url)
        self.assertNotContains(r, '987654')
        self.assertNotContains(r, 'id="cost-estimate-box"')
        self.assertContains(r, 'Double')
        self.assertEqual(definition['parameters'][0]['option_pricing']['Double'],987654)

    def test_assignment_endpoint_cannot_be_called_by_client(self):
        url=reverse('dashboard:admin_assign_billing',args=[self.req.pk])
        self.assertEqual(self.client.get(url).status_code,405)
        self.client.post(url, {'billing_channel':'BAD'})
        self.req.refresh_from_db()
        self.assertEqual(self.req.billing_channel,'GENOCLAB')
        self.client.post(url, {'billing_channel':'OHB'})
        self.req.refresh_from_db()
        self.assertEqual(self.req.billing_channel,'OHB')
        self.client.force_login(self.customer)
        self.assertEqual(self.client.post(url, {'billing_channel':'GENOCLAB'}).status_code,403)

    def test_tariff_period_api_and_effective_price(self):
        url=reverse('dashboard:pricing_add_api',args=[self.service.pk])
        data={'name':'OHB tariff','channel':'OHB','pricing_type':'BASE','amount':'700','is_active':'on',
              'valid_from':timezone.localdate().isoformat(),'valid_until':timezone.localdate().isoformat()}
        r=self.client.post(url,data)
        self.assertEqual(r.status_code,201)
        self.assertEqual(r.json()['config']['valid_until'],data['valid_until'])
        self.assertEqual(resolve_cost(self.service,'OHB',sample_table=[{'id':1}])['total'],700)
        data['valid_until']='invalid'
        self.assertEqual(self.client.post(url,data).status_code,400)
        data['valid_until']=(timezone.localdate()-timedelta(days=1)).isoformat()
        self.assertEqual(self.client.post(url,data).status_code,400)
