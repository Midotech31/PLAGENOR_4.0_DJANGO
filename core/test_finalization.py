"""Financial release, disclosure and SMTP integration regressions."""
from datetime import timedelta
from decimal import Decimal
from email import policy
from email.parser import BytesParser
from pathlib import Path
import socketserver
import tempfile
import threading
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document

from accounts.models import User
from core.models import FinancialSettings, FinancialAudit, Invoice, Request, Service, ServicePricing
from core.exceptions import InvalidTransitionError, AuthorizationError, PricingConfigurationError, FinancialValidationError
from core.pricing import resolve_cost
from core.workflow import transition, force_transition, _send_transition_emails
from documents.financial_snapshot import capture_snapshot
from documents.generators import generate_quote, generate_invoice_document
from notifications.emails import notify_submission_confirmation, send_email_notification


def document_text(path):
    doc = Document(path)
    def contents(container):
        text = [p.text for p in container.paragraphs]
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    text.extend(contents(cell))
        return text
    return '\n'.join(contents(doc))


@override_settings(STORAGES={'default': {'BACKEND':'django.core.files.storage.FileSystemStorage'}, 'staticfiles': {'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}}, PRIVILEGED_MFA_ENFORCEMENT=False, EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class FinalizationTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        setting = override_settings(MEDIA_ROOT=self.temp.name)
        setting.enable(); self.addCleanup(setting.disable)
        self.ops = User.objects.create_user(username='ops-final', role='PLATFORM_ADMIN')
        self.owner = User.objects.create_user(username='client-final', role='CLIENT', email='client@example.test', first_name='Client')
        self.service = Service.objects.create(code='FINAL', name='Prestation', genoclab_price=1000, ibtikar_price=500)
        self.req = Request.objects.create(display_id='FINAL-1', channel='GENOCLAB', status='QUOTE_DRAFT',
            requester=self.owner, service=self.service, quote_amount=1000,
            billing_assigned_by=self.ops, billing_assigned_at=timezone.now(),
            quote_detail={'items':[{'label':'Prestation', 'unit_price':1000, 'quantity':1, 'total':1000}], 'vat_rate':0})

    def test_ohb_assignment_is_exclusive_to_ops_and_drafts_are_recalculated(self):
        url = reverse('dashboard:assign_billing_channel', args=[self.req.pk])
        for role in ['CLIENT', 'SUPER_ADMIN', 'FINANCE', 'MEMBER']:
            user = User.objects.create_user(username='routing-' + role, role=role)
            self.client.force_login(user)
            self.assertEqual(self.client.post(url, {'billing_channel':'OHB'}).status_code, 403)
        self.client.force_login(self.ops)
        generate_quote(self.req)
        self.assertEqual(self.client.post(url, {'billing_channel':'OHB'}).status_code,302)
        self.req.refresh_from_db()
        self.assertEqual(self.req.billing_channel,'OHB')
        self.assertEqual(self.req.billing_assigned_by,self.ops)
        self.assertEqual(self.req.quote_detail,{})
        self.assertEqual(self.req.quote_number,'')
        self.assertTrue(FinancialAudit.objects.filter(action='INVOICE_CHANNEL_ASSIGNED').exists())
        self.client.force_login(self.owner)
        response=self.client.get(reverse('dashboard:client_request_detail',args=[self.req.pk]))
        self.assertNotContains(response,'billing_channel')
        self.assertNotContains(response,'Opérations Hors Budget')

    def test_ohb_quote_invoice_and_identity_are_zero_vat_and_immutable(self):
        self.req.billing_channel='OHB';self.req.save()
        self.client.force_login(self.ops)
        payload={'item_label_0':'Analyse', 'item_unit_price_0':'1234.56', 'item_quantity_0':'2',
                 'admin_fees':'10', 'report_fees':'20', 'vat_rate':'19', 'action':'send'}
        response=self.client.post(reverse('dashboard:admin_prepare_quote',args=[self.req.pk]),payload)
        self.assertEqual(response.status_code,302)
        self.req.refresh_from_db(); self.assertEqual(self.req.status,'QUOTE_SENT')
        self.assertEqual(self.req.quote_detail['vat_rate'],0)
        self.assertEqual(self.req.quote_amount,Decimal('2499.12'))
        quote_text=document_text(generate_quote(self.req))
        self.assertIn('École Supérieure en Sciences Biologiques d’Oran',quote_text)
        self.assertIn('ESSBO-DEV',quote_text)
        self.assertIn('TVA non applicable',quote_text)
        self.client.post(reverse('dashboard:assign_billing_channel',args=[self.req.pk]),{'billing_channel':'GENOCLAB'})
        self.req.refresh_from_db(); self.assertEqual(self.req.billing_channel,'OHB')
        transition(self.req,'QUOTE_VALIDATED_BY_CLIENT',self.owner)
        self.req.order_file='orders/test.pdf';self.req.save(update_fields=['order_file'])
        transition(self.req,'ORDER_UPLOADED',self.owner)
        page=self.client.get(reverse('dashboard:admin_request_detail',args=[self.req.pk]))
        self.assertContains(page,reverse('dashboard:admin_generate_invoice',args=[self.req.pk]))
        self.client.post(reverse('dashboard:admin_generate_invoice',args=[self.req.pk]))
        invoice=self.req.invoice_set.get()
        self.assertEqual(invoice.billing_channel,'OHB');self.assertEqual(invoice.vat_amount,0)
        self.assertEqual(invoice.total_ttc,Decimal('2499.12'))
        text=document_text(generate_invoice_document(invoice))
        self.assertIn('École Supérieure en Sciences Biologiques d’Oran',text)
        self.assertIn('ESSBO-INV',text)
        self.req.billing_channel='GENOCLAB'
        with self.assertRaises(ValidationError):self.req.save()

    def test_estimator_uses_current_channel_tariff_and_hides_amount_when_disabled(self):
        url=reverse('dashboard:estimate',args=[self.service.code])
        self.assertEqual(self.client.post(url,{'channel':'GENOCLAB'}).json()['total'],'1000.00')
        self.assertEqual(self.client.post(url,{'channel':'IBTIKAR'}).json()['total'],'500.00')
        self.assertEqual(self.client.post(url,{'channel':'OHB'}).status_code,400)
        FinancialSettings.objects.create(show_estimates=False)
        response=self.client.post(url,{'channel':'GENOCLAB'})
        self.assertEqual(response.json(),{'visible':False})
        self.assertEqual(response['Cache-Control'],'private, no-store')

    def test_payment_order_never_confirms_payment_and_private_upload_is_owned(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        Invoice.objects.create(invoice_number='PO-1',request=self.req,client=self.owner,vat_rate=0)
        self.req.status='INVOICE_GENERATED';self.req.save()
        self.client.force_login(self.owner)
        response=self.client.post(reverse('dashboard:client_upload_payment_order',args=[self.req.pk]), {
            'payment_order_reference':'OP-2026-1','payment_order_date':timezone.localdate().isoformat(),
            'payment_order_file':SimpleUploadedFile('order.pdf',b'%PDF-1.4\n%%EOF',content_type='application/pdf')})
        self.assertEqual(response.status_code,302)
        self.req.refresh_from_db()
        self.assertEqual(self.req.payment_order_reference,'OP-2026-1')
        self.assertEqual(self.req.status,'INVOICE_GENERATED')
        self.assertIsNone(self.req.payment_verified_at)
        self.assertEqual(self.req.invoice_set.get().payment_status,'PENDING')
        other=User.objects.create_user(username='other-client',role='CLIENT')
        self.client.force_login(other)
        response=self.client.get('/media/'+self.req.payment_order_file.name)
        self.assertIn(response.status_code,[403,404])

    def test_revision_retires_old_number_and_preserves_old_content(self):
        transition(self.req,'QUOTE_SENT',self.ops)
        number=self.req.quote_number
        transition(self.req,'QUOTE_REJECTED_BY_CLIENT',self.owner)
        transition(self.req,'QUOTE_DRAFT',self.ops,notes='Prix révisé')
        self.assertEqual(self.req.quote_number,'')
        self.assertEqual(FinancialAudit.objects.get(action='QUOTE_REVISION').details['number'],number)
        transition(self.req,'QUOTE_SENT',self.ops)
        self.assertNotEqual(self.req.quote_number,number)

    def test_draft_finances_and_downloads_hidden_from_client(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse('dashboard:client_request_detail', args=[self.req.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, '1000')
        for route in ['dashboard:download_quote', 'documents:quote']:
            self.assertEqual(self.client.get(reverse(route, args=[self.req.pk])).status_code, 403)
        self.req.refresh_from_db()
        self.assertEqual(self.req.quote_amount, Decimal('1000'))

    def test_policy_hides_pricing_attributes_without_removing_questions(self):
        self.service.code='egtp_imt'; self.service.save()
        FinancialSettings.objects.create(show_estimates=False)
        response=self.client.get(reverse('dashboard:service_form_fragment', args=[self.service.code]))
        self.assertEqual(response.status_code,200)
        self.assertNotContains(response, 'data-pricing-value=')
        self.assertNotContains(response, 'data-option-pricing=')
        self.assertNotContains(response, 'id="cost-estimate-box"')
        self.assertContains(response, 'Estimation en cours')
        self.assertEqual(response['Cache-Control'],'private, no-store')

    def test_policy_and_tariff_management_require_admin_and_record_history(self):
        url=reverse('dashboard:financial_settings')
        self.client.force_login(self.owner)
        self.assertEqual(self.client.post(url, {}).status_code,403)
        self.client.force_login(self.ops)
        self.assertEqual(self.client.get(url).status_code,200)
        self.assertEqual(self.client.post(url, {'invoice_payment_method':'Virement bancaire','invoice_payment_days':'30'}).status_code,302)
        self.assertFalse(FinancialSettings.objects.get().show_estimates)
        self.assertTrue(FinancialAudit.objects.filter(action='ESTIMATE_POLICY').exists())
        self.assertEqual(self.client.get(reverse('dashboard:superadmin_service_edit',args=[self.service.pk])).status_code,200)
        self.assertEqual(self.client.get(reverse('dashboard:superadmin')).status_code,403)

    def test_quote_number_is_stable_and_issued_document_identity_frozen(self):
        first=generate_quote(self.req)
        self.req.refresh_from_db(); number=self.req.quote_number
        self.assertTrue(number)
        generate_quote(self.req); self.req.refresh_from_db()
        self.assertEqual(self.req.quote_number,number)
        transition(self.req,'QUOTE_SENT',self.ops)
        self.req.refresh_from_db()
        self.owner.first_name='Changed'; self.owner.save()
        from core.models import PlatformContent
        PlatformContent.objects.create(key='genoclab_issuer_name',lang='fr',value='Changed issuer')
        self.req.refresh_from_db()
        text=document_text(generate_quote(self.req))
        self.assertIn(number,text); self.assertNotIn('Changed issuer',text); self.assertNotIn('Changed',text)
        self.client.force_login(self.ops)
        self.assertEqual(self.client.post(reverse('dashboard:admin_prepare_quote',args=[self.req.pk]),{}).status_code,403)
        self.client.post(reverse('dashboard:admin_adjust_cost',args=[self.req.pk]),{'admin_price':'5','cost_justification':'Correction documentée'})
        self.req.refresh_from_db(); self.assertIsNone(self.req.admin_validated_price)

    def test_zero_vat_and_invoice_immutability(self):
        invoice=Invoice.objects.create(invoice_number='INV-FINAL',client=self.owner,request=self.req,
            line_items=[{'description':'Prestation','quantity':1,'unit_price':1000,'total':1000}],
            subtotal_ht=1000,vat_rate=0,vat_amount=0,total_ttc=1000,document_snapshot=capture_snapshot(self.req))
        text=document_text(generate_invoice_document(invoice))
        self.assertNotIn('19%',text.replace(' ',''))
        self.assertNotIn('1 190',text)
        invoice.total_ttc=1190
        with self.assertRaises(ValidationError): invoice.save()
        invoice.refresh_from_db(); invoice.payment_status='COMPLETED'; invoice.save()
        self.assertEqual(invoice.total_ttc,1000)

    def test_expired_rates_never_fall_back_to_stale_flat_price(self):
        tier=ServicePricing.objects.create(service=self.service,channel='GENOCLAB',name='Ancien',amount=90,
            valid_until=timezone.localdate()-timedelta(days=1))
        with self.assertRaises(PricingConfigurationError): resolve_cost(self.service,'GENOCLAB')
        tier.valid_until=timezone.localdate();tier.save()
        self.assertEqual(resolve_cost(self.service,'GENOCLAB')['total'],90)
        self.assertEqual(resolve_cost(self.service,'IBTIKAR')['total'],500)

    def test_expired_quote_and_unauthorized_force_are_rejected(self):
        self.req.status='QUOTE_SENT';self.req.quote_valid_until=timezone.localdate()-timedelta(days=1);self.req.save()
        with self.assertRaises(InvalidTransitionError): transition(self.req,'QUOTE_VALIDATED_BY_CLIENT',self.owner)
        with self.assertRaises(AuthorizationError): force_transition(self.req,'COMPLETED',self.ops)

    def test_dedicated_email_is_not_duplicated_by_generic_email(self):
        with patch('notifications.emails.notify_appointment') as appointment, patch('notifications.emails.notify_status_change') as generic:
            _send_transition_emails(self.req,'APPOINTMENT_PROPOSED','APPOINTMENT_CONFIRMED')
        appointment.assert_called_once(); generic.assert_not_called()
        with patch('notifications.emails.notify_report_delivery') as report:
            _send_transition_emails(self.req,'REPORT_UPLOADED','REPORT_VALIDATED')
        report.assert_not_called()

    def test_real_smtp_transport_contains_absolute_links_and_plain_text(self):
        captured=[]
        class Handler(socketserver.StreamRequestHandler):
            def handle(self):
                self.wfile.write(b'220 localhost ESMTP\r\n')
                while True:
                    line=self.rfile.readline()
                    if not line: break
                    verb=line.split(b' ',1)[0].strip().upper()
                    if verb in (b'EHLO',b'HELO'): self.wfile.write(b'250 localhost\r\n')
                    elif verb==b'DATA':
                        self.wfile.write(b'354 End with dot\r\n')
                        data=[]
                        while True:
                            row=self.rfile.readline()
                            if row in (b'.\r\n',b''):break
                            data.append(row[1:] if row.startswith(b'..') else row)
                        captured.append(b''.join(data)); self.wfile.write(b'250 accepted\r\n')
                    elif verb==b'QUIT': self.wfile.write(b'221 bye\r\n');break
                    else:self.wfile.write(b'250 OK\r\n')
        with socketserver.TCPServer(('127.0.0.1',0),Handler) as server:
            thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
            try:
                with self.settings(EMAIL_BACKEND='django.core.mail.backends.smtp.EmailBackend',EMAIL_HOST='127.0.0.1',
                    EMAIL_PORT=server.server_address[1],EMAIL_USE_TLS=False,EMAIL_USE_SSL=False,
                    EMAIL_HOST_USER='',EMAIL_HOST_PASSWORD='',EMAIL_TIMEOUT=3,PUBLIC_BASE_URL='https://plagenor.example.test'):
                    notify_submission_confirmation(self.req)
            finally:server.shutdown();thread.join(timeout=3)
        self.assertEqual(len(captured),1)
        message=BytesParser(policy=policy.default).parsebytes(captured[0])
        self.assertEqual(message['To'],'client@example.test')
        self.assertTrue(message.get_body(preferencelist=('plain',)).get_content().strip())
        html=message.get_body(preferencelist=('html',)).get_content()
        self.assertIn('https://plagenor.example.test/dashboard/client/request/',html)
        self.assertNotIn('href=""',html)

    def test_live_estimator_applies_canonical_field_modifiers_and_rejects_invalid_values(self):
        from core.models import ServiceFormField
        field = ServiceFormField.objects.create(service=self.service,name='mode',label='Option',field_type='enum',
            affects_pricing=True, price_modifier_type='multiply',price_modifier_value=2,
            option_pricing={'triple': 3})
        self.assertEqual(resolve_cost(self.service, 'GENOCLAB',service_params={'mode':'triple'})['total'], 3000)
        field.price_modifier_type='add';field.save()
        self.assertEqual(resolve_cost(self.service,'GENOCLAB',service_params={'mode':'other'})['total'],1002)
        field.price_modifier_type='set';field.save()
        self.assertEqual(resolve_cost(self.service,'GENOCLAB',service_params={'mode':'other'})['total'],2)
        field.field_type='boolean';field.save()
        self.assertEqual(resolve_cost(self.service,'GENOCLAB',service_params={'mode':'false'})['total'],1000)
        field.option_pricing={'bad':-1};field.save()
        with self.assertRaises(PricingConfigurationError):resolve_cost(self.service,'GENOCLAB',service_params={'mode':'bad'})
        field.option_pricing={};field.price_modifier_type='unknown';field.save()
        with self.assertRaises(PricingConfigurationError):resolve_cost(self.service,'GENOCLAB',service_params={'mode':'true'})

    def test_fractional_vat_rounding_survives_database_and_document_generation(self):
        from core.financial import compute_invoice_totals
        items=[{'description':'Arrondi','quantity':3,'unit_price':'9.99','total':'29.97'}]
        totals=compute_invoice_totals(items,vat_rate=Decimal('.195'))
        self.assertEqual(Decimal(str(totals['subtotal_ht'])),Decimal('29.97'))
        self.assertEqual(Decimal(str(totals['vat_amount'])),Decimal('5.84'))
        self.assertEqual(Decimal(str(totals['total_ttc'])),Decimal('35.81'))
        inv=Invoice.objects.create(invoice_number='INV-FRACTIONAL',client=self.owner,request=self.req,
            line_items=items,subtotal_ht=totals['subtotal_ht'],vat_rate=Decimal('.195'),
            vat_amount=totals['vat_amount'],total_ttc=totals['total_ttc'],document_snapshot=capture_snapshot(self.req))
        inv.refresh_from_db();self.assertEqual(inv.vat_rate,Decimal('.1950'))
        self.assertIn('19.5%',document_text(generate_invoice_document(inv)).replace(' ',''))
        for rate in ['NaN','-0.1','1.1','0.12345']:
            with self.assertRaises(FinancialValidationError):compute_invoice_totals([],vat_rate=rate)

    def test_legacy_documents_gain_dated_archives_and_cannot_change_amounts(self):
        from documents.financial_snapshot import preserve_legacy_snapshot
        self.req.status='QUOTE_SENT';self.req.save(update_fields=['status'])
        original=generate_quote(self.req)
        self.req.refresh_from_db()
        self.assertIn('legacy_captured_at',self.req.quote_snapshot)
        self.req.quote_amount=5
        with self.assertRaises(ValidationError):self.req.save()
        self.req.refresh_from_db()
        inv=Invoice.objects.create(invoice_number='INV-LEGACY',client=self.owner,request=self.req,
            line_items=[{'description':'Legacy','quantity':1,'unit_price':1000,'total':1000}],
            subtotal_ht=1000,vat_rate=0,vat_amount=0,total_ttc=1000)
        preserve_legacy_snapshot(inv);inv.refresh_from_db()
        self.assertIn('legacy_captured_at',inv.document_snapshot)
        self.owner.first_name='New identity';self.owner.save()
        self.assertNotIn('New identity',document_text(generate_invoice_document(inv)))
        self.assertEqual(FinancialAudit.objects.filter(action='INVOICE_LEGACY_SNAPSHOT').count(),1)

    def test_issuer_configuration_is_separate_and_does_not_change_existing_snapshots(self):
        from core.models import PlatformContent
        from dashboard.views.financial_settings import IssuerForm
        from documents.genoclab_layout import CMS_DEFAULTS
        self.client.force_login(self.ops)
        for channel in ['genoclab','ohb']:
            payload={'action':channel}
            payload.update({channel+'-'+key:CMS_DEFAULTS.get('genoclab_'+key,'') or 'Valeur de test' for key in IssuerForm.base_fields})
            payload[channel+'-issuer_name']='Issuer '+channel
            response=self.client.post(reverse('dashboard:financial_settings'),payload)
            self.assertEqual(response.status_code,302)
        self.assertEqual(PlatformContent.objects.get(key='ohb_issuer_name',lang='fr').value,'École Supérieure en Sciences Biologiques d’Oran')
        self.assertEqual(PlatformContent.objects.get(key='genoclab_issuer_name',lang='fr').value,'Issuer genoclab')
        self.assertEqual(FinancialAudit.objects.filter(action='INVOICE_ISSUER_SETTINGS').count(),2)
        bad=self.client.post(reverse('dashboard:financial_settings'),{'action':'genoclab','genoclab-issuer_name':''})
        self.assertEqual(bad.status_code,200)
        self.assertTrue(bad.context['issuer_forms'][0]['form'].errors)

    def test_admin_email_editor_and_real_templates_in_three_languages(self):
        from django.core import mail
        from notifications.emails import notify_status_change
        self.client.force_login(self.owner)
        route=reverse('dashboard:email_templates')
        self.assertEqual(self.client.get(route).status_code,403)
        self.client.force_login(self.ops)
        for language in ['fr','en','ar']:
            url=route+'?language='+language+'&event=request_status_change'
            self.assertEqual(self.client.get(url).status_code,200)
            response=self.client.post(url,{'subject':'Notification '+language,'body':'Texte '+language+' <script>alert(1)</script>'})
            self.assertEqual(response.status_code,200)
            self.assertFalse(response.context['form'].errors)
            self.owner.preferred_language=language;self.owner.save()
            self.assertTrue(notify_status_change(self.req,'REQUEST_CREATED','QUOTE_SENT'))
            message=mail.outbox[-1]
            self.assertIn('Notification '+language,message.subject)
            html=message.alternatives[0].content
            self.assertIn('lang="'+language+'"',html)
            self.assertNotIn('<script>',html)
            self.assertIn('https://',html)
        self.assertEqual(FinancialAudit.objects.filter(action='NOTIFICATION_TEMPLATE_UPDATED').count(),3)
        self.assertEqual(self.client.get(route+'?language=xx').status_code,403)
        self.assertEqual(self.client.get(route+'?event=bad').status_code,403)
        response=self.client.post(route,{'subject':'Bad\nSubject','body':'Test'})
        self.assertTrue(response.context['form'].errors)

    def test_email_failure_deduplication_and_role_appropriate_links(self):
        from notifications.emails import _email_ctx
        from django.core import mail
        self.assertFalse(send_email_notification([], 'Subject','<p>Test</p>'))
        self.assertTrue(send_email_notification([' client@example.test ','client@example.test'],'Subject','<p>Text</p>'))
        self.assertEqual(mail.outbox[-1].to,['client@example.test'])
        with patch('notifications.emails.send_mail',return_value=0):
            self.assertFalse(send_email_notification('client@example.test','Subject','Test'))
        with patch('notifications.emails.send_mail',side_effect=OSError('SMTP unavailable')):
            self.assertFalse(send_email_notification('client@example.test','Subject','Test'))
        self.assertIn('/dashboard/ops/',_email_ctx(self.req,recipient=self.ops)['dashboard_url'])
        self.owner.email='';self.owner.save()
        self.assertFalse(notify_submission_confirmation(self.req))

    def test_ibtikar_home_scope_and_revision_header(self):
        from core.models import PlatformContent
        for language,word in [('fr','tous les étudiants algériens'),('en','all Algerian students'),('ar','لجميع الطلبة الجزائريين')]:
            self.assertIn(word,PlatformContent.objects.get(key='ibtikar_description',lang=language).value)
        with patch.dict('os.environ',{'RENDER_GIT_COMMIT':'a'*40}):
            response=self.client.get('/readyz')
            self.assertEqual(response.status_code,200)
            self.assertEqual(response['X-PLAGENOR-Revision'],'a'*40)
        with patch.dict('os.environ',{'RENDER_GIT_COMMIT':'invalid'}):
            self.assertNotIn('X-PLAGENOR-Revision',self.client.get('/healthz'))

    @override_settings(ROOT_URLCONF='plagenor.urls_e2e')
    def test_financial_browser_fixture_is_loopback_only_and_absent_in_production(self):
        self.ops.username='admin_ops';self.ops.save()
        self.owner.username='client';self.owner.save()
        self.client.force_login(self.ops)
        url='/__e2e__/financial-request/'
        self.assertEqual(self.client.post(url,REMOTE_ADDR='203.0.113.1').status_code,404)
        self.assertEqual(self.client.get(url).status_code,404)
        response=self.client.post(url)
        self.assertEqual(response.status_code,200)
        self.assertEqual(Request.objects.get(pk=response.json()['id']).status,'REQUEST_CREATED')
        with self.settings(ROOT_URLCONF='plagenor.urls'):
            self.assertEqual(self.client.post(url).status_code,404)
